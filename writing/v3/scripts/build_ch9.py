#!/usr/bin/env python3
"""Build writing/v3/Chapter_9_Conclusion.docx.

v3 rewrite per writing/v3/STRUCTURE.md: the Conclusion is short and general.
The v2 chapter's per-chapter recap and specific result numbers are gone (they
all remain stated in Chapters 6 and 7); the systems-engineering argument is
kept in qualitative form; the closing remarks are kept. No subsections.
Per-chapter References section dropped (skill rule 12); the references dict
is kept for build_thesis.py.
"""
from docx import Document
from docx.shared import Pt

H1, H2, H3, P = "h1", "h2", "h3", "p"

content = [
(H1, "Chapter 9: Conclusion"),

(P, "This thesis set out to reconstruct, in a metric virtual scene, a person's upper body together with the object their two hands manipulate, and to know how far the reconstruction can be trusted. The constraint was deliberate: one commodity depth camera [9], three printed paper markers, and no instrumentation worn on the body. Everything else is software. The system that closes the thesis meets that goal. A levelled, metrically calibrated virtual scene contains an avatar that tracks the person's upper body while a carried object follows its independently measured trajectory, and every held, bridged, or unmeasured quantity is declared as such all the way to the screen."),

(P, "The route to that result runs through the chapters as a chain of design decisions rather than a collection of parts. A kinematic model built entirely in rotation matrices turns measured landmarks into joint angles, with Euler angles admitted only at the display boundary. Two processing pipelines, one for the person and one for the scene and object, were kept strictly independent, sharing no code and no intermediate data, and were joined only in a calibrated common world frame anchored to the desk. That independence repaid its cost at evaluation time. Because the two chains rest on different physics, their agreement on the carried object is a genuine measurement, and the system was able to grade its own accuracy, at the centimetre level where sight lines are clean, without any ground-truth equipment. The same system then ran causally at the recording's own pace, meeting the frame budget with margin, and the price of causality was measured rather than assumed. The specific results, and the limits with their causes and sizes, are stated in Chapters 6 through 8."),

(P, "The wider argument of this thesis concerns how such systems should be built now. For half a century, computing improved mainly because the hardware underneath it improved, and system designers could treat next year's hardware as a free performance upgrade. That era is ending. As transistor miniaturization approaches its physical limits, the consensus in the architecture community is that future performance must come from the top of the computing stack, from software, algorithms, and system architecture, rather than from the bottom [51]. Industry has said the same thing more bluntly: NVIDIA's chief executive declared in 2022 that Moore's law is dead [52]."),

(P, "The consequence for anyone building a measurement system is simple to state. Capability is no longer delivered automatically by better components; it has to be engineered out of the components already available. This is the core concern of systems engineering: under a fixed resource budget, find the arrangement of parts whose whole exceeds the sum of the parts, the arrangement that makes one plus one larger than two. Given components A, B, and C, the design question is no longer which more powerful component D to add. It is which arrangement of A, B, and C exposes a capability that none of them has alone."),

(P, "Human motion tracking offers a clear instance of the two competing philosophies. The instrumentation route adds hardware until the problem becomes easy: inertial sensors strapped to every body segment [53], instrumented gloves on the hands [18], or multi-sensor laboratory installations [19]. Each added sensor measures directly what would otherwise have to be inferred, and each adds cost, setup time, and intrusion on the movement being studied. This thesis took the other route. Every component is a commodity or free: a depth camera sold for a few hundred dollars [9], markers printed on paper, an open source pose estimation network [39], an open source computer vision library, and a game engine. None of these tools was designed for the task at hand. The engineering effort went not into acquiring stronger components but into the arrangement: independent pipelines, a calibrated common frame, causal variants of every filter, and explicit declaration of everything not measured. The claim that the whole exceeded the sum of its parts is not rhetorical. The evaluation of Chapter 6, where the two pipelines grade each other, and the real-time study of Chapter 7, where careful structure outperformed added parallel hardware, are both measured instances of it."),

(P, "The trade this system makes is equally plain. It does not reach the accuracy of a marker-based capture studio, and it does not need to: it degrades in ways whose laws are derived and whose sizes are measured, and it announces its own failures. Low cost with acceptable, quantified accuracy is not a lesser version of the instrumented approach. It is a different point on the design curve, and as hardware improvement slows, it is the point that careful arrangement of available tools can keep pushing forward."),

(P, "The specific measurements belong to this sensor, this room, and this recording. Two things in the work are more portable. The first is the method: derive each error's law, measure its size on real data, and make the system announce what it does not know. The second is the design stance just argued: when better hardware can no longer be assumed, capability must be composed from what is available, and the composition itself, done carefully, is where the capability comes from. This thesis is one worked example of that stance, and the part this work would most hope to see reused."),
]

references = {
9:  'L. Keselman, J. Iselin Woodfill, A. Grunnet-Jepsen, and A. Bhowmik, "Intel RealSense stereoscopic depth cameras," in Proc. IEEE Conf. Computer Vision and Pattern Recognition Workshops (CVPRW), 2017.',
18: 'R. Y. Wang and J. Popovic, "Real-time hand-tracking with a color glove," ACM Transactions on Graphics, vol. 28, no. 3, pp. 1-8, 2009.',
19: 'C. Diaz and S. Payandeh, "Multimodal sensing interface for haptic interaction," Journal of Sensors, vol. 2017, art. 2072951, pp. 1-24, 2017, doi: 10.1155/2017/2072951.',
39: 'V. Bazarevsky, I. Grishchenko, K. Raveendran, T. Zhu, F. Zhang, and M. Grundmann, "BlazePose: On-device real-time body pose tracking," arXiv:2006.10204, 2020.',
51: 'C. E. Leiserson, N. C. Thompson, J. S. Emer, B. C. Kuszmaul, B. W. Lampson, D. Sanchez, and T. B. Schardl, "There\'s plenty of room at the Top: What will drive computer performance after Moore\'s law?," Science, vol. 368, no. 6495, art. eaam9744, 2020.',
52: 'D. Takahashi, "Jensen Huang Q&A: Why Moore\'s Law is dead, and smart design is replacing it," GamesBeat/VentureBeat, Sep. 21, 2022. [Online]. Available: https://gamesbeat.com/jensen-huang-qa-why-moores-law-is-dead-and-smart-design-is-replacing-it/',
53: 'I. H. Lopez-Nava and A. Munoz-Melendez, "Wearable inertial sensors for human motion analysis: A review," IEEE Sensors Journal, vol. 16, no. 22, pp. 7821-7834, 2016.',
}

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

for kind, text in content:
    if kind == H1:
        doc.add_heading(text, level=1)
    elif kind == H2:
        doc.add_heading(text, level=2)
    elif kind == H3:
        doc.add_heading(text, level=3)
    else:
        doc.add_paragraph(text)

out = "/home/luo/Desktop/New_SandBox/writing/v3/Chapter_9_Conclusion.docx"
doc.save(out)
print("saved", out, "| words:", sum(len(t.split()) for k, t in content if k == "p"))
