#!/usr/bin/env python3
"""Build writing/v4/Thesis_V4.docx: the merged full thesis.

Order per writing/v4/STRUCTURE.md: Chapters 1-9, then one consolidated
References section, then the Appendix. This script performs the single
citation renumbering pass decided on 2026-07-28:
  - citations are renumbered by order of first appearance in the final text;
  - the meaning of an old number is chapter-local (each chapter script's
    references dict), which resolves the old [15] collision: the Chapter 1
    paper and the Chapter 3 paper are different entries and end up with
    different final numbers;
  - the two author-list variants of the old [26] (MediaPipe Hands) are
    unified to the full-author form.
It also prints the old-to-new mapping used to update references.md.

Chapter content is imported from the sibling build_chN.py / build_appendix.py
scripts (importing them rebuilds the per-chapter docx as a side effect, which
is harmless and keeps everything in sync).
"""
import importlib.util
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from eqn import add_display_eq, add_display_math

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = "/home/luo/Desktop/New_SandBox/writing/v4/Thesis_V4.docx"
OUT_P1 = "/home/luo/Desktop/New_SandBox/writing/v4/Thesis_V4_Part1_Ch1-3.docx"
OUT_P2 = "/home/luo/Desktop/New_SandBox/writing/v4/Thesis_V4_Part2_Ch4-9_App.docx"

H1, H2, H3, P, IMG, CAP, TBL, EQ = "h1", "h2", "h3", "p", "img", "cap", "tbl", "eq"
MATH = "mth"  # displayed, unnumbered worked-example mathematics

MODULES = [f"build_ch{n}" for n in range(1, 10)] + ["build_appendix"]

# The old [26] appears in two author-list variants; canonicalize to the full form.
MP_HANDS_SHORT = 'F. Zhang et al., "MediaPipe Hands: On-device real-time hand tracking," arXiv:2006.10214, 2020.'
MP_HANDS_FULL = ('F. Zhang, V. Bazarevsky, A. Vakunov, A. Tkachenka, G. Sung, C.-L. Chang, '
                 'and M. Grundmann, "MediaPipe Hands: On-device real-time hand tracking," '
                 'arXiv:2006.10214, 2020.')


def load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def canonical(ref):
    return MP_HANDS_FULL if ref == MP_HANDS_SHORT else ref


CITE = re.compile(r"\[(\d+)\]")

chapters = [load(m) for m in MODULES]

new_num = {}      # canonical ref string -> final number
final_refs = {}   # final number -> ref string
old_to_new = []   # (module, old, new) for the report

merged = []
for mod in chapters:
    refs = {n: canonical(s) for n, s in mod.references.items()}
    remap = {}

    def sub(match):
        old = int(match.group(1))
        assert old in refs, f"{mod.__name__}: cites [{old}] but its references dict has no entry"
        key = refs[old]
        if key not in new_num:
            new_num[key] = len(new_num) + 1
            final_refs[new_num[key]] = key
        if old not in remap:
            remap[old] = new_num[key]
            old_to_new.append((mod.__name__, old, new_num[key]))
        return f"[{remap[old]}]"

    for item in mod.content:
        kind = item[0]
        if kind in (P, CAP, H1, H2, H3):
            merged.append((kind, CITE.sub(sub, item[1])))
        elif kind == TBL:
            merged.append((TBL, [[CITE.sub(sub, c) for c in row] for row in item[1]]))
        else:
            merged.append(item)
    if mod.__name__ == "build_ch9":
        merged.append((H1, "References"))
        # placeholder; the reference paragraphs are inserted after the scan of
        # the appendix module completes (its first-appearance numbers must be
        # known), so remember the insertion point.
        merged.append(("REFS_HERE",))

# splice the finished reference list into the placeholder position
idx = merged.index(("REFS_HERE",))
ref_items = [(P, f"[{n}] {final_refs[n]}") for n in sorted(final_refs)]
merged[idx:idx + 1] = ref_items

def render(items, out):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for item in items:
        kind = item[0]
        if kind == H1:
            doc.add_heading(item[1], level=1)
        elif kind == H2:
            doc.add_heading(item[1], level=2)
        elif kind == H3:
            doc.add_heading(item[1], level=3)
        elif kind == P:
            doc.add_paragraph(item[1])
        elif kind == EQ:
            add_display_eq(doc, item[1], item[2])
        elif kind == MATH:
            add_display_math(doc, item[1])
        elif kind == IMG:
            doc.add_picture(item[1], width=Inches(item[2]))
        elif kind == CAP:
            p = doc.add_paragraph(item[1])
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.italic = True
        elif kind == TBL:
            rows = item[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    cell = t.cell(i, j)
                    cell.text = c
                    for par in cell.paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(10)
                            if i == 0:
                                run.font.bold = True

    doc.save(out)
    print("saved", out)


render(merged, OUT)

# Two-part delivery split (user request, 2026-07-28): Part 1 is Chapters 1-3,
# Part 2 is Chapters 4-9 with References and the Appendix. Citation numbers
# stay global (first appearance over the whole thesis), so the parts read
# identically to the merged document. Part 1 gets its own copy of the full
# consolidated References so its citations resolve.
ch4_idx = next(i for i, it in enumerate(merged)
               if it[0] == H1 and it[1].startswith("Chapter 4"))
part1 = merged[:ch4_idx] + [(H1, "References")] + ref_items
part2 = merged[ch4_idx:]
render(part1, OUT_P1)
render(part2, OUT_P2)
print("final reference count:", len(final_refs))
print("\nold-to-new map (first use per module):")
for name, old, new in old_to_new:
    print(f"  {name}: [{old}] -> [{new}]")
