#!/usr/bin/env python3
"""Build writing/v4/Chapter_7_Real_Time_Feasibility.docx.

v3 port: content carried per the approved v2-to-v3 map; per-chapter References
section dropped (skill rule 12); references dict kept for build_thesis.py.

All numbers come from the pinned real-time artifacts (realtime/REALTIME.md,
realtime/dataset/validate_realtime_output.txt, bench_solver.csv/png,
dataset README); nothing is invented and nothing is recomputed. Structure
follows ToC v3 sections 7.1-7.5.
"""
from docx import Document
from docx.shared import Pt, Inches

FIG = "/home/luo/Desktop/New_SandBox/writing/v2/figures/"
H1, H2, H3, P, IMG, CAP, TBL = "h1", "h2", "h3", "p", "img", "cap", "tbl"

content = [
(H1, "Chapter 7: Real-Time Feasibility and Trade-off Analysis"),

(P, "Every result so far was computed offline, with the luxury of seeing the whole recording at once. This chapter removes that luxury: the same two pipelines are run frame by frame, at camera rate, with output appearing while the motion is still happening, and the accuracy cost of the change is quantified. The offline pipelines are left untouched; a separate real-time stack is built beside them that reuses their validated mathematics, replaces only the parts that are illegal in a live setting, and streams into the same Unity receiver as Chapter 5. Because both stacks process the identical recording, every difference between their outputs is attributable to causality alone."),

(P, "One framing point should be stated before anything else. Real time here means the evaluation recording played back at its recorded pacing as a stand-in for a live camera: frames become available one at a time, at 30 per second, and each must be processed before the next arrives, with no access to the future and no second pass. This is exactly the constraint a live sensor imposes, applied to data whose offline ground truth is already pinned, and that pinned ground truth makes the accuracy comparison of Section 7.3 possible. Section 7.1 describes the architecture, Section 7.2 measures the per-stage latency and the value of further acceleration, Section 7.3 quantifies the accuracy cost of causal processing and the latency dial behind it, Section 7.4 examines the smoothing substitutions and the two stability failure modes that only exist live, and Section 7.5 draws the offline-versus-real-time comparison together."),

(H2, "7.1 Real-Time Pipeline Architecture"),

(P, "The real-time stack is three processes plus Unity, shown in Figure 7.1. Process A is the person pipeline: for each arriving frame it aligns depth to colour, runs the pose detector, deprojects the gated landmarks, applies the causal filter of Section 7.4, and solves the thirteen joint angles, publishing the result to a shared memory channel as an 84-byte packet. The packet type is new in name only: it is exactly the person half of the integrated packet of Chapter 5, the frame index, timestamp, pelvis position, thirteen angles, and per-joint live mask. Process B is the object pipeline: marker detection with corner refinement, planar pose estimation, its own causal filter, and the map into the desk world, published as the unchanged 44-byte object packet. The two processes share no code and no data, preserving the independence that Chapter 6's evaluation depends on; even the causal filter class is deliberately duplicated into each side rather than imported from a common module. A third process, the merger, is the only component allowed to read both channels: it publishes the frozen scene calibration once and then, each frame, pairs the newest person and object records into the same 108-byte integrated packet as the offline replay. Unity runs the Chapter 5 receiver without a single change, which is the point: from the display's perspective, offline replay and live processing are indistinguishable contracts."),

(IMG, FIG + "ch7_fig1_rt_arch.png", 6.3),
(CAP, "Figure 7.1. The real-time architecture: two independent per-frame pipelines publish to shared memory, a merger pairs their outputs into the integrated packet of Chapter 5, and the Unity receiver is unchanged. The start barrier releases both pipelines at the same instant, emulating one shared camera."),

(P, "One synchronization detail matters more than it looks. The two pipelines each open the recording independently, and their initialization times differ; the pose model alone takes about 0.6 seconds to load. Started naively, the two playbacks begin at different wall-clock times, and the measured person-to-object skew was 567 milliseconds, seventeen frames of one pipeline's output being paired with the wrong frames of the other. The fix is a start barrier: each process finishes its initialization, signals readiness through a file, and only opens the recording when the launcher releases both together, the same release discipline a single shared camera would enforce. With the barrier the measured skew has a median of zero and a maximum of 33.4 milliseconds, exactly one frame, the irreducible pairing quantum of two unsynchronized 30 frames-per-second loops. The merger measures and reports this skew on every run, so a regression would be caught rather than absorbed."),

(H2, "7.2 Computational Bottlenecks and Latency"),

(P, "Every timing in this chapter comes from one machine, the machine the experiments ran on: an AMD Ryzen 9 7950X processor with 16 cores and 32 threads, 64 GB of memory, and an NVIDIA GeForce RTX 3080 graphics card with 10 GB of video memory. The division of labour is fixed throughout: the GPU's only job is neural network inference for the pose detector, and everything else, including all of the kinematic mathematics, runs on the CPU. The absolute numbers below should be read against this configuration; the structure of the budget, the dominant stage and the stages not worth accelerating, transfers to other hardware."),

(P, "Table 7.1 gives the person pipeline's measured per-stage latency over the full run, with both pipelines running concurrently. The distribution of cost is lopsided. The entire mathematical content of the thesis, deprojection, filtering, and the thirteen-angle solve, costs under one millisecond per frame; the shared memory write costs ten microseconds; the depth-to-colour alignment about 2.3 milliseconds. The solve row times the deployed solver, which is the same per-frame code the offline pipeline runs; no faster variant is substituted live, and the benchmark below shows that none is needed. Everything else is the pose detector: 9.2 milliseconds at the median and 27.6 at the 99th percentile, which is 90 percent of the frame budget's worst case on its own. The object pipeline is flatter: marker detection with corner refinement [46] takes 12.2 to 14.5 milliseconds, and everything after it, pose estimation, filtering, and the world map, adds about 0.3, for a 99th percentile total of 14.9 milliseconds. Both pipelines fit inside the 33.3 millisecond budget of 30 frames per second, the person side with about 2.6 milliseconds of headroom at the 99th percentile. The sustained end-to-end rate measured over the run is 25.7 frames per second, and the limiter is not computation: both processes report the identical 35.0 second duration for the 30 second recording, which identifies the playback pacing of the recording library as the bottleneck. Should the inference budget ever tighten, the detector's smaller model variant buys roughly a factor of three [39]."),

(TBL, [["stage", "p50 (ms)", "p95 (ms)", "p99 (ms)"],
       ["depth to colour alignment", "2.3", "2.5", "2.7"],
       ["MediaPipe inference (GPU, heavy model)", "9.2", "22.1", "27.6"],
       ["deprojection and causal filtering", "0.5", "0.6", "0.7"],
       ["13-angle kinematic solve", "0.2", "0.2", "0.3"],
       ["shared memory write", "0.01", "0.01", "0.01"],
       ["total, person pipeline", "12.2", "24.9", "30.7"]]),
(CAP, "Table 7.1. Measured per-stage latency of the person pipeline over the full run, both pipelines concurrent. The frame budget at 30 fps is 33.3 ms. The object pipeline totals 14.9 ms at p99, dominated by marker detection."),

(P, "The solve deserves one more level of scrutiny, because the same mathematics runs in two regimes with very different budgets: one frame at a time in the live loop, and whole recordings at once when a dataset is reprocessed offline. The system carries exactly two implementations. The deployed one, used by the offline pipeline and the live loop alike, is the per-frame solver of Chapter 3, eight landmarks in, thirteen angles out. The second is a vectorized numpy implementation of the same closed-form chain that evaluates every frame of a recording in one call, built for offline reprocessing; before any timing it was verified to agree with the per-frame solver to within ten billionths of a degree, so the benchmark compares implementations of identical mathematics. Table 7.2 and Figure 7.2 give the timings in both regimes, and the result separates cleanly. One frame at a time the two are practically indistinguishable, at 91 and 99 microseconds, because at batch size one the cost is Python and library call overhead, not arithmetic; a variant that splits the two arm solves across two threads was also measured and is slower, at 151 microseconds, since sub-100-microsecond work is smaller than the cost of handing it to a thread, so the deployed solve stays single threaded. In the batch regime the difference is decisive: at ten thousand frames the vectorized solver costs 0.51 microseconds per frame against the per-frame solver's 87.6, roughly 170 times the throughput."),

(TBL, [["implementation", "one frame at a time (us)", "batch of 10,000 (us per frame)"],
       ["per-frame solver (deployed, offline and live)", "99", "87.6"],
       ["vectorized numpy (offline batch)", "91", "0.51"]]),
(CAP, "Table 7.2. The solver benchmark, restricted to the two implementations the system carries: the per-frame solver deployed in both the offline and the real-time pipelines, and the vectorized numpy solver for offline batch reprocessing. Both agree to within ten billionths of a degree. Per frame the choice is immaterial; in batch the vectorized solver delivers roughly 170 times the throughput."),

(IMG, "/home/luo/Desktop/New_SandBox/writing/v4/figures/ch7_fig2_bench_v4.png", 6.5),
(CAP, "Figure 7.2. The solver benchmark: per-frame latency at batch size one (left) and per-frame cost against batch size (right). At batch size one the two implementations are separated only by call overhead; in the batch regime they separate by more than two orders of magnitude."),

(P, "Three conclusions follow. First, the live solve needs no accelerating. Table 7.1 reports the in-pipeline cost of the deployed per-frame solver, measured with the pose detector and both pipelines running concurrently; Table 7.2 reports the same call timed alone in a tight loop, and the two numbers differ in measurement context, not in implementation. Either number is under one percent of the 33.3 millisecond budget, so no choice of solver implementation can be visible in the live loop, and the simple, validated per-frame solver stays deployed. Second, vectorization pays offline: the batch solver reprocesses recordings at roughly 170 times the per-frame solver's throughput, and its self-test reproduces the per-frame solver over 10,785 real frames to within a tenth of a picodegree. Third, threads help only where the work is coarse enough: inside the sub-millisecond solve they add overhead, but at the frame level of the object pipeline, where the vision library releases Python's global lock, processing alternate frames on two threads measured 6.2 milliseconds per frame against 11.9 serial. Even that is kept out of the running system, because it adds a frame of latency and the serial pipeline already fits the budget; it is recorded as headroom, not deployed. The parallelism that carries the system is at the process level, the person and object pipelines running concurrently on separate cores by design."),

(H2, "7.3 Accuracy vs. Latency Trade-off"),

(P, "The real-time stack must accept two kinds of degradation relative to the offline one: it smooths causally, so its output lags and jitters slightly more, and it cannot bridge gaps retrospectively. Both costs were measured directly, by running the full real-time stack on the evaluation recording and comparing its per-frame output against the pinned offline solve of the same frames. Table 7.3 gives the angle comparison after warm-up. The joint groups differ by medians of 0.0 to 0.64 degrees, with the worst group, the left shoulder, at a median of 0.64 and a 95th percentile of 3.24 degrees; the left side is worst for the same reason it was worst throughout Chapter 6, occlusion by the carried box. The root orientation tracks to a median of 0.24 degrees, the pelvis position to a median of 2.4 millimetres, and the object position differs from the offline filtered track by a median of 6.2 millimetres, 95th percentile 13.9, which is the direct signature of causal versus zero-phase smoothing of the same underlying measurements."),

(TBL, [["quantity", "median", "p95", "max"],
       ["root orientation (deg)", "0.24", "1.08", "2.89"],
       ["right shoulder angles (deg)", "0.38", "1.72", "4.13"],
       ["right elbow angle (deg)", "0.00", "0.96", "2.24"],
       ["left shoulder angles (deg)", "0.64", "3.24", "15.27"],
       ["left elbow angle (deg)", "0.00", "2.18", "9.97"],
       ["pelvis position (mm)", "2.4", "7.4", ""],
       ["object position (mm)", "6.2", "13.9", ""]]),
(CAP, "Table 7.3. The measured cost of causality: per-frame difference between the real-time stack and the offline solve of the same recording, after warm-up. The left arm is worst for the same occlusion reason as in Chapter 6."),

(P, "The latency side of the trade-off has one dominant dial: the causal filter's cutoff. The One Euro filter [40] smooths with an adaptive low-pass whose minimum cutoff sets how aggressively slow motion is flattened, and lag is the price of aggression; Table 2.3 measured that price at 28.4 millimetres of deviation during fast motion under an offline-flavoured tuning. The first real-time configuration used a minimum cutoff of 0.05 hertz and produced a cross-correlation lag of nine frames against the offline track, 300 milliseconds, visible as the avatar trailing the motion. Retuning the minimum cutoff to 1 hertz with the speed coefficient at 1 brought the lag to three frames, 100 milliseconds, at a correlation of 0.998, while still holding the jitter to the residuals of Table 7.3. The trade is explicit and continuous: a lower cutoff buys stiller rest poses at the price of lag, a higher one buys responsiveness at the price of visible noise, and the chosen point keeps the total motion-to-display delay near the threshold where interactive use remains comfortable. Nothing else in the pipeline offers a comparable lever, because Section 7.2 showed the computation itself already fits the frame budget."),

(H2, "7.4 Smoothing and Stability Considerations"),

(P, "Chapter 2 argued that zero-phase filtering is the correct offline choice precisely because it is non-causal; running live forfeits it, and every smoothing stage needs a causal replacement. Table 7.4 lists the substitutions. The centred despike window becomes a trailing one; rejected samples become missing samples, handled by the same occlusion fallback as always, rather than being interpolated from a future that does not exist yet. Gap interpolation and edge backfill disappear entirely, because holding the last value under a lowered live mask is the online semantics of a gap, and the data integrity machinery of the earlier chapters already renders it. The landmark smoother becomes the One Euro filter per axis, and the object track's polynomial smoothing [44] becomes a One Euro position filter plus a complementary filter on rotation that blends each new measurement into the state along the geodesic, with the blend factor chosen to match the same 3 hertz cutoff as the offline design."),

(TBL, [["offline stage (zero-phase)", "real-time replacement (causal)"],
       ["centred Hampel despike", "trailing-window Hampel; a rejected sample becomes a missing sample"],
       ["interior gap interpolation and edge backfill", "not replaced: hold-last under a lowered live mask is the online semantics"],
       ["Butterworth 3 Hz, forward-backward (landmarks)", "One Euro per axis, minimum cutoff 1 Hz, speed coefficient 1"],
       ["Savitzky-Golay and rolling chordal mean (object)", "One Euro position filter plus a complementary geodesic rotation filter matched to a 3 Hz cutoff"]]),
(CAP, "Table 7.4. The causal filter substitutions. The data integrity conventions are unchanged: everything held rather than measured travels with a lowered live flag and renders flagged."),

(P, "Two failure modes surfaced that have no offline counterpart, and both are worth recording because they are generic to causal robust filtering, not quirks of this system. The first is rejection lockout. A trailing despike window that excludes rejected samples from its own history goes stale after genuinely fast motion: the window still describes the old position, so every new sample looks like a spike, is rejected, and never enters the window, locking the filter out permanently. On the evaluation recording this collapsed the object's live coverage from 877 frames to 541 before the cause was found. The fix encodes a physical insight: a spike that persists is not a spike. After three consecutive rejections the filter accepts the sample and reseeds its window from it, restoring coverage to 863 of 899 frames. The offline filter never met this failure because its centred window always contains the future side of the motion."),

(P, "The second is motion-onset false rejection. The despike threshold scales with the local scatter of the window, with an absolute floor for windows at rest; offline, with a centred window, a floor of 20 millimetres works well. A trailing window at the onset of motion has near-zero scatter, because the motion has not entered it yet, so the first genuinely fast samples exceed the floor and are wrongly rejected, exactly at the moments that matter most. Raising the live floor to 35 millimetres, still comfortably below the 54 millimetre class of real depth spikes the filter exists to catch, lifted the fraction of frames with all thirteen joints live from 92.4 to 97.3 percent. Both failure modes carry the same lesson: causal robust statistics need an escape hatch for legitimate novelty, because unlike their offline counterparts they cannot look ahead to see that the world really did change."),

(H2, "7.5 Comparison: Offline vs. Real-Time"),

(P, "Table 7.5 draws the two stacks side by side. The shared parts are deliberate: the same recording, the same calibration, the same kinematic mathematics, the same packet contract, the same transport, and the same Unity receiver, so that the comparison isolates exactly one variable, causal versus retrospective processing. Only the parts that must differ do: filtering, gap handling, and pacing. The accuracy cost is the content of Table 7.3, fractions of a degree at the median and single millimetres of position; the latency cost is three frames of filter lag plus the processing time of Table 7.1; the coverage cost is 863 live object frames against 877, and 97.3 percent fully live person frames against 100, because the real-time stack may not bridge a gap it has not yet seen the far side of."),

(TBL, [["aspect", "offline pipeline", "real-time stack"],
       ["input", "whole recording, random access", "one frame at a time, at recorded pacing"],
       ["landmark smoothing", "Butterworth 3 Hz, zero phase", "One Euro, causal"],
       ["object smoothing", "Savitzky-Golay + chordal mean", "One Euro + complementary geodesic filter"],
       ["gap handling", "interior interpolation, edge backfill", "hold-last under a lowered live mask"],
       ["lag", "zero (zero-phase)", "3 frames (100 ms)"],
       ["angles vs the other stack", "reference", "worst joint-group median 0.64 deg, p95 3.24"],
       ["pelvis / object position", "reference", "median 2.4 mm / 6.2 mm"],
       ["object live frames", "877 of 899", "863 of 899"],
       ["person all-joints-live", "100% after warm-up", "97.3% after warm-up"],
       ["per-frame compute, p99", "not constrained", "person 30.7 ms, object 14.9 ms (budget 33.3)"],
       ["transport and Unity receiver", "shared memory, Chapter 5 receiver", "identical, unchanged"]]),
(CAP, "Table 7.5. Offline versus real-time, same recording and same mathematics. Every difference is attributable to causality; the accuracy rows are the medians of Table 7.3 and the compute rows are the pinned validator bounds."),

(P, "The comparison is enforced, not merely tabulated: an automated validator runs the full real-time stack and applies eleven checks, all passing on the pinned run, covering frame coverage on both sides, the causal-versus-offline accuracy bounds, the three-frame lag bound, live-coverage floors, the per-frame compute budgets, and the integrity of the integrated packet region. Figure 7.3 shows the live demonstration: the recording streamed through the three-process stack into the unchanged scene, with the avatar tracking the motion as it happens, one hundred milliseconds behind the person. The conclusion of the feasibility question is therefore quantified: the system runs at camera rate on modest hardware with roughly a tenth of a second of latency, at a cost of well under a degree of median joint-angle accuracy relative to its own offline output, and the offline pipeline retains its role as the reference-quality pass, with the real-time stack as a faithful, slightly lagged approximation of it. The remaining distance to a true live deployment, a camera instead of a recording, and the limitations that pacing imposes on the measured frame rate, are discussed in Chapter 8."),

(IMG, FIG + "ch7_fig3_rt_still.png", 6.3),
(CAP, "Figure 7.3. Still from the pinned live demonstration: the recording processed frame by frame through the real-time stack and rendered in the unchanged Unity scene, with the sensor-view inset at the top right. Motion appears from the first second of playback."),

]

references = {
39: 'V. Bazarevsky, I. Grishchenko, K. Raveendran, T. Zhu, F. Zhang, and M. Grundmann, "BlazePose: On-device real-time body pose tracking," arXiv:2006.10204, 2020.',
40: 'G. Casiez, N. Roussel, and D. Vogel, "1 euro filter: A simple speed-based low-pass filter for noisy input in interactive systems," in Proc. SIGCHI Conf. Human Factors in Computing Systems (CHI), 2012, pp. 2527-2530.',
44: 'A. Savitzky and M. J. E. Golay, "Smoothing and differentiation of data by simplified least squares procedures," Analytical Chemistry, vol. 36, no. 8, pp. 1627-1639, 1964.',
46: 'J. Wang and E. Olson, "AprilTag 2: Efficient and robust fiducial detection," in Proc. IEEE/RSJ Int. Conf. Intelligent Robots and Systems (IROS), 2016, pp. 4193-4198.',
}

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == H3:
        doc.add_heading(item[1], level=3)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == IMG:
        doc.add_picture(item[1], width=Inches(item[2]))
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

out = ("/home/luo/Desktop/New_SandBox/writing/v4/"
       "Chapter_7_Real_Time_Feasibility.docx")
doc.save(out)
print("saved", out)
