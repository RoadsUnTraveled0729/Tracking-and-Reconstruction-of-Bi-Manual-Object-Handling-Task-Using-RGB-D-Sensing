"""Build Thesis_V5_Changes.docx: the change summary that accompanies each
delivered version. Delivery protocol (user decision 2026-08-05): every round
ships Thesis_Vx.docx, Thesis_Vx_commented.docx, and Thesis_Vx_Changes.docx.

Content is drawn from the recorded fix log (REVIEW_FIXES.md items 32-34);
scope is V4 (the reviewed draft) to V5 (current). Audience: the supervisor.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/home/luo/Desktop/New_SandBox/writing/v5/Thesis_V5_Changes.docx"

SECTIONS = [
("Chapter 1", [
 "The chapter was shortened from ten pages to seven.",
 "The literature review is condensed, and the reviewed work with its "
 "limitations is summarized in the new Table 1.1.",
 "The research gaps and the objectives are now numbered lists.",
 "The new Section 1.4, Organization of the Thesis, holds the chapter "
 "overview that previously opened Chapter 2.",
]),
("Chapter 2", [
 "The chapter was rewritten in plainer, more direct language.",
 "New figures show real data: the capture scene from the camera's "
 "viewpoint (Figure 2.2), the MediaPipe landmarks (Figure 2.3), and the "
 "ArUco detections (Figure 2.4), all on the running example frame.",
 "New short subsections introduce the pose landmark detector and the "
 "fiducial markers; their internals moved to the appendices, and each "
 "subsection points to its own appendix.",
 "The filter comparison was removed as out of scope; the text now "
 "describes only the Butterworth filter that is used.",
 "The three error kinds are itemized as Cases A, B, and C, and Case is "
 "used consistently.",
 "The Data Integrity rule is defined once at the start of the chapter "
 "and later sections simply apply it.",
 "The sampling rate section was deleted; its details sit in one line "
 "where the recordings are introduced.",
]),
("Chapter 3", [
 "Section 3.4 is renamed Joint Angle Computation and restricted to the "
 "model and its worked examples. The new Section 3.5 shows how the "
 "computed angles are converted for the Unity avatar: the torso first, "
 "carried through with the numbers of the running example, then the arm "
 "chain.",
 "The new Figure 3.5 shows the T-pose on the avatar, and the T-pose is "
 "defined as the reference configuration of the model.",
 "Symbols and subscripts are typeset in math mode throughout the thesis.",
 "Worked numbers are printed to two decimals everywhere, with the "
 "convention stated once at the first worked example.",
 "The nomenclature is now consistent: p marks a measured landmark point, "
 "v a segment vector, and a hat a unit direction; the matrix entry "
 "notation is defined before its first use.",
 "Confusing passages were rewritten in plain language, including the "
 "handedness discussion, gimbal lock, and the decode convention, and "
 "the worked examples are titled Worked Example.",
]),
("Chapters 4 to 9", [
 "The same rule as Part 1 was applied: unnecessary comparisons and "
 "repeated exactness figures removed, scope narrowed, while the "
 "development history is kept, including the change of transport from "
 "UDP to shared memory, the anchoring bug, and the filter retunes.",
 "The new Figure 5.2 draws the streamed record as a block of memory with "
 "the byte layout and the writer and reader protocols.",
 "Chapter 7 opens its performance section by stating the hardware, and "
 "the solver discussion keeps only the two implementations the system "
 "carries.",
 "The lever arm figures of Chapter 8 were corrected against the "
 "calibration source and now agree with Table 8.1.",
]),
("Appendices", [
 "The appendices were reordered to match their first mention in the body "
 "and expanded. Appendix A covers the acquisition of the recordings: the "
 "capture procedure with a code listing, the structure of the recorded "
 "bag file with a diagram and its measured contents, and replay in the "
 "vendor's viewer, whose 3D views are compared against the camera view "
 "in Figure A.2.",
 "Appendices B, C, and D hold depth sensing and deprojection, pose "
 "landmark detection, and fiducial marker detection.",
 "The new Appendix E lists the full hardware and software platform "
 "behind every timing and processing result.",
]),
("Front matter", [
 "The thesis now opens in the department format: title page, approval "
 "page, abstract, acknowledgments, table of contents, list of figures, "
 "list of tables, and list of equations.",
]),
("References", [
 "The reference list is unchanged: the numbering [1] to [52] is "
 "identical to the reviewed draft.",
]),
]

doc = Document()
for sname in ("Normal", "List Bullet"):
    st = doc.styles[sname]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Summary of Changes: Thesis V4 to V5")
r.font.size = Pt(14)
r.font.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Lanqing Luo, August 2026").font.size = Pt(11)

intro = doc.add_paragraph(
    "This document summarizes the changes between the reviewed draft and "
    "the current version. In addition, every review comment has been placed "
    "back into the commented copy of the thesis at the location of the "
    "change, each with a short reply stating what was done.")
intro.paragraph_format.space_after = Pt(10)

for heading, items in SECTIONS:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.add_run(heading).font.bold = True
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(3)

doc.save(OUT)
print("saved", OUT)
