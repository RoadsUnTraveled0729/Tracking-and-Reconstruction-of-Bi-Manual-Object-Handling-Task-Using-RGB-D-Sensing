"""Build Thesis_V6_Changes.docx: the change summary that accompanies each
delivered version. Delivery protocol (user decision 2026-08-05): every round
ships Thesis_Vx.docx, Thesis_Vx_commented.docx, and Thesis_Vx_Changes.docx.

Scope: V5 (delivered to the supervisor 2026-08-05) to V6 (current).
Audience: the supervisor. Grow the SECTIONS list as V6 rounds accumulate.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/home/luo/Desktop/New_SandBox/writing/v6/Thesis_V6_Changes.docx"

SECTIONS = [
("Appendix A", [
 "Figure A.2 is completed with screenshots of the vendor viewer replaying "
 "the evaluation recording in 3D: panel (a) is the recorded camera view, "
 "panel (b) the point cloud textured with the colour stream, and panel (c) "
 "the same cloud coloured by depth.",
 "All three panels show the same moment of the recording; the camera frame "
 "of panel (a) was extracted from the recorded bag file at the instant the "
 "viewer was paused for the 3D views.",
]),
]

doc = Document()
for sname in ("Normal", "List Bullet"):
    st = doc.styles[sname]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Summary of Changes: Thesis V5 to V6")
r.font.size = Pt(14)
r.font.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Lanqing Luo, August 2026").font.size = Pt(11)

intro = doc.add_paragraph(
    "This document summarizes the changes since the previously delivered "
    "version. The review comments remain in the commented copy of the "
    "thesis at the location of each change, each with a short reply "
    "stating what was done.")
intro.paragraph_format.space_after = Pt(10)

for heading, items in SECTIONS:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.add_run(heading).font.bold = True
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(3)

doc.save(OUT)
print("saved", OUT)
