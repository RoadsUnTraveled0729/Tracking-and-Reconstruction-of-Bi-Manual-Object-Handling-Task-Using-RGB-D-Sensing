#!/usr/bin/env python3
"""Verify D-082's Chapter 7 Scene conversion without rewriting pinned evidence.

Run from the repository root with /home/luo/anaconda3/bin/python. Full-precision
coordinates are used throughout; tolerances come from the existing independent
Chapter 7 verifier, not from displayed rounding or a physical error allowance.
"""
import hashlib
import io
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from lxml import etree
from scipy.spatial.distance import pdist

HERE = Path(__file__).resolve().parents[2]
REPO = Path(__file__).resolve().parents[4]
OUT = Path(__file__).resolve().parent
EVIDENCE = HERE / 'audit_evidence/ch7_restructured'
sys.path.insert(0, str(HERE / 'scripts'))
import make_ch7_object_waypoint_figures as plot
import evaluate_ch7_object as pinned_object
from verify_ch7_restructured import TOL_CM, TOL_M, TOL_UNIT

BASELINE = 'e69dfa6'
INPUTS = set()
REPORT = {'decision': 'D-082', 'baseline_commit': BASELINE, 'checks': [],
          'recordings': {}, 'coordinate_pairs': {}}


def source(path):
    INPUTS.add(Path(path).resolve())
    return Path(path)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def check(name, difference, tolerance):
    maximum = float(np.max(np.abs(difference)))
    assert maximum <= tolerance, (name, maximum, tolerance)
    REPORT['checks'].append({'name': name, 'maximum_absolute_difference': maximum,
                             'tolerance': tolerance, 'status': 'PASS'})


def load_csv(name):
    return pd.read_csv(source(EVIDENCE / name), float_precision='round_trip')


def coordinates(rows, prefix):
    return rows[[prefix + '_' + a for a in 'xyz']].to_numpy(float)


def summarize(distance_cm):
    return {'n': len(distance_cm), 'median': float(np.median(distance_cm)),
            'p95': float(np.percentile(distance_cm, 95, method='linear')),
            'maximum': float(np.max(distance_cm))}


def fit_scatter(points):
    centroid = points.mean(axis=0)
    _, sigma, axes = np.linalg.svd(points - centroid, full_matrices=False)
    direction = axes[0]
    residual = (points - centroid) - np.outer((points - centroid) @ direction,
                                            direction)
    distance = np.linalg.norm(residual, axis=1) * 100
    return dict(**summarize(distance), centroid_m=centroid.tolist(),
                direction=(direction if direction[0] >= 0 else -direction).tolist(),
                tilt_from_horizontal_deg=float(np.degrees(np.arcsin(abs(direction[1])))),
                vertical_component_median_cm=float(np.median(abs(residual[:, 1])) * 100),
                pca_sigma_cm=(sigma / np.sqrt(len(points)) * 100).tolist()), distance


def translation_check(rows, first, second, offset, name, error_column=None):
    a, b = coordinates(rows, first), coordinates(rows, second)
    finite = np.isfinite(a).all(axis=1) & np.isfinite(b).all(axis=1)
    before = np.linalg.norm(a[finite] - b[finite], axis=1) * 100
    after = np.linalg.norm((a[finite] + offset) - (b[finite] + offset), axis=1) * 100
    check(name, after - before, TOL_CM)
    if error_column:
        check(name + ' against pinned per-row distances',
              before - rows.loc[finite, error_column].to_numpy(float), TOL_CM)
    return {'rows': len(rows), 'finite_pairs': int(finite.sum()),
            'unavailable_coordinate_pairs': int((~finite).sum()),
            'before_cm': summarize(before), 'after_cm': summarize(after),
            'maximum_distance_change_cm': float(np.max(abs(after - before)))}


def main():
    object_data = json.loads(source(EVIDENCE / 'object.json').read_text())
    human_data = json.loads(source(EVIDENCE / 'human.json').read_text())
    waypoint_pairs = load_csv('object_pairs.csv')
    offsets = {}
    for alias, record in object_data['recordings'].items():
        calibration = json.loads(source(REPO / record['calibration']).read_text())
        track = pd.read_csv(source(REPO / record['track']), float_precision='round_trip')
        stored_all = track[['unity_px', 'unity_py', 'unity_pz']].to_numpy(float)
        accepted = (np.isfinite(stored_all).all(axis=1)
                    & (track.detected.to_numpy() == 1)
                    & (track.filled.to_numpy() != 1))
        stored = stored_all[accepted]
        frames = track.frame.to_numpy(int)[accepted]
        gravity = np.asarray(calibration['scene_geometry']['gravity_up_unity'], float)
        gravity /= np.linalg.norm(gravity)
        G = plot.from_to_rotation(gravity, np.array([0., 1., 0.]))
        offset = plot.floor_translation(calibration)
        offsets[alias] = offset
        before = (G @ stored.T).T
        after = before + offset
        check(alias + ' G orthogonality', G.T @ G - np.eye(3), TOL_UNIT)
        check(alias + ' G proper determinant', np.linalg.det(G) - 1., TOL_UNIT)
        check(alias + ' G maps gravity to Scene y', G @ gravity - [0., 1., 0.], TOL_UNIT)

        stream_alias = 'r6b_trails' if alias == 'r6b' else alias
        stream = pd.read_csv(source(REPO / 'eval/output' / ('unity_check_' + stream_alias)
                                    / 'integrated_stream.csv'), float_precision='round_trip')
        plane = -gravity * calibration['scene_geometry']['origin_above_tabletop_m']
        def project(q):
            return q - gravity * np.dot(q - plane, gravity)
        center = (project(np.zeros(3)) + project(stream[['opx', 'opy', 'opz']].iloc[0].to_numpy())) / 2
        receiver_offset = np.array([0., -(G @ (center - gravity *
                                             (plot.DESK_THICK + plot.LEG_H)))[1], 0.])
        check(alias + ' direct floor sum versus receiver geometry', offset - receiver_offset, TOL_M)
        S = np.array([[1., 0., 0.], [0., 0., 1.], [0., 1., 0.]])
        camera_to_world = np.linalg.inv(np.asarray(calibration['T_cam_desk'], float))
        expected_translation = G @ S @ camera_to_world[:3, 3] + offset
        check(alias + ' final Scene origin versus existing human evaluator',
              expected_translation - human_data[alias]['camera_to_unity_translation'], TOL_M)

        plot_cm, plot_frames, _, plot_offset = plot.scene_samples(alias)
        assert np.array_equal(plot_frames, frames)
        check(alias + ' plotted samples include one floor translation', plot_cm / 100 - after, TOL_M)
        check(alias + ' plotted floor translation', plot_offset / 100 - offset, TOL_M)
        check(alias + ' all unique marker pair distances after G and floor placement',
              pdist(after) - pdist(stored), TOL_M)
        check(alias + ' all unique marker pair distances before versus after floor placement',
              pdist(after) - pdist(before), TOL_M)

        detector_track = {'xyz': before, 't': track.time_s.to_numpy(float)[accepted]}
        detector_before = pinned_object.detect(detector_track)
        detector_after = pinned_object.detect(dict(detector_track, xyz=after))
        assert detector_before['index'] == detector_after['index']
        waypoint_records, means_before, means_after = {}, {}, {}
        plotted_waypoints = plot.waypoints(plot_cm, plot_frames, record, plot_offset)
        for name, entry in record['waypoints'].items():
            selected_frame = int(frames[detector_after['index'][name]])
            expected_frame = entry['last_frame'] if name == 'W1' else entry['first_frame']
            assert selected_frame == expected_frame, (alias, name, selected_frame, expected_frame)
            inside = (frames >= entry['first_frame']) & (frames <= entry['last_frame'])
            assert int(inside.sum()) == entry['n_samples']
            old_mean, scene_mean = before[inside].mean(axis=0), after[inside].mean(axis=0)
            means_before[name], means_after[name] = old_mean, scene_mean
            check(alias + ' ' + name + ' original waypoint mean', old_mean - entry['position_m'], TOL_M)
            check(alias + ' ' + name + ' common translation of mean', scene_mean - old_mean - offset, TOL_M)
            check(alias + ' ' + name + ' plotted waypoint', plotted_waypoints[name] / 100 - scene_mean, TOL_M)
            pinned_rows = waypoint_pairs[(waypoint_pairs.recording == alias) & (waypoint_pairs.waypoint == name)]
            assert np.array_equal(pinned_rows.frame.to_numpy(int), frames[inside])
            check(alias + ' ' + name + ' all pinned waypoint samples',
                  pinned_rows[['x_m', 'y_m', 'z_m']].to_numpy(float) - before[inside], TOL_M)
            waypoint_records[name] = {'first_frame': entry['first_frame'], 'last_frame': entry['last_frame'],
                                      'n_samples': entry['n_samples'], 'pinned_before_floor_m': old_mean.tolist(),
                                      'scene_m': scene_mean.tolist()}
        segments = []
        for entry in record['segments']:
            first, last = entry['segment'].split(' to ')
            delta = means_after[last] - means_after[first]
            length = float(np.linalg.norm(delta) * 100)
            absolute = abs(length - entry['physical_cm'])
            values = {'reconstructed_cm': length, 'absolute_error_cm': absolute,
                      'relative_error_pct': absolute / entry['physical_cm'] * 100}
            for key, value in values.items():
                check(alias + ' ' + entry['segment'] + ' ' + key, value - entry[key], TOL_CM)
            check(alias + ' ' + entry['segment'] + ' component displacement', delta - entry['delta_m'], TOL_M)
            assert 'xyz'[int(np.argmax(abs(delta)))] == entry['dominant_axis']
            segments.append(dict(segment=entry['segment'], physical_cm=entry['physical_cm'], **values))
        scatter = {}
        for name, entry in record['scatter'].items():
            first, last = entry['span_frames']
            inside = (frames >= first) & (frames <= last)
            prior, prior_distances = fit_scatter(before[inside])
            scene, scene_distances = fit_scatter(after[inside])
            assert scene['n'] == entry['n']
            check(alias + ' ' + name + ' every line-scatter distance', scene_distances - prior_distances, TOL_CM)
            for key in ('median', 'p95', 'maximum', 'vertical_component_median_cm', 'pca_sigma_cm',
                        'direction', 'tilt_from_horizontal_deg'):
                check(alias + ' ' + name + ' ' + key, np.asarray(scene[key]) - entry[key], TOL_CM)
            check(alias + ' ' + name + ' centroid translation',
                  np.asarray(scene['centroid_m']) - entry['centroid_m'] - offset, TOL_M)
            scatter[name] = {'before_floor': prior, 'scene': scene}
        REPORT['recordings'][alias] = {
            'accepted_samples': int(accepted.sum()), 'all_unique_marker_pairs': len(pdist(after)),
            'G': G.tolist(), 'translation_m': offset.tolist(),
            'floor_sources': {'calibrated_origin_above_tabletop_m': calibration['scene_geometry']['origin_above_tabletop_m'],
                              'desk_thickness_m': plot.DESK_THICK, 'leg_height_m': plot.LEG_H},
            'receiver_geometry_translation_m': receiver_offset.tolist(),
            'human_scene_origin_check_m': expected_translation.tolist(),
            'waypoint_selection_before': detector_before['index'],
            'waypoint_selection_after': detector_after['index'],
            'waypoint_boundary_frames': {name: int(frames[index]) for name, index in detector_after['index'].items()},
            'waypoints': waypoint_records, 'segments': segments, 'scatter': scatter}

    spatial = load_csv('spatial_check_pairs.csv')
    for alias, rows in spatial.groupby('recording'):
        for first, second, column in (('model_wrist', 'marker', 'distance_cm'),
                                      ('measured_wrist', 'cube_centre', None)):
            name = alias + ' spatial ' + first + ' to ' + second
            REPORT['coordinate_pairs'][name] = translation_check(rows, first, second, offsets[alias], name, column)
        a, b = coordinates(rows, 'measured_wrist'), coordinates(rows, 'cube_centre')
        substituted = ~np.isfinite(a).all(axis=1) | (rows.wrist_src.to_numpy(int) != 0)
        assert np.array_equal(substituted, rows.substituted.to_numpy(bool))
        a[substituted] = coordinates(rows, 'model_wrist')[substituted]
        holds = rows.copy()
        for index, axis in enumerate('xyz'):
            holds['hold_' + axis] = a[:, index]
        name = alias + ' spatial original hold test point to cube_centre'
        REPORT['coordinate_pairs'][name] = translation_check(holds, 'hold', 'cube_centre',
                                                              offsets[alias], name, 'hold_distance_cm')
        from verify_ch7_restructured import HOLD_RADIUS_CM
        before_gate = np.linalg.norm(a - b, axis=1) * 100 < HOLD_RADIUS_CM
        after_gate = np.linalg.norm((a + offsets[alias]) - (b + offsets[alias]), axis=1) * 100 < HOLD_RADIUS_CM
        assert np.array_equal(before_gate, after_gate)
        finite_distance = np.isfinite(coordinates(rows, 'model_wrist')).all(axis=1) & np.isfinite(coordinates(rows, 'marker')).all(axis=1)
        included = before_gate & (rows.carried.to_numpy(int) == 1) & (rows.detected.to_numpy(int) == 1) & finite_distance
        assert np.array_equal(included, rows.included.to_numpy(bool))
        REPORT['coordinate_pairs'][alias + ' spatial selection'] = {
            'rows': len(rows), 'included_rows': int(rows.included.sum()),
            'hold_radius_cm': HOLD_RADIUS_CM, 'changed_hold_radius_decisions': 0,
            'model_wrist_substitutions_in_original_hold_gate': int(substituted.sum()),
            'other_selection_fields': 'Unchanged pinned carried, detected, wrist source and substitution fields.'}
        for (part, side), selected in rows[rows.included == 1].groupby(['part', 'side']):
            name = alias + ' included spatial ' + part + ' ' + side
            REPORT['coordinate_pairs'][name] = translation_check(selected, 'model_wrist', 'marker',
                                                                   offsets[alias], name, 'distance_cm')

    # These existing files already store final Scene positions (human/bare) or
    # Camera' positions (synthetic). Zero means no extra floor placement.
    for filename, first, second, column in (
            ('human_pairs.csv', 'rig', 'measured', 'error_cm'),
            ('bare_model_pairs.csv', 'model', 'measured', 'error_cm'),
            ('bare_model_pairs.csv', 'model', 'measured_filtered', 'error_vs_filtered_cm'),
            ('synthetic_pairs.csv', 'reconstructed', 'reference', 'error_cm')):
        rows = load_csv(filename)
        name = filename + ' ' + first + ' to ' + second
        REPORT['coordinate_pairs'][name] = translation_check(rows, first, second, np.zeros(3), name, column)

    preserved = sorted(p for p in EVIDENCE.iterdir() if p.is_file())
    preserved += [HERE / 'figures/ch7_valid_joints_r6b.png', HERE / 'figures/ch7_valid_joints_r7.png']
    REPORT['preserved_baseline_sha256'] = {}
    for path in preserved:
        relative = path.relative_to(REPO).as_posix()
        baseline = subprocess.run(['git', 'show', BASELINE + ':' + relative], cwd=REPO,
                                  check=True, capture_output=True).stdout
        current = source(path).read_bytes()
        assert current == baseline, relative
        REPORT['preserved_baseline_sha256'][relative] = hashlib.sha256(current).hexdigest()

    chapter = HERE / 'Chapter_7_Evaluation.docx'
    baseline_docx = subprocess.run(['git', 'show', BASELINE + ':' + chapter.relative_to(REPO).as_posix()],
                                   cwd=REPO, check=True, capture_output=True).stdout
    original, current = Document(io.BytesIO(baseline_docx)), Document(chapter)
    def tables(doc):
        return [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert tables(original) == tables(current)
    REPORT['chapter_tables'] = {'unchanged_from_baseline': len(current.tables)}
    def math_xml(doc):
        return [etree.tostring(node, method='c14n') for node in doc.element.xpath('.//m:oMath')]
    assert math_xml(original) == math_xml(current)
    REPORT['chapter_equations'] = {'unchanged_omml_displays_from_baseline': len(math_xml(current))}
    text = '\n'.join(p.text for p in current.paragraphs)
    assert 'levelled' not in text.lower() and 'leveled' not in text.lower()
    REPORT['chapter_printed_removed_frame_occurrences'] = 0
    for path in (HERE / 'scripts/make_ch7_object_waypoint_figures.py',
                 HERE / 'scripts/build_ch7.py', HERE / 'scripts/verify_ch7_restructured.py',
                 HERE / 'scripts/evaluate_ch7_object.py', REPO / 'eval/unity_check/check_unity_log.py',
                 REPO / 'Unity/Assets/Scripts/ArucoSceneReceiver.cs', Path(__file__)):
        source(path)
    REPORT['tolerance_sources'] = {'source': 'scripts/verify_ch7_restructured.py',
                                   'TOL_CM': TOL_CM, 'TOL_M': TOL_M, 'TOL_UNIT': TOL_UNIT,
                                   'meaning': 'Double-precision arithmetic only; no physical or display-rounding allowance.'}
    REPORT['input_sha256'] = {path.relative_to(REPO).as_posix(): sha256(path) for path in sorted(INPUTS)}
    REPORT['status'] = 'PASS'
    output = OUT / 'ch7_scene_invariance.json'
    output.write_text(json.dumps(REPORT, indent=2, allow_nan=False) + '\n')
    print('PASS:', len(REPORT['checks']), 'numeric checks; all tables and pinned evidence unchanged.')
    print('PASS:', output)


if __name__ == '__main__':
    main()
