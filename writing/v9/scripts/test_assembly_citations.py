#!/usr/bin/env python3
"""Regression checks for split-run citation remapping and reference identity."""
import copy
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

import assembly_citations as citations


class CitationTests(unittest.TestCase):
    def test_split_runs_keep_formatting_and_math(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Source [").bold = True
        p.add_run("4").italic = True
        p.add_run("2] and [5], then [42].")
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = "x"
        run.append(text)
        math.append(run)
        p._p.append(math)
        properties = [etree.tostring(r._r.rPr) if r._r.rPr is not None else None for r in p.runs]
        math_before = etree.tostring(math)
        citations.rewrite_element(p._p, {42: 1, 5: 12})
        self.assertEqual(p.text, "Source [1] and [12], then [1].")
        self.assertEqual(properties, [etree.tostring(r._r.rPr) if r._r.rPr is not None else None for r in p.runs])
        self.assertEqual(math_before, etree.tostring(math))

    def test_first_appearance_identity_and_source_preservation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            chapter = Document()
            chapter.add_paragraph("Chapter 1")
            chapter.add_paragraph("First [42], then [1], again [42].")
            table = chapter.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Table source [5]."
            appendix = Document()
            appendix.add_paragraph("Appendix A")
            appendix.add_paragraph("Appendix-only [51], then [1].")
            paths = [root / "Chapter.docx", root / "Appendix.docx"]
            for source, path in zip((chapter, appendix), paths):
                source.save(path)
            originals = [path.read_bytes() for path in paths]
            entries = [(1, "Reference one"), (5, "Reference five"),
                       (42, "Craig fourth edition"), (50, "Uncited"),
                       (51, "Appendix reference")]
            plan = citations.make_plan(paths, entries)
            self.assertEqual(plan["mapping"], {42: 1, 1: 2, 5: 3, 51: 4})
            master = Document()
            for source_index, source in enumerate((chapter, appendix)):
                if source_index:
                    master.add_paragraph("References")
                    for number, text in citations.final_entries(plan):
                        master.add_paragraph(f"[{number}] {text}")
                for element in source.element.body:
                    if element.tag == qn("w:sectPr"):
                        continue
                    copied = copy.deepcopy(element)
                    citations.rewrite_element(copied, plan["mapping"])
                    master.element.body.sectPr.addprevious(copied)
            saved = root / "Assembly.docx"
            master.save(saved)
            result = citations.verify_assembly(Document(saved), plan)
            self.assertEqual(result["citation_occurrences"], 6)
            self.assertEqual(result["removed_uncited_source_numbers"], [50])
            self.assertEqual(originals, [path.read_bytes() for path in paths])
            # Contiguous but wrong citations must fail; target existence alone
            # cannot detect a wrong source-to-final assignment.
            broken = Document(saved)
            broken.paragraphs[1].text = "First [2], then [1], again [2]."
            with self.assertRaises(AssertionError):
                citations.verify_assembly(broken, plan)
            broken = Document(saved)
            for p in broken.paragraphs:
                if p.text == "[1] Craig fourth edition":
                    p.text = "[1] Reference one"
            with self.assertRaises(AssertionError):
                citations.verify_assembly(broken, plan)


if __name__ == "__main__":
    unittest.main()
