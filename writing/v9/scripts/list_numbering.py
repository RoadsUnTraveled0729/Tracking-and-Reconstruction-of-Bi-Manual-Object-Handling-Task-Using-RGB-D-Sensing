"""Shared numbering-restart helper for "List Number 2" runs (NUM item kind).

Companion to the BUL bullet-item kind added in commit 1d5521a (see, e.g.,
writing/v9/scripts/build_ch1.py, "bullet item (supervisor comment C54,
2026-09-16)"): NUM is the numbered-list sibling of BUL, same tuple
contract, `(NUM, "Label.", "text")` or `(NUM, "", "text")`. python-docx's
built-in "List Number" style (style_id ListNumber) is reserved:
verify_submission.py asserts exactly 8 paragraphs carry that style_id for
Appendix A.1, and submission_layout.apply() re-pins every one of them to
numId 7. NUM paragraphs use the sibling built-in style "List Number 2"
(style_id ListNumber2) instead, so the two counts never collide.

Word numbers a run of List Number 2 paragraphs by the numId each
paragraph's <w:numPr> points at: paragraphs sharing a numId are one
running count, and a fresh numId restarts the count at 1. This module
walks doc.paragraphs, finds every contiguous run of ListNumber2
paragraphs (a run ends at any paragraph of a different style - "the next
run, separated by any other paragraph, restarts at 1"), and for each run
adds one new <w:num> to the document's own numbering part. Each new
<w:num> points at the abstractNum the ListNumber2 style already resolves
to (found via styles.xml's w:pPr/w:numPr/w:numId and numbering.xml's
matching w:num/w:abstractNumId) and carries a
<w:lvlOverride ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>, then
every paragraph in that run is pointed at the new numId, ilvl 0.

Call restart_numbered_lists(doc) once per document that carries NUM
paragraphs:
  - in every chapter builder, just before doc.save, so the standalone
    chapter .docx already numbers correctly on its own; and
  - again in submission_layout.apply(doc), on the assembled master, after
    build_thesis.py has merged every chapter's body into it.
The second call is required, not a redundant safety net. build_thesis.py
merges chapter body XML into a master built from a bare python-docx
Document(), so the master's numbering.xml is the default template's, not
any chapter's: none of the numId values a chapter wrote into its
paragraphs' numPr resolve to anything in the master's numbering part.
Calling this function again on the master creates fresh <w:num> entries
in the master's own numbering part and repoints every ListNumber2
paragraph at one of them - exactly as submission_layout.apply already
re-pins ListNumber paragraphs to numId 7 for the same reason. normalize()
in build_thesis.py does not touch w:numPr (it is not a heading, caption,
placeholder run, table cell or bare-math/image paragraph), so the numPr
merged in from each chapter survives the merge unchanged; this function
then overwrites it with the master's own numId regardless.
"""
from docx.oxml.ns import qn

STYLE_ID = "ListNumber2"


def _abstract_num_id(doc):
    """The abstractNumId the ListNumber2 style's own numbering resolves to."""
    style_el = None
    for candidate in doc.styles.element.findall(qn("w:style")):
        if candidate.get(qn("w:styleId")) == STYLE_ID:
            style_el = candidate
            break
    if style_el is None:
        raise ValueError(f"styles.xml has no style with styleId {STYLE_ID!r}")

    pPr = style_el.find(qn("w:pPr"))
    numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
    numId_el = numPr.find(qn("w:numId")) if numPr is not None else None
    if numId_el is None:
        raise ValueError(
            f"style {STYLE_ID!r} carries no w:pPr/w:numPr/w:numId in "
            "styles.xml; it cannot be restarted")
    style_num_id = int(numId_el.get(qn("w:val")))

    numbering = doc.part.numbering_part.element
    try:
        num_el = numbering.num_having_numId(style_num_id)
    except KeyError:
        raise ValueError(
            f'numbering.xml has no <w:num w:numId="{style_num_id}"> for '
            f"style {STYLE_ID!r}")
    abstract_id = int(num_el.abstractNumId.get(qn("w:val")))
    have = {a.get(qn("w:abstractNumId"))
            for a in numbering.findall(qn("w:abstractNum"))}
    if str(abstract_id) not in have:
        raise ValueError(
            f'numbering.xml\'s <w:num w:numId="{style_num_id}"> points at '
            f"abstractNumId {abstract_id}, which has no <w:abstractNum> "
            f"definition; style {STYLE_ID!r} cannot be restarted")
    return abstract_id


def restart_numbered_lists(doc):
    """Restart numbering at 1 for every contiguous run of ListNumber2
    paragraphs (NUM items) in doc.paragraphs."""
    abstract_id = None
    numbering = None
    run = []

    def flush():
        nonlocal abstract_id, numbering
        if not run:
            return
        if numbering is None:
            numbering = doc.part.numbering_part.element
            abstract_id = _abstract_num_id(doc)
        num_el = numbering.add_num(abstract_id)
        override = num_el.add_lvlOverride(ilvl=0)
        override.add_startOverride(1)
        num_id = num_el.numId
        for paragraph in run:
            pPr = paragraph._p.get_or_add_pPr()
            numPr = pPr.get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = num_id
        run.clear()

    for paragraph in doc.paragraphs:
        if paragraph.style.style_id == STYLE_ID:
            run.append(paragraph)
        else:
            flush()
    flush()
