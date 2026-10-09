#!/usr/bin/env python3
"""Build writing/v9/Abstract_V9.docx: the abstract on its own, for sending
separately (supervisor round 7, 2026-09-10, C51). The text is the ABSTRACT
list of scripts/build_frontmatter.py, so this file and the
front matter of the assembled thesis cannot drift apart.

Run: python3 writing/v9/scripts/build_abstract.py
"""
import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
FM = HERE / "build_frontmatter.py"
OUT = HERE.parent / "Abstract_V9.docx"

# Load only the module constants: executing the module would rebuild the
# TOC, so read ABSTRACT and TITLE_MIXED from the source text instead.
ns = {}
src = FM.read_text()
for name in ("TITLE_MIXED", "ABSTRACT"):
    start = src.index(f"\n{name} = ")
    end = start + 1
    depth = 0
    for i in range(start + 1, len(src)):
        c = src[i]
        if c in "([": depth += 1
        elif c in ")]":
            depth -= 1
            if depth == 0:
                end = i + 1; break
    exec(src[start + 1:end], ns)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
doc.add_heading(ns["TITLE_MIXED"], level=1)
doc.add_heading("Abstract", level=2)
for para in ns["ABSTRACT"]:
    doc.add_paragraph(para)
doc.save(OUT)
words = sum(len(p.split()) for p in ns["ABSTRACT"])
print("saved", OUT, "| abstract words:", words, "| paragraphs:", len(ns["ABSTRACT"]))
assert 250 <= words <= 400, words  # the supervisor's text of 2026-09-14 is 370 words
