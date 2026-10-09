#!/usr/bin/env python3
"""Build writing/v8/Chapter_3_Kinematic_Modeling.docx (V8 rewrite round).

Chapter 3 against the V8 TOC: 3.1 Coordinate Systems and Transformations
(old 3.1.1 + 3.1.2 merged), 3.2 Torso Reference Frame, 3.3 The Arm Model
with the NEW 3.3.1 kinematic schematic (supervisor comment C9), 3.4 Joint
Angle Computation with 3.4.4 rewritten as "The Wrist as a Tracked Point"
(C10), and 3.5 Worked Example consolidating the three old per-stage
examples into one continuous solve that cites the schematic labels (C11).
The old Unity-mapping section moves to Chapter 6 and is not built here.

Round of 2026-08-31 (supervisor comments C15-C21, PROF_COMMENTS_ROUND3):
the homogeneous T matrix shown wherever R and t appear (general form
eq 3.2/3.4, chain definitions eq 3.5/3.6, root eq 3.12, shoulder eq
3.13, elbow eq 3.14, all three numerically in Section 3.5), and the
information-flow diagram as the new Figure 3.1 (make_ch3_flow_fig.py);
equations and figures renumbered accordingly.

Worked example pinned to frame 533 of the rail recording
recording_20260831_065553 (rail-only round 2026-09-01; replaced R5
frame 270); every number printed here is produced by
writing/v7/scripts/ch3_numbers.py. The hip landmarks of that frame
take their depth from the rail (0.99 m against 1.33 m at the
shoulders), so the root frame decodes to a 32.1 degree pitch on an
upright subject; the worked example states that plainly.
Figures: writing/v8/figures/ (regenerated from the same frame; see
make_ch3_figs.py, make_ch3_chain_fig.py, make_ch3_fig2.py).

2026-09-06 (MATH_LOGIC_REVIEW.md M11, K4, K5): the gimbal-lock branch of
equation (3.17) is printed as an unnumbered display after it; the
worked-example rounding sentence says several degrees, not the last
digit; the right upper arm points outward, not across the body; the
root-frame construction names its non-collinearity condition.

Condensed build 2026-09-06 (writing/v9/CONDENSE_BRIEF.md): the
chapter is cut from about 7,500 to about 4,700 words of prose, captions
and tables. Kept: the coordinate conventions and the person-space flip,
the T notation wherever R and t appear (C15 to C21), Figure 3.1, the
PUMA-style schematic, every equation 3.1 to 3.24, the observability
limits of 3.4.3 and 3.4.4, and the two numeric checks that close the
worked example. Cut: the chapter preview, the standard-mathematics
explanations of Section 3.1, the old Figure 3.6 (seven-DoF arm, standard
robotics) and Figure 3.9 (the Euler order built one rotation at a time),
panel (b) of the old Figure 3.5 (avatar T-pose, duplicates Figure 2.4b),
and most of the printed arithmetic of Section 3.5, which now shows the
inputs, one intermediate step per construction and every result value.
The old Figures 3.3 and 3.8 are merged into one two-panel Figure 3.3;
figures are renumbered 3.1 to 3.7 and the redrawn ones are written to
writing/v9/figures by make_ch3_condensed_figs.py. Equation
numbers are unchanged. Notes: writing/v9/notes_ch3.md.

Humanizer pass 2026-09-06 (writing/WRITING_SKILL.md = blader/humanizer,
user direction): prose redrafted against the 35 patterns plus the kept
supervisor constraints; every EQ/MATH/IMG/TBL item, number, citation,
cross-reference and the ten comment anchors (C9-C11, C15-C21) kept;
review loop in writing/reviews/Chapter_3_humanizer_round*.txt. The
worked-frame cross-reference "Figures 3.4 and 3.9" was a renumbering
leftover and now reads 3.4 and 3.8. Body spelling is Canadian
(modelling); the chapter title keeps the TOC spelling.

Revision of 2026-09-11, Phase 3 (REVISION_2026-09-11_BRIEF.md), continuity
only: no method, parameter, equation, figure, table or number changed. The
opening states what Chapter 2 provides (metric landmarks behind the
visibility gate, the filtering and the hip depth preparation) and what this
chapter makes of them, and adds the section roadmap. Section 3.3.1 carries
the model scope in one place, worded to agree with Chapter 9 Section 9.6
(four of the seven arm rotations; the wrist a point with no hand
orientation, Section 3.4.4; no fingers, contact or force; the rigid torso
plate of Section 3.2.1; legs not tracked) plus the effective segment
lengths of DECISIONS.md D-004/D-005 and skill_set/subject-arm-lengths.md
(medians of the landmark track, not anatomical bone lengths, set against
the subject's measurement in Section 6.3; no length value is printed here).
The closing states what the chapter established and hands on to Chapter 4
(the object observed independently), Chapter 5 (missing landmarks) and
Chapter 6, and sends the solver check to Appendix H instead of "Chapter 7
evaluates the accuracy of the angles" (2026-09-12: D-029 removed solver
verification from Chapter 7, where this closing had pointed at Section
7.5; the author moved that record to Appendix H, so the pointer now goes
there). The
side-subscript paragraph of Section 3.1 is split into sentences that open
with words, which clears the one check_style long-sentence hit (the check
reads text runs only, so a sentence opening with inline math read as a
continuation). Restatement cut to pay for the additions: the chaining
gloss before equation (3.3), the forward/backward gloss after equation
(3.4), the roadmap sentence at the end of Section 3.1, the repeated "to
the subject's right" in Section 3.2.2, the dot-product gloss in Section
3.3.3, and the rigid-plate repeat in Section 3.3.1. The worked-example
frame and Figure 3.3 now name the rail recording.

Craig notation pass of 2026-09-12 (DECISIONS.md D-054, D-056, D-057, D-058,
D-061; FRAME_INVENTORY.md section 8.9.3, the 42 apply-ready rows C3-02 to
C3-51 in line order). No number, equation number, figure, table or
transformation direction changed. Every frame-relative rotation and every
homogeneous transform now carries both frames, the reference frame as a
left superscript and the described frame as a left subscript, under the
semantic names D-058 fixes: Sensor, Person, and the landmark names L24,
L12, L11, L14 and L13 for the link frames. The old subscript convention
went three ways at once (R_root named the described frame, R_sh,r the
reference frame, R_arm,r neither), which is audit item C3-13; the Craig
labels replace it and no side subscript survives. Translations are
written as Craig origin positions, so the translation meaning of t
retires while the angle labels t_y, t_z, t_t are untouched. Equation
(3.22) stays a class-3 operator per D-057, carrying no frame pair, with
its operands and the components equation (3.23) reads labelled in L12 and
the prose saying so. The class-4 maps F of equation (3.1) and M of
Section 3.4.2 keep their own symbol class and are never written in Craig
notation. Not touched, and listed for the author: Figure 3.5's caption
(FRAME_INVENTORY.md section 9, S-6), the unqualified "no matrix is ever
inverted numerically" of Section 3.1 (S-10), and the fixed-angle versus
Euler terminology of equations (3.15) to (3.17).

Round 1 fix pass of 2026-09-11 (writing/reviews/final_2026-09-11/round1/
Chapter_3.txt and Thesis_wide.txt, under FIX_DECISIONS.md; log in
FIXES_ch1-3.md). No number, equation, figure or table changed. Section 3.5
now states that its values are taken before the hip depth preparation of
Section 2.6, at Table 3.1 and again beside the 32.1 degree root pitch,
which is the qualification Chapter 9 Table 9.1 cites. Section 3.1 says what
the object track does with the handedness instead of attributing a choice
to Chapter 4, which makes none. Section 3.3.1 states the missing hand
orientation once. Section 3.4.3 names Sections 9.4 and 9.6 instead of
"Chapter 9".

Prose condensation of 2026-09-14 (user direction; the build, prose and
language rules of writing/v9/REVISION_2026-09-14_BRIEF.md section 1 and
.agent/CHAPTER_REVIEW.md). The prose is cut 15 percent, from 4,469 to 3,797
words. Captions (438 words), headings (69) and Table 3.1 (75) are untouched,
and so are the 24 numbered equations 3.1 to 3.24, their numbers, the
unnumbered gimbal-lock and worked-example displays, the seven figures and
every printed number. Every cross-reference target is kept: the only
reference count that falls is the pair in the cut sentence "Chapter 2 reports
every measurement in the sensor frame of Section 2.2", and both targets are
still referenced from Section 3.1. Kept in full: the Craig frame pairs and
every frame definition, the person-space flip and the T notation of C15 to
C21, the model scope of Section 3.3.1 with the effective segment lengths of
D-004/D-005, the observability limits of Sections 3.4.3 and 3.4.4, the
gimbal-lock branch of equation (3.17), the non-collinearity condition of the
root construction, the rounding sentence and the outward right upper arm of
the worked example, the hip depth qualification of Section 3.5 at Table 3.1
and beside the 32.1 degree pitch, and the two numeric checks that close the
solve.
Cut: the restatement of Chapter 2 in Section 3.1 and of the four solved
angles in the opening, both of which Section 3.3.1 carries; the
standard-mathematics glosses of Section 3.1 (the bottom row of the
homogeneous matrix and the appended 1, the offset-then-transpose reading of
equation 3.4, the quadrant rule of the two-argument arctangent, the
perpendicularity remark after equation 3.9, the sign-correction and
unambiguity remarks after equation 3.7); the third-column vector written
inline in Section 3.4, which equation (3.16) prints; the Section 3.1 preview
of the rotation blocks that Sections 3.2.2 and 3.3.3 build, now a pointer to
Section 3.3.3; the wrist coordinates quoted beside Figure 3.2, which Table
3.1 prints unchanged in both frames; the avatar gloss of the reference
configuration; the "root frame rides the torso" remark of Section 3.1, which
Section 3.2.1 carries, and the animation-system clause beside it, the Unity
world and the camera pose keeping their sentence; the transitional sentences ("The model stops
there", "That has two consequences", "The labels stay the same from here on",
"The same equations run on every frame and on both arms"); and the second
statements of the parent-frame reading of an elbow angle, of the out-of-plane
freedom of the elbow, of the second segment that makes the twist observable
and of the elbow elevation property in the worked example. Sentences
elsewhere are tightened without losing content.
The mapping paragraph after equation (3.6) is written as sentences that open
with words, as the side-subscript paragraph already was, so check_style
reports no long sentence. check_refs reports no unresolved reference. Note for
the author: item S-10 of the 2026-09-12 note above quotes wording ("no matrix
is ever inverted numerically") that the chapter no longer contained before
this round; Section 3.1 now says only that the transpose is the inverse and
that the transformations of the chapter use it directly.

Condensation round 2 of 2026-09-15 (writing/v9/CONDENSE2_BRIEF.md, the
Chapter 3 lines of its "Do" list). Prose 3,812 to 3,599 words; captions,
headings, Table 3.1, all seven figures, all 24 numbered equations with
their numbers, the two unnumbered displays and every printed number are
untouched, so nothing is renumbered. Cut: the chapter opener's section
listing (the Figure 3.1 sentence is kept as the one roadmap sentence) and
its preview of the frame chain; the closer's recap sentence, the forward
pointers to Appendix H and Chapters 4, 5 and 6 staying; in Section 3.1 the
restatement of the transformation notation that Section 2.2.1 carries (the
notation sentence is now a pointer to it), the frame-label list that
Section 2.2.1 also carries, the generic handedness and composition glosses,
the second "right and left refer to the subject" before Section 3.2.1, and
the per-transformation naming after equation (3.5), shortened to the root
transformation plus the two that follow; in Section 3.4 the textbook
extraction steps between equations (3.16) and (3.17). Kept because Section
2.2.1 does not carry them: the columns of the rotation part, the transpose
as the inverse, and the C15 scoping of equations (3.2) to (3.4). The left
counterpart symbols of equation (3.5) stay in print in the shortened naming
sentence (orchestrator decision of 2026-09-15 under comment C16: a
transformation named in the prose carries its T matrix).
Notes: writing/v9/notes_condense2_ch3.md.

"""
# D-073: current frame labels are Camera and Camera'; historical notes
# above retain the terminology of their original decisions. No numeric map changes.
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

from eqn import r, nor, sub, sup, hat, frac, mat, d, eqArr, \
    add_display_eq, add_display_math, add_inline_math

REPO = Path(__file__).resolve().parents[3]
FIG = REPO / "writing" / "v9" / "figures"
FIGC = REPO / "writing" / "v9" / "figures"

H1, H2, H3, P, IMG, CAP, TBL, EQ = "h1", "h2", "h3", "p", "img", "cap", "tbl", "eq"
MATH = "mth"  # displayed, unnumbered worked-example mathematics
PM = "pm"     # paragraph mixing text runs and inline math
BUL = "bul"   # bullet item (supervisor comment C54, 2026-09-16)
NUM = "num"  # numbered item (C54 mechanics follow-up, 2026-09-16; see list_numbering.py)


def T(t):
    return ("t", t)


def X(f):
    return ("m", f)


# --- equation-fragment shorthands ---
def Rot(axis, arg=None):
    base = sub(r("R"), r(axis))
    return base + (d(arg) if arg is not None else "")


def t_(axis):
    return sub(r("t"), r(axis))


def cs(letter, axis):
    return sub(r(letter), r(axis))


# --- Craig frame notation (DECISIONS.md D-054, D-056, D-058) ---
# Class 1 (frame-relative proper rotation) and class 2 (homogeneous frame
# transform) symbols carry both frames: the reference frame as a left
# superscript and the described frame as a left subscript. The labels are
# the system's own names, fixed by D-058 and recorded per frame in
# FRAME_INVENTORY.md section 2: Camera and Camera' for the two frames of
# Chapter 2, and the landmark names L24, L12, L11, L14 and L13 for the
# link frames of the arm chain. The side that the old "sh,r" / "sh,l"
# subscripts carried (user direction 2026-08-31) now comes from the
# landmark itself, so no side subscript survives on a relabelled symbol.
# Class 3 operators (the elementary rotations, and the un-swing of
# equation 3.22) carry no frame pair, and the class-4 maps F and M are
# never written in this notation.
_SLOT = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'


def pre(base, ref, desc=None):
    """Craig left indices: reference frame above, described frame below."""
    lower = desc if desc is not None else _SLOT
    return ('<m:sPre>'
            f'<m:sub>{lower}</m:sub><m:sup>{ref}</m:sup>'
            f'<m:e>{base}</m:e></m:sPre>')


FSEN, FPER = nor("Camera"), nor("Camera'")
FL24, FL12, FL11, FL14, FL13 = (nor("L24"), nor("L12"), nor("L11"),
                                nor("L14"), nor("L13"))
FA, FB = nor("A"), nor("B")


def Ppt(frame, idx):
    """Craig point: the landmark position expressed in one frame."""
    return pre(sub(r("P"), r(idx)), frame)


def Porg(frame, org):
    """Craig origin position: the origin of one frame expressed in another."""
    return pre(sub(r("P"), nor(org + "ORG")), frame)


def Ax(frame, letter, of):
    """Craig axis: one axis direction of a frame expressed in another."""
    return pre(sub(hat(r(letter)), of), frame)


P23, P24, P12 = Ppt(FPER, "23"), Ppt(FPER, "24"), Ppt(FPER, "12")
P14, P16 = Ppt(FPER, "14"), Ppt(FPER, "16")
RAB, TAB = pre(r("R"), FA, FB), pre(r("T"), FA, FB)
XAB, YAB, ZAB = Ax(FA, "X", FB), Ax(FA, "Y", FB), Ax(FA, "Z", FB)
PA, PB = pre(r("P"), FA), pre(r("P"), FB)
RABt = sup(d(RAB), r("T"))
RBA, TBA = pre(r("R"), FB, FA), pre(r("T"), FB, FA)
RPL24 = pre(r("R"), FPER, FL24)
RPL24t = sup(d(RPL24), r("T"))
RL24L12 = pre(r("R"), FL24, FL12)
RL12L14 = pre(r("R"), FL12, FL14)
RPL14 = pre(r("R"), FPER, FL14)
RPL14t = sup(d(RPL14), r("T"))
TPL24 = pre(r("T"), FPER, FL24)
TL24L12 = pre(r("T"), FL24, FL12)
TL12L14 = pre(r("T"), FL12, FL14)
TL24L11 = pre(r("T"), FL24, FL11)
TL11L13 = pre(r("T"), FL11, FL13)
TPL14 = pre(r("T"), FPER, FL14)
TL24P = pre(r("T"), FL24, FPER)
TL12P = pre(r("T"), FL12, FPER)
TL14P = pre(r("T"), FL14, FPER)


def blockT(rot, trans):
    """4x4 homogeneous transformation as a 2x2 block matrix."""
    return mat([[rot, trans], [r("0"), r("1")]])
XL24 = Ax(FPER, "X", FL24)
YL24 = Ax(FPER, "Y", FL24)
ZL24 = Ax(FPER, "Z", FL24)
XL14rest = Ax(FL14, "X", FL14)
XL14arm = Ax(FL12, "X", FL14)
aL12 = pre(hat(r("a")), FL12)
gL14 = pre(hat(r("g")), FL14)
fL12 = pre(r("f"), FL12)
fpL12 = pre(r("f′"), FL12)
sP = pre(r("s"), FPER)
vP = pre(r("v"), FPER)
vL24 = pre(r("v"), FL24)

content = [
(H1, "Chapter 3: Kinematic Modelling"),

(P, "Chapter 2 delivers eight body landmarks per frame as metric 3D positions in one calibrated scene (Appendix C, Appendix D). Points carry no orientation. This chapter converts them into a kinematic state: a root pose, four angles per arm, and the wrist as a point. The model is closed form: every angle comes from the landmark geometry of the current frame, with no optimization and no state carried between frames."),

(P, "Figure 3.1 draws the information flow."),

(IMG, FIGC / "ch3_fig_flow.png", 6.3),
(CAP, "Figure 3.1. The chapter's information flow. Landmarks enter from Chapter 2 on the left, the stored thirteen-angle pose leaves on the right, and the section numbers locate each construction."),

(H2, "3.1 Coordinate Systems and Transformations"),

(P, "This chapter attaches a coordinate frame to each body segment. Unity uses a left-handed frame with y upward, so the kinematics uses the matching axis convention of {Camera'}."),

(P, "Converting a position into {Camera'} changes the handedness itself rather than relabelling the axes. If p is a landmark position in the sensor frame, its position q in the y-up camera frame is"),

(EQ, pre(r("q"), FPER) + r(" = ") + r("F") + r(" ") + pre(r("p"), FSEN) + r(",        ") + r("F") + r(" = ") + nor("diag") + d(r("1, −1, 1")), "3.1"),

(P, 'The determinant of F is -1, the sign of a reflection. F applies to points only, never to rotation matrices, so every rotation in the pipeline comes out proper. Figure 3.2 draws the two frames with the right wrist of Section 3.5. Every landmark position in this chapter is a y-up camera-frame value unless stated otherwise.'),

(IMG, FIGC / "ch3_fig_colocated.png", 5.0),
(CAP, "Figure 3.2. The raw camera frame {Camera} and y-up camera frame {Camera'} drawn from one optical centre. Only the y axis differs: {Camera} reads y downward and {Camera'} upward."),

(P, "The frame {Camera'} is not person-centred, and it is not the frame of the finished Unity scene of Figure 2.4(b), which Chapter 6 builds from the calibrated camera pose and the desk marker of Section 2.2."),

(P, "The left-handed choice changes no physical result. The object track keeps the right-handed world frame instead, and Chapter 6 converts it for Unity with the axis swap of equation (6.1)."),

(PM, [T("Section 2.2 fixes the notation of a rigid body transformation [42]. The reference frame A is a left superscript and the described frame B a left subscript, so "),
      X(RAB), T(" is the orientation of a child frame B in a parent frame A. The columns of the rotation part "),
      X(RAB), T(" are the axis directions of B in the coordinates of A, "),
      X(RAB + r(" = ") + mat([[XAB, YAB, ZAB]])),
      T(". They are perpendicular unit vectors, so the transpose of "),
      X(RAB), T(" is its inverse, which the transformations of this chapter use directly. With the origin of B expressed in A as the translation part "),
      X(Porg(FA, "B")), T(", the two combine into the 4 by 4 homogeneous transformation matrix "),
      X(TAB), T(":")]),

(EQ, TAB + r(" = ") + blockT(RAB, Porg(FA, "B")) + r(",        ") +
     mat([[PA], [r("1")]]) + r(" = ") + TAB + r(" ") +
     mat([[PB], [r("1")]]), "3.2"),

(P, "Running the rotation the other way expresses a measured segment direction in the coordinates of the child frame B:"),

(EQ, pre(r("v"), FB) + r(" = ") + RBA + r(" ") + pre(r("v"), FA) + r(",        ") +
     RBA + r(" = ") + RABt, "3.3"),

(PM, [T("Throughout, "), X(pre(r("P"), FA)),
      T(" is a measured landmark point carrying the frame it is expressed in, and v a segment vector between two points. Equation (3.3) is the rotation part acting alone, so it applies to directions. The chain frames sit at different places, so a landmark moves between them by the full transformation of equation (3.2), whose inverse is")]),

(EQ, eqArr(
     PB + r(" = ") + RABt + r(" ") + d(PA + r(" − ") + Porg(FA, "B")),
     TBA + r(" = ") + sup(d(TAB), r("−1")) + r(" = ") +
     blockT(RABt, r("−") + RABt + r(" ") + Porg(FA, "B"))), "3.4"),

(P, "Figure 3.3(a) draws the chain on the worked-example frame of Section 3.5."),

(IMG, FIGC / "ch3_fig_chain_segments.png", 6.3),
(CAP, "Figure 3.3. The worked-example frame, front view. (a) The chain of frames, the root at right hip L24, shoulder at L12, elbow at L14. (b) The four arm segments."),

(P, "Each neighbouring pair in Figure 3.3(a) is related by its own homogeneous transformation:"),

(EQ, eqArr(
     TPL24 + r(" = ") + blockT(RPL24, Porg(FPER, "L24")),
     TL24L12 + r(" = ") + blockT(r("I"), Porg(FL24, "L12")),
     TL12L14 + r(" = ") + blockT(RL12L14, Porg(FL12, "L14"))), "3.5"),

(PM, [T("A frame that exists once per arm takes its label from its landmark, so the side needs no further subscript: L12 and L14 on the subject's right, L11 and L13 on the left. The root transformation "),
      X(TPL24),
      T(" is the pose of the root frame with respect to the y-up camera frame. The two that follow carry the shoulder frame with respect to the root and the elbow frame with respect to the shoulder, with "),
      X(TL24L11),
      T(" and "), X(TL11L13),
      T(" their left counterparts. Section 3.2 builds the root pose, Section 3.3.3 places both origins, and Section 3.4.2 solves the shoulder rotation that fills the last block and mirrors the construction for the left arm. Chaining the three describes the right elbow frame directly in the y-up camera frame,")]),

(EQ, TPL14 + r(" = ") + TPL24 + r(" ") + TL24L12 + r(" ") + TL12L14, "3.6"),

(PM, [T("Every landmark is converted to the y-up camera frame by equation (3.1), and the inverse transformations of equation (3.4) carry it into the frame where its joint is solved. The shoulder landmark goes into the root frame by "),
      X(TL24P), T(". The elbow goes into the shoulder frame by "),
      X(TL12P), T(" for the angles of Section 3.4.2. The wrist goes into the elbow frame by "),
      X(TL14P), T(" for the flexion of Section 3.4.3. Section 3.5 prints all three as numbers.")]),
(H2, "3.2 Torso Reference Frame"),

(H3, "3.2.1 Landmark Selection"),

(P, "The right hip L24, the left hip L23, and the right shoulder L12, three MediaPipe landmarks [43], define the torso root. Left and right always refer to the subject, who faces the sensor, so the right side appears on the left of the image."),

(P, "These points move more slowly than the hands and forearms, the object rarely hides them, and they form a broad triangle across the trunk, so a small position error changes the angles little."),

(P, 'With three torso points, the model cannot separate the pelvis from the upper trunk. The torso is therefore one rigid plate with no twist between shoulders and hips, and a forward lean pitches the root frame with it. Chapter 9 lists the simplification among the limitations.'),

(H3, "3.2.2 Derivation of the Torso Frame"),

(PM, [T("The flip F of equation (3.1) converts the three positions to the y-up camera frame, giving "),
      X(P23), T(", "), X(P24), T(" and "), X(P12),
      T(". Figure 3.4(a) shows the two input vectors on frame 1400 of the two-hand rail recording, and Figure 3.4(b) the finished frame at L24.")]),

(IMG, FIGC / "ch3_fig_torso_video.png", 6.3),
(CAP, "Figure 3.4. Torso-frame construction. (a) The measured shoulder L12 and hips L23 and L24 define the hip vector and the spine-side vector s. (b) The derived axes at L24."),

(P, "The x axis points along the hip line toward the subject's right, and the hat marks a unit vector."),
(EQ, XL24 + r(" = ") + frac(P24 + r(" − ") + P23,
     d(P24 + r(" − ") + P23, "|", "|")), "3.7"),
(P, "The second input is the spine-side vector from the right hip up to the right shoulder."),
(EQ, sP + r(" = ") + P12 + r(" − ") + P24, "3.8"),
(PM, [T("This vector is not an axis, because it is generally not perpendicular to the hip line: in the worked-example frame the angle between them is 83.8 degrees. It only selects the torso plane, the one plane through the hip line that also contains "),
      X(sP), T(".")]),

(P, "The z axis is the normalized cross product of the two inputs:"),
(EQ, ZL24 + r(" = ") + frac(XL24 + r(" × ") + sP,
     d(XL24 + r(" × ") + sP, "|", "|")), "3.9"),
(PM, [T("The order of the factors sets the sign: with "),
      X(XL24), T(" to the subject's right and "), X(sP),
      T(" up the trunk, the product points forward, out of the chest toward the camera.")]),

(P, "The y axis closes the frame with a second cross product."),
(EQ, YL24 + r(" = ") + ZL24 + r(" × ") + XL24, "3.10"),
(PM, [T("The vectors "), X(ZL24), T(" and "), X(XL24),
      T(" are unit vectors at right angles, so their product needs no normalization. The construction gives an orthonormal frame whenever the two hips are distinct and the shoulder does not lie on the hip line.")]),

(P, "The three axes assemble into the root rotation matrix, its columns in the y-up camera-frame order right, up, forward."),
(EQ, RPL24 + r(" = ") + mat([[XL24, YL24, ZL24]]), "3.11"),
(PM, [T("The right hip is the origin of the root frame, "),
      X(Porg(FPER, "L24") + r(" = ") + P24),
      T(", so the three axes fill in the first transformation of equation (3.5).")]),

(EQ, TPL24 + r(" = ") + blockT(RPL24, P24), "3.12"),

(PM, [T("Figure 3.5 shows the reference configuration that fixes the meaning of the finished matrix: the subject upright, facing the sensor, arms straight out to the sides, the T-pose that is the zero of every arm angle in Section 3.3. Such a subject has the right side toward the camera's left and the chest toward the negative z of the y-up camera frame, giving "),
      X(RPL24), T(" the columns (-1, 0, 0), (0, 1, 0), (0, 0, -1). These decode in the convention of Section 3.4 to the Euler angles (0, 180, 0).")]),

(IMG, FIGC / "ch3_fig_tpose_a.png", 4.6),
(CAP, "Figure 3.5. The reference configuration: the subject in the T-pose with a frame at each landmark. All frames share the root orientation, right arm along +x, left arm along -x."),
(H2, "3.3 The Arm Model"),

(H3, "3.3.1 Kinematic Schematic and Solved Angles"),

(P, "A human arm is described in [42] with seven rotational degrees of freedom, three at the shoulder, one at the elbow and three at the wrist. This model solves the first four: the two shoulder swing angles, the shoulder twist and the elbow flexion. The chain ends at the wrist as a tracked point, since the three wrist rotations cannot be observed from the eight body landmarks."),

(P, "The wrist orientation is not solved. Section 3.4.4 gives the reason no algebra over these landmarks reaches it. Fingers, contact points and grip force lie outside the landmark set, the torso is the single rigid plate of Section 3.2.1, and the legs are not tracked. The recovery of Chapter 5 and the avatar of Chapter 6 use segment lengths to rebuild positions. These are effective lengths, medians of a segment over the landmark track rather than anatomical bone lengths, and Section 6.3 sets them against the subject's own measurement."),

(P, "Figure 3.6 places those four joints on the body as a manipulator diagram. The shoulder at L12 is a cluster of three revolute joints with one shared centre, and a single hinge sits at the elbow L14."),

(IMG, FIG / "ch3_fig_schematic.png", 6.3),
(CAP, "Figure 3.6. The kinematic schematic of the right arm, a chain of revolute-joint cylinders in the Programmable Universal Machine for Assembly (PUMA) style. Each dash-dot line is a rotation axis."),

(PM, [T("The azimuth "),
      X(t_("y")), T(" turns the arm about the torso's y axis, the elevation "),
      X(t_("z")), T(" raises it about the z axis, the twist "),
      X(t_("t")), T(" rolls the upper arm about its own length. The flexion "),
      X(sub(r("e"), r("y"))),
      T(" bends the forearm relative to the upper arm.")]),

(H3, "3.3.2 Definition of Shoulder, Elbow, and Wrist"),

(P, "Each arm is a chain of two rigid segments, the upper arm from the shoulder to the elbow and the forearm from the elbow to the wrist. Figure 3.3(b) shows the four segments, the right arm on L12, L14 and L16 and the left arm on L11, L13 and L15."),

(H3, "3.3.3 Establishing Local Arm Frames"),

(P, "Every joint rotation is measured relative to its parent frame, the shoulder relative to the torso and the elbow relative to the rotated upper arm."),

(PM, [T("The parent frame of the shoulder is L12, the root orientation "),
      X(RPL24),
      T(" moved up the trunk to the shoulder landmark. No separate orientation is derived there: a second basis from the shoulder line and the upper arm would add a shoulder-girdle degree of freedom the data cannot support. With the shoulder angles at zero, the right upper arm points along +x. This fills in the second transformation of equation (3.5):")]),

(EQ, TL24L12 + r(" = ") + blockT(r("I"), RPL24t + d(P12 + r(" − ") + P24)), "3.13"),

(PM, [T("The rotation block is the identity, "),
      X(RL24L12 + r(" = ") + r("I")), T(", and the origin "),
      X(Porg(FL24, "L12") + r(" = ") + RPL24t + d(P12 + r(" − ") + P24)),
      T(" is the measured shoulder landmark in torso coordinates, by the mapping of equation (3.4).")]),

(P, 'The parent of the elbow is the upper arm, so the elbow frame L14 is the root orientation composed with the solved shoulder rotation, with its origin at the elbow landmark. This follows the PUMA robot model, where each new frame is the previous one turned by the joint rotation [42]. In the T-pose, the forearm continues the arm axis along the local +x. The third transformation of equation (3.5) is then'),

(EQ, TL12L14 + r(" = ") + blockT(RL12L14, RPL24t + d(P14 + r(" − ") + P12)), "3.14"),

(PM, [T("with the solved shoulder rotation "), X(RL12L14),
      T(" of Section 3.4.2 in the rotation block and the elbow landmark in shoulder-frame coordinates, "),
      X(Porg(FL12, "L14")), T(", as its origin. The shoulder frame shares the root orientation, so the same "),
      X(RPL24t),
      T(" maps it.")]),

(PM, [T("A measured segment reaches its parent frame in two steps. It is formed as a vector between two landmarks, for the right upper arm "),
      X(vP + r(" = ") + P14 + r(" − ") + P12),
      T(", and equation (3.3) pulls it into the parent frame as "),
      X(vL24 + r(" = ") + RPL24t + r(" ") + vP),
      T(". Section 3.4 works backwards from it to the angles.")]),
(H2, "3.4 Joint Angle Computation"),

(P, "The root orientation is encoded by three Euler angles in the Unity convention (Chapter 6), applied about the parent-frame axes in the order z, then x, then y. The arm angles are joint coordinates; Section 3.4.2 defines their separate composition. In the matrix product the rotation applied first sits nearest the column vector, so the convention reads"),

(EQ, r("R") + r(" = ") + Rot("y", r("y")) + r(" ") + Rot("x", r("x")) + r(" ") + Rot("z", r("z")), "3.15"),
(PM, [T("where "), X(sub(r("R"), r("x"))), T(", "), X(sub(r("R"), r("y"))), T(", "), X(sub(r("R"), r("z"))),
      T(" are the elementary rotations about the coordinate axes. Multiplied out, with c for cosine and s for sine, the product is the matrix this section reads its angles from:")]),
(EQ, r("R") + r(" = ") + mat([
    [cs("c","y")+cs("c","z")+r(" + ")+cs("s","y")+cs("s","x")+cs("s","z"),
     r("−")+cs("c","y")+cs("s","z")+r(" + ")+cs("s","y")+cs("s","x")+cs("c","z"),
     cs("s","y")+cs("c","x")],
    [cs("c","x")+cs("s","z"), cs("c","x")+cs("c","z"), r("−")+cs("s","x")],
    [r("−")+cs("s","y")+cs("c","z")+r(" + ")+cs("c","y")+cs("s","x")+cs("s","z"),
     cs("s","y")+cs("s","z")+r(" + ")+cs("c","y")+cs("s","x")+cs("c","z"),
     cs("c","y")+cs("c","x")]]), "3.16"),

(PM, [T("Write "),
      X(sub(r("m"), r("ij"))),
      T(" for the entry of a measured numerical matrix "), X(r("m")),
      T(" in row i and column j. The three angles follow from its entries:")]),
(EQ, r("x") + r("=") + nor("asin") + d(r("−")+sub(r("m"),r("23"))) + r(", ") +
     r("y") + r("=") + nor("atan2") + d(sub(r("m"),r("13"))+r(",")+sub(r("m"),r("33"))) + r(", ") +
     r("z") + r("=") + nor("atan2") + d(sub(r("m"),r("21"))+r(",")+sub(r("m"),r("22"))), "3.17"),
(PM, [T("The arcsine restricts x to -90 to +90 degrees. At either end of that range ("),
      X(d(sub(r("m"), r("23")), "|", "|") + r(" = 1")),
      T(") the y and z rotations turn about the same axis and only their combined angle is defined, the configuration called gimbal lock. The z angle is set to zero there by convention, so only the split between the two angles is a choice. The entries "),
      X(sub(r("m"),r("13"))), T(" and "), X(sub(r("m"),r("33"))),
      T(" of the third column both vanish there, so the locked branch reads y from "),
      X(sub(r("m"),r("12"))), T(" and "), X(sub(r("m"),r("11"))),
      T(" instead:")]),
(MATH, eqArr(
    r("x = +90°:   y") + r("=") + nor("atan2") + d(sub(r("m"),r("12"))+r(",")+sub(r("m"),r("11"))) + r(",   z = 0"),
    r("x = −90°:   y") + r("=") + nor("atan2") + d(r("−")+sub(r("m"),r("12"))+r(",")+sub(r("m"),r("11"))) + r(",   z = 0"))),

(H3, "3.4.1 Torso Pose"),

(PM, [T("The torso is not a joint. Its pose is the pair "),
      X(RPL24 + r(", ") + Porg(FPER, "L24")),
      T(", and equation (3.17) applied to "), X(RPL24),
      T(" gives its three angles, as Section 3.5 shows on measured data.")]),


(H3, "3.4.2 Shoulder Orientation"),

(PM, [T("A shoulder pose must describe the direction of the upper arm and its roll about its own axis, the three angles labelled at the shoulder of Figure 3.6. They separate into swing and twist [44], with the twist innermost:")]),
(EQ, RL12L14 + r(" = ") + Rot("y", t_("y")) + r(" ") + Rot("z", t_("z")) + r(" ") + Rot("x", t_("t")), "3.18"),
(PM, [T("where "), X(t_("y")), T(" and "), X(t_("z")),
      T(" are the swing angles and "), X(t_("t")),
      T(" the twist about the arm axis +x of the T-pose. A rotation about x leaves the x axis fixed, so with the twist innermost the first column of "),
      X(RL12L14), T(" is free of the twist factor. This column is the upper arm direction "), X(aL12),
      T(" of the elbow frame described in the shoulder frame,")]),
(EQ, aL12 + r(" = ") + XL14arm + r(" = ") + RL12L14 + r(" ") + XL14rest + r(" = ") +
     Rot("y", t_("y")) + r(" ") + Rot("z", t_("z")) + r(" ") + XL14rest + r(" = ") +
     d(cs("c","y")+cs("c","z") + r(",  ") + cs("s","z") + r(",  −") + cs("s","y")+cs("c","z")), "3.19"),
(P, "so the upper-arm direction depends on the swing alone. The twist is invisible in the shoulder-to-elbow vector and has to be measured from a second segment, the forearm. The joint order is not an Euler order, so Chapter 6 rebuilds the rotation from the three angles as three factors in this order. Figure 3.7 shows both halves of the decomposition."),

(IMG, FIGC / "ch3_fig_swing_twist.png", 6.3),
(CAP, "Figure 3.7. The swing-twist decomposition. (a) The swing aims the rest direction +x of the arm into the observed direction â. (b) The twist rolls about the arm axis."),

(PM, [T("The swing angles follow from the components of "), X(aL12),
      T(". Its y component is "), X(cs("s", "z")), T(", so")]),
(EQ, t_("z") + r(" = ") + nor("asin") + d(sub(r("a"),r("y"))), "3.20"),
(PM, [T("the elevation of the arm, positive upward and restricted to -90 to +90 degrees. The x and z components share the factor "),
      X(cs("c", "z")), T(", which cancels in the two-argument arctangent:")]),
(EQ, t_("y") + r(" = ") + nor("atan2") + d(r("−")+sub(r("a"),r("z"))+r(", ")+sub(r("a"),r("x"))), "3.21"),
(PM, [T("the azimuth of the arm, positive when sweeping backward. It is undefined for the arm pointing straight up or straight down ("),
      X(t_("z") + r(" = ±90°")),
      T("), where the x and z components both vanish. The convention "),
      X(t_("y") + r(" = 0")),
      T(" is applied there, and any axial rotation folds into the twist.")]),

(P, "A rotation about the arm axis is only visible in the plane perpendicular to it, so the procedure undoes the swing first, making that plane the y-z plane of the shoulder frame. The matrix product below is a rotation operator given by its axes and its angles, not a frame-relative rotation, so it carries no pair of frame labels. Its input and output are described in the shoulder frame L12:"),
(EQ, fpL12 + r(" = ") + Rot("z", r("−")+t_("z")) + r(" ") + Rot("y", r("−")+t_("y")) + r(" ") + fL12, "3.22"),
(PM, [T("where "), X(fL12),
      T(" is the forearm vector in the parent frame of Section 3.3.3. Once the swing is undone, the nonzero perpendicular component of "), X(fpL12),
      T(" points along the +z axis of L12 at zero twist. Under a twist "),
      X(t_("t")), T(", its unit direction in the (y, z) plane is "),
      X(d(r("−") + nor("sin ") + t_("t") + r(", ") + nor("cos ") + t_("t"))),
      T(", so the twist reads")]),
(EQ, t_("t") + r(" = ") + nor("atan2") + d(r("−")+sub(r("f′"),r("y"))+r(", ")+sub(r("f′"),r("z"))), "3.23"),
(PM, [T("The two components the arctangent reads are components of the rotated vector "),
      X(fpL12),
      T(" in the shoulder frame L12. The twist is positive for internal rotation, where the forearm rotates downward from forward.")]),

(PM, [T("The left arm is solved on mirrored inputs. Both left segment vectors are expressed in the root basis and their x components flipped with "),
      X(r("M") + r(" = ") + nor("diag") + d(r("−1, 1, 1"))),
      T(", the sagittal mirror. The same right-arm solver then runs on the result, so a mirror-symmetric pose gives identical angles on both sides with the same anatomical meaning.")]),

(H3, "3.4.3 Elbow Flexion and Extension"),

(PM, [T("The elbow is solved in the elbow frame L14 of Section 3.3.3. Its orientation in the y-up camera frame is "),
      X(RPL14 + r(" = ") + RPL24 + r(" ") + RL24L12 + r(" ") + RL12L14),
      T(", where the middle factor is the identity of equation (3.13) and the last the solved shoulder rotation of equation (3.18). The unit forearm direction there is "),
      X(gL14 + r(" = ") + frac(RPL14t + r(" ") + d(P16 + r(" − ") + P14),
                                   d(P16 + r(" − ") + P14, "|", "|"))),
      T(", and its swing is read from the components of "),
      X(gL14), T(" as at the shoulder:")]),
(EQ, sub(r("e"),r("z")) + r(" = ") + nor("asin") + d(sub(r("g"),r("y"))) + r(",      ") +
     sub(r("e"),r("y")) + r(" = ") + nor("atan2") + d(r("−")+sub(r("g"),r("z"))+r(", ")+sub(r("g"),r("x"))), "3.24"),
(PM, [T("Here "), X(sub(r("e"), r("y")) + r(" = 0")), T(" is a straight arm and "),
      X(sub(r("e"), r("y")) + r(" = −90")), T(" is a right angle bend, the hinge angle labelled in Figure 3.6.")]),

(PM, [T("The angle "), X(sub(r("e"), r("z"))),
      T(" is always zero. The shoulder twist aligns the perpendicular component of the forearm with local +z, so the flexion plane is already the x-z plane of the elbow frame by the time the elbow is solved. That closes the dimension count: two measured segment directions give four observable degrees of freedom and the model extracts four independent angles. The angle "),
      X(sub(r("e"), r("z"))),
      T(" stays among the streamed angles for generality but holds no independent information.")]),

(PM, [T('A three-landmark arm breaks down in one case. When the arm is straight, the forearm is parallel to the upper arm, its perpendicular component vanishes, and no algebra can recover a roll about an axis from two collinear segments. The solver detects this when the forearm lies along the upper arm axis and the perpendicular part of the forearm vanishes against its length, applies '),
      X(t_("t") + r(" = 0")),
      T(" so the elbow frame stays defined, and flags the twist as unobservable; later stages hold the last valid value. A nearly straight elbow leaves the flexion well conditioned but makes the twist sensitive to noise, a limitation Chapter 9 discusses.")]),

(H3, "3.4.4 The Wrist as a Tracked Point"),

(P, "The wrist L16 is not a solved joint. It enters as a tracked 3D point, as Figure 3.6 draws it, and does two jobs: it gives the forearm direction the elbow flexion is read from, and it makes the shoulder twist observable."),

(P, "A full three-degree-of-freedom wrist orientation would need a landmark beyond the wrist. Forearm pronation, the roll of the forearm about its own axis, does not move the wrist point, so no algebra over these landmarks can observe it. The tracked position is still kept and transmitted, and when the landmark is lost during a grip the recovery in Chapter 5 rebuilds it with an additional geometric constraint from the object pose. A hand model such as [41] would attach at the same point."),
(H2, "3.5 Worked Example"),

(P, "Frame 533 of the rail recording, drawn in Figure 3.3, is taken while the cube slides along the rail in the task of Chapter 2. All eight landmarks are measured on it, none flagged by the despiking stage or filled by the gap-bridging stage of Section 2.5, and no failure detector of Chapter 5 fires on it. Values are computed from the unrounded positions and printed rounded to two decimals, so a step recomputed by hand from the printed inputs can differ by several degrees where the arguments are small. The two-argument arctangent of the printed 0.06 and 0.03 gives 63.4 degrees, while the unrounded values give the 59.9 degrees printed below."),

(P, "Table 3.1 lists the eight landmark positions in both frames, after the filtering of Section 2.5 and before the hip depth preparation of Section 2.6. The same frame runs through the worked examples of Appendix C, Appendix D and Appendix E."),

(TBL, [["Landmark", "Camera (m)", "Camera' (m)"],
       ["L11 left shoulder", "(0.22, -0.34, 1.36)", "(0.22, 0.34, 1.36)"],
       ["L12 right shoulder", "(-0.16, -0.33, 1.31)", "(-0.16, 0.33, 1.31)"],
       ["L13 left elbow", "(0.28, -0.03, 1.38)", "(0.28, 0.03, 1.38)"],
       ["L14 right elbow", "(-0.20, -0.04, 1.19)", "(-0.20, 0.04, 1.19)"],
       ["L15 left wrist", "(0.24, 0.20, 1.25)", "(0.24, -0.20, 1.25)"],
       ["L16 right wrist", "(-0.20, 0.11, 1.06)", "(-0.20, -0.11, 1.06)"],
       ["L23 left hip", "(0.12, 0.18, 0.99)", "(0.12, -0.18, 0.99)"],
       ["L24 right hip", "(-0.06, 0.19, 0.99)", "(-0.06, -0.19, 0.99)"]]),
(CAP, "Table 3.1. The eight landmarks of the worked-example frame in the sensor frame and in the y-up camera frame, printed to two decimals. Each column carries the frame its entries are expressed in, the left superscript of the position symbols used from here on. The conversion flips only the sign of y, and the hips read negative in the y-up camera frame because they sit below the camera height."),

(P, "The torso construction of Section 3.2 needs three of these points:"),
(MATH, eqArr(
    P23 + r(" = (0.12, −0.18, 0.99)"),
    P24 + r(" = (−0.06, −0.19, 0.99)"),
    P12 + r(" = (−0.16, 0.33, 1.31)"))),

(P, "Equations (3.7) and (3.8) give the hip line, 0.18 m long, and the spine-side vector, 0.63 m long. The cross products of equations (3.9) and (3.10) close the frame:"),
(MATH, eqArr(
    XL24 + r(" = (−1.00, −0.05, −0.04)"),
    ZL24 + r(" = (0.01, 0.53, −0.85)"),
    YL24 + r(" = (−0.06, 0.85, 0.53)"))),
(PM, [T("The x component of "), X(XL24),
      T(" is close to -1, and "),
      X(ZL24),
      T(" is tipped upward by the hip-to-shoulder depth difference. Equation (3.11) assembles the columns:")]),
(MATH, RPL24 + r(" = ") + mat([
    [r("−1.00"), r("−0.06"), r("0.01")],
    [r("−0.05"), r("0.85"), r("0.53")],
    [r("−0.04"), r("0.53"), r("−0.85")]])),
(PM, [T("Appending the origin "), X(Porg(FPER, "L24") + r(" = ") + P24),
      T(" gives the root pose of equation (3.12) as numbers:")]),
(MATH, TPL24 + r(" = ") + mat([
    [r("−1.00"), r("−0.06"), r("0.01"), r("−0.06")],
    [r("−0.05"), r("0.85"), r("0.53"), r("−0.19")],
    [r("−0.04"), r("0.53"), r("−0.85"), r("0.99")],
    [r("0"), r("0"), r("0"), r("1")]])),
(P, "Equation (3.13) needs the shoulder landmark in the torso's own coordinates, which equation (3.4) gives as (0.07, 0.63, 0.00). The forward offset is zero, since the spine-side vector lies in the torso plane by construction:"),
(MATH, TL24L12 + r(" = ") + mat([
    [r("1"), r("0"), r("0"), r("0.07")],
    [r("0"), r("1"), r("0"), r("0.63")],
    [r("0"), r("0"), r("1"), r("0.00")],
    [r("0"), r("0"), r("0"), r("1")]])),
(P, "an identity rotation with a pure origin shift. Equation (3.17) decodes the root matrix:"),
(MATH, eqArr(
    r("x") + r(" = ") + nor("asin") + d(r("−0.53")) + r(" = −32.1°"),
    r("y") + r(" = ") + nor("atan2") + d(r("0.01, −0.85")) + r(" = 179.5°"),
    r("z") + r(" = ") + nor("atan2") + d(r("−0.05, 0.85")) + r(" = −3.3°"))),
(P, "Compared with the T-pose (0, 180, 0), the subject faces almost straight at the sensor and the hip line is almost level. The frame is pitched by 32.1 degrees because the hips sit nearer the camera than the shoulders. The two hip landmarks fall on the pixels where the pelvis meets the rail, so the depth sampled there is the depth of the rail, 0.99 m against 1.31 and 1.36 m at the shoulders. Chapter 5 describes this kind of failure. The pipeline replaces the hip depth with the preparation of Section 2.6 before it solves, and the numbers of this section are taken before that replacement. The arm angles below are read against the pitched frame, and these three values are the torso pose of Section 3.4.1."),

(P, "The y-up camera-frame upper arm vector of the right arm is"),
(MATH, vP + r(" = ") + P14 + r(" − ") + P12 + r(" = (−0.04, −0.29, −0.12)")),
(P, "with length 0.32 m. Equation (3.3) pulls it into the parent frame and normalizing gives the local arm direction:"),
(MATH, eqArr(
    vL24 + r(" = ") + RPL24t + r(" ") + vP + r(" = (0.05, −0.31, −0.05)"),
    aL12 + r(" = (0.17, −0.97, −0.16)"))),
(P, "The upper arm points slightly outward, strongly downward and slightly behind the pitched torso plane, since +x for this right arm points outward to the subject's right. Figure 3.3(b) shows that pose, the right arm reaching down to the cube. The same two steps on the forearm give"),
(MATH, fL12 + r(" = ") + RPL24t + r(" ") + d(P16 + r(" − ") + P14) + r(" = (0.01, −0.20, 0.03)")),
(P, "with length 0.20 m. Equations (3.20) and (3.21) read the two swing angles of the right arm (Figure 3.6):"),
(MATH, eqArr(
    t_("z") + r(" = ") + nor("asin") + d(r("−0.97")) + r(" = −76.6°"),
    t_("y") + r(" = ") + nor("atan2") + d(r("0.16, 0.17")) + r(" = 43.0°"))),
(P, "The arm points 76.6 degrees below horizontal, and equation (3.22) undoes this swing on the forearm vector:"),
(MATH, fpL12 + r(" = ") + Rot("z", r("76.6°")) + r(" ") + Rot("y", r("−43.0°")) + r(" ") + fL12 + r(" = (0.19, −0.06, 0.03)")),
(PM, [T("Its perpendicular part points mostly along negative y. Equation (3.23) gives the twist:")]),
(MATH, t_("t") + r(" = ") + nor("atan2") + d(r("0.06, 0.03")) + r(" = 59.9°")),
(P, 'The forearm turns inward across the front of the body, an internal rotation. In the elbow frame, the forearm direction is'),
(MATH, gL14 + r(" = (0.94, 0.00, 0.34)")),
(PM, [T("Its y component is zero, the "),
      X(sub(r("e"), r("z")) + r(" = 0")),
      T(" property of Section 3.4.3 on the measured frame. The flexion of equation (3.24) is")]),
(MATH, sub(r("e"), r("y")) + r(" = ") + nor("atan2") + d(r("−0.34, 0.94")) + r(" = −19.8°")),
(PM, [T("The solved shoulder rotation and the local upper-arm vector "),
      X(vL24),
      T(" complete the last transformation of equation (3.5), the elbow pose of equation (3.14):")]),
(MATH, TL12L14 + r(" = ") + mat([
    [r("0.17"), r("0.95"), r("−0.27"), r("0.05")],
    [r("−0.97"), r("0.12"), r("−0.20"), r("−0.31")],
    [r("−0.16"), r("0.30"), r("0.94"), r("−0.05")],
    [r("0"), r("0"), r("0"), r("1")]])),
(PM, [T("Its first column is exactly "),
      X(aL12),
      T(" and its last column is the elbow position in the shoulder frame. The numeric chain closes with two checks. The product of the three matrices of equation (3.6), applied to the origin, "),
      X(mat([[Porg(FPER, "L14")], [r("1")]]) + r(" = ") + TPL14 + r(" ") + sup(d(r("0, 0, 0, 1")), r("T"))),
      T(", lands on (-0.20, 0.04, 1.19), the measured elbow position of Table 3.1, to machine precision. Equation (3.4) applied three times carries the measured wrist through the inverse chain into the elbow frame, "),
      X(mat([[pre(sub(r("P"), r("16")), FL14)], [r("1")]]) + r(" = ") + sup(d(TPL14), r("−1")) + r(" ") + mat([[P16], [r("1")]])),
      T(". The result is (0.19, 0.00, 0.07), whose length is the forearm length 0.20, and dividing by that length reproduces the "),
      X(gL14),
      T(" used for the flexion.")]),
(P, "The left arm runs through the same mirrored solver and gives a shoulder azimuth of 76.1 degrees, an elevation of -53.5 degrees, a twist of 108.9 degrees and an elbow flexion of -37.3 degrees, the arm hanging at the side in Figure 3.3(b). The stored pose has thirteen angular values: three root Euler angles, two shoulder swing angles, shoulder twist and elbow flexion for each arm, plus the two redundant elbow elevation fields, eleven of them independent."),

(P, "Appendix H records the verification of the chapter's mathematics against synthetic exact references. A missing or misplaced landmark leaves that arm without a state. Chapter 5 supplies an additional geometric constraint from the object Chapter 4 tracks. Chapter 6 carries the finished state into the reconstruction."),
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

# D-106: unnumbered display ordinals continuing into the following clause.
DISPLAY_COMMAS = {8, 10, 6}
math_display_index = 0

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == H3:
        doc.add_heading(item[1], level=3)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == PM:
        p = doc.add_paragraph()
        for typ, val in item[1]:
            if typ == "t":
                p.add_run(val)
            else:
                add_inline_math(p, val)
    elif kind == EQ:
        add_display_eq(doc, item[1], item[2])
    elif kind == MATH:
        math_display_index += 1
        add_display_math(doc, item[1],
                         punctuation="," if math_display_index in DISPLAY_COMMAS else ".")
    elif kind == IMG:
        path = Path(item[1])
        assert path.exists(), f"missing figure: {path}"
        doc.add_picture(str(path), width=Inches(item[2]))
        if path.name in ("ch3_fig_flow.png", "ch3_fig_torso_video.png"):
            # D-076/D-087: keep these figures with their following captions.
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
    elif kind == BUL:
        # bullet item (supervisor comment C54, 2026-09-16)
        p = doc.add_paragraph(style="List Bullet")
        if item[1]:
            p.add_run(item[1]).font.bold = True
            p.add_run(" " + item[2])
        else:
            p.add_run(item[2])
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
    elif kind == NUM:
        # numbered item, restarted per contiguous run (C54 mechanics
        # follow-up, 2026-09-16; see list_numbering.py)
        p = doc.add_paragraph(style="List Number 2")
        if item[1]:
            p.add_run(item[1]).font.bold = True
            p.add_run(" " + item[2])
        else:
            p.add_run(item[2])
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row_ in enumerate(rows):
            for j, c in enumerate(row_):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

out = REPO / "writing" / "v9" / "Chapter_3_Kinematic_Modeling.docx"
from list_numbering import restart_numbered_lists
restart_numbered_lists(doc)

doc.save(out)
print("saved", out, "| items:", len(content))
