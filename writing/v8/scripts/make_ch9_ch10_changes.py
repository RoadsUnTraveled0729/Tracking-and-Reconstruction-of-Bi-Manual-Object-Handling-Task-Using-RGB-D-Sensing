#!/usr/bin/env python3
"""Build the concise change log for the supervisor for round 7 (2026-09-10):
writing/v8/Thesis_V8_Changes_Abstract_Ch9_Ch10.docx. It compares this
version with the version he last saw (9 September, his verdict "Everything
looks good") for the abstract and the old Chapter 9, which is now
Chapters 9 and 10. Word counts are read from the built parts
(count_words.py: prose plus captions plus tables); the 9 September counts
are the old Chapter 9 build of commit de74302. Same format as
make_ch7_ch8_changes.py."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "writing" / "v8" / "condensed" / "scripts"))
from count_words import count  # noqa: E402

C = REPO / "writing" / "v8" / "condensed"
c9 = count(str(C / "Chapter_9_Discussion.docx")); W9 = c9[0] + c9[1] + c9[3]
c10 = count(str(C / "Chapter_10_Conclusions_Future_Work.docx")); W10 = c10[0] + c10[1] + c10[3]
OLD9_WORDS = 2239   # old Chapter 9 of 9 September (count_words on the de74302 build)
OLD_ABSTRACT_WORDS = 349
ABSTRACT_WORDS = int(sys.argv[1]) if len(sys.argv) > 1 else 349
PAGES_OLD = 166
PAGES_NEW = sys.argv[2] if len(sys.argv) > 2 else "166"


def new_doc(title, intro):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    doc.add_heading(title, level=1); doc.add_paragraph(intro); return doc


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, txt in enumerate(r):
            cell = t.cell(i, j); cell.text = txt
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10); run.font.bold = (i == 0)


def numbered(doc, items):
    for head, body in items:
        p = doc.add_paragraph(style="List Number"); r = p.add_run(head + " "); r.bold = True; p.add_run(body)


doc = new_doc(
    "Changes from the version of 9 September to this version: abstract, Chapter 9 and Chapter 10",
    "The version you commented on (9 September, \"Everything looks good.\") asked for a rewritten abstract and for the last chapter to become two: a discussion of the results and a conclusions and future work chapter. This version answers both. The three parts are also sent as separate files: Abstract_V8.docx, Chapter_9_Discussion.docx and Chapter_10_Conclusions_Future_Work.docx. Everything else in the thesis is the same apart from the sentences that point at the new chapters.")
doc.add_heading("How much changed", level=2)
table(doc, [["", "9 September", "This version"],
            ["Abstract words", str(OLD_ABSTRACT_WORDS), str(ABSTRACT_WORDS)],
            ["Abstract paragraphs", "3", "4"],
            ["Chapter 9 words", str(OLD9_WORDS), str(W9)],
            ["Chapter 9 sections", "3", "7"],
            ["Chapter 9 figures and tables", "2 and 1", "2 and 1"],
            ["Chapter 10 words", "(part of Chapter 9)", str(W10)],
            ["Chapter 10 sections", "", "2"],
            ["Thesis pages", str(PAGES_OLD), str(PAGES_NEW)]])
doc.add_heading("Abstract", level=2)
numbered(doc, [
 ("Four paragraphs in your order (your comment).",
  "The first says what the problem is and why: two-handed handling of objects, its uses in robot imitation learning, avatar animation and clinical assessment, and the cost of capture arrays and wearables against one consumer RGB-D camera and three printed markers. The second names the challenges: a landmark returned whether or not the body part is there, wrists hidden by the object and the other hand, hips covered by the hands and taking the desk depth, no wrist orientation in the landmark set, and the need to run causally. The third says how they were addressed: the two measurement chains that share only the frames and the world frame, so one stands in for the other; the wrist rebuilt from the object pose and the elbow by two-link inverse kinematics; the geometric detectors; the remembered hip depth; the measured, held or constrained label on every quantity. The fourth says what was accomplished, with the same numbers as before: the rail as the physical straight-line reference, 1.0 and 2.4 centimetres on the first recording, 0.3 and 1.6 on the handover recording, the recovery graded against the unmasked solve and the hand-placed labels, and the causal version at 29.0 frames per second."),
 ("Unchanged.",
  "The ground-truth statement you asked for earlier (the tracked object trajectory compared with the straight rail defined in the laboratory) and every number. The abstract stays within 350 words."),
])
doc.add_heading("Chapter 9: Discussion", level=2)
numbered(doc, [
 ("A discussion by result (your comment).",
  "The old Chapter 9 listed the limitations in one section. The new chapter discusses each result of Chapters 7 and 8 in turn and says what it means: 9.1 the object track against the rail (what the residual across the fitted line does and does not show, the physical cross-checks, the marker scale, the blocked marker); 9.2 the arm pose and the recovery (why the synthetic windows test only the slow case, the label floor, where the recovery helps and where the frozen grip offset defeats it); 9.3 the second arm and the handover (both arms measured, the cost on the torso, the hip-hold alternative with its numbers on both recordings); 9.4 the torso and the hidden hips; 9.5 landmark failures at high confidence; 9.6 the live structure (what the causal path gives up, shown on the frames); 9.7 the limits of the model and the evaluation with Table 9.1."),
 ("Same evidence.",
  "Both failure-mode figures and Table 9.1 stay; the marker figure now comes first because Section 9.1 cites it, so it is Figure 9.1 and the landmark figure 9.2. Every number in the chapter was already printed in Chapters 7 and 8 or in the old Chapter 9; the new sentences are interpretation, not new results."),
])
doc.add_heading("Chapter 10: Conclusions and Future Work", level=2)
numbered(doc, [
 ("Conclusions (your comment).",
  "Section 10.1 answers the five objectives of Chapter 1 in order, as the old Section 9.3 did, followed by the real-time result, the torso preparation caveat, the systems framing and the scope statement."),
 ("Future work.",
  "Section 10.2 carries the old Section 9.2: a live-sensor session, more subjects, a hand-specific chain, the chain-level Bayesian recovery, the hip-hold rule (its numbers now in Section 9.3) and the hidden regrasp."),
 ("Sentences elsewhere.",
  "Chapter 1 (objective 4 and the chapter roadmap), Chapter 3, Chapter 7 (the end of Section 7.7), Chapter 8 (the end of Section 8.4) and Appendices E and G now point at Chapter 10 or at the numbered section of Chapter 9 that discusses the item."),
])
doc.add_paragraph("Unchanged: every number, figure and table of Chapters 1 to 8 and the appendices, and every recording.")
OUT = REPO / "writing" / "v8" / "Thesis_V8_Changes_Abstract_Ch9_Ch10.docx"
doc.save(OUT); print("saved", OUT, W9, W10)
