#!/usr/bin/env python3
"""Build writing/v8/Thesis_V8_Changes_Ch2_Ch3.docx: the concise change log
for the supervisor covering Chapters 2 and 3, V7 (as sent) to V8, organised
by his round 4 comments (C43-C47). Counts come from the built chapter files
(word count of the prose, figures, tables, numbered equations)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[3]
OUT = REPO / "writing" / "v8" / "Thesis_V8_Changes_Ch2_Ch3.docx"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)

doc.add_heading("Changes from V7 to V8: Chapters 2 and 3", level=1)
doc.add_paragraph("V7 is the version you commented on (2026-09-05). This note lists what changed in the two chapters and which of your comments each change answers. Other chapters are unchanged in this round.")

doc.add_heading("How much changed", level=2)
rows = [["", "Chapter 2 (V7)", "Chapter 2 (V8)", "Chapter 3 (V7)", "Chapter 3 (V8)"],
        ["Words", "1803", "3562", "7617", "7477"],
        ["Sections", "4", "5 (2.2 new, with 2.2.1 and 2.2.2)", "14", "14"],
        ["Figures", "8", "11", "11", "10"],
        ["Tables", "1", "3", "1", "1"],
        ["Numbered equations", "0", "1", "24", "24"]]
t = doc.add_table(rows=len(rows), cols=5); t.style = "Table Grid"
for i, r in enumerate(rows):
    for j, c in enumerate(r):
        cell = t.cell(i, j); cell.text = c
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(10); run.font.bold = (i == 0)

doc.add_heading("Chapter 2: Experimental Setup and Measurements", level=2)
doc.add_paragraph("The chapter is retitled (it was Experimental Setup and Data Acquisition) because it now defines the frames and the calibration as well as the sensors. Sections 2.2 to 2.4 of V7 are now 2.3 to 2.5.")
items = [
 ("Frames, T-matrix notation and calibration in Chapter 2 (your comment: define the coordinate frames and the T matrices here, like Chapter 2 of the robotics book, and include the calibration section).",
  "New Section 2.2 Coordinate Frames and Scene Calibration. Section 2.2.1 defines every frame: the camera frame C, the world frame W on the desk marker, the wall marker frame, the object frame O, person space P, the body frames of Chapter 3, and the Unity frame U. The T notation is given once, and Table 2.2 lists every transformation of the thesis: what it carries a point from and into, whether it is fixed by the calibration or measured on every frame, and where it is derived. Section 2.2.2 is the static scene calibration moved from Chapter 4 (old Section 4.1): the chordal-mean anchor with equation (2.1), Table 2.3 with the frozen scene description, and Figure 2.6 with the calibrated scene. Chapter 3 Section 3.1 and Chapter 4 now refer back to it instead of defining the frames again. The camera-frame photograph (old Figure 3.2) moved here as Figure 2.4."),
 ("Overall picture of the scene with the frames drawn (your comment: draw a frame on each ArUco marker, the sensor frame, and arrows for the T matrices).",
  "New Figure 2.5. Panel (a) is the colour frame of the recording with the wall marker frame, the world frame W on the desk marker, the object frame O on the cube and the camera frame C, with an arrow for each transformation (T_C,wall, T_CW, T_CO, T_WO). Panel (b) is a top view of the same frames with the object path and the subject at the working distance. The world frame stays on the desk marker rather than the wall marker: the desk marker sits inside the working area next to the object path and the tabletop, so heights and object positions read directly in it, while the wall marker only supplies the gravity direction. The data-flow diagram (Figure 2.3) now shows the scene calibration block in Pipeline B."),
 ("Rename Section 2.4 and say where the filtering goes next.",
  "The section is now 2.5 Measurement Filtering. Its closing paragraph states where each output goes: the filtered landmarks are what Chapter 3 solves the joint angles from, the object track is cleaned in its own stage in Chapter 4, and Chapter 5 uses the cleaned object pose to recover a held wrist."),
 ("Order of the stages in the data-flow diagram (follow-up to the Chapter 5 rewrite).",
  "Figure 2.3 placed the pose recovery of Chapter 5 after the kinematic solver. Chapter 5 places it between the filtering and the solve, because the layer rebuilds the failed landmarks before the solve runs and only the twist hold and the rate limit act on the angles after it. The recovery block now sits between the filter block and the solver block, with the object-pose arrow from Pipeline B into it, and the sentence that walks through Pipeline A gives the same order."),
 ("Two decimals at most (your comment of 2026-09-07 on Chapter 4, applied to every chapter).",
  "One number in Section 2.5 had more: the factor that scales the median absolute deviation, 1.4826, is now written as about 1.48. Appendix F carries the same change. Nothing else in Chapters 2 and 3 had more than two decimals; the T matrices of Chapter 3 were already printed that way."),
 ("Use AI for clarity but keep my level of English and vocabulary.",
  "Sections 2.3 to 2.5 keep my own sentences. The edits there are limited to updated cross-references, a few short linking sentences (for example that the landmark coordinates are in the camera frame C of Section 2.2), and small clarity fixes such as naming who does what. Sections 2.1 and 2.2 are written in the same plain style."),
]
for head, body in items:
    p = doc.add_paragraph(style="List Number"); r = p.add_run(head + " "); r.bold = True; p.add_run(body)

doc.add_heading("Chapter 3: Kinematic Modeling", level=2)
doc.add_paragraph("No new comments on this chapter in this round. The changes follow from Chapter 2 and from the clarity direction:")
items3 = [
 "Section 3.1 no longer defines the camera frame and the scene frames. It refers to Section 2.2 and to Figures 2.4 and 2.5, and keeps only the person-space flip (equation (3.1)), the homogeneous T (equations (3.2) to (3.6)) and the body chain.",
 "The camera-frame photograph (old Figure 3.2) moved to Chapter 2, so the remaining figures are renumbered 3.2 to 3.10. All 24 equations and their numbers are unchanged, including the T matrices added for your round 3 comments.",
 "The prose is revised for readability in the same plain style as Chapter 2: shorter and plainer sentences, the actor named, no change to the content, the equations, the figures or the worked-example numbers.",
 "Figure 3.1 and its caption now say that the landmarks enter after the recovery of Chapter 5 has replaced any that failed, so that the figure agrees with Figure 2.3 and with Chapter 5. The flow itself is unchanged.",
 "One cross-reference corrected: the worked-example frame is the one shown in Figures 3.4 and 3.8 (it said 3.9 after the renumbering).",
]
for b in items3:
    doc.add_paragraph(b, style="List Bullet")
doc.save(OUT); print("saved", OUT)
