#!/usr/bin/env python3
"""Build the concise change logs for the supervisor for Chapters 7 and 8:
writing/v8/Thesis_V8_Changes_Ch7.docx and Thesis_V8_Changes_Ch8.docx.
Chapter 7: the version he last saw (8 September 2026) to this version (the
handover recording, round 6 C50); Chapter 8: V7 as sent (2026-09-05) to the
condensed V8 of 2026-09-07, by the author's restructuring brief (decision
D-008), since he made no comment on it. Word counts are read from the built chapters (count_words.py: prose plus
captions plus tables); the V7 counts are those CONDENSE_SUMMARY.md records.
Same format as make_ch4_ch5_changes.py."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "writing" / "v8" / "condensed" / "scripts"))
from count_words import count  # noqa: E402

C = REPO / "writing" / "v8" / "condensed"
c7 = count(str(C / "Chapter_7_Evaluation.docx")); W7 = c7[0] + c7[1] + c7[3]
c8 = count(str(C / "Chapter_8_Real_Time_Feasibility.docx")); W8 = c8[0] + c8[1] + c8[3]


def new_doc(title, intro):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    doc.add_heading(title, level=1); doc.add_paragraph(intro); return doc


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, txt in enumerate(r):
            cell = t.cell(i, j); cell.text = txt
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10); run.font.bold = (i == 0)


def numbered(doc, items):
    for head, body in items:
        p = doc.add_paragraph(style="List Number"); r = p.add_run(head + " "); r.bold = True; p.add_run(body)


# ------------------------------------------------------------------ Chapter 7
# Since 2026-09-09 the Chapter 7 log compares THIS version with the version
# the supervisor last saw (sent 2026-09-08, his round-6 comment C50 on it),
# at the user's direction ("be concise, just the comparison between this
# version and last version"). The earlier V7-to-V8 list of the 2026-09-07
# restructure is in the git history of this script (commit ec0467c).
doc = new_doc(
    "Changes from the version of 8 September to this version: Chapter 7",
    "The version you commented on (8 September, \"It looks good. Just one comment.\") asked for the second arm in the wooden block experiment. This version answers it with a new recording and a new section; everything else in the chapter is the same apart from renumbering.")
doc.add_heading("How much changed", level=2)
table(doc, [["", "Chapter 7 (8 September)", "Chapter 7 (this version)"],
            ["Words", "4796", str(W7)],
            ["Sections", "8", "9"],
            ["Figures", "12", "15"],
            ["Tables", "4", "5"],
            ["Thesis pages", "161", "166"]])
doc.add_heading("Chapter 7: Evaluation", level=2)
numbered(doc, [
 ("New Section 7.7, Handover on the Rail (your comment).",
  "I recorded the rail task again on the same setup with the handover you described: the right hand carries the cube to the rail and slides it, the left hand takes it over on the rail and slides it to the far end. The section reports the recording as it came out of the same pipeline, errors included. Figure 7.11 shows four colour frames (the right-hand slide, both hands at the cube, the left-hand slide, the far end). Table 7.1 gives the deviation of the marker origin from the rail line per part of the slide: right hand 29.7 cm of travel with a median of 0.7 cm, the 5.3 s handover with a median of 0.3 cm and at most 1.7, the left hand with a median of 0.1 cm, together with the frames on which the torso detector fires and the hip preparation replaces the hips. Figure 7.12 draws the marker origin along the rail with the hand at the cube marked, and the slide from above and from the side with both model wrists. Figure 7.13 shows the trails inside the Unity scene with both wrists, one panel per part."),
 ("What went wrong, stated in the section.",
  "The wall card had shifted between the two recordings, so the levelling of the new recording uses the gravity calibrated on the first one; the camera and the desk marker had not moved. The hands in front of the hips take the hip depth for 14 seconds, so the torso rests on the hip depth preparation of Section 2.6, and the model wrist jumps by up to 23 cm at the three moments the hips enter or leave that repair. The section ends with the alternative of holding the hips at their last accepted position, with its numbers; the thesis keeps the existing preparation, and Section 9.2 discusses the assumption."),
 ("Renumbering.",
  "The recovery section is now 7.8 and the solver check 7.9; the label figures are 7.14 and 7.15; the rail masking table is 7.2 and the label and synthetic tables 7.3 to 7.5. The loop recording is called the loop recording throughout, no longer the second recording."),
 ("Sentences elsewhere.",
  "The abstract, Chapter 1 (scope, Section 1.3, objective 5), Section 2.1 and Chapter 9 (9.1, Table 9.1, 9.3, objective 5) no longer describe the rail task as one-handed only; Section 9.2 gains one paragraph on the hip-hold assumption; Chapters 5, 6 and 9 now cite the recovery section as 7.8."),
])
doc.add_paragraph("Unchanged: every number of the first recording, Sections 7.1 to 7.6 apart from the sentences that name the new recording, the recovery results, the solver check, and the appendices.")
doc.save(REPO / "writing" / "v8" / "Thesis_V8_Changes_Ch7.docx"); print("saved Ch7 log", W7)

if "--chapter7-only" in sys.argv:
    sys.exit(0)

# ------------------------------------------------------------------ Chapter 8
doc = new_doc(
    "Changes from V7 to V8: Chapter 8",
    "V7 is the version you have (sent 2026-09-05). You made no comment on Chapter 8. The chapter was rewritten so that it explains how the offline pipeline was carried over to a live system, what that system looks like, and what the transfer costs, instead of reporting timing tables and live-versus-offline error statistics. How accuracy is measured stays in Chapter 7.")
doc.add_heading("How much changed", level=2)
table(doc, [["", "Chapter 8 (V7)", "Chapter 8 (V8)"],
            ["Words", "3717", str(W8)],
            ["Sections", "3", "4"],
            ["Figures", "2", "3 (all new)"],
            ["Tables", "4", "1"]])
doc.add_heading("Chapter 8: Real-Time Feasibility", level=2)
doc.add_paragraph("The section order is now 8.1 The Live Structure, 8.2 What the Transfer Changes, 8.3 What Is Lost, 8.4 Cost.")
numbered(doc, [
 ("System design diagram.",
  "New Figure 8.1 sets the offline pass and the live structure side by side, each live stage under the offline stage it replaces, colour-coded: unchanged, adapted for causality, new wiring, and analysis tools with no live form. Section 8.1 describes the structure: one capture process owning the sensor or the recording and a shared frame buffer, the two branch processes with the one-way object link, the merger rendering at a declared delay of two frames, and the Unity receiver."),
 ("What the transfer changes.",
  "Section 8.2 keeps the feature table (Table 8.1) and explains the causal replacements: the trailing despike, gap bridging only inside the merger's delay, the One Euro filter in place of the zero-phase smoother, the cleaning rule ahead of the causal filter, and the grip offset warmed up over the first clean frames with bounded episode entry and exit."),
 ("What is lost, shown on frames.",
  "New Section 8.3 names four losses and shows them: the smoother's view of the future (the live trace lags a moving joint), the reproducibility of carried states (the two paths froze shoulder twists about 40 degrees apart because the opening frames were treated differently, which moves the rendered wrist by a few centimetres even where every landmark is measured), the grip offset (episode mean offline, running mean live, so the rebuilt wrists and elbows differ inside the gap), and the edges of episodes and gaps (a release confirmed late, long gaps held). New Figure 8.2 draws the elbow flexion and shoulder twist of both paths in three windows, and new Figure 8.3 draws both chains on the colour frame at frames 18, 548 and 595. The evaluation tools have no live form, and a live camera session is still untested."),
 ("Numbers reduced.",
  "The per-stage timing tables, the latency figure and the live-versus-offline angle statistics are out of the chapter. Section 8.4 states only that the replay ran at about 29 frames per second with detection taking most of the budget, and that the reconstruction is two frames behind capture by design."),
])
doc.save(REPO / "writing" / "v8" / "Thesis_V8_Changes_Ch8.docx"); print("saved Ch8 log", W8)
