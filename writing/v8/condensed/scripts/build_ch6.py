#!/usr/bin/env python3
"""Build writing/v7/Chapter_6_System_Integration.docx (V7 rewrite round).

Chapter 6 "System Integration and Graphical Reconstruction" against the
V7 TOC: 6.1 Communication Architecture, 6.2 Coordinate System
Conversion, 6.3 Avatar Reconstruction in Unity, 6.4 Scene Reconstruction
in Unity.

METHOD ONLY (user rule 2026-08-28, skill_set/thesis-structure-rules.md):
no accuracy or error numbers (Chapter 7), no latency, cadence or frame
rate numbers (Chapter 8), no validation counts. The Unity-vs-video check
of eval/reports/r5_unity_check.md is internal and is not cited; only its
captured stills are reused as imagery.

Sources for every statement, in the order the sections use them:
  6.1  v2/V2.md, v2/reports/r5_realtime_probe.md,
       v2/integration/run_v2.py, v2/broker/capture_broker.py,
       v2/common/shm_ring.py, v2/person/v2_person.py,
       v2/object/v2_object.py, v2/common/person_shm_v2.py (PSR2),
       v2/common/object_link.py (B->A link), v2/integration/
       v2_integrate.py (merger, PSI2), v1/aruco/send_scene_poses.py
       (PSB2, PSB3).
  6.2  v1/KINEMATIC_MODEL.md sec 3 and v1/kinematics/root_frame.py
       (the y flip F), v1/ARUCO_MODEL.md sec 5 and v1/aruco/frames.py
       (the swap, written S here), v1/INTEGRATION.md sec 1 (the anchor),
       Unity/Assets/Scripts/IntegratedSceneReceiverV2.cs (anchor node),
       Unity/Assets/Scripts/ArucoSceneReceiver.cs (gravity levelling).
  6.3  Unity/Assets/Scripts/IntegratedSceneReceiver.cs (bone driving,
       AlignArm T-pose forcing, MatchArm subject proportions, uniform
       height scale), IntegratedSceneReceiverV2.cs (tags to colours),
       v2/common/display_lpf.py + eval/DECISIONS.md E-018 (display
       smoother), v1/kinematics/shoulder.py (the 13-angle layout),
       eval/reports/recording_20260825_222315_offset_fit.json (subject
       segment lengths).
  6.5  worked example on frame 548 of the rail recording: v2/output/
       v2_person_dump_r6b_full.csv, v2_object_dump_r6b_full.csv,
       v2_integrate_dump_r6b_full.csv; the mapping recomputed from
       eval/output/scene_calibration_r6bc.json (scratch, 2026-09-07).
  6.4  Unity/Assets/Scripts/ArucoSceneReceiver.cs (BuildStaticScene,
       BuildDeferredScene), v1/aruco/send_scene_poses.py (PSB3 fields).

Figures: writing/v7/figures/ch6_fig_arch.png (make_ch6_arch_fig.py),
ch6_fig_frames.png (make_ch6_frames_fig.py),
writing/v8/condensed/figures/ch6_fig_buffer.png and ch6_fig_records.png
(make_ch6_memory_fig.py,
layouts from v2/common/shm_ring.py and v2/integration/v2_integrate.py,
frame 548 trace from v2/output/v2_*_dump_r6b_full.csv), ch6_fig_unity.png
(make_ch6_unity_fig.py, sensor-view captures of the rail recovery stream,
eval/unity_check/run_unity_capture.py, 2026-09-07; rig joints logged).
Citations: writing/v7/references.md ([1]-[57] frozen numbering).
Condensed round (2026-09-06, writing/v8/condensed/CONDENSE_BRIEF.md):
prose+captions+tables cut from 4,760 to 3,184 words. Kept whole: the four
processes and the shared frame index of 6.1, Table 6.1, the merger's four
states, all of 6.2 (equations 6.1 to 6.3 and the anchor argument), and the
avatar mapping of 6.3 with equations (6.4) and (6.5). Compacted: the
sequence-counter protocol, the cross-branch inversion, startup ordering,
and the scene-building detail of 6.4. Stale references repointed for V8:
the twist hold is Section 5.5 (was 5.6), the output-state tags point at
the Chapter 5 opening, and the static scene calibration is Section 2.2.2
with the frames in Section 2.2.1 (was "Chapter 4"). No figure, table or
equation removed. See writing/v8/condensed/notes_ch6.md.
Audit pass (2026-09-06, MATH_LOGIC_REVIEW.md M16, M17): the merger is
the stage that pairs the two streams (the recovery link also reads the
object record); one packet means one render time, not one observation
time, and the hold flag exposes the older stream; the avatar lengths are
fixed in the rig (Unity IntegratedSceneReceiver.cs) and are loop-derived
even when the rail recording is replayed.
Trim pass (2026-09-06, evening): about 270 prose words of wording,
restatements and figure-reading sentences removed from 6.1, 6.3 and 6.4;
Section 6.2 untouched; no claim, number, equation, figure or table lost.
Review pass (2026-09-07, writing/reviews/Chapter_6_v8_round*.txt plus the
technical and cross-reference checks in notes_ch6.md section (h)):
  - code facts corrected: the receiver composes the transforms of 6.2 to
    6.4 (not 6.2 only); the landmark branch reuses the marker branch's
    swap and Euler routines through the link; the merger's flags mark
    interpolated and blended only, a held person shows in its group
    states and a held object lowers its live bit after a staleness limit
    (v2_integrate.py); only the root and the object travel as Euler
    triples, the arm angles as joint coordinates (shoulder.py); the
    Figure 6.5 captures come from the receiver that applies the T-pose
    and proportion corrections (IntegratedSceneReceiver.cs), the receiver
    on the merger's record applies the height scale only (V2.cs).
  - frame readings to the precision rule: 14.4 deg / 9.9 cm on frame 114,
    31.0 deg / 2.4 cm on frame 700 (angles_recovery.csv, fk with the rig
    lengths 0.256/0.252).
  - the calibrated lengths are attributed to Section 5.1 (Chapter 3 uses
    none); the subject's direct arm measurement (user, 2026-09-07: about
    25 cm for both segments) is stated beside the two calibrations.
  - notation (C15/C46): equation (6.2) also written as the homogeneous
    T_UP with R_cam and t_cam named as the blocks of T_WC of Table 2.2;
    the display point is p_U as in (6.1).
  - prose: passives with a hidden actor given their actor, long sentences
    split, project labels (skew guard, person anchor, state tag, live
    mask) replaced by plain words, "state" instead of "tag".
Round 5 (2026-09-07, supervisor C48 and the user's direction of the same
day): no number carries more than two decimals, so the session-clock
times of 6.5 and Table 6.2 print to hundredths and the two person
samples around the render time are placed by their offsets in whole
milliseconds; and the worked example now shows the calculation of the
coordinate transformation instead of stating its results: R_cam, t_cam,
A = S R_cam F and S t_cam with their numbers, the pelvis through
equation (6.2), the levelling rotation and the raise of Section 6.4,
and the object's inverse path through the link, all printed to two
decimals from ch6_numbers.py (which also corrects two double-rounded
prose values: the camera at 43.1 cm, not 43.2, and the cube at
-20.6 cm, not -20.7).
Phase 3 of the revision of 2026-09-11 (REVISION_2026-09-11_BRIEF.md):
continuity and wording only, no method, parameter, equation, figure,
table or number changed. New opening paragraph (what Chapters 2 to 5
leave and the roadmap) and new closing paragraph (what the chapter
established, and why Chapters 7 and 8 follow). Real-time vocabulary (brief section 1, locked
fact 9): the merger's two-frame offset is named a configured buffering
and render delay and is "the one Chapter 8 configures (Section 8.2)"
(was "declares"); the worked example names its replay as the recording
at its recorded pace (Section 8.4). No causal or live analysis was
absorbed from Chapter 8; the passages that sit closest to Chapter 8 are
listed in notes_revision_ch6.md as relocation candidates and left in
place. Section 6.3 keeps the effective segment lengths, their medians
and the subject's about 25 cm per segment (D-005), which Chapters 3, 7
and 9 cite. Dropped one stale clause: Chapter 7 no longer quotes the
two clean-frame segment medians the 6.3 sentence pointed at.

Chapter 7 restructure of 2026-09-12 (D-029, D-037; CH7_INTEGRATION_FIXLIST.md
items 6-1 to 6-5). The three cross-references the previous round repointed at
the new Chapter 7 all resolved to different content and were repaired:
 - 6-1, the Figure 6.{F_SCENE} caption. It sent the reader to Section 7.1 for
   "the task and its phases". Section 7.1 is now the evaluation method, and the
   protocol drawing and the colour frame strip went with the old Sections 7.1
   and 7.2. The caption now names the moment itself and points at Section 2.1
   for the setup and the task.
 - 6-2, Section 6.4. The sentence "Section 7.4.1 measures on synthetic windows
   what a nearly straight arm costs the recovery" was deleted. Section 7.4.1
   now selects three handover windows by reference wrist excursion and
   characterizes no arm straightness; the mechanism survives in Section 5.5.
 - 6-3, Section 6.3. The "(Section 7.1)" pointer beside the loop recording was
   dropped. The rebuilt Section 7.1 names that recording but no longer
   describes it, so the pointer promised what it could not deliver.
 - 6-4, the chapter closing. "Physical, manual and synthetic references" became
   measured landmarks, manual wrist labels and within-recording comparisons,
   and the synthetic exact references now point at Appendix H, which is where
   D-037 put the solver verification. The physical tape reference of Section
   7.2 is held back while Sections 7.2.1 and 7.2.2 are Data Required (D-036).
No number, figure, table or equation of this chapter changed.

Follow-up pass of 2026-09-12, after Section 7.2 was completed (D-044, D-045):
 - The Figure 6.{F_SCENE} caption points at Section 2.1 for the task phases,
   which D-044 defines there. That is the single phase-vocabulary pointer of
   this chapter, placed at the first use ("sliding it along").
 - The chapter closing carries the physical reference again. Sections 7.2.1
   and 7.2.2 now compare reconstructed marker-centre segment lengths with the
   author's tape measurements of the physical route, so the clause removed
   under D-036 is restored. No quantity of Section 7.2 is named here.
 - D-045: no claim of this chapter rests on the dropped rail-height
   cross-check, and the 3.8 centimetre rail height is printed only in
   Chapter 2 Table 2.1 as an apparatus dimension.
Still no number, figure, table or equation of this chapter changed.

Frame notation pass of 2026-09-12 (D-054, D-056, D-058, D-060, D-061, D-062,
D-063; FRAME_INVENTORY.md sections 8.5 and 8.9.6, the 23 APPLY rows C6-05 to
C6-34; BONE_ROTATION_CLASSIFICATION.md for equations (6.4) and (6.5)).
Notation only: no number, no transformation direction and no code path
changed, and the one new printed quantity is the floor drop d, which was
already in the source comments.
 - Every frame-relative proper rotation and every homogeneous frame transform
   now carries both of its frames, the reference frame as a left superscript
   and the described frame as a left subscript; positions carry the frame they
   are expressed in. The two handedness changing maps S and F keep their bare
   letters (D-054 class 4), are never given Craig notation, and no frame is
   cancelled through either. Equation (6.1) and the bare S of the return leg
   of 6.5 are left exactly as they were (the IMPROPER rows C6-01 to C6-04 and
   C6-26), so the point symbols of 6.5 changed while that swap did not.
 - The historical pass introduced a separate intermediate frame for gravity
   alignment before the scene frame (D-061). D-082 supersedes that choice:
   equation (6.6) now prints the full Unity-to-Scene transform with the
   gravity operator G and translation column (0, d, 0); d remains a
   recording-specific scalar, 71.6 cm on
   the rail recording and 70.9 cm on the handover recording, never as a
   constant of the system. Equation (6.6) is the last equation of the chapter,
   so nothing renumbered. The back reference required by D-063 names the frame
   drawn in Figure 2.4(b) as the scene frame.
 - Equations (6.4) and (6.5) follow D-062's frame branch: R_0, R_1, R_2,
   R_chain, R_rest and R_bone are class-1 frame relations and carry both
   frames. The degenerate spine line of (6.4) is carried by the sentence
   before it. The rig's reference alignment is printed as a stated assumption,
   not as a fact. The A of (6.2) and the A of (6.5) were different quantities
   under one symbol and are now separated by their labels.
 - Left untouched and still open, as recorded in FRAME_INVENTORY.md section 9:
   the two mutually exclusive re-expression rules (S-3), the levelling factor
   that equation (6.5) does not account for in prose (S-4), the direction of F
   against equation (3.1) (S-9), the "standard chaining of homogeneous
   transformations" sentence after equation (6.2), the earlier shared-frame
   sentence of Section 6.2 (S-12), and Figure 6.4's "the
   four coordinate frames" caption.

D-082 (2026-09-13): eliminate the intermediate coordinate frame. Gravity
alignment is the proper rotation operator G; floor placement is t_f. Their
composition is the full Unity-to-Scene transform in equation (6.6). Figure
6.4 distinguishes the five retained frame boxes from the operations between
them. Equation (6.4), rig configuration, the unverified spawn-axis assumption,
and all final numerical coordinates remain unchanged.

D-079/D-083 (2026-09-13): the tabletop used for heights is the horizontal
model through the depth-derived point with normal set by calibrated gravity,
not the depth-fitted plane. Preserve every height. Describe the frame-548
forearm discrepancy while twist is held without assigning its whole cause.
"""
# D-073: current frame labels are Camera and Camera'; historical notes
# above retain the terminology of their original decisions. No numeric map changes.
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

from eqn import r, nor, sub, sup, mat, d, eqArr, add_display_eq, add_display_math, add_inline_math

REPO = Path(__file__).resolve().parents[4]
FIG = REPO / "writing" / "v7" / "figures"
FIGC = REPO / "writing" / "v8" / "condensed" / "figures"

H1, H2, H3, P, IMG, CAP, TBL, EQ, PM = "h1", "h2", "h3", "p", "img", "cap", "tbl", "eq", "pm"
MATH = "mth"  # displayed, unnumbered worked-example mathematics


def T(t):
    return ("t", t)


def X(f):
    return ("m", f)


def num_mat(rows):
    return mat([[r(c) for c in row] for row in rows])


F_ARCH, F_BUF, F_REC, F_FRAMES, F_SCENE = 1, 2, 3, 4, 5

# --- frame notation ---------------------------------------------------
# D-054, D-056, D-058, D-060, D-061.  A frame-relative proper rotation and
# a homogeneous frame transform carry the reference frame as a left
# superscript and the described frame as a left subscript; a position
# carries the frame it is expressed in as a left superscript, and the
# origin of another frame as an ORG subscript.  The two handedness
# changing maps S and F are class 4 and keep their bare letters: they are
# never written in this form and no frame is ever cancelled through them.
def pre(base, ref, described=""):
    """Leading superscript (reference frame) and leading subscript."""
    # An empty OMML script slot renders as a visible placeholder box in the
    # LibreOffice path that produces the thesis PDF, so an absent described
    # frame carries a blank run instead of nothing.
    _BLANK = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'
    return ("<m:sPre><m:sub>" + (nor(described) if described else _BLANK)
            + "</m:sub><m:sup>" + nor(ref) + "</m:sup><m:e>" + base
            + "</m:e></m:sPre>")


def Rf(ref, described):
    """Class 1: the orientation of one frame relative to another."""
    return pre(r("R"), ref, described)


def Tf(ref, described):
    """Class 2: a homogeneous frame transform."""
    return pre(r("T"), ref, described)


def Pf(ref, org=None):
    """A position expressed in a frame, or the origin of another frame."""
    return pre(sub(r("P"), nor(org + "ORG")) if org else r("P"), ref)


# --- equation-fragment shorthands -------------------------------------
Sm = r("S")
Fm = sup(r("F"), r("-1"))
Aup = Rf("Unity", "Camera'")           # was A in (6.2) and (6.3) (C6-08)
Rcam = Rf("World", "Camera")          # was R_cam (C6-05)
tcam = Pf("World", "Camera")          # was t_cam (C6-06)
StCam = Pf("Unity", "Camera")         # was S t_cam (C6-07)
RuMc = Rf("Unity", "MappedCamera")    # first factor of (6.3) (C6-12, C6-30)
RmcP = Rf("MappedCamera", "Camera'")   # second factor of (6.3) (C6-12)
Gm = r("G")                           # proper gravity alignment operator (D-082)
tf = sub(r("t"), nor("f"))             # recording-specific floor translation
TsU = Tf("Scene", "Unity")            # full gravity alignment and floor placement
RsW = Rf("Camera", "World")           # was R = R_cam transposed (C6-25)
tsW = Pf("Camera", "World")           # was t (C6-25)
Rx90 = sub(r("R"), r("x")) + d(nor("-90 deg"))
# Equations (6.4) and (6.5) under D-062: R_0, R_1, R_2, R_chain, R_rest and
# R_bone are all class-1 frame-relative rotations, so each carries both of
# its frames (BONE_ROTATION_CLASSIFICATION.md sections 0, 2, 4).
RpL24 = Rf("Camera'", "L24")           # was R_0 (the root rotation)
RL24L14 = Rf("L24", "L14")            # was R_1 (the shoulder rotation)
RL14L16 = Rf("L14", "L16")            # was R_2 (the elbow rotation)
RpL14 = Rf("Camera'", "L14")           # was R_chain of the upper arm
RpL16 = Rf("Camera'", "L16")           # was R_chain of the forearm
RsBone = Rf("Scene", "Bone")          # was R_bone
RsPerson = Rf("Scene", "Camera'")      # was A in (6.5); NOT (6.2)'s anchor
RpSeg = Rf("Camera'", "Segment")       # was R_chain, one bone at a time
RsegBone = Rf("Segment", "Bone")      # was R_rest
pU = Pf("Unity")                      # was p_U (C6-23)
pS = Pf("Scene")                      # was p_scene (C6-20)
pW = Pf("World")                      # was p_W (C6-31)
pC = Pf("Camera")                     # was p_C (C6-31)
qP = pre(r("q"), "Camera'")            # was q (C6-23)
dcol = mat([[r("0")], [r("d")], [r("0")]])
TWC = Tf("World", "Camera")           # was T_WC (C6-09)
TUP = Tf("Unity", "Camera'")           # was T_UP (C6-10)

content = [
(H1, "Chapter 6: System Integration and Graphical Reconstruction"),

(P, "Chapters 2 to 5 leave two measurement branches in one calibrated scene. The landmark branch of Chapters 3 and 5 delivers, for every frame, a pelvis position in the y-up camera frame {Camera'}, thirteen joint angles, and the state the recovery layer gave each joint group, measured, held or constrained. The marker branch of Chapters 2 and 4 delivers a cleaned track of the carried object in the world frame. The two branches run as separate processes in different coordinate frames, so neither yet shows the person and the object together. This chapter integrates them into one rendered Unity scene. Section 6.1 gives the communication architecture, which pairs the branches by frame index and session time. Sections 6.2 to 6.4 convert the coordinate frames, drive the avatar with the output states of Chapter 5, and build the measured room, and Section 6.5 closes with a worked example."),

(H2, "6.1 Communication Architecture"),

(P, "Unity performs no measurements and runs no solver. The receiver script uses the computed values: it composes the transforms of Sections 6.2 to 6.4 from the streamed values, so Unity holds no second implementation of the model."),

(P, f"Figure 6.{F_ARCH} shows the processes at run time: a capture process, one process per branch, and a merger that feeds the display."),

(IMG, FIG / "ch6_fig_arch.png", 6.4),
(CAP, f"Figure 6.{F_ARCH}. The integration architecture: the capture process publishes aligned frames into the shared frame buffer, each branch reads it and writes its own record, the landmark branch also reads the object record, and the merger combines both for the display."),

(P, "The two branches run as independent processes. The marker branch imports nothing from the landmark branch. The landmark branch accesses the marker branch only through the link described below, which reuses the marker branch's swap and Euler routines to read the object record. On the landmark side run detection, filtering, the chain of Chapter 3 and the recovery of Chapter 5. The marker branch detects markers, computes the world pose of the object and applies the cleaning step of Chapter 4."),

(P, "Shared memory has two purposes here. The frame buffer carries the images from the capture process to the two branches; Unity never reads it. The four records of Table 6.1 carry the computed numbers between the branches, the merger and Unity, which reads the scene record once and the combined record on every rendered frame."),

(H3, "6.1.1 The Frame Buffer: Capture Process to the Branches"),

(P, "Only the capture process opens the source, since two processes cannot share one physical camera and both branches need every frame. It aligns the depth frame to the colour frame once and stamps each frame with a session timestamp derived from the sensor's hardware timestamps. It then publishes the pair into a shared frame buffer with one writer and any number of readers."),

(P, f"The buffer is eight fixed-size slots of about 1.5 megabytes each, one aligned colour and depth pair per slot. Ahead of the images, each slot carries its own sequence counter, the frame index, the hardware timestamp and the arrival time on the computer's clock; the session timestamp is the hardware timestamp less the clock origin held in the header. A 128-byte header in front of the slots carries the image size, the intrinsics, the clock origin and the count of frames published so far (Figure 6.{F_BUF}). The capture process numbers the frames as it publishes them, and that frame index travels in every record that describes the frame. Frame n lands in slot n modulo 8. A reader therefore has eight frames before the writer comes round to that slot again, about a quarter of a second at the sensor's 30 frames per second."),
# v2/common/shm_ring.py: HDR_SIZE 128, DEFAULT_NSLOTS 8, slot 32 + 640*480*3
# + 640*480*2 = 1,536,032 bytes, total 12,288,384; lap protection ~266 ms.

(P, f"Each slot runs a lock-free counter protocol, drawn in Figure 6.{F_BUF}. Before it copies a frame into a slot, the capture process raises that slot's counter to an odd value. When the copy is complete, it raises the counter to the next even value, and only then raises the count of published frames in the header. A branch finds the newest published slot from that count, copies the slot, and reads the slot counter again. The branch keeps the copy when that counter is even and equal to the counter inside the copy. Otherwise, the capture process wrote to the slot while the branch was reading it, so the branch discards the copy and reads again. No lock is involved, and the capture process never waits for a branch. A branch that falls behind loses whole frames and never reads a half-written one."),

(IMG, FIGC / "ch6_fig_buffer.png", 5.6),
(CAP, f"Figure 6.{F_BUF}. The frame buffer as a memory block, written by the capture process and read by the two branches: a 128-byte header and eight slots of one aligned colour and depth pair each, with the byte layout of one slot. The slot's own sequence counter precedes the frame index, the hardware timestamp, the arrival time and the images. Frame 548 of the rail recording, traced in Section 6.5, sits in slot 4. Beneath the block the figure notes the capture process's write steps and a branch's read steps."),

(P, "Both branches see the same aligned frame, with the same frame index and session timestamp, so a shared frame index means the same captured instant on both sides. This one process also selects the source, a recording or the camera, so no stage downstream needs to know which it is; Section 8.5 lists what a live session would still need."),

(H3, "6.1.2 The Records: Branches to the Merger and Unity"),

(P, f"The person, object, scene and combined records are blocks of shared memory of one shape: a four-character code naming the type, a sequence counter, the frame or tick described, and the payload in a fixed byte order. Table 6.1 lists them beside the frame buffer, whose header also opens with a four-character code. The counter follows the lock-free protocol of Section 6.1.1, drawn again beside the record in Figure 6.{F_REC}. The writer and the reader take the roles the capture process and a branch play there. A late reader misses whole records and never reads part of one. The person record is 84 bytes, the object record 44, the scene record 100 and the combined record 112; Table 6.1 lists their contents and Figure 6.{F_REC} the layout of the largest. The merger writes the scene record once, from the calibration, before its first tick."),
# sizes: PSR2 84 (person_shm_v2.py), PSB2 44 and PSB3 100 (send_scene_poses.py),
# PSI2 112 (v2_integrate.py); SeqReader in v1/realtime/integration/realtime_integrate.py

(IMG, FIGC / "ch6_fig_records.png", 6.2),
(CAP, f"Figure 6.{F_REC}. The combined record as a memory block, written by the merger and read by Unity: 112 bytes with its byte offsets, the merger's write protocol on the left and the receiver's read protocol on the right. The sequence counter is odd while a write is in progress and even when the record is stable. The person, object and scene records follow the same pattern."),

(TBL, [["Record", "Code", "Written by", "Rate", "Contents"],
       ["frame buffer", "PSF1", "capture process", "per frame",
        "per slot the aligned colour and depth images with the frame index and timestamps; in the header the image size, the sensor intrinsics, the clock origin and the count of frames published"],
       ["person record", "PSR2", "landmark branch", "per frame",
        "frame index and session timestamp, pelvis position in the y-up camera frame, thirteen joint angles, and a live bit and a state per joint group"],
       ["object record", "PSB2", "marker branch", "per frame",
        "frame index and session timestamp, object position and orientation in display axes, and a live bit lowered when the marker is not measured"],
       ["scene record", "PSB3", "merger", "once",
        "the frozen scene description of Table 2.3, already in display axes"],
       ["combined record", "PSI2", "merger", "per output tick",
        "tick index and render time, pelvis, the thirteen angles, the object pose, the per-group live bits of the person and the object live bit, the interpolation flags and the group states"]]),
(CAP, "Table 6.1. The shared memory blocks. Every block is fixed size and begins with its four-character code; the four records carry the sequence counter in their header and the frame buffer in each slot. The scene record is written once before the stream begins, so its counter is never contended."),

(P, "Per joint group the person record carries a live bit that says whether the group was measured on this frame, and a state, measured, held or constrained, in the sense of Chapter 5."),

(P, "A single link crosses between the branches, in one direction: the landmark branch reads the object record, because the wrist recovery of Chapter 5 needs the object pose. Bringing that pose back into the camera frame undoes the axis swap of Section 6.2 into the world frame, then applies the calibrated pose of the world frame in the camera frame from Section 2.2.2. The round trip returns the measured pose. The link rejects a record whose frame index sits more than three frames from the frame being solved."),

(P, 'The merger pairs the two output streams. It keeps a short buffer of recent samples from each stream and emits 30 ticks per second. Each tick renders the world at a render time on the session clock two frame intervals behind the tick, not at the newest sample. Those two frame intervals are the configured buffering and render delay described in Section 8.2. By that render time both branches usually hold a sample, so the merger can pair them. Only samples whose live bit is up enter the object buffer.'),

(P, 'Each stream leaves the merger in one of four states: measured, interpolated, held or blended. A stream is measured when a sample lands on the render time and interpolated when the render time falls between two samples no more than three frame intervals apart. The merger holds it across a longer gap. When a sample returns after such a gap, the merger blends the held pose into that sample over two ticks. The merger interpolates the angles component by component, and the object pose in display axes by a straight line in position with a shortest-arc path in orientation.'),

(P, "The record's flags say whether a stream was interpolated or blended, and whether an interpolation bridged a dropout or only resampled between two adjacent frames. A held person shows in its group states, and a held object lowers the object live bit of the combined record once its last sample is more than half a second older than the render time. The group states travel through the same buffer, ordered measured, held, constrained. On an interpolated tick each group takes the higher of the two states in that order. On a held tick the merger raises every group to at least held. The person and object share one record evaluated at the same render time. When one stream is held across a gap, the record says so through the group states of the person or, after that half second, the live bit of the object."),

(H2, "6.2 Coordinate System Conversion"),

(P, f"Figure 6.{F_FRAMES} shows the four frames used in coordinate conversion before gravity alignment."),

(IMG, FIGC / "ch6_fig_frames.png", 6.4),
(CAP, f"Figure 6.{F_FRAMES}. The four frames in the coordinate conversion before gravity alignment. The landmark branch reaches the display through the inverse flip, calibrated camera pose and swap; the marker branch reaches it through the swap. Section 6.4 completes the mapping to Scene with gravity alignment and floor placement."),

(P, f"The landmark branch works in the y-up camera frame {{Camera'}}, which shares the optical centre of {{Camera}} and flips only its y axis through F of Section 3.1. The marker branch works in the desk-anchored world frame of Section 2.2.1. Unity uses left-handed coordinates [47]. The swap S exchanges the second and third world coordinates into the display frame {{Unity}}; its second axis follows the desk marker normal until gravity alignment is applied,"),

(EQ, eqArr(
     pU + r(" = ") + Sm + r(" ") + pW,
     Rf("Unity", "MappedMarker") + r(" = ") + Sm + r(" ") + Rf("World", "Object") + r(" ") + Sm,
     Sm + r(" = ") + mat([[r("1"), r("0"), r("0")],
                          [r("0"), r("0"), r("1")],
                          [r("0"), r("1"), r("0")]])), "6.1"),

(P, "Here S changes both the reference axes from {World} to {Unity} and the marker axes from {Object} to {MappedMarker}. Changing both bases requires conjugation. Changing only the reference basis requires left multiplication, as in the subsequent levelling and rig mapping. The swap is its own inverse, so Section 6.1.2 undoes it by applying it again. Its determinant is minus one, as is that of the flip F, so each changes the handedness of the frame."),

(P, "The two results do not land in the same place, because a y-up camera-frame point is expressed relative to the camera and a world point relative to the desk marker. The calibration of Section 2.2.2 froze the camera's own pose in the world, which Section 4.1 obtains by inverting the anchoring transformation, and that pose closes the gap. The inverse flip, equal to F because F is self-inverse, returns a y-up camera-frame point q to raw sensor coordinates, and the calibrated camera orientation rotates it into the world. The swap then carries it into the display frame, and the calibrated camera position, mapped by the same swap, supplies the offset. Collecting the three linear maps into one matrix, and the whole map into one homogeneous transformation,"),

(EQ, eqArr(
    pU + r(" = ") + Aup + r(" ") + qP + r(" + ") + StCam + r(",        ") +
    Aup + r(" = ") + Sm + r(" ") + Rcam + r(" ") + Fm + r(","),
    mat([[pU], [r("1")]]) + r(" = ") + TUP + r(" ") + mat([[qP], [r("1")]]) + r(",        ") +
    TUP + r(" = ") + mat([[Aup, StCam], [r("0"), r("1")]])), "6.2"),

(PM, [T("where "), X(Rcam), T(" and "), X(tcam), T(" are the rotation block and the translation column of "), X(TWC),
      T(", the calibrated camera pose in the world of Table 2.2. "), X(TUP),
      T(" is the same map as one homogeneous transformation matrix, the product of the 4 by 4 forms of the swap, "), X(TWC),
      T(" and the flip. The determinant of "), X(Aup),
      T(" is minus one times plus one times minus one, so the two handedness changes cancel. The product combines two handedness-changing basis maps with the calibrated rigid transform [42]. Its net linear block is proper, and the person arrives in {Unity} without a mirror image. The y-up camera frame and the display axes are both left-handed, so this rotation relates two left-handed frames. The algebra of the reference carries over, and only its assumption of right-handed frames does not hold here.")]),

(P, "The composition splits into a factor the display already has and a constant. Inserting the swap twice in the middle changes nothing. Regrouping then gives the anchor rotation as the calibrated camera orientation conjugated into display axes times the fixed pair formed by the swap of equation (6.1) and the flip of equation (3.1). Multiplying that pair out gives the matrix of a rotation of minus ninety degrees about the x axis,"),

(EQ, eqArr(
    Aup + r(" = ") + Sm + r(" ") + Rcam + r(" ") + Fm + r(" = ") + Sm + r(" ") + Rcam + r(" ") + d(Sm + r(" ") + Sm) + r(" ") + Fm + r(" = ") + d(Sm + r(" ") + Rcam + r(" ") + Sm) + r(" ") + d(Sm + r(" ") + Fm) + r(","),
    RuMc + r(" = ") + Sm + r(" ") + Rcam + r(" ") + Sm + r(",        ") + RmcP + r(" = ") + Sm + r(" ") + Fm + r(","),
    Sm + r(" ") + Fm + r(" = ") + mat([[r("1"), r("0"), r("0")],
                                        [r("0"), r("0"), r("1")],
                                        [r("0"), r("1"), r("0")]]) + r(" ") +
    mat([[r("1"), r("0"), r("0")],
         [r("0"), r("−1"), r("0")],
         [r("0"), r("0"), r("1")]]) + r(" = ") +
    mat([[r("1"), r("0"), r("0")],
         [r("0"), r("0"), r("1")],
         [r("0"), r("−1"), r("0")]]) + r(" = ") + Rx90), "6.3"),

(P, "The first factor is the orientation the display already uses to place the sensor body in the scene, so the receiver needs no new mathematics. It creates one static anchor node at the calibrated camera pose composed with that constant rotation. That first factor relates the display axes to the mapped camera frame, which is the sensor node's own system in the display, with the optical axis upward and the downward image direction forward. Every quantity from the person record passes through this node, positions as points and rotations by multiplying on the left. Conjugating a body rotation by the anchor would cancel the camera orientation out of the pose and turn the figure the wrong way."),

(P, "Section 6.4 levels and places the display. The object pose is applied locally beneath the parent node. The rig is driven in global scene coordinates through a parented anchor, which supplies the same levelling rotation and floor translation."),

(P, "The root orientation and the object pose travel as Euler angles in the fixed order the animation system uses [47], which is the order Section 3.4 decodes a rotation matrix into. The receiver rebuilds those two rotations directly from the three numbers. The arm angles travel as the joint coordinates of Section 3.4, and Section 6.3 says how the receiver reassembles them."),

(H2, "6.3 Avatar Reconstruction in Unity"),

(P, "The avatar is a humanoid character supplied by the laboratory. The receiver drives five of its bones. The spine bone at the base of the trunk takes the root rotation and the position, and the upper arm and forearm of each side take the arm rotations. Everything else, the legs included, keeps its rest pose, since the eight tracked landmarks measure no leg."),

(P, "The thirteen streamed angles reassemble five rotations. The root Euler triple rebuilds the trunk rotation directly. The shoulder swing pair and twist of Section 3.4 follow as axis-angle factors about the up, forward and right axes in that order. The elbow adds two angles about up and forward. The left arm repeats the pattern with every component except the twist reversed in sign, the sagittal mirror of Chapter 3."),

(P, "Each reconstructed rotation relates two frames of the chain of Chapter 3: the root frame at landmark L24 in the y-up camera frame, the upper arm frame at landmark L14 in the root frame, and the forearm frame in the upper arm frame. The forearm frame has its origin at the wrist landmark L16 and its first axis along the forearm. Chapter 3 uses the elbow rotation as an operator acting on a vector while here it relates two frames, because the notation follows how a rotation is used and not the shape of its symbol. The spine bone takes the root rotation alone, and the chain composition adds one factor per joint,"),

(EQ, eqArr(
    RpL14 + r(" = ") + RpL24 + r(" ") + RL24L14 + r(","),
    RpL16 + r(" = ") + RpL24 + r(" ") + RL24L14 + r(" ") + RL14L16), "6.4"),

(P, "The upper arm therefore adds the shoulder and the forearm the elbow, one factor per joint of the forward kinematics of Chapter 3. The shoulder frame at landmark L12 inherits the orientation of the root frame, so the shoulder factor is equally the upper arm frame taken in that shoulder frame."),

(P, "The figure needs three corrections: two in the avatar setup when the scene starts, and one on the processing side on every output tick before the record leaves it."),

(P, "The rest pose is the first correction, since the receiver composes the chain product against it. Each driven bone carries two frames: the segment frame that the chain of equation (6.4) ends on, and the bone frame, the axes the rig's author gave that bone. Nothing is ever expressed in the bone frame, and it is named here because the rotation written to the bone relates it to the scene frame of Section 6.4. That rotation, taken one bone at a time, is"),

(EQ, RsBone + r(" = ") + RsPerson + r(" ") + RpSeg + r(" ") + RsegBone, "6.5"),

(PM, [T("The leading factor includes the proper gravity alignment operator "), X(Gm), T(" of Section 6.4:")]),
(MATH, RsPerson + r(" = ") + Gm + r(" ") + Aup),
(P, "The receiver reads this global anchor rotation. The bone and segment axes are unchanged by the anchor, so the chain is left-multiplied rather than conjugated. The floor translation changes joint positions but not this orientation."),

(P, "In equation (6.5), the leading factor is the orientation of the y-up camera frame in the scene frame, and the chain product of equation (6.4) sits in the middle. The trailing factor is the bone's rest rotation, captured when the scene starts, which absorbs the axis conventions built into the rig as a constant. Reading it as the relation between the segment frame and the bone's own axes rests on an assumption this thesis states rather than establishes."),

(P, "The avatar is taken to be authored so that with every driven bone at its identity rotation it stands upright, its right side along the scene's first axis and its chest along the third. It is taken to be placed with no rotation, so that the model's axes coincide with the scene axes when the scene starts. The model's zero configuration is then the reference configuration of Section 3.2.2, the T-pose that is the zero of every arm angle, and the composition is correct only under that alignment. The setup therefore forces the T-pose first, rotating each arm segment so that it points straight out to the side. It captures the five rest rotations once the setup has aligned the arms and sized the rig."),

(P, "Proportions are the second correction. The landmark branch sends joint angles and one point, so the rendered hand ends up wherever the rig's bone lengths put it. The avatar setup scales the character uniformly to 170 centimetres. It then matches the trunk and arms to effective segment lengths taken from the subject's landmark data: a hip-to-mid-shoulder length of 47.9 centimetres, a right upper arm of 25.6 centimetres and a left of 26.2, a right forearm of 25.2 and a left of 24.7. Each is the median of that segment over the whole landmark track of the loop recording. The setup fixes the arm lengths, so a replay of the rail recording uses the same arms."),

(P, "The trunk length is configured separately, because it runs from the pelvis point the record carries to the mid-shoulder. The rail capture in Chapter 7 retains the receiver's default trunk length of 57.6 centimetres. Its pelvis is the corrected point from the hip depth preparation of Section 2.6, which sits deeper than the raw hips. The loop capture uses a trunk length of 47.9 centimetres and the raw hip midpoint as its pelvis."),

(P, 'The recovery of Chapter 5, however, ran on the lengths calibrated on the rail recording as Section 5.1 describes, 31.7 and 20.5 centimetres for the right arm. Both sets are segment lengths taken from landmark data. They differ mostly in the split between upper arm and forearm: the totals are 50.8 centimetres for the loop values and 52.2 for the rail values.'),

(P, "A direct measurement of the subject gives about 25 centimetres for both the upper arm and the forearm, so the loop values are close to the anatomy. The rail split is consistent with the elbow landmark sitting farther along the arm on that recording. The recovery keeps the rail lengths because they describe where that recording's landmarks sit, and the avatar keeps the loop lengths. Scaling the trunk first matters because the trunk moves the shoulder joints. Uniform scaling leaves every rotation unchanged."),

(P, "A display smoother on the angles and the pelvis is the third correction; the processing side applies it as they enter the output record. The solver of Chapters 3 and 5 changes state in short ramps because its rate limits cap how far a joint may move between frames. The eye reads a ramp as a jump. A linear low-pass alone cannot remove a ramp, only delay it, so the smoother pairs a first-order low-pass with a cap on the angle change per output tick. It changes only the displayed values; the processing side writes the tables Chapters 7 and 8 analyse before the smoother runs."),

(P, "Each tick the receiver maps the streamed pelvis point from {Camera'} into {Scene} through the anchor, including the levelling and floor translation of Section 6.4. It translates the root so that the spine bone lands on that point, and the root rotation orients the trunk about it."),

(P, f"The reconstruction shows where each value came from. A joint group not measured on a tick has a small sphere at its joint: at the pelvis for the root, at the shoulder for the shoulder swing and at the elbow for the elbow. The sphere sits a little in front of the body along the line of sight of the camera that renders it, so that the mesh does not hide it. Red marks a held value and blue a constrained one, the states of Section 6.1.2; the twist group has no sphere. A smaller amber sphere at the pelvis shows a pose interpolated across a gap in the stream or ramping back after a long gap."),

(P, f"The captures of Figure 6.{F_SCENE} come from the receiver that applies both start-up corrections; it reads only the live bits, so every sphere there is red. The receiver that reads the merger's record draws all three colours. It applies the uniform height scale but neither the forced T-pose nor the measured segment lengths, and it captures the rest rotations of equation (6.5) from the rig as authored."),

(H2, "6.4 Scene Reconstruction in Unity"),

(P, "The receiver builds the room from the static scene record, which arrives once before the stream begins and holds the frozen scene description of Table 2.3, already in display axes by the swap of Section 6.2."),

(P, "The receiver levels the scene first. The world frame's third axis is the desk marker's face normal, and the card sits on a leaning stand. Raw world axes would therefore draw a tilted room on a level screen. The receiver therefore creates a parent node for the scene geometry, object and person anchor. It rotates that node by the shortest rotation that carries the gravity direction of Section 2.2.2 onto the screen's vertical. The wall then comes out vertical and the desk top horizontal, and the object pose is applied beneath the node as a local pose. The rig remains outside that hierarchy and receives global poses through the anchor."),

(PM, [T("Write this gravity alignment operation as "), X(Gm), T(", a proper rotation matrix. It acts on the swapped coordinates from {Unity}, preserving lengths and handedness. The receiver then adds the translation "), X(tf), T(" to raise the reconstruction until the drawn floor sits at zero height.")]),

(P, "Together these operations define the final frame {Scene}, the coordinate system Unity itself draws in. Its second axis follows the calibrated upward gravity direction, and its origin lies on the drawn floor directly below the desk marker origin. The translation raises that marker origin by the recording-specific distance d. The full transformation from {Unity} to {Scene}, and its application to a world point after the swap, are"),

(EQ, eqArr(
    TsU + r(" = ") + mat([[Gm, tf], [r("0"), r("1")]]) + r(",        ") + tf + r(" = ") + dcol,
    pS + r(" = ") + Gm + r(" ") + Sm + r(" ") + pW + r(" + ") + tf), "6.6"),
# d = origin_above_tabletop_m + DeskThick + LegH: 0.716215 m from
# eval/output/scene_calibration_r6bc.json (the rail recording) and 0.709312 m
# from scene_calibration_r7c.json (the handover recording), printed to the
# millimetre as 71.6 and 70.9 cm; DeskThick + LegH = 0.72 m at
# eval/unity_check/check_unity_log.py:41. Recorded in FRAME_INVENTORY.md
# section 3 and UNITY_FRAME_RESOLUTION.md section 5; the transform itself is
# defined by D-082, retaining the recording-specific treatment of d in D-061.

(P, "The offset d belongs to the recording and not to the system. It is the height of the drawn desk top above the floor, the desk thickness plus the leg height, less the small offset between the horizontal tabletop model and the origin of the world frame. That gives 71.6 centimetres on the rail recording and 70.9 centimetres on the handover recording. The final coordinate system drawn in Figure 2.4(b) is the scene frame."),

(P, "The wall is a thin slab in the plane of the calibrated wall marker, and the desk a horizontal slab on the tabletop model of Section 2.2.2. That model passes through the tabletop point derived from depth, with its normal set by calibrated gravity. No measurement sets the desk width, the slab thicknesses, the legs or the floor. The rail is not part of the scene record, so the receiver does not draw it."),

(P, "A plate printed with the marker pattern sits at each calibrated marker pose. The receiver draws the desk card over the desk slab. Its calibrated centre lies 0.4 centimetres below the horizontal tabletop model, so the slab would otherwise cut the tilted card in half. The depth image puts that centre 1 to 2 centimetres above the desk around it, a difference within the pose error of a 45 millimetre marker seen from 58 centimetres. The object is a cube of the measured edge length, half an edge behind its own plate along the face normal. It turns red for a held pose once the object live bit drops, and amber for a pose interpolated across a dropout or blended back. A model of the sensor sits at the calibrated camera pose, and a camera attached to it renders an inset at the calibrated vertical field of view for comparison with the recorded colour frame."),
# desk card: notes_ch6.md section (j), decision D-009 (calibrated centre 0.4 cm
# below the horizontal tabletop model, depth 1 to 2 cm above the desk around it);
# 58 cm is the length of the anchor translation (0.023, 0.187, 0.552) m, 58.4 cm
# (eval/output/scene_calibration_r6bc.json).  Marker size 45 mm: Table 2.1 and
# locked fact 1 of REVISION_2026-09-11_BRIEF.md section 3; the legacy 50 mm
# ("a 5 centimetre marker") was not used in the reported experiments.

(P, f"Figure 6.{F_SCENE} shows the reconstruction on two frames of the rail recording, rendered from that camera at the sensor's pose so the view matches the real sensor's. Every sphere whose joint group is not marked measured in the record is red; the twist groups have no sphere. The root has one at the pelvis on both frames, because the depth pixels of the hips land on the rail in front of the subject and the hip depth preparation of Section 2.6 corrects them before the solve. The idle left arm hangs at the subject's side, partly out of the frame, so its groups are not measured and it has one at its shoulder and one at its elbow."),

(P, "On frame 114 the right arm reaches forward and down to the cube on the desk, and the rendered hand lands beside the cube, 14.1 centimetres from the measured wrist. The elbow is bent by 14.4 degrees there, so the rule of Section 5.5 holds the shoulder twist. The rendered discrepancy includes the solver state, shoulder placement and rig proportions; this capture does not isolate their individual contributions. On frame 700 the elbow is bent by 31.0 degrees, the twist is measured, and the hand sits on the cube, 6.4 centimetres from the measured wrist."),
# Unity person log of the 2026-09-07 capture (rig joints logged), mapped
# through the calibration as in eval/unity_check/check_unity_log.py:
# frame 114 rendered hand to measured wrist 14.1 cm; rig directions with the
# calibrated split anchored at the measured shoulder 8.9 cm; anchored at the
# rig shoulder 13.0; rig shoulder 5.7 cm deeper than the landmark; rig upper
# arm within 0.1 deg of the measured direction. frame 700: 6.4 cm. Rel_y
# -14.4 / -31.0 deg from eval/output/recovery_r6b/angles_recovery.csv.

(IMG, FIG / "ch6_fig_unity.png", 6.4),
# Fix list item 6-1 (D-029 restructure): the caption pointed at Section 7.1
# for "the task and its phases". The rebuilt Section 7.1 is the evaluation
# method, and the four-step protocol drawing and the colour frame strip were
# removed with the old Sections 7.1 and 7.2. Under D-044 the task phases are
# defined in Chapter 2 Section 2.1, so the caption points there for the
# phase names, and "sliding it along" is the first use of the vocabulary in
# this chapter. One pointer serves the chapter; later uses are not tagged.
(CAP, f"Figure 6.{F_SCENE}. The reconstruction as Unity draws it, seen from the virtual sensor at the calibrated camera pose. (a) Frame 114, with the hand at the cube on the desk, and the markers and the tracked cube in the measured room; each sphere whose joint group is not marked measured in the record is red. (b) Frame 700, with the cube on the rail and the hand sliding it along. Section 2.1 describes the recording setup and defines the task phases."),

(H2, "6.5 Worked Example: Frame 548"),

(P, "This section takes one frame of the rail recording, so that the records of Section 6.1 and the maps of Sections 6.2 to 6.4 appear with their numbers, as Section 5.6 did for the recovery. The frame is 548, replayed through the causal structure of Section 8.2 at the recording's own pace (Section 8.4), at a moment when the cube slides along the rail. Table 6.2 collects the values."),

(P, "The capture process published the frame at 18.28 seconds on the session clock into slot 4 of the frame buffer, 548 modulo 8. It raised that slot's counter to an odd value, copied the colour and depth bytes, raised the counter to the next even value, and only then raised the count of published frames in the header. Each branch saw the new count, copied slot 4 and checked the slot counter before using its copy."),

(P, "The landmark branch then wrote the person record: frame 548, the pelvis and the thirteen angles. Its live bits are down for the root and the right shoulder twist. Its states say the root is constrained and the twist held. The record for frame 547 carried the same two states. The solve did not draw on the object record on this frame, because the right wrist was measured. The link reads that record on every frame to follow the grip episode, taking the newest stable one within three frames."),

(P, "The marker branch wrote the object record with its live bit up. For frame 547 it had written a record with the live bit down, because the marker was not detected there, and the merger keeps such a sample out of its buffer. Table 6.2 lists the processing time stored in each branch's record for this frame. Section 8.4 reports processing times per frame over the recording."),
# v2/output/v2_person_dump_r6b_full.csv frame 548: time_s 18.27995, pel_z
# 0.9917, mask 122 (root bit 0 and R twist bit 2 down), tag_0 2 CONSTRAINED,
# tag_2 1 HELD (547 identical), compute_ms 8.36, rec_right 0;
# v2_object_dump_r6b_full.csv 548: live 1, compute_ms 13.66; 547: live 0.

(P, "The merger's tick at 18.32 seconds on the session clock, with frame 549 just published, rendered the world as of 18.25 seconds, two frame intervals earlier. The person samples of frames 547 and 548 lay on either side of that render time, 4 milliseconds before it and 29 milliseconds after it, so the merger interpolated between them. For the object stream frame 547 was missing, so the merger interpolated between frames 546 and 548 and raised the flag that says the interpolation bridged a dropout. Into the combined record went the tick, the render time, both poses, the live bits, the flags and the group states, with the root still constrained and the twist still held. On its next rendered frame the receiver in Unity copied that record under the same counter protocol."),
# v2_integrate_dump_r6b_full.csv tick 548: tau 18.2508, p_f0/p_f1 547/548,
# o_f0/o_f1 546/548, flags 35 (person interp, object interp, object bridge),
# tags 18 (root 2, R twist 1); frame 547 time_s 18.24659; --delay-frames 2;
# tau - t547 = 4.2 ms, t548 - tau = 29.1 ms (ch6_numbers.py).

(P, "The pelvis in the person record is (3.0, −17.9, 99.2) centimetres in the y-up camera frame. Equation (6.2) carries it into display axes, and the calculation runs here with its numbers. The rotation block and the translation column of the calibrated camera pose of Table 2.2 are"),
# src: ch6_numbers.py from eval/output/scene_calibration_r6bc.json T_cam_desk
#      (R_cam = R^T, t_cam = -R^T t = (-0.013549, -0.393115, 0.431492))
(MATH, eqArr(
    Rcam + r(" = ") + num_mat([
        ["1.00", "−0.02", "−0.01"],
        ["0.00", "−0.48", "0.88"],
        ["−0.02", "−0.88", "−0.48"]]),
    tcam + r(" = (−0.01, −0.39, 0.43) m"))),
(P, "The swap, this rotation and the flip multiply into the anchor rotation, and the swap carries the camera position into display axes:"),
# src: ch6_numbers.py: A = S R_cam F (det 1.0), S t_cam = (-0.013549, 0.431492, -0.393115)
(MATH, eqArr(
    Aup + r(" = ") + Sm + r(" ") + Rcam + r(" ") + Fm + r(" = ") + num_mat([
        ["1.00", "0.02", "−0.01"],
        ["−0.02", "0.88", "−0.48"],
        ["0.00", "0.48", "0.88"]]),
    StCam + r(" = ") + Sm + r(" ") + tcam + r(" = (−0.01, 0.43, −0.39) m"))),
(P, "The pelvis point of the record then maps as"),
# src: ch6_numbers.py: q = (0.030224, -0.179074, 0.991679) -> A q =
#      (0.015887, -0.636577, 0.781618) -> p_U = (0.002338, -0.205084, 0.388503)
(MATH, eqArr(
    qP + r(" = (0.03, −0.18, 0.99) m"),
    pU + r(" = ") + Aup + r(" ") + qP + r(" + ") + StCam + r(" = (0.00, −0.21, 0.39) m"))),
(P, "The matrices and vectors are printed to two decimals. The pelvis sits at (0.2, −20.5, 38.9) centimetres in display axes, before the parent levelling and floor translation of Section 6.4. The three named steps give the same point. Undoing the flip gives the raw camera point (3.0, 17.9, 99.2), within 2 millimetres of the midpoint of the two hip landmarks, which the hip depth preparation left unchanged on this frame. The calibrated camera orientation and position carry it into the world frame, (0.2, 38.9, −20.5), and the swap exchanges the last two coordinates into display axes. The camera itself sits at the swapped camera position, (−1.4, 43.1, −39.3) centimetres in the same axes."),
# src: ch6_numbers.py: F q = (0.030224, 0.179074, 0.991679); R_cam F q + t_cam
#      = (0.002338, 0.388503, -0.205084); hip landmark midpoint (0.0295,
#      0.1776, 0.9911); S t_cam in cm (-1.4, 43.1, -39.3) (V7 printed 43.2
#      by rounding the four-decimal 0.4315 a second time).
(P, "The parent node of Section 6.4 applies the full transformation of equation (6.6). Its gravity alignment carries the calibrated gravity direction, (−0.02, 0.85, 0.53) in display axes, onto the screen's vertical, a turn of 32.0 degrees. Its translation uses the offset d at 71.6 centimetres on this recording, so that the drawn floor sits at zero height:"),
# src: ch6_numbers.py: gravity_up_unity (-0.021236, 0.847708, 0.530038),
#      32.04 deg; G = FromToRotation(g, up) (ArucoSceneReceiver.cs); the
#      floor is DeskThick + LegH = 0.72 m below the tabletop plane, which
#      lies 0.003785 m above the world origin, so the raise is 0.716215 m
#      and the drawn desk top sits at 0.72 m; the operation G p_U yields
#      (0.000349, 0.032020, 0.438149) -> scene (0.000349, 0.748235,
#      0.438149), printed to two decimals as (0.00, 0.03, 0.44) and
#      (0.00, 0.75, 0.44) m.
(MATH, eqArr(
    Gm + r(" = ") + num_mat([
        ["1.00", "0.02", "0.01"],
        ["−0.02", "0.85", "0.53"],
        ["0.01", "−0.53", "0.85"]]),
    Gm + r(" ") + pU + r(" = (0.00, 0.03, 0.44) m"),
    mat([[pS], [r("1")]]) + r(" = ") + TsU + r(" ") + mat([[pU], [r("1")]]),
    pS + r(" = ") + Gm + r(" ") + pU + r(" + ") + tf + r(" = (0.00, 0.75, 0.44) m"))),
(P, "The spine bone lands at (0.0, 74.8, 43.8) centimetres in the scene frame, 2.8 centimetres above the drawn desk top at 72.0. In the replay of Section 6.3 the receiver's log places the spine bone on the mapped point to within a millimetre on every frame."),
# check_unity_log.py check 2: hip bone on the mapped pelvis point, max
# 0.0008 mm over 899 frames.

(P, "The object record holds the marker origin of the cube already in display axes, (−23.4, −18.8, 42.8) centimetres with Euler angles (−62.0, 2.8, −9.7) degrees. The link of Section 6.1.2 undoes the swap and applies the calibrated anchor in the direction from the world to the camera, the inverse of the step above. The anchor's rotation block is the transpose of the one above and its translation column is the desk marker seen from the camera:"),
# src: ch6_numbers.py from scene_calibration_r6bc.json T_cam_desk:
#      R = R_cam^T, t = (0.023458, 0.187377, 0.552492)
(MATH, eqArr(
    RsW + r(" = ") + sup(d(Rcam), r("T")) + r(" = ") + num_mat([
        ["1.00", "0.00", "−0.02"],
        ["−0.02", "−0.48", "−0.88"],
        ["−0.01", "0.88", "−0.48"]]),
    tsW + r(" = (0.02, 0.19, 0.55) m"))),
(P, "The cube's marker origin then returns to the camera frame as"),
# src: ch6_numbers.py: object dump 548 (ux,uy,uz) (-0.233985, -0.188405,
#      0.428127), euler (-62.02, 2.81, -9.68); S p_U = (-0.233985, 0.428127,
#      -0.188405); R p_W + t = (-0.206489, 0.150504, 1.020790) (V7 printed
#      -20.7 cm by rounding the four-decimal -0.2065 a second time).
(MATH, eqArr(
    pW + r(" = ") + Sm + r(" ") + pU + r(" = ") + Sm + r(" (−0.23, −0.19, 0.43) = (−0.23, 0.43, −0.19) m"),
    pC + r(" = ") + RsW + r(" ") + pW + r(" + ") + tsW + r(" = (−0.21, 0.15, 1.02) m"))),
(P, "The result, (−20.6, 15.1, 102.1) centimetres in the camera frame, is the pose the marker branch reported, with the marker's face normal turned to the camera 6.7 degrees off the optical axis. The depth image reads 98.5 centimetres at the marker's pixel, 3.6 centimetres nearer than the pose, within the pose error of a 45 millimetre marker at that distance."),
# marker size 45 mm: Table 2.1 and locked fact 1 of REVISION_2026-09-11_BRIEF.md
# section 3 (the legacy 50 mm was not used in the reported experiments)
# face normal in camera frame (0.115, 0.017, -0.993) -> 6.7 deg; aruco raw
# 548: m1 pixel (202, 335), depth_z 0.985.

(P, f"The thirteen angles become bone rotations by equation (6.5), composed here from the record's angles as the receiver of Figure 6.{F_SCENE} composes them. The right upper arm takes the root Euler angles (−13.6, 176.2, −1.3) degrees, the shoulder azimuth and elevation (−44.3, −77.2) degrees and the held twist of −15.1 degrees. After the anchor and the levelling they place the bone's axis along (−0.14, −0.93, −0.34), 0.2 degrees from the measured shoulder-to-elbow direction carried through the same maps. The elbow is bent by 22.7 degrees about the up axis and 0.0 about the forward axis. The forearm lands 6.9 degrees from the measured direction while the shoulder twist is held under the rule of Section 5.5."),
# fk_arm_dirs (eval/inspect/check_v1_overlay.py) with the frame 548 angles,
# root recompose_zxy, then A and G: upper (-0.1437, -0.929, -0.3409) vs
# measured (-0.1466, -0.9288, -0.3403), 0.2 deg; forearm 6.9 deg.

(P, f"The receiver that reads the merger's record draws the pelvis sphere blue, because the root's state is constrained, and no sphere for the twist. The merger marks both streams as interpolated and flags the object interpolation as bridging the missing frame 547. The cube is therefore tinted amber on this tick. The person stream crossed no dropout, so the pelvis carries no amber sphere. The receiver of Figure 6.{F_SCENE}, which reads the live bits alone, would draw the pelvis sphere red."),

(TBL, [["Stage", "Record or map", "Frame 548"],
       ["Capture process", "frame buffer, slot 4", "18.28 s on the session clock; colour and depth, 1,536,032 bytes"],
       ["Landmark branch", "person record", "8.4 ms; pelvis (3.0, −17.9, 99.2) cm; live bits down for the root and the right twist; root constrained, twist held"],
       ["Marker branch", "object record", "13.7 ms; (−23.4, −18.8, 42.8) cm, (−62.0, 2.8, −9.7) deg; live; frame 547 not live"],
       ["Merger", "combined record", "tick at 18.32 s, render time 18.25 s; person 547 to 548, object 546 to 548, dropout bridged"],
       ["Link", "swap undone, anchor applied", "cube marker origin (−20.6, 15.1, 102.1) cm in the camera frame"],
       ["Receiver, position", "equations (6.2) and (6.6)", "pelvis (0.2, −20.5, 38.9) cm in display axes; (0.0, 74.8, 43.8) cm in the scene frame; 32.0 deg"],
       ["Receiver of Figure 6.5, rotation", "equation (6.5)", "right upper arm 0.2 deg from the measured direction; forearm 6.9 deg"]]),
(CAP, "Table 6.2. Frame 548 of the rail recording through the causal structure and the receiver. Positions in centimetres; clock times on the session clock, durations in milliseconds."),

# Fix list item 6-4 (D-029, D-037): "synthetic references" named the solver
# verification, which left Chapter 7 under D-029 and is now Appendix H. The
# reference kinds are taken from the rebuilt Section 7.1: the physical tape
# reference of Section 7.2, measured 3D landmarks (7.3), manual wrist
# proxies (7.4.2) and the unmasked reconstruction as a within-recording
# comparison (7.4.1). The physical reference was held back while Sections
# 7.2.1 and 7.2.2 were Data Required under D-036; both now report
# reconstructed marker-centre segment lengths against the author's tape
# measurements of the physical route, so the clause is restored. No
# quantity of Section 7.2 is named here.
(P, "The chapter closes the construction of the system. Both branches now reach one rendered scene. The person stands in the measured room beside the tracked object. Every joint group carries the state the recovery layer of Chapter 5 gave it. A held joint and a rebuilt joint are therefore visible as such. The chapter does not establish how near that reconstruction is to the references, or whether it can be produced while the motion happens. Chapter 7 evaluates the reconstruction offline against physical references, measured landmarks and manual wrist labels, and against within-recording comparisons where no independent reference exists. Section 7.2 sets the reconstructed object segment lengths against tape measurements of the physical route. Appendix H holds the verification of the kinematic solve of Chapter 3 against synthetic exact references. Chapter 8 asks whether the same architecture runs causally."),


]

references = {
    42: "J. J. Craig, Introduction to Robotics: Mechanics and Control, 3rd ed. "
        "Upper Saddle River, NJ, USA: Pearson Prentice Hall, 2005.",
    47: 'Unity Technologies, "Unity scripting reference: Transform, '
        'Quaternion," Unity Documentation. [Online]. Available: '
        "https://docs.unity3d.com/ScriptReference/",
}

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

# D-106: unnumbered display ordinals continuing into the following clause.
DISPLAY_COMMAS = set()
math_display_index = 0

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == H3:
        doc.add_heading(item[1], level=3)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == PM:
        p = doc.add_paragraph()
        for typ, val in item[1]:
            if typ == "t":
                p.add_run(val)
            else:
                add_inline_math(p, val)
    elif kind == EQ:
        add_display_eq(doc, item[1], item[2])
    elif kind == MATH:
        math_display_index += 1
        add_display_math(doc, item[1],
                         punctuation="," if math_display_index in DISPLAY_COMMAS else ".")
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == IMG:
        path = Path(item[1])
        assert path.exists(), f"missing figure: {path}"
        doc.add_picture(str(path), width=Inches(item[2]))
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row_ in enumerate(rows):
            for j, c in enumerate(row_):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

out = REPO / "writing" / "v8" / "condensed" / "Chapter_6_System_Integration.docx"
doc.save(out)
print("saved", out, "| items:", len(content))
