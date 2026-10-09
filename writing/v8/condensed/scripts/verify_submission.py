#!/usr/bin/env python3
"""Check actual V8 submission artifacts against Checklist 2.8 and baseline.

This validates document structure and rendered geometry, not experiment
accuracy. The baseline is the committed pre-edit manuscript inventory.
"""
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree as E

import assembly_citations as citations
import build_thesis as assembly

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'audit_evidence/submission_checklist_2026-09-13'
DOCX = ROOT.parent / 'Thesis_V8_Condensed.docx'
PDF = DOCX.with_suffix('.pdf')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
results = []


def check(label, condition, detail=None):
    results.append({'check': label, 'status': 'PASS' if condition else 'FAIL', 'detail': detail})


def plain(node):
    return ''.join(node.xpath('.//w:t/text()', namespaces=NS))


def normalized(text):
    return re.sub(r'\s+', ' ', text).strip()


def main():
    baseline = json.loads((EVIDENCE / 'baseline.json').read_text())
    doc = Document(DOCX)
    with ZipFile(DOCX) as archive:
        body = E.fromstring(archive.read('word/document.xml')).find('w:body', NS)
        math = [[t.text for t in m.findall('.//m:t', NS)] for m in body.findall('.//m:oMath', NS)]
        same, punctuated = 0, 0
        for before, after in zip(baseline['math_tokens'], math):
            same += before == after
            punctuated += before == after[:-1] and after[-1:] in [['.'], [',']]
        check('Native mathematical content preserved', len(math) == len(baseline['math_tokens']) == same + punctuated,
              {'objects': len(math), 'unchanged': same, 'punctuated': punctuated})
        check('All displays punctuated', punctuated == 107)
        media = sorted(hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/'))
        check('Embedded image bytes preserved', media == sorted(baseline['media'].values()), len(media))
    tables = [t for t in body.findall('w:tbl', NS) if t.xpath('./w:tblPr/w:tblStyle[@w:val="TableGrid"]', namespaces=NS)]
    data = [[plain(p) for p in table.findall('.//w:p', NS)] for table in tables[1:]]
    allowed = {'ArUco ID': 'ArUco identifier', '64 GB': '64 gigabytes (GB)',
               'Ubuntu 24.04 LTS': 'Ubuntu 24.04 long-term support (LTS)'}
    def canonical(text):
        return re.sub(r'\[\d+\]', '[REF]', allowed.get(text, text))
    check('All result-table content preserved',
          [[canonical(t) for t in row] for row in baseline['data_tables']] ==
          [[canonical(t) for t in row] for row in data], len(data))
    eq_tables = [t for t in body.findall('w:tbl', NS) if any(re.fullmatch(r'\([A-H0-9]+\.\d+\)', plain(c)) for c in t.findall('.//w:tc', NS))]
    check('Symmetric six-inch numbered equation layout', len(eq_tables) == 56 and all(
        [int(c.get(qn('w:w'))) for c in t.findall('w:tblGrid/w:gridCol', NS)] == [648, 7344, 648] for t in eq_tables), len(eq_tables))
    check('All table grids within text area', all(sum(int(c.get(qn('w:w'))) for c in t.findall('w:tblGrid/w:gridCol', NS)) <= 8640 for t in body.findall('w:tbl', NS)))
    sizes = [int(n.get(qn('w:val'))) for root in (doc.element, doc.styles.element) for n in root.iter(qn('w:sz'))]
    check('Text font sizes at least 12 point', min(sizes) >= 24)
    check('Section page geometry', all(s.page_width.inches == 8.5 and s.page_height.inches == 11 and
          s.left_margin.inches == s.right_margin.inches == 1.25 and
          s.top_margin.inches == s.bottom_margin.inches == 1 for s in doc.sections))
    terms = [plain(row.findall('w:tc', NS)[0]) for row in tables[0].findall('w:tr', NS)[1:]]
    check('Selected glossary terms sorted', len(terms) == 12 and terms == sorted(terms, key=str.casefold), terms)
    steps = [p for p in doc.paragraphs if p.style.style_id == 'ListNumber']
    check('Native numbered acquisition steps', len(steps) == 8 and all(p._p.find('w:pPr/w:numPr', NS) is not None for p in steps))
    headings = [p.text for p in doc.paragraphs if p.style.style_id.startswith('Heading')]
    check('References after Appendix H', headings[-1] == 'References' and headings.index('References') > next(i for i,h in enumerate(headings) if h.startswith('Appendix H:')))
    navigation = json.loads((EVIDENCE / 'navigation.json').read_text())
    lengths = [len(i['rows']) for i in navigation['indexes']]
    check('Four populated navigation fields', lengths == [118, 50, 27, 56], lengths)
    check('No title/contents self-entry', not any(re.match(r'(?:Title Page|Table of Contents)\t', r['text']) for r in navigation['indexes'][0]['rows']))
    check('Cached Word reopen matches PDF', navigation.get('cached_reopen_check', '').startswith('PASS'), navigation.get('cached_reopen_check'))
    check('Dot-leader tabs in navigation caches', sum(1 for p in doc.paragraphs if p.style.style_id in ('TOC1','TOC2','TOC3') and p._p.xpath('./w:pPr/w:tabs/w:tab[@w:leader="dot"]')) == sum(lengths))
    plan = citations.make_plan([ROOT / f for _, f in assembly.PARTS + [assembly.APPENDIX_PART]], assembly.fm.parse_references())
    check('Literature source identities verified', citations.verify_assembly(doc, plan)['status'] == 'PASS')
    # Check introductions against the source order, including plural label lists.
    seen, captions, missing = set(), [], []
    for element in body:
        if element.xpath('./w:pPr/w:pStyle[starts-with(@w:val,"TOC") or @w:val="TableofFigures"]', namespaces=NS):
            continue
        text = plain(element)
        match = re.match(r'^(Figure|Table) ([A-H0-9]+\.\d+)\. ', text)
        if match:
            key = (match[1], match[2]); captions.append((key, text))
            if key not in seen: missing.append(key)
        elif element.tag == qn('w:p') and not element.xpath('./w:pPr/w:pStyle[starts-with(@w:val,"TOC")]', namespaces=NS):
            for match in re.finditer(r'\b(Figures?|Tables?)\s+([A-H0-9]+\.\d+(?:(?:,?\s+and\s+|,\s*|\s+to\s+)[A-H0-9]+\.\d+)*)', text):
                for num in re.findall(r'[A-H0-9]+\.\d+', match[2]):seen.add((match[1].rstrip('s'),num))
    check('All figure/table numbered introductions', not missing, missing)
    # Poppler exposes rendered image and text coordinates directly.
    with tempfile.TemporaryDirectory(prefix='submission-verify-') as temp:
        xml_path = Path(temp, 'layout.xml')
        subprocess.run(['pdftohtml','-xml','-hidden','-noroundcoord','-zoom','1',str(PDF),str(xml_path)], stdout=subprocess.DEVNULL, check=True)
        pages = E.parse(str(xml_path)).findall('page')
        pdf_texts = [normalized(' '.join(''.join(t.itertext()) for t in page.findall('text'))) for page in pages]
        body_start = next(i for i,text in enumerate(pdf_texts) if text.startswith('Chapter 1: Introduction'))
        glossary_start = next(i for i,text in enumerate(pdf_texts) if text.startswith('Glossary'))
        check('One-page glossary before Chapter 1', body_start == glossary_start + 1, glossary_start + 1)
        image_caption_pages, caption_errors, heading_errors = {}, [], []
        for i,page in enumerate(pages):
            if i < body_start:continue
            texts = page.findall('text')
            for j,text in enumerate(texts):
                content = normalized(''.join(text.itertext()))
                match = re.match(r'^(Figure) ([A-H0-9]+\.\d+)\. ',content)
                if match:
                    images = [image for image in page.findall('image') if float(image.get('top')) + float(image.get('height')) <= float(text.get('top')) + 2]
                    image_caption_pages[match[2]] = i+1
                    if not images:caption_errors.append([match[2],i+1])
                if content in headings:
                    below = [t for t in texts[j+1:] if float(t.get('top')) < 720]
                    if not below and not page.findall('image'):heading_errors.append([content,i+1])
        check('Every figure shares a page with its caption', len(image_caption_pages) == 50 and not caption_errors, caption_errors)
        check('No headings stranded at page foot', not heading_errors, heading_errors)
        # A 1-point allowance accounts for glyph-side bearings and Poppler's
        # text-box rounding, not document-margin shrinkage.
        bbox_path = EVIDENCE / 'layout.html'
        subprocess.run(['pdftotext','-bbox-layout',str(PDF),str(bbox_path)],check=True)
        bx = E.parse(str(bbox_path));xn={'x':'http://www.w3.org/1999/xhtml'}
        margin_errors=[]
        for i,page in enumerate(bx.findall('.//x:page',xn)):
            for word in page.findall('.//x:word',xn):
                if float(word.get('yMin')) > 60 and float(word.get('yMax')) < 735 and (float(word.get('xMin'))<89 or float(word.get('xMax'))>523):margin_errors.append([i+1,word.text])
        check('Rendered text respects side margins', not margin_errors, margin_errors)
        summary={'pages':len(pages),'figure_pages':image_caption_pages,'checks':results,
                 'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest(),
                 'pdf_sha256':hashlib.sha256(PDF.read_bytes()).hexdigest()}
    (EVIDENCE/'validation.json').write_text(json.dumps(summary,indent=2)+'\n')
    for result in results:print(result['status'],result['check'],result['detail'] if result['status']=='FAIL' else '')
    print('PDF pages:',summary['pages'])
    raise SystemExit(any(r['status']=='FAIL' for r in results))


if __name__=='__main__':main()
