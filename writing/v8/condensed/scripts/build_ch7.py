#!/usr/bin/env python3
"""Build Chapter 7 from verified coordinate evidence (D-029 to D-038).

Object segment results are computed under D-035 and D-036, which supersede the
D-030 block by supplying the physical reference and the phase-boundary waypoint
definitions. Section 7.3.3 carries the wrist-to-marker check under D-038.
Experimental numbers come from audit_evidence/ch7_restructured and from the
author's own measurements, never from old chapter prose.
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from eqn import add_display_eq, r, nor, sub, sup, frac, d, eqArr


# Craig's left superscript names the frame a quantity is expressed in
# (D-054). An absent left subscript is a blank run, not an empty element:
# an empty one draws a placeholder box in some renderers. The same
# construction is used in build_ch2.py.
BLANK = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'


def presup(base, above, below=BLANK):
    """A quantity expressed in one frame: Craig left superscript only."""
    return ('<m:sPre><m:sub>' + below + '</m:sub><m:sup>' + above
            + '</m:sup><m:e>' + base + '</m:e></m:sPre>')

HERE = Path(__file__).resolve().parent.parent
EVIDENCE = HERE / 'audit_evidence/ch7_restructured'
human = json.loads((EVIDENCE / 'human.json').read_text())
synthetic = json.loads((EVIDENCE / 'synthetic.json').read_text())
natural = json.loads((EVIDENCE / 'natural.json').read_text())['summary']
# Section 7.2 object result, pinned under D-035 and D-036. Every object number
# printed below is formatted straight out of this file, so the prose cannot
# drift from the evidence. audit_evidence/ch7_restructured/object.json.
object_json = json.loads((EVIDENCE / 'object.json').read_text())
obj = object_json['recordings']
HALF = object_json['detection_rule']['parameters']['half_window_samples']
# Bare kinematic-model wrist error, recomputed against the accepted raw
# measured landmark of D-040 and reported at the two scopes of D-041.
# audit_evidence/ch7_restructured/bare_model.json.
bare = json.loads((EVIDENCE / 'bare_model.json').read_text())['recordings']
# Loop-recording object-marker detection coverage, recomputed and pinned in
# audit_evidence/ch7_restructured/loop_detection.json by
# scripts/evaluate_ch7_loop_detection.py. Section 7.4.2 prints the detected
# criterion, detected == 1 on the cleaned object track, which the ArUco
# extractor sidecar independently confirms; the stricter accepted criterion of
# Section 7.2, which also drops despiked and interpolated rows, gives 2046 of
# 2099 and is recorded in the same file. The location of the missing frames
# comes from the pinned two-hand grip episodes of the recovery output.
loop_detected = json.loads((EVIDENCE / 'loop_detection.json').read_text())
loop_gap = loop_detected['location_of_missing_frames']
assert len(loop_gap['missing_outside_handover_span']) == 1
doc = Document()
doc.styles['Normal'].font.name = 'Times New Roman'
doc.styles['Normal'].font.size = Pt(12)
for name, size in (('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
for section in doc.sections:
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
text = []


def norm(value):
    return d(value, l='&#x2016;', rr='&#x2016;')


def summation(value):
    return ('<m:nary><m:naryPr><m:chr m:val="&#x2211;"/>'
            '<m:limLoc m:val="subSup"/></m:naryPr>'
            '<m:sub>' + r('i=1') + '</m:sub><m:sup>' + r('n')
            + '</m:sup><m:e>' + value + '</m:e></m:nary>')


pj = sub(r('p'), r('i'))
delta = d(pj + r('-c'))
projection = d(r('I-') + r('u') + sup(r('u'), r('T'))) + delta
lr = sub(nor('Lrec'), r('j'))
lp = sub(r('L'), r('j'))
absolute = sub(r('A'), r('j'))
MATH = {
    '7.1': eqArr(lr + r('=') + norm(sub(r('q'), r('j+1')) + r('-') + sub(r('q'), r('j'))),
                 absolute + r('=') + d(lr + r('-') + lp, l='|', rr='|'),
                 sub(r('R'), r('j')) + r('=100') + frac(absolute, lp) + nor(' (%)')),
    '7.2': eqArr(r('c=') + frac(r('1'), r('n')) + summation(pj),
                 r('C=') + frac(r('1'), r('n')) + summation(delta + sup(delta, r('T')))),
    '7.3': eqArr(sub(nor('min'), eqArr(r('c,u'), norm(r('u')) + r('=1')))
                 + summation(sup(norm(projection), r('2'))),
                 nor('line') + d(r('s')) + r('=c+su')),
    '7.4': sub(r('d'), r('i')) + r('=') + norm(projection),
    # C7-07a (FRAME_INVENTORY 8.9.7 row 5, D-060): both operands are in the
    # Unity scene frame, so the frame moves to Craig's left superscript and the
    # right superscript keeps only the provenance of the position.
    '7.5': sub(r('e'), r('k,t')) + r('=') + norm(
        presup(sup(sub(r('p'), r('k,t')), nor('rig')), nor('Scene')) + r('-')
        + presup(sup(sub(r('p'), r('k,t')), nor('measured')), nor('Scene'))),
    # C7-07b (FRAME_INVENTORY 8.9.7 row 8): both operands are y-up camera-frame
    # positions; Section 7.4 enters no display frame.
    '7.6': sub(r('e'), r('m,k,t')) + r('=') + norm(
        presup(sup(sub(r('p'), r('m,k,t')), nor('masked')), nor("Camera'")) + r('-')
        + presup(sup(sub(r('p'), r('k,t')), nor('unmasked')), nor("Camera'"))),
}


def heading(value, level):
    doc.add_heading(value, level=level)
    text.append('\n' + value + '\n')


def p(value):
    doc.add_paragraph(value)
    text.append(value + '\n')


def equation(value, number):
    # Structured OMML avoids LibreOffice interpreting ASCII norm bars as
    # logical operators or placeholders during the PDF conversion.
    add_display_eq(doc, MATH[number], number)
    text.append(value + '  (' + number + ')\n')


def table(rows, caption):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    if len(rows[0]) == 6:
        t.autofit = False
        for column, width in zip(t.columns, (1.25, 1.35, 0.45, 1.15, 1.15, 1.3)):
            column.width = Inches(width)
            for cell in column.cells:
                cell.width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            t.cell(i, j).text = str(value)
            for par in t.cell(i, j).paragraphs:
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
        props = t.rows[i]._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
    par = doc.add_paragraph(caption)
    par.runs[0].font.size = Pt(10)
    for row in rows:
        text.append(' | '.join(map(str, row)))
    text.append(caption + '\n')


def figure(name, width, caption, context=None):
    if context is not None:
        doc.add_picture(str(HERE / 'figures' / context), width=Inches(width))
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    doc.add_picture(str(HERE / 'figures' / name), width=Inches(width))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    par = doc.add_paragraph(caption)
    par.paragraph_format.keep_together = True
    par.runs[0].font.size = Pt(10)
    text.append(caption + '\n')


def row_stats(s):
    return [str(s['n'])] + [f"{s[k]:.1f}" for k in ('median', 'p95', 'maximum')]


def caveat(alias, item):
    # object.json records the disclosures the prose must carry as a caveats
    # list; this picks one by its stored name rather than by position.
    return next(c for c in obj[alias]['caveats'] if c['item'] == item)


def object_rows(alias):
    # object.json recordings.<alias>.segments: physical reference, reconstructed
    # endpoint separation, absolute and relative error of equation 7.1.
    rows = [['Segment', 'Physical length (cm)', 'Reconstructed length (cm)',
             'Absolute error (cm)', 'Relative error (%)']]
    for s in obj[alias]['segments']:
        rows.append([s['segment'], f"{s['physical_cm']:.1f}",
                     f"{s['reconstructed_cm']:.2f}",
                     f"{s['absolute_error_cm']:.2f}",
                     f"{s['relative_error_pct']:.2f}"])
    return rows


def bands(alias):
    # object.json caveat "detector-selection band": the spread of each length
    # over the detection parameters and over boundary shifts of two and five
    # accepted samples, phase-pattern failures excluded.
    seg = caveat(alias, 'detector-selection band')['segments']
    parts = [f"{v['minimum_cm']:.2f} to {v['maximum_cm']:.2f} centimetres for {k}"
             for k, v in seg.items()]
    return ', '.join(parts[:-1]) + ' and ' + parts[-1]


def bare_stat(alias, joint, scope, key='vs_raw_measured'):
    # bare_model.json recordings.<alias>.joints.<joint>.<scope>.<key>:
    # measured wrist against the wrist the kinematic model places directly,
    # before Unity rig mapping and display filtering.
    return bare[alias]['joints'][joint][scope][key]


def held_share(alias, joint, scope):
    # Share of the frame set on which the forearm twist output is held.
    block = bare[alias]['joints'][joint][scope]
    return 100 * block['twist_held_frames'] / block['n']


def twist_span(alias, state):
    # Range of the bare-model wrist median over both arms and both scopes,
    # split by whether the forearm twist output is measured or held.
    key = 'vs_raw_measured_twist_' + state
    values = [bare[alias]['joints'][j][scope][key]['median']
              for j in ('right_wrist', 'left_wrist')
              for scope in ('section_7_3_frame_set', 'recording_wide_accepted')
              if bare[alias]['joints'][j][scope][key]['n']]
    return f"{min(values):.1f} and {max(values):.1f}"


def human_rows(alias, sides):
    rows = [['Arm', 'Joint', 'n', 'Median (cm)', 'p95 (cm)', 'Maximum (cm)']]
    for side in sides:
        for joint in ('elbow', 'wrist'):
            rows.append([side.capitalize(), joint.capitalize(), *row_stats(human[alias]['joints'][side + '_' + joint])])
    return rows


heading('Chapter 7: Experimental Evaluation', 1)
heading('7.1 Evaluation Method', 2)
p('The evaluation uses the single-hand rail recording, the handover recording on the rail, and the loop recording for natural occlusion. Tape measurements provide the reference for object segment lengths once waypoint correspondence is established. Rendered elbow and wrist positions are compared with synchronized MediaPipe landmarks deprojected using aligned depth. Synthetic removal uses valid unmasked kinematic reconstruction as its reference. Natural occlusion uses manually annotated wrist labels, with their observed separation from accepted wrist measurements evaluated on clean frames.')
p('Physical length error is defined as absolute and relative error. Positional differences and trajectory scatter are summarized by the median, the 95th percentile (p95), and the maximum. Each joint and method uses its stated frame set; n is the number of evaluated frames. Percentiles use linear interpolation between ordered observations.')
# D-040: one evaluation reference for every quantitative reconstruction error
# in the thesis, the accepted raw MediaPipe and depth wrist before temporal
# filtering. bare_model.json records both variants, so the cost of the
# convention is measured rather than assumed.
p('The reference wrist for quantitative reconstruction error is the accepted raw MediaPipe and depth three-dimensional wrist measurement, taken before temporal filtering. The same reference applies to the kinematic-model wrist error, to the rendered Unity rig wrist error, and to any clean-frame baseline that stands for reconstruction error against a measured wrist. A filtered wrist is used elsewhere for algorithmic purposes; wherever it appears it is named as a filtered measurement and is not the evaluation reference. No comparison table mixes the two references.')
# D-047: two object coverage terms, defined once here and used at every site
# that reports an object frame count. The two counts answer different
# questions and are deliberately not harmonized into one convention. No count
# belongs in the definition itself.
p('The object frame counts reported below use two terms. A detected frame is a frame on which the ArUco extractor returned a marker pose. An accepted object sample is a returned pose that remained usable after the cleaning rules of Section 7.2. The two answer different questions, so each count names the term it reports.')

heading('7.2 Object Reconstruction Accuracy', 2)
p('The route has four physical waypoint roles: W1 is the parked start on the desk, W2 is the foot of the lift, W3 is the start of the rail, and W4 is the far endpoint. The segments are W1 to W2, W2 to W3, and W3 to W4. A segment comparison requires the same named point on the object at both physically measured endpoints and at their corresponding reconstructed positions.')
p('Let q_j and q_(j+1) be reconstructed positions at the independently identified waypoint events, and let L_j be the physically measured endpoint separation. Reconstructed length, absolute error, and relative error are defined by')
equation('Lrec_j = ||q_(j+1) - q_j||;  A_j = |Lrec_j - L_j|;  R_j = 100 A_j / L_j (%).', '7.1')
p('Here L_j is positive. Reconstructed length is endpoint separation, not accumulated sample-to-sample travel or the longest span of a fitted line. A stationary-frame interval, when available, defines an endpoint by the mean of its accepted reconstructed samples; an identified single-frame event uses that sample.')
p('For reconstructed samples p_i, i = 1, ..., n, within one segment and expressed in {Scene} of Section 6.4, orthogonal least-squares line fitting uses the sample mean c and covariance C:')
equation('c = (1/n) sum_i p_i;  C = (1/n) sum_i (p_i - c)(p_i - c)^T.', '7.2')
p('Principal component analysis (PCA) gives a unit eigenvector u of C associated with its largest eigenvalue. This direction and the line through c minimize the squared perpendicular distances:')
equation('min_(c,u), ||u||=1  sum_i ||(I - uu^T)(p_i - c)||^2;  line(s) = c + su.', '7.3')
p('The parameter s ranges over the real numbers and I is the identity matrix. The perpendicular distance of sample p_i is')
equation('d_i = ||(I - uu^T)(p_i - c)||.', '7.4')
p('The distances d_i describe scatter of the reconstructed trajectory about its own fitted line. They do not measure physical accuracy. A degenerate sample set with no unique principal direction does not define a fitted segment direction. Physical length error and trajectory scatter are reported separately.')

p('The measured point is the ArUco marker centre on the cube face, at both physically measured endpoints and in the reconstruction. The physical separations are tape measurements made by the author between marker-centre positions, at a resolution of 1 millimetre, and the physical setup is the same in both recordings. The physical route is not registered into the world frame of the desk-marker calibration or into {Scene}, so the trajectory plots report reconstructed samples rather than surveyed physical paths or physical waypoint coordinates.')

heading('7.2.1 Single-Hand Rail Task', 3)
# object.json recordings.r6b.samples: 900 frames, 889 accepted, 11 excluded.
s = obj['r6b']['samples']
p(f"The rail recording holds {s['frames_total']} frames, of which {s['accepted']} supply accepted object samples of the marker centre. The other {s['frames_total'] - s['accepted']} frames are excluded because the marker was not detected, because the stored sample was filled by the tracker, or because the recorded position was not finite. No exclusion uses the size of any error.")
# object.json detection_rule.estimator and recordings.r6b.waypoints.
w = obj['r6b']['waypoints']
p(f"Waypoint timing comes from the marker-centre motion of this recording alone. Each axis velocity is a centred difference over {HALF} accepted samples on either side. A transition is bracketed at the half-peak of the velocity of the moving axis, and the boundary is the last sample within 5 millimetres of the local pre-transition level of that axis. W1 covers frames {w['W1']['first_frame']} to {w['W1']['last_frame']}, an opening interval of {w['W1']['n_samples']} samples whose mean gives the position. W2 is frame {w['W2']['first_frame']} and W3 is frame {w['W3']['first_frame']}, each a single transition sample, because the cube does not pause on either side of the lift and no dwell exists there to average. W4 covers frames {w['W4']['first_frame']} to {w['W4']['last_frame']}, a closing interval of {w['W4']['n_samples']} samples whose mean gives the final position.")
p('Table 7.1 compares the reconstructed segment lengths with the physical reference, using equation (7.1). Figure 7.1 places recorded views of the task above the accepted samples and the four waypoint positions.')
table(object_rows('r6b'), 'Table 7.1. Rail recording: reconstructed marker-centre segment lengths against the physical reference, with absolute and relative error.')
figure('ch7_object_waypoints_r6b.png', 6.1, 'Figure 7.1. Rail recording: selected video frames show object acquisition, lift, slide and arrival at the far end. Below, accepted marker-centre samples and the four reconstructed waypoints are shown in {Scene}: x-y at upper left, x-z at lower left and a three-dimensional view on the right. The x-y view expands the height scale to show the lift; the other views use equal metric scale. All positions include the recording-specific floor translation. The video frames do not define the waypoint estimates.', context='ch7_video_context_r6b.png')
p(f"Each reconstructed length is reported with a range obtained by varying the detection parameters and shifting each boundary by two and by five accepted samples. Variants that no longer reproduce the depth, vertical and lateral phase pattern are excluded. The ranges are {bands('r6b')}. The point estimate alone overstates the precision of the result.")
# object.json caveats "single-sample transition waypoints" (local depth spread
# within plus or minus 7 accepted samples) and "boundary resolved only to the
# velocity half-peak".
depth = caveat('r6b', 'single-sample transition waypoints')['local_depth_spread_mm']
p(f"W2 and W3 are single accepted samples, so each carries the instantaneous depth error of one frame rather than an averaged value. Within {HALF} frames on either side the depth coordinate spreads {depth['W2']:.1f} millimetres at W2 and {depth['W3']:.1f} millimetres at W3. The criterion based on departure from the pre-transition level also fails to separate either boundary from the half-peak of its axis velocity, so the reported sample is the latest the search allows and the transition lies at or before it.")
# object.json caveat "W1 and W4 averaging intervals". The closing duration of
# about 0.8 s is 25 accepted samples at the recorded frame interval; the
# remaining sentence is the author's reading of the same trajectory, recorded
# in the writer brief of 2026-09-12.
w4 = caveat('r6b', 'W1 and W4 averaging intervals')
p(f"The W4 interval holds {w['W4']['n_samples']} samples, about 0.8 seconds. Its depth coordinate spans {w4['W4_spread_mm'][2]:.1f} millimetres, and its largest single-sample step is {w4['W4_largest_step_mm']:.1f} millimetres, between frames {w4['W4_largest_step_between_frames'][0]} and {w4['W4_largest_step_between_frames'][1]}. The closing interval is therefore not a single settled cluster. The recording ends about ten frames after the cube settles, and no longer settled interval is available to average.")
# object.json recordings.r6b.scatter.rail_W3_to_W4.
sc = obj['r6b']['scatter']['rail_W3_to_W4']
p(f"Scatter about the fitted line is a separate within-recording comparison, defined by equations (7.2) to (7.4), between the reconstructed trajectory and a line fitted to the same samples. Over the {sc['n']} samples between W3 and W4 the median perpendicular distance is {sc['median']:.2f} centimetres, the p95 is {sc['p95']:.2f} centimetres and the maximum is {sc['maximum']:.2f} centimetres. The fitted line lies {sc['tilt_from_horizontal_deg']:.1f} degrees from horizontal. These distances describe the spread of the reconstructed samples and are not comparable with the errors of Table 7.1.")

heading('7.2.2 Handover Task', 3)
# object.json recordings.r7.samples and recordings.r7.waypoints.
s = obj['r7']['samples']
w = obj['r7']['waypoints']
p(f"The handover recording holds {s['frames_total']} frames, of which {s['accepted']} supply accepted object samples of the marker centre, and the other {s['frames_total'] - s['accepted']} are excluded under the same rule. Its waypoints are detected from its own marker-centre motion, with the estimator and the parameters used for the rail recording.")
p(f"W1 covers frames {w['W1']['first_frame']} to {w['W1']['last_frame']}, an opening interval of {w['W1']['n_samples']} samples. W2 is frame {w['W2']['first_frame']} and W3 is frame {w['W3']['first_frame']}, again single transition samples across the lift. W4 covers frames {w['W4']['first_frame']} to {w['W4']['last_frame']}, a closing interval of {w['W4']['n_samples']} samples, because the cube stays on the rail while the recording continues after the transfer. Table 7.2 reports the resulting lengths against the same physical reference, and Figure 7.2 shows the trajectory.")
table(object_rows('r7'), 'Table 7.2. Handover recording: reconstructed marker-centre segment lengths against the physical reference, using the columns of Table 7.1.')
figure('ch7_object_waypoints_r7.png', 6.1, 'Figure 7.2. Handover recording: selected video frames show the right-hand slide, transfer, left-hand slide and arrival at the far end. Below, accepted marker-centre samples and the four reconstructed waypoints are shown in {Scene}: x-y at upper left, x-z at lower left and a three-dimensional view on the right. The plots use this recording\'s floor translation and the scale conventions of Figure 7.1. The selected video moments provide task context.', context='ch7_video_context_r7.png')
depth = caveat('r7', 'single-sample transition waypoints')['local_depth_spread_mm']
p(f"The ranges of reconstructed lengths under the tested detection settings are {bands('r7')}. W2 and W3 are again single accepted samples. Within {HALF} frames on either side the depth coordinate spreads {depth['W2']:.1f} millimetres at W2 and {depth['W3']:.1f} millimetres at W3, so the local depth variation at both transition endpoints is larger than in the rail recording.")
sc = obj['r7']['scatter']['rail_W3_to_W4']
p(f"Over the {sc['n']} samples between W3 and W4 the median perpendicular distance from the fitted line is {sc['median']:.2f} centimetres, the p95 is {sc['p95']:.2f} centimetres and the maximum is {sc['maximum']:.2f} centimetres. The fitted line lies {sc['tilt_from_horizontal_deg']:.1f} degrees from horizontal. This remains a within-recording comparison and is reported separately from the length errors.")
# Sign invariance verified over object.json recordings.*.sensitivity.variants:
# every phase-passing variant of the band family (52 for the rail recording, 46
# for the handover recording) keeps the sign of all three errors. The wider
# ten-frame shift, which the band excludes, flips W3 to W4 in the rail
# recording and W2 to W3 in the handover recording.
p('The two recordings agree on the sign of every error. The depth leg from W1 to W2 reconstructs short in both, while the two legs after the lift reconstruct long in both. The handover recording disagrees with the physical reference more than the rail recording on all three segments, in absolute and in relative terms. Every detection variant included in the ranges above keeps these signs, so the threshold changes the error magnitude within each range but does not produce the pattern. A wider sweep that shifts a boundary by ten accepted samples lies outside that set of variants and reverses one sign in each recording, on W3 to W4 in the rail recording and on W2 to W3 in the handover recording.')
p('The proper gravity-alignment rotation G and the common floor translation of Section 6.4 preserve Euclidean distances, including segment lengths and perpendicular distances from a trajectory to its fitted line. They therefore leave Tables 7.1 and 7.2 and the scalar scatter statistics unchanged. Absolute coordinates depend on the origin. Displacement components, fitted-line tilt and the checks used to confirm task phases depend on the axis orientation. Those checks identify which axis has the largest motion. The floor translation shifts the plotted positions without changing those displacements, tilts or checks. The trajectory panels of Figures 7.1 and 7.2 use the final Scene coordinates throughout.')

heading('7.3 Human Reconstruction Accuracy under Valid Landmark Measurements', 2)
# D-049: the human acceptance concept of this section is named here, once.
# It stays distinct from the two object terms defined in Section 7.1 under
# D-047 (detected frame, accepted object sample); no count is shared.
p('For each arm, the reference contains raw MediaPipe elbow and wrist coordinates with valid aligned depth. A valid-landmark frame is a frame in which the shoulder, elbow, wrist and required torso inputs used by the reconstruction of this section satisfy the landmark-validity checks. Such a frame requires observed, finite, positive-depth shoulder, elbow, wrist and torso inputs, no filled or rejected filtered input, and acceptance by the landmark checks applied to individual frames and detected events. Frames outside the detector evaluation interval are excluded. These checks define when a measurement counts as valid; they do not establish an anatomical reference.')
p('The corresponding Unity joints are the forearm bone head for the elbow and the hand bone head for the wrist. Their saved world positions include display filtering and rig mapping. The calibrated camera mapping, gravity alignment and recording-specific floor translation carry the measured landmarks into {Scene}, where the logged rig joints are expressed. Frame identifiers match exactly, and timestamp and pelvis correspondence checks verify the saved capture. For joint k at frame t, positional error is')
equation('e_(k,t) = ||^Scene p_(k,t)^rig - ^Scene p_(k,t)^measured||.', '7.5')
p('Both positions in equation (7.5) are expressed in the Unity scene frame of Section 6.4. The left superscript names that frame, and the right superscript names where the position comes from, either the rendered rig or the measured landmark. The statistics describe reconstruction error relative to measured 3D landmarks. Detector-accepted inputs can still produce constrained root or held twist output; the comparison retains those outputs as part of the rendered reconstruction. It does not isolate an entirely measured-output kinematic solve.')
# D-041: the two quantities are reported separately and never interchangeably.
# bare_model.json quantity block defines both.
p('The chapter distinguishes two reconstruction quantities throughout, and neither stands for the other. The kinematic-model error compares the measured joint with the position that the kinematic model and forward kinematics place directly from the solved angles, before Unity rig mapping and display filtering. The rendered rig error compares the same measured joint with the final Unity rig joint, after the complete display and rendering path. Equation (7.5) applies to both, with the corresponding reconstructed position.')
# D-043: capture configuration, stated as fact. bare_model.json records the
# calibrated segment lengths against the avatar lengths actually used: both
# captures carried the arm lengths of the loop recording, and the rail capture
# also kept the receiver's default torso length (rig_dimensions.csv beside each
# capture and the receiver's serialized configuration). No share of the
# rendered-rig difference is attributed to it here.
p('Both Unity captures ran the avatar with the arm segment lengths of the loop recording rather than the calibrated lengths of the recording being replayed, and the rail capture also kept the default torso length of the receiver. Tables 7.3 and 7.4 are therefore end-to-end rendered results under the capture configuration used, not tests of rig mapping and display filtering with matched body dimensions. Display filtering, rig mapping and the remaining stages of the display path are not separated by this comparison, so no part of the rendered difference is attributed to any one of them.')

heading('7.3.1 Single-Hand Rail Task', 3)
s = human['r6b']['joints']['right_wrist']
p(f"The working right arm supplies {s['n']} matched valid-landmark frames from the saved capture. Table 7.3 reports the elbow and wrist errors on this same frame set. The captured frames are not a complete or uniform sample of the task. Every retained frame has a constrained root output, so these results do not establish reconstruction performance without root constraints.")
# D-042: no bare-model value is printed for this recording. bare_model.json
# records the right arm here as held twist and constrained root on every
# accepted frame, at both scopes, so the distance would not isolate the
# kinematic model. The assertions keep that statement tied to the evidence.
for scope in ('section_7_3_frame_set', 'recording_wide_accepted'):
    block = bare['r6b']['joints']['right_wrist'][scope]
    assert block['twist_held_frames'] == block['n']
    assert block['output_tags']['root'].get('2') == block['n']
p('The right forearm twist output is held on every valid-landmark frame of this recording, alongside the constrained root output just noted. A kinematic-model wrist distance computed here would include the effects of the held twist and the constrained root. It would not isolate the agreement between the kinematic model and the measured landmarks. No such distance is reported for this recording or compared with the handover results. The quantitative kinematic-model accuracy comes from the handover recording.')
table(human_rows('r6b', ['right']), 'Table 7.3. Single-hand task: rendered joint error relative to accepted measured 3D landmarks on the captured subset of valid-landmark frames.')
p('Figure 7.3 shows the marker-centre and model-wrist trajectories through the single-hand rail task. These views provide spatial context for the captured reconstruction.')
figure('ch7_unity_trajectories_r6b.png', 6.3, 'Figure 7.3. Single-hand rail task in Unity: marker-centre and model-wrist trajectories during carry, lift and slide. Wrist colours identify the recorded state. The avatar and cube show the final frame of each interval. The dashed line is a schematic scene guide. These interval views provide spatial context; Table 7.3 evaluates rendered rig joints on a separate matched valid-landmark subset.')

heading('7.3.2 Handover Task', 3)
p('Table 7.4 evaluates both arms separately. The elbow and wrist of each arm share a valid-landmark frame set; the two arms have different coverage because their measurement validity differs. Figure 7.4 shows the saved Unity trajectories through the handover for spatial context. The quantitative capture is incomplete and includes constrained root output on part of each valid-landmark set. Differences from Table 7.3 cannot be attributed solely to the handover condition because the recorded subsets and rig states also differ.')
table(human_rows('r7', ['right', 'left']), 'Table 7.4. Handover task: rendered joint error relative to accepted measured 3D landmarks, reported separately for each arm on its captured subset of valid-landmark frames.')
figure('ch7_unity_trajectories_r7.png', 6.3, 'Figure 7.4. Handover task in Unity: marker-centre and both model-wrist trajectories during the right-hand, transfer and left-hand intervals, with the state colours and schematic scene guide shown in the legend. The avatar and cube show each interval\'s final frame. Table 7.4 evaluates rendered rig joints on separate matched valid-landmark frame sets for each arm.')
# D-041 scope one: the capture subset, frame-matched to Table 7.4 and used only
# for the contrast with the rendered rig. bare_model.json recordings.r7.joints.
# *_wrist.section_7_3_frame_set.vs_raw_measured.
sub_r = bare_stat('r7', 'right_wrist', 'section_7_3_frame_set')
sub_l = bare_stat('r7', 'left_wrist', 'section_7_3_frame_set')
p(f"The kinematic-model wrist error separates the model from the display path. It uses the same {sub_r['n']} right-arm and {sub_l['n']} left-arm valid-landmark frames as Table 7.4. On those frames the kinematic-model wrist lies a median of {sub_r['median']:.2f} centimetres from the measured wrist on the right and {sub_l['median']:.2f} centimetres on the left, against the rendered rig medians in that table. The additional discrepancy arises after the kinematic model, along the path described above.")
# D-041 scope two: recording-wide accepted frames, the representative bare-model
# performance. The subset over-represents held-twist frames, so the two scopes
# are stated separately and are never compared with Table 7.4 across frame sets.
wide_r = bare_stat('r7', 'right_wrist', 'recording_wide_accepted')
wide_l = bare_stat('r7', 'left_wrist', 'recording_wide_accepted')
p(f"That capture subset is not representative of the full valid-landmark frame population, and least so on the left arm. It over-represents frames on which the forearm twist output is held, at {held_share('r7', 'left_wrist', 'section_7_3_frame_set'):.0f} per cent of the subset against {held_share('r7', 'left_wrist', 'recording_wide_accepted'):.0f} per cent of the valid-landmark frames of the whole recording. Across all valid-landmark frames the kinematic-model wrist median is {wide_r['median']:.2f} centimetres on the right, over {wide_r['n']} frames, and {wide_l['median']:.2f} centimetres on the left, over {wide_l['n']} frames. These recording-wide values are the representative performance of the kinematic model, and they are not set against Table 7.4, whose frame set is different.")
# Supporting structure, all from bare_model.json: the twist split of the
# bare-model wrist medians and the elbow medians on the Table 7.4 frame sets.
elbow_r = bare_stat('r7', 'right_elbow', 'section_7_3_frame_set')
elbow_l = bare_stat('r7', 'left_elbow', 'section_7_3_frame_set')
p(f"The state of the forearm twist output separates two cases. Where the twist is measured the kinematic-model wrist median lies between {twist_span('r7', 'measured')} centimetres; where it is held it lies between {twist_span('r7', 'held')} centimetres. The elbow is almost unaffected, because the twist rotates about the upper-arm axis: on the frame sets of Table 7.4 the kinematic-model elbow median is {elbow_r['median']:.2f} centimetres on the right and {elbow_l['median']:.2f} centimetres on the left.")

heading('7.3.3 Human-Object Spatial Check', 3)
# Author's physical measurement, REVISION_2026-09-11_BRIEF.md section 3 fact 4
# (approximately 16 cm, normal single-hand grip, MediaPipe wrist landmark
# location to ArUco marker centre), retained in Chapter 7 by D-038.
p('The separation between the wrist and the held cube has one external physical reference. The author measured approximately 16 centimetres between the physical location corresponding to the MediaPipe wrist landmark and the object ArUco marker centre, for a normal single-hand grip. This is an approximate scalar physical reference, not a per-frame constraint.')
# D-089: place the recorded/Unity pairs beside the grip discussion.
p('Figure 7.5 pairs recorded and rendered views of the single-hand grip during acquisition, lift and slide. Red spheres identify joint groups that the stream does not report as measured. The views illustrate the task and retain the captured pose and object-placement discrepancies; the numerical tables use their own logged frame sets and capture configurations.')
figure('ch7_visual_examples_r6b.png', 6.5, 'Figure 7.5. Rail recording: object acquisition, lift and slide, with recorded video on the left and the saved Unity sensor view on the right. The visualization capture was made on 7 September 2026. The selected moments illustrate the task and do not define an evaluation sample.')
# Reconstructed medians, REVISION_2026-09-11_BRIEF.md section 3 fact 5, pinned
# to eval/reports/r7_handover.json (*_fk_wrist_to_marker_at_cube_cm medians);
# printed as the author gave them. D-038 restricts the comparison to the two
# single-hand intervals.
p('Reconstruction places the wrist a median of 15.77 centimetres from the marker centre while the right hand holds the cube alone, and 14.08 centimetres while the left hand holds it alone. Both reconstructed positions are expressed in {Scene}. The common floor translation cancels in their difference, and the proper gravity-alignment rotation preserves its norm, so these distances are unchanged. Both medians sit within about two centimetres of the physical reference, so the reconstructed human-object separation is of the correct magnitude.')
p('The two medians recorded during the transfer itself, 17.91 centimetres for the right hand and 15.85 centimetres for the left, are not evaluated against the same reference. Grip geometry changes while both hands are on the cube, so the single-hand measurement does not describe that configuration. Those two values describe the reconstructed grip geometry during transfer.')
p('Figure 7.6 shows the change of holding hand in the recorded scene and the corresponding saved Unity views.')
figure('ch7_visual_examples_r7.png', 6.5, 'Figure 7.6. Handover recording: right-hand slide, transfer and left-hand slide, using the paired views of Figure 7.5. This visualization was captured on 9 September 2026 with a torso-length override. The visualization supplies qualitative examples; Table 7.4 uses a separate capture.')
p('This is an approximate external physical check. It is not an anatomical reference for the wrist, because the physical endpoint is the location corresponding to a detector landmark rather than a joint centre. It also does not validate the three-dimensional hand-object offset that the recovery of Chapter 5 uses, because one scalar distance constrains neither the direction of that offset nor its behaviour from frame to frame.')

heading('7.4 Reconstruction During Landmark Failure', 2)
heading('7.4.1 Synthetic Landmark Removal', 3)
p('The experiment removes the right elbow and wrist inputs after offline filtering, while retaining the measured shoulder and torso. The reference is the unmasked kinematic reconstruction of the same frames. Every selected reference frame has valid observed arm and torso inputs, accepted landmark checks, an established grip, and measured root, shoulder swing, shoulder twist and elbow output groups. The handover recording provides suitable windows; the single-hand recording does not provide a fully measured root-and-arm reference window under these conditions.')
p('All eligible windows contain 45 frames, matching the recovery limit. Motion magnitude is the maximum distance of the unmasked wrist from its first position within a window. Selection precedes recovery scoring. The first window has the smallest excursion. The second has the largest excursion among windows that do not overlap the first. The third has the excursion closest to the median of all candidates and overlaps neither earlier selection. Table 7.5 lists the selections. These are contrasting windows from one recording, not independent experimental trials.')
rows = [['Window', 'Frame range', 'n', 'Reference wrist excursion (cm)']]
for w in synthetic['selection']['selected']:
    rows.append([w['window'].capitalize(), f"{w['start']} to {w['stop']}", str(w['stop'] - w['start'] + 1), f"{w['excursion_cm']:.1f}"])
table(rows, 'Table 7.5. Synthetic removal windows selected using valid unmasked wrist motion before method errors were calculated.')
p('The methods are hold-last joint angles, direction-memory recovery, and object-assisted recovery. Each method replays the recording from a fresh initial state and agrees with the reference immediately before removal. Masked landmarks do not update the hand-object offset, which is already established and remains frozen during the gap. Each method uses the same retained shoulder position and calibrated segment lengths to reconstruct the elbow and wrist. For method m, the deviation from the unmasked reference is')
equation("e_(m,k,t) = ||^Camera' p_(m,k,t)^masked - ^Camera' p_(k,t)^unmasked||.", '7.6')
p("Both positions in equation (7.6) are expressed in the y-up camera frame {Camera'}, the left-handed frame of Section 3.1 in which the kinematic chain is built, and no display frame enters this section. Tables 7.6 and 7.7 use identical frames and metrics for every method. This is a test of missing solver inputs with fixed offline preprocessing and calibration. It does not measure anatomical accuracy or the effects of removing observations before filtering.")
for number, joint in ((6, 'elbow'), (7, 'wrist')):
    rows = [['Window', 'Method', 'n', 'Median (cm)', 'p95 (cm)', 'Maximum (cm)']]
    for s in synthetic['results']:
        if s['joint'] == joint:
            rows.append([s['window'].capitalize(), s['method'].capitalize(), *row_stats(s)])
    table(rows, f'Table 7.{number}. Synthetic {joint} reconstruction error relative to the valid unmasked model reconstruction.')
p('No method has the smallest error across every window and joint. Object assistance reduces the largest wrist deviation in the intermediate- and higher-motion windows, but increases elbow deviation in the higher-motion window. The lower-motion window favours direction memory. The results therefore support a window-dependent comparison rather than a general recovery ranking.')

heading('7.4.2 Natural Occlusion', 3)
p("Natural occlusion is evaluated on the loop recording. A single annotator marked the bead bracelet on the left wrist and the watch on the right wrist. The same visible feature defines the manual wrist label on clean and occluded frames. Each click is deprojected using the aligned depth and colour intrinsics. Table 7.8 compares the label with a filtered wrist measurement in the camera frame {Camera}, and Table 7.9 compares it with the reconstructed wrist in the y-up camera frame {Camera'}. The two frames share an origin and differ by the fixed axis flip of Section 3.1, which preserves distance, so every reported separation is the same in either frame.")
# Object-marker availability bounds when object-assisted recovery can have
# object information at all, so the coverage is stated here before the methods
# are compared. Source: audit_evidence/ch7_restructured/loop_detection.json.
# D-047: the count is detection coverage, named with the term defined in
# Section 7.1, and the closing sentence keeps a reader from reading it as a
# statement about sample quality. The informal plausibility observation of
# loop_detection.json, including its implied speed threshold, stays out of the
# prose because the thesis defines no such evaluation rule.
p(f"The loop recording holds {loop_detected['frames_total']} frames, of which {loop_detected['detected']['n_frames']} are detected frames for the object marker. Of the {loop_detected['detected']['n_missing']} frames without a detected marker, {loop_gap['n_missing_in_handover_span']} lie in the two-hand transfer on the desk, where a hand crosses the border of the printed square, and the remaining frame sits at its near edge. The object therefore supplies no geometric constraint across that stretch, while every frame labelled below carries a detected marker. A detected frame is not necessarily an accepted object sample or a reliable trajectory sample, because detection coverage records whether the extractor returned an object observation and not whether that observation survived the later cleaning.")
p('Figure 7.7 shows four retained video frames from the natural-failure sample. The visible wrist features allow manual wrist labelling even when the corresponding landmark input fails.')
figure('ch7_failure_context.png', 6.1, 'Figure 7.7. Loop recording: natural-failure frames 1427, 1462, 1777 and 1890, with the affected arm identified above each view. The bracelet or watch remains visible for manual wrist labelling. These recorded examples are separate from the synthetic input removal of Section 7.4.1.')
p('Table 7.8 first measures how far the label sits from accepted, filtered MediaPipe-plus-depth wrist measurements on clean wrist frames. That filtered measurement locates the detected landmark for this annotation comparison and is not the evaluation reference of Section 7.1; the table reports a label separation and not a reconstruction error. The visible feature is displaced from the landmark definition of the wrist; its separation is not an anatomical calibration.')
rows = [['Arm', 'Label', 'n', 'Median (cm)', 'p95 (cm)', 'Maximum (cm)']]
for side, feature in (('left', 'Bracelet'), ('right', 'Watch')):
    rows.append([side.capitalize(), feature, *row_stats(natural['clean'][side])])
table(rows, 'Table 7.8. Clean-frame distance between the manual wrist label and an accepted filtered wrist measurement.')
p('During occlusion, the plain solve uses the reported landmarks, hold-last retains the last accepted joint angles, and object-assisted recovery reconstructs the missing inputs. Each method is evaluated against the same retained label points for that arm. The metric is the Euclidean distance from the reconstructed wrist to the label, summarized in Table 7.9; the clean-frame scalar offset is not subtracted.')
p('Figure 7.8 illustrates a left-wrist failure with no MediaPipe wrist observation. The object-derived wrist estimate and the wrist placed by the reconstructed arm are distinct points.')
figure('ch7_natural_proxy_left.png', 6.1, 'Figure 7.8. Loop frame 1462, left wrist: the square marks the manual wrist label, the blue chain is the reconstructed arm, and the cross is the object-derived wrist estimate. The plain-solve and object-assisted distances to the label are 14.0 and 13.7 centimetres. The panels illustrate one retained failure frame.')
rows = [['Arm', 'Method', 'n', 'Median (cm)', 'p95 (cm)', 'Maximum (cm)']]
for side in ('left', 'right'):
    for key, label in (('plain_fk', 'Plain solve'), ('hold_fk', 'Hold-last'), ('recovery_fk', 'Object-assisted')):
        rows.append([side.capitalize(), label, *row_stats(natural['failure'][side][key])])
table(rows, 'Table 7.9. Reconstructed wrist distance to the same manual wrist label on retained natural-occlusion frames, separately for each arm.')
p('Figure 7.9 shows a right-wrist example in which object assistance places the reconstructed wrist closer to the manual wrist label. Together with Figure 7.8, it illustrates why the two arms are reported separately.')
figure('ch7_natural_proxy_right.png', 6.1, 'Figure 7.9. Loop frame 1890, right wrist, with the symbols of Figure 7.8. The object-assisted reconstructed wrist lies 5.0 centimetres from the manual wrist label, against 14.0 centimetres for the plain solve. The cross marks the object-derived estimate, not the final reconstructed wrist. This is one favourable example within the labelled sample.')
p('Object assistance is closer to the right-wrist label on this small labelled set, while the left-wrist comparison shows no corresponding improvement. Label-to-wrist separation can change with arm orientation, clothing or jewellery motion, and depth selection. Repeated-click and depth uncertainties were not quantified. Frames on which the marked feature was completely hidden were excluded, so the results describe selectively visible samples rather than all occlusion frames. These distances do not establish anatomical accuracy. Checking the reconstructed wrist against the hand-object offset used by recovery tests only whether the wrist satisfies the recovery constraint. The offset is not an independent accuracy reference.')

out = HERE / 'Chapter_7_Evaluation.docx'
doc.save(out)
(HERE / 'CH7_RESTRUCTURED.txt').write_text('\n'.join(text))
print('PASS:', out)
print('PASS:', HERE / 'CH7_RESTRUCTURED.txt')
