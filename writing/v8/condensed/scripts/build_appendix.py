#!/usr/bin/env python3
"""Build the condensed appendices.

2026-09-12: Appendix H carries the kinematic solver verification. The
restructure of D-029 removed it from Chapter 7 (it was Section 7.5 with
Table 7.5), which left Chapter 3 pointing at a section that no longer
exists. The author placed the record here so the evidence stays
available without re-entering the evaluation narrative. The prose,
the table and every source comment are the ones from Section 7.5 at
commit 347f873; only the connective sentences changed. No number was
recomputed.

2026-09-12, follow-up to the same restructure: Appendix G.4 no longer
names "the plain, masked and recovery solves of Chapter 7". The rebuilt
Chapter 7 has no masked solve; Section 7.4.1 compares hold-last,
direction memory and object-assisted recovery, and Section 7.4.2
compares the plain solve, hold-last and object-assisted recovery. The
sentence now matches Chapter 2 Section 2.6 and Chapter 5 Section 5.1
word for word. The claim itself is unchanged and was re-verified against
the variant definitions; the comment above the paragraph gives the
sources. No number changed.

2026-09-07: Appendix G documents the approved D2/M06 method; C38
display precision and the remaining filter-event overclaim corrected.
Exact edits and source locations: notes_followup_ch2.md.

2026-09-06 (MATH_LOGIC_REVIEW.md M12, M13, M14 and smaller items):
Appendix F states what zero phase does and does not preserve and what a
median filter does to ramps and spikes; Appendix E.2 calls the 40 degree
switch a heuristic; E.4 quotes the reconstruction to 0.0001; Table E.1
separates range from depth; Appendix D notes that registration is
geometric.

2026-09-07 (supervisor round 5, C48): every displayed number carries
two decimals at most, matrix entries included; E.4 no longer prints the
trace arithmetic (the rounded trace would give 180 degrees, not 177.7);
the C.3 normalized-coordinate interval is stated in words. Full-precision
values stay in the source comments and in appendix_numbers.py.

Eight appendices against the condensed table of contents, lettered in the order
they are first referenced in the body (Chapter 1, Section 1.4 lists them
in this order, and Chapters 2 and 3 reference them in place):

  A  RGB-D Data Acquisition        A.1 procedure, A.2 session file
                                   structure, A.3 replay and inspection
  B  Hardware and Software Platform
  C  Pose Landmark Detection       C.1 BlazePose, C.2 extraction,
                                   C.3 worked example
  D  Depth Sensing and Deprojection D.1 calibration, D.2 alignment,
                                   D.3 pixel-to-world, D.4 worked example
  E  Fiducial Marker Detection and Pose Recovery
                                   E.1 detection, E.2 planar pose and the
                                   two-solution ambiguity, E.3 Rodrigues,
                                   E.4 worked example
  F  Characteristics of the Candidate Filters
  G  Hip Depth Preparation
  H  Kinematic Solver Verification  H.1 synthetic datasets, H.2 measured
                                   input, H.3 scope of the checks

Appendices are method and reference material only: procedures,
definitions, derivations and worked examples. Validation, coverage and
accuracy figures belong to the evaluation chapter and are not repeated
here (skill_set/thesis-structure-rules.md, user rule 2026-08-28).
Appendix H is the author's stated exception to that rule: after D-029
removed solver verification from Chapter 7, the author directed that it
be kept as a verification record in an appendix rather than restored to
the evaluation chapter. It verifies the mathematics of Chapter 3, not
the reconstruction of any recording.

Every number printed here is produced by
writing/v7/scripts/appendix_numbers.py from the frozen artifacts of the
rail recording (recording_20260831_065553: 900 frames, 29.99 s,
640 by 480 at 30 frames per second; rail-only round 2026-09-01, which
replaced the R5 loop recording and its frame 270). The worked-example
frame is 533, the Chapter 3 worked-example frame, which carries all
three marker detections and all eight landmarks. Figures come from
make_appA_figs.py, make_appC_figs.py, make_appD_fig.py,
make_appE_figs.py and make_appF_fig.py.
Citations follow writing/v7/references.md ([1]-[57]).
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches

from eqn import r, nor, sub, sup, frac, mat, d, eqArr, \
    add_display_eq, add_display_math

REPO = Path(__file__).resolve().parents[4]
FIG = REPO / "writing" / "v7" / "figures"

H1, H2, P, IMG, CAP, TBL, EQ = "h1", "h2", "p", "img", "cap", "tbl", "eq"
CODE = "code"   # monospace listing
MATH = "mth"    # displayed, unnumbered worked-example mathematics

# Craig's left superscript names the frame a quantity is expressed in and
# the left subscript the frame described (D-054). An absent script is a
# blank run, not an empty element: an empty one draws a placeholder box in
# some renderers. The same construction is used in build_ch2.py.
BLANK = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'


def pre(base, lsub, lsup):
    """Craig pre-sub-superscript: left superscript over left subscript."""
    # An empty OMML script slot renders as a placeholder box in the PDF
    # render path, so an absent frame carries a blank run instead.
    _BLANK = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'
    return ('<m:sPre><m:sub>' + (lsub or _BLANK) + '</m:sub><m:sup>'
            + (lsup or _BLANK) + '</m:sup><m:e>' + base + '</m:e></m:sPre>')


def inf(ref, sym):
    """A quantity expressed in one frame: left superscript only."""
    return pre(sym, BLANK, nor(ref))


def cR(ref, desc):
    """Class-1 frame-relative proper rotation, both frames named."""
    return pre(r("R"), nor(desc), nor(ref))


# --- equation-fragment shorthands ---
pS = inf("Camera", r("p"))          # A-D02: a deprojected point, camera frame
wS = inf("Camera", r("w"))          # A-E05: the recovered axis, camera frame
RSO = cR("Camera", "Object")        # A-E04: the object marker pose of E.4
fx, fy = sub(r("f"), r("x")), sub(r("f"), r("y"))
cx, cy = sub(r("c"), r("x")), sub(r("c"), r("y"))
wx = sub(d(r("w"), "[", "]"), r("×"))
wx2 = sup(sub(d(r("w"), "[", "]"), r("×")), r("2"))

STEP = "step"

content = [

# ==========================================================================
(H1, "Appendix A: RGB-D Data Acquisition"),

(H2, "A.1 Acquisition Procedure"),

(P, "Prepare and record the session in the following order."),

(STEP, 'Draw the working area on grid paper on the desk. Use the RealSense D435 sensor [10] with the workstation of Appendix B for one continuous take.'),

(STEP, 'Place the rail inside the working area and set it level.'),

(STEP, 'Fix the three markers at the back wall, at the world origin on the desk and on one face of the cube.'),

(STEP, 'Mount the camera on its support at the desk edge, at about chest height. Aim it to keep the subject, desk, rail and all three markers in view.'),

(STEP, 'Request the colour stream in 8-bit-per-channel bgr8 format and the depth stream in 16-bit z16 format. Set both streams to 640 by 480 pixels at 30 frames per second.'),

(STEP, 'Verify the negotiated stream profiles against the requested profiles before keeping any frame. Stop capture with a message naming the stream if they differ. A slower bus can cause the device to select a lower resolution or frame rate without warning.'),

(STEP, 'Start recording with the streams running once the subject is at the desk. Perform the task once: lift the cube from the desk onto the rail, then slide it to the far end. Write everything the device delivers into the single session file described in Section A.2.'),

(STEP, 'Stop recording after the subject steps back. Retain the complete take, including the approach and retreat spans, so later processing stages can segment it from the recorded data.'),

(H2, "A.2 Structure of the Recorded Bag File"),

(P, "The sensor stores each session as a single bag file, a rosbag container in which named topics carry messages held in compressed chunks. The layout separates description from data, as Figure A.1 draws and Table A.1 lists. The file stores its version, the device and sensor descriptions, the value of every sensor option at record time, and the calibration of each stream once at the start. The camera intrinsics of Appendix D come from that calibration. The rest is a sequence of per-frame messages, an image message and a metadata message per captured frame."),

(IMG, FIG / "appA_fig_bag.png", 6.3),
(CAP, "Figure A.1. Structure of the recorded session file: the one-time description topics on the left, the per-frame image and metadata messages on the right."),

(P, "Replay reads the same structure back, so a recording carries its own calibration and needs no separate calibration file."),

(TBL, [["Property", "Value, read from the recording"],
       ["container", "rosbag 2.0, compressed chunks"],
       ["streams", "depth (z16, 16-bit) and colour (bgr8, 24-bit)"],
       ["resolution and rate", "640 by 480 at 30 frames per second, both streams"],
       ["duration", "29.99 seconds, first to last frame"],
       ["frames delivered on replay", "900 depth and 900 colour"],
       ["raw pixel payload", "1.38 GB"],
       ["file size on disk", "0.82 GB"],
       ["one-time topics", "file version, device and sensor descriptions, sensor options, per-stream intrinsics and pose"],
       ["per-frame topics", "image data and frame metadata, per stream"]]),
(CAP, "Table A.1. Contents of the recording's session file."),

(H2, "A.3 Replay and Visual Inspection"),

(P, "The vendor's viewer application opens the same session file and replays it as though a live sensor were attached, which makes it the quickest way to audit a recording before any processing runs. Figure A.2 places the camera's own view of one moment beside the point cloud the viewer deprojects from it."),

(IMG, FIG / "appA_fig_cloud.png", 6.3),
(CAP, "Figure A.2. (a) The recorded colour frame of the worked-example moment. (b) The same moment as a metric point cloud with the colour stream mapped onto it. (c) The same cloud coloured by distance."),

(P, "A second, automated pass reads the raw outputs of the two extractors and reports what the recording contains before any modelling starts. It lists the frames in which the subject is present, the segmentation into approach, manipulation and retreat, the frames on which each marker and each landmark is available, and an inventory of the occlusion episodes. The pass reports these conditions without rejecting the recording, with a single exception: because the desk marker defines the world origin, a frame on which it is missing stamps the recording as rejected and stops every stage after it."),

# ==========================================================================
(H1, "Appendix B: Hardware and Software Platform"),

(P, "Every processing result and every timing in this thesis was produced on the one workstation of Table B.1, and the cost measurements of Chapter 8 should be read against this configuration. The division of labour is fixed throughout: the pose detector's neural network inference runs on the graphics card through the framework's GPU delegate. Everything else, including the depth sampling, the marker pose recovery, the kinematic mathematics and the filtering, runs on the processor in a single thread."),

(TBL, [["Component", "Value"],
       ["processor", "AMD Ryzen 9 7950X, 16 cores, 32 threads"],
       ["memory", "64 gigabytes (GB)"],
       ["graphics card", "NVIDIA GeForce RTX 3080, 10 GB"],
       ["sensor", "Intel RealSense D435 RGB-D camera [10]"],
       ["operating system", "Ubuntu 24.04 long-term support (LTS)"],
       ["Python", "3.11.9"],
       ["RealSense library", "2.55.1"],
       ["MediaPipe", "0.10.15"],
       ["OpenCV", "4.10.0"],
       ["NumPy", "1.26.4"],
       ["SciPy", "1.13.1"],
       ["Unity editor", "6000.3.19f1"]]),
(CAP, "Table B.1. The workstation and the software versions behind every result in this thesis."),

# ==========================================================================
(H1, "Appendix C: Pose Landmark Detection"),

(H2, "C.1 MediaPipe BlazePose Overview"),

(P, "The detector this work uses is the BlazePose model [43], deployed through the MediaPipe framework [28], whose landmark network regresses 33 keypoints inside the region of interest that a person detector localizes. For each of them it returns normalized image coordinates, a depth estimate inferred from appearance, and a visibility score between 0 and 1 for its confidence that the landmark is present and unoccluded. The pipeline runs the heavy variant in video mode, which tracks the person across frames rather than detecting from scratch on each one. It expects a single person and leaves the detection, presence and tracking confidence thresholds at 0.30. The appearance-derived depth is not used, because it is unreliable during fast motion and object interaction [3], [13]. Metric depth comes instead from the aligned depth image of Appendix D, which is a measurement rather than an inference."),

(H2, "C.2 Landmark Coordinate Extraction"),

(P, "Of the 33 landmarks the model returns, this thesis uses the eight of Table C.1, covering the torso and the two arms."),

(TBL, [["ID", "Landmark", "ID", "Landmark"],
       ["L11", "left shoulder", "L12", "right shoulder"],
       ["L13", "left elbow", "L14", "right elbow"],
       ["L15", "left wrist", "L16", "right wrist"],
       ["L23", "left hip", "L24", "right hip"]]),
(CAP, "Table C.1. The eight MediaPipe landmarks used in this work, with the model's indices."),

(P, "The extractor takes three steps for each landmark on each frame. The normalized coordinates are scaled to the colour image and truncated to whole pixels, so the pixel u is the integer part of the horizontal coordinate times 640 and v the integer part of the vertical coordinate times 480. The aligned depth image is sampled there by the 5 by 5 median of Section D.3, and the pixel with its depth is back-projected by equation (D.2). Figures C.1 and C.2 show the eight landmarks of the worked-example frame, at the sampled pixels and as position vectors from the sensor origin."),

(IMG, FIG / "appC_fig_landmarks.png", 5.2),
(CAP, "Figure C.1. The eight upper-body landmarks on the worked-example frame, drawn at the pixels the extractor sampled, with the segments between them."),

(IMG, FIG / "appC_fig_vectors.png", 6.0),
(CAP, "Figure C.2. The same eight landmarks as metric position vectors from the sensor origin, with the camera axes drawn at the origin."),

(P, "The model returns coordinates for all 33 landmarks whether it can see them or not, so an occluded wrist still receives a position. The visibility score is then the only signal that separates a measurement from a guess. A landmark whose visibility falls below 0.50 is therefore written as an empty record, and the same applies when the 5 by 5 depth window holds no valid return. Every exported landmark carries a code recording which of three cases produced it: a measured sample, a sample blocked by low visibility, or a sample blocked by missing depth. Later stages can therefore tell a gap from a value."),

(H2, "C.3 Worked Example"),

(P, "Take the right wrist, landmark L16, on frame 533 of the recording, the frame of Figure C.1. The detector reports it with a visibility of 0.58, above the 0.50 gate. The extractor stores the sampled pixel, not the normalized pair. Multiplying the normalized coordinates by 640 and 480 and truncating gives the pixel (209, 308)."),  # src: appendix_numbers.py, vis 0.5778; x_n in [0.3266, 0.3281), y_n in [0.6417, 0.6438)
(P, "The 5 by 5 median depth at pixel (209, 308) is 1.06 metres. Back-projection through equation (D.2), with the intrinsics of Table D.1, gives"),
(MATH, eqArr(  # src: appendix_numbers.py, x = -0.200529, y = +0.104745 at depth 1.060
    r("x") + r(" = 1.06 · ") + frac(r("209 − 323.94"), r("607.56")) + r(" = −0.20"),
    r("y") + r(" = 1.06 · ") + frac(r("308 − 248.02"), r("607.02")) + r(" = +0.10"))),
(P, "so the right wrist sits at (−0.20, 0.10, 1.06) metres in the camera frame. Table C.2 carries the same three steps through the other two landmarks of the right arm, and Section 3.5 continues from the filtered form of these numbers into the kinematic solve."),

(TBL, [["Landmark", "Visibility", "Pixel", "Depth (m)", "Camera-frame position (m)"],
       # src: appendix_numbers.py; vis 0.9987/0.7031/0.5778, depth 1.311/1.196/1.060,
       #      xyz (-0.163859, -0.334799, 1.311) (-0.200667, -0.045351, 1.196) (-0.200529, 0.104745, 1.060)
       ["L12 right shoulder", "1.00", "(248, 93)", "1.31", "(−0.16, −0.33, 1.31)"],
       ["L14 right elbow", "0.70", "(222, 225)", "1.20", "(−0.20, −0.05, 1.20)"],
       ["L16 right wrist", "0.58", "(209, 308)", "1.06", "(−0.20, 0.10, 1.06)"]]),
(CAP, "Table C.2. The right-arm landmarks of the worked-example frame, from the visibility score through the sampled pixel and depth to the back-projected position. These are raw extractor positions, before the filtering of Section 2.5; the filtered values are in Table 3.1."),

# ==========================================================================
(H1, "Appendix D: Depth Sensing and Deprojection"),

(P, "The Intel RealSense D435 is a stereo-depth camera that produces a per-pixel depth measurement alongside a colour image [10]. Depth comes from triangulating between two infrared imagers separated by a known baseline, assisted by an infrared projector that adds texture to surfaces too uniform for stereo matching. The scene spans depths from about 0.53 metres at the desk marker to about 3.2 metres at the wall marker, inside the sensor's usable range. Stereo depth noise grows roughly with the square of the distance [37], [38]."),

(H2, "D.1 RGB Camera Calibration"),

(P, "Calibration establishes the relationship between a 3D point in the scene and the pixel it lands on. The pinhole model describes it through four numbers: the focal lengths in pixels along the two image axes, and the principal point, the pixel where the optical axis meets the image [51]. A point p with coordinates x, y and z in the camera frame projects, as shown in Figure D.1(a), to the pixel"),
(EQ, r("u") + r(" = ") + fx + r(" ") + frac(r("x"), r("z")) + r(" + ") + cx + r(",        ") +
     r("v") + r(" = ") + fy + r(" ") + frac(r("y"), r("z")) + r(" + ") + cy, "D.1"),

(P, "This work performs no manual calibration. The vendor's library exposes the factory calibration held in each device, and the recording carries those values in its own calibration topic, so the intrinsics used are the ones the sensor reports for the colour stream at 640 by 480, listed in Table D.1."),

(TBL, [["Parameter", "Value", "Meaning"],
       # src: calibration topic, fx 607.56128 fy 607.01508 cx 323.93756 cy 248.01741
       ["fx", "607.56 px", "focal length along the horizontal image axis"],
       ["fy", "607.02 px", "focal length along the vertical image axis"],
       ["cx", "323.94 px", "principal point, horizontal (image centre 320)"],
       ["cy", "248.02 px", "principal point, vertical (image centre 240)"],
       ["distortion model", "inverse Brown-Conrady", "the model the stream declares"],
       ["distortion coefficients", "all zero", "the recorded values for this stream"]]),
(CAP, "Table D.1. Factory colour-camera intrinsics, read from the calibration topic of the recording."),

(P, "The table connects back to the model in three ways. The two focal lengths differ by less than 0.1 percent, so the pixels are square to that precision. The principal point sits a few pixels from the geometric image centre, so the deprojection of Section D.3 measures every pixel offset from that point rather than from the centre. The distortion coefficients are all zero, so the vendor's general deprojection routine reduces to the closed form of equation (D.2) for this data. Ignoring the small principal-point offset, the focal lengths and image size give a nominal field of view of 55.6 by 43.1 degrees, which guided the camera placement in Chapter 2."),

(H2, "D.2 Depth-to-RGB Alignment"),

(P, "The depth imager and the colour camera sit at a small offset inside the sensor housing, and each has its own intrinsics [10], so a raw depth pixel does not observe the same scene point as the colour pixel at the same location. The alignment step drawn in Figure D.1(b) resolves this. Each depth pixel is lifted to a 3D point through the depth intrinsics, moved into the colour camera's frame through the known transformation between the two imagers, and projected back onto the colour image grid through the colour intrinsics of Table D.1. That transformation is held inside the vendor library and is never applied by this pipeline, so only the aligned colour optical frame, the camera frame, is used downstream. The registration is geometric: a colour pixel and its aligned depth now share one ray, but the depth belongs to the nearest surface along that ray. A landmark on a body part behind the cube therefore takes the cube's depth, the case Section 5.2 deals with."),

(IMG, FIG / "appD_fig_alignment.png", 6.3),
(CAP, "Figure D.1. (a) The pinhole geometry the intrinsics describe: a scene point, the ray to the camera centre, and the pixel where it crosses the image plane. (b) Depth-to-colour alignment: the depth image is reprojected into the colour camera's geometry so that both buffers share one pixel grid."),

(P, "Both buffers are 640 by 480, so the aligned depth image shares the pixel grid of the colour image and the two values at one pixel describe the same physical point. That aligned pair is the only input the rest of the pipeline sees, and both detectors work in colour pixel coordinates that index the depth image directly."),

(H2, "D.3 Pixel-to-Camera Mapping"),

(P, "A pixel with a valid depth is lifted to a 3D point in the camera frame by running the projection backwards:"),
(EQ, r("x") + r(" = ") + r("z") + r(" ") + frac(r("u") + r(" − ") + cx, fx) + r(",      ") +
     r("y") + r(" = ") + r("z") + r(" ") + frac(r("v") + r(" − ") + cy, fy) + r(",      ") +
     pS + r(" = ") + d(r("x, y, z")), "D.2"),
(P, "The camera frame is right-handed, with x to image right, y downward and z forward along the optical axis, and the depth measurement is the z coordinate itself. The depth image stores integer counts, one count being one millimetre for this device. The fixed y-axis flip of Section 3.1 expresses the same point in the left-handed y-up camera frame {Camera'}, used for body kinematics. Its origin is still the camera origin; it is not person-centred. Chapter 6 gives the further mapping into Unity."),

(P, 'Reading depth at a single pixel fails in two ways. The pixel can give a zero return, which is common at object edges and on dark or reflective surfaces, or it can carry single-pixel speckle. The pipeline therefore takes the median of the nonzero depths in a 5 by 5 window centred on the target pixel, counting a zero return as an absent sample rather than as zero metres. When the window holds no valid return, the sample is reported as missing.'),

(H2, "D.4 Worked Example"),

(P, "Take the object marker on frame 533, the frame of Appendix C. The marker detector of Appendix E places its centre at pixel (200, 341), where the 5 by 5 median depth is 0.99 metres, and equation (D.2) with the intrinsics of Table D.1 gives"),
(MATH, eqArr(  # src: appendix_numbers.py, depth 0.986, x = -0.201136, y = +0.151036
    r("x") + r(" = 0.99 · ") + frac(r("200 − 323.94"), r("607.56")) + r(" = −0.20"),
    r("y") + r(" = 0.99 · ") + frac(r("341 − 248.02"), r("607.02")) + r(" = +0.15"))),
(P, "so the marker centre sits at (−0.20, 0.15, 0.99) metres in the camera frame. Running the result forward through equation (D.1) returns the pixel the detection started from, so the two equations are inverses."),

# ==========================================================================
(H1, "Appendix E: Fiducial Marker Detection and Pose Recovery"),

(P, "The landmark detections of Appendix C measure the subject. The objectives of Chapter 1 also require a measurement of the manipulated object in a frame that does not move with the camera, and a scene geometry against which the reconstruction can be read. Fiducial markers supply both, because they are printed patterns whose geometry is known exactly, so their pose follows from a single image with no training data and no appearance model. Section 2.4 gives the three markers and their roles. This work uses ArUco markers because they detect reliably, cost little and encode their own identity [29]; the size and distance relationship of [30] set their printed sizes."),

(H2, "E.1 Detection and Sub-Pixel Corners"),

(P, "Detection runs in four stages. Adaptive thresholding turns the image into binary regions by comparing each pixel with the mean of its own neighbourhood, which keeps the detector working under uneven lighting. Contour extraction finds candidate quadrilaterals. The interior of each candidate is perspective-corrected and matched against the dictionary, which either rejects it or returns a marker identity, and the four corner positions are refined below the pixel grid. Figure E.1 shows the result on the worked-example frame."),

(IMG, FIG / "appE_fig_detect.png", 5.6),
(CAP, "Figure E.1. Marker detection on the worked-example frame. (a) The full frame with the three detected markers outlined and the recovered axes at each centre. (b) to (d) Each marker enlarged, with its four refined corners numbered in detection order."),

(P, "Detection itself is standard and is not a contribution of this work [29]. The corner refinement that follows it is a design decision, because corner precision limits pose precision. With the detector's default settings the coordinates land on whole pixels, so a marker that sits still is reported at identical corners frame after frame, which is the detector returning to the same pixel grid, not a stable measurement. The pipeline therefore enables the refinement based on the AprilTag method [48], which takes each corner as the intersection of straight lines fitted to the two marker edges that meet there. Every frame of the recording uses the refined corners."),

# The clause "and Section 9.1 uses the pair" stood here until 2026-09-12.
# The author's instruction of that date dropped the depth-implied marker
# size diagnostic from the thesis, so Section 9.1 no longer sets this
# depth sample against the marker-geometry translation and the
# justification was false. The sentence now describes what the branch
# records and where the appendix lists it, and claims no use the thesis
# does not contain. The author settled its status on 2026-09-12 (D-050):
# the field stays in Appendix E.1 as part of the recorded-data description,
# with documentary status only. It was stored by the pipeline, no analysis
# or result in this thesis uses it, and it must NOT be cited as evidence
# for any claim in Chapters 7 to 10. Do not remove it merely because it has
# no downstream consumer, unless Appendix E is later redefined to document
# only analysed quantities.
(P, "The marker branch also records, at detection time, the median of a 5 by 5 patch of aligned depth values at each marker's centre pixel. The pose recovery never reads it. The value comes from the other measurement principle, the stereo depth stream rather than the marker geometry, and it remains in the recorded data, although no analysis of this thesis reads it. Table E.1 lists it beside the recovered translation for the worked-example frame. A frame on which a marker is not detected stays an empty record."),

(H2, "E.2 Planar Pose Recovery and the Two-Solution Ambiguity"),

(P, "A detected marker gives four corner points in the image. The printed marker gives the same four corners in its own frame, at plus and minus half the printed side length L in the marker plane where the third coordinate is zero, in the order the detector returns them. Recovering the pose means finding the orientation of the marker's own frame relative to the camera frame, and the position t of the marker origin in the camera frame. Together they carry these four known corners onto the four detected pixels through the projection of equation (D.1). The marker's own frame follows the rule of Section 2.2.1, so the pose recovered here is that of the world frame for the desk marker, of the wall marker frame, or of the object frame, each relative to the camera frame. The corners carry noise, so the solver finds the orientation and position that minimize the summed squared image distance between the detected and projected corners. The four corners provide eight equations for the six unknowns."),

(P, "A planar target has a structural weakness that a target with depth does not have. Because the four corners lie in one plane, a plane seen nearly face-on can be tilted toward the camera or away from it by the same small angle and produce the same four pixels to within the corner noise. The two readings are mirror images of the marker normal about the camera ray. The solver used here, the infinitesimal plane-based method [52], returns both locally optimal poses, called lobes below, rather than choosing one of them. Figure E.2 draws the geometry and the pair returned for the wall marker on the worked-example frame. The branch records its near-degenerate flag there, as on almost every frame. The two normals differ by 36.4 degrees, and both readings reproject with a largest corner gap of 0.22 pixels on sides averaging 30.03 pixels."),

(IMG, FIG / "appE_fig_lobes.png", 6.3),
(CAP, "Figure E.2. (a) A nearly face-on plane admits two pose readings, tilted toward the camera or away from it, that are mirror images about the camera ray. (b) Both readings the solver returns for the wall marker, reprojected over the detection they come from, with the corner where they separate most magnified in the inset."),

(P, "The pipeline selects its policy by the angular separation of the two lobes, not by how well either lobe reprojects. That choice is a heuristic of the implementation. A marker close to the camera shows visible perspective across its own width. The lobes then sit far apart, and the wrong one places its corners visibly away from the detection, so the closer-fitting lobe is taken. A marker whose projection is nearly an affine copy of the printed square gives almost no perspective cue. The two lobes can then stay far apart in angle while their reprojections coincide, since a plane tilted toward the camera and one tilted away by the same angle project to the same affine image. Angular separation therefore does not say how well the two readings can be told apart on the image. The threshold between the two cases is a lobe separation of 40 degrees, and the pair of Figure E.2 sits just below it, so the discriminators below decide instead."),

(P, "Below that separation the discriminator becomes physical and temporal instead. On the first such frame the plumb wall decides: the lobe whose in-plane up axis better agrees with an estimate of the gravity direction in the camera frame is taken. That estimate is seeded from the plane fitted to the depth returns around the desk marker and needs to be right only to within about 10 degrees. On every later frame the lobe closest to a reference pose is kept, and the reference updates only when the chosen pose lies within 30 degrees of it, so one corrupted detection cannot corrupt the frames that follow. After five consecutive frames far from the reference the tracker re-seeds. A per-frame flag records how each pose was resolved."),

(H2, "E.3 From the Rodrigues Form to a Rotation Matrix"),

(P, "The pose solver returns its rotation in the compact Rodrigues form, a single 3-vector whose direction is the rotation axis and whose length is the rotation angle in radians. The rest of the system works in rotation matrices, so the vector is converted on receipt, and the same formula reappears in the object-track cleaning of Chapter 4."),

(P, "Write w for the unit axis and theta for the angle, so the Rodrigues vector is theta times w. The skew-symmetric matrix of w, written with the cross subscript, has rows (0, -w3, w2), (w3, 0, -w1) and (-w2, w1, 0), and multiplying it by a vector v gives w cross v. The conversion is Rodrigues' rotation formula [42]:"),
(EQ, r("R") + d(r("w, θ")) + r(" = ") + r("I") + r(" + ") + nor("sin") + d(r("θ")) + r(" ") + wx +
     r(" + ") + d(r("1 − ") + nor("cos") + d(r("θ"))) + r(" ") + wx2, "E.1"),
(P, "Equation (E.1) is a rotation operator. It builds a rotation from an axis and an angle, so it carries no frame labels, and a frame is named only where the rotation it returns is used. The formula holds for every angle and produces a proper rotation matrix by construction. Its inverse, recovering the axis and the angle from a matrix, is the matrix logarithm the gap bridging of Chapter 4 uses, so the pipeline moves between the two forms without introducing quaternions."),

(H2, "E.4 Worked Example: Frame 533 Detections"),

(P, "Frame 533 of the recording, shown in Figure E.1, carries detections of all three markers. Table E.1 collects what the marker branch records for each of them: the printed side length, the centre pixel, the recovered marker origin in the camera frame, the range it implies, and the independent 5 by 5 median depth at the same pixel. The wall marker is printed at 150 millimetres because it sits farthest from the camera, and the object marker occupies one face of the 70 millimetre cube. The desk and object translations listed rest on the 45 millimetre size of Table 2.1, and the wall translation on the 150 millimetre size of the same table."),

# src: frame 533 of eval/output/recording_20260831_065553_aruco_raw_scaled.csv
#      (the scaled track every later stage reads); ranges are the norms of
#      the tabulated translations
(TBL, [["Marker", "Role", "Black square (mm)", "Centre pixel", "Marker origin t, camera frame (m)", "Range (m)", "Depth sample (m)"],
       #      wall t (-1.030905, -0.470312, 3.052835) range 3.2563 depth 3.185;
       #      desk t (0.023456, 0.187143, 0.551804) range 0.5831 depth 0.532;
       #      object t (-0.209474, 0.155957, 1.022892) range 1.0557 depth 0.986
       ["id 0", "wall", "150", "(119, 155)", "(−1.03, −0.47, 3.05)", "3.26", "3.19"],
       ["id 2", "desk", "45", "(350, 455)", "(0.02, 0.19, 0.55)", "0.58", "0.53"],
       ["id 1", "object", "45", "(200, 341)", "(−0.21, 0.16, 1.02)", "1.06", "0.99"]]),
(CAP, "Table E.1. The three marker detections of the worked-example frame, in the order of the panels of Figure E.1. The marker origin and the range come from the corner geometry alone and carry the printed-size scaling. The depth sample comes from the stereo depth stream and does not enter the pose recovery. The range is the length of that position vector, and the depth sample is the coordinate along the optical axis, so the two columns are different quantities."),

(P, "Take the object marker through the rotation conversion. The pose recovery returns the rotation of the object frame relative to the camera frame"),
# src: appendix_numbers.py, R = [[0.998259, -0.008185, 0.058415], [-0.005833,
#      -0.999170, -0.040308], [0.058697, 0.039897, -0.997478]], trace -0.998390,
#      angle 177.7007 deg, axis (0.9996, -0.0035, 0.0293), sin 0.0401, 1-cos 1.9992
(MATH, RSO + r(" = ") + mat([
    [r("1.00"), r("−0.01"), r("0.06")],
    [r("−0.01"), r("−1.00"), r("−0.04")],
    [r("0.06"), r("0.04"), r("−1.00")]])),
(P, "The entries are printed to two decimals and the angle to a tenth of a degree. The angle follows from the trace:"),
(MATH, r("θ") + r(" = ") + nor("acos") + d(frac(nor("tr") + d(RSO) + r(" − 1"), r("2"))) + r(" = 177.7°")),
(P, "and the axis from the skew-symmetric part, each component being one off-diagonal difference divided by twice the sine of the angle:"),
(MATH, wS + r(" = (1.00, 0.00, 0.03)")),
(P, "The marked face of the cube is turned almost 180 degrees about an axis close to the camera's x axis, the cube sitting on the rail with its face square to the camera. The skew-symmetric matrix of those components is"),
(MATH, wx + r(" = ") + mat([
    [r("0.00"), r("−0.03"), r("0.00")],
    [r("0.03"), r("0.00"), r("−1.00")],
    [r("0.00"), r("1.00"), r("0.00")]])),
(P, "and its square is"),
(MATH, wx2 + r(" = ") + mat([
    [r("0.00"), r("0.00"), r("0.03")],
    [r("0.00"), r("−1.00"), r("0.00")],
    [r("0.03"), r("0.00"), r("−1.00")]])),
(P, "The two scalar factors of equation (E.1) at this angle are the sine, 0.04, and one minus the cosine, 2.00, so the three terms assemble to"),
(MATH, RSO + r(" = ") + r("I") + r(" + 0.04 ") + wx + r(" + 2.00 ") + wx2 + r(" = ") + mat([  # rebuilt matrix rounded from full precision (largest difference from R 2.27e-07)
    [r("1.00"), r("−0.01"), r("0.06")],
    [r("−0.01"), r("−1.00"), r("−0.04")],
    [r("0.06"), r("0.04"), r("−1.00")]])),
(P, "The recomposed matrix agrees with R to the displayed rounding. This is a numerical consistency check, not a measurement of pose accuracy. The third column is the object frame's own outward normal, its third axis expressed in the camera frame, (0.06, −0.04, −1.00), pointing almost straight back at the camera. That pose is the input the world anchoring of Chapter 4 carries forward."),

# ==========================================================================
(H1, "Appendix F: Characteristics of the Candidate Filters"),

(P, "The filtering stage of Section 2.5 can select any of four smoothers. This appendix describes what each smoother computes and how it behaves. This provides a common basis for the choice in Chapter 2 and the real-time comparison in Chapter 8. Every parameter quoted is the one the filter run of the recording used, at the recorded rate of 29.98 samples per second, the nominal 30 frames per second."),

(P, "Section 2.5 gives the two stages ahead of the smoothers. The rolling Hampel test comes first, with an implemented robust scale based on rolling absolute residuals [53], [54]. Each sample is first compared with its own seven-frame rolling median; a second rolling median of those absolute residuals is multiplied by about 1.48. This differs from taking every deviation about the current window's single median. Linear interpolation then fills gaps of up to 5 frames. Their rules matter here in two respects: leading and trailing gaps are not backfilled, and each smoother runs on each unbroken run of valid samples separately. No filter therefore smooths across a gap it cannot see into."),

(P, "The zero-phase Butterworth low-pass is the candidate the results use, chosen for a passband that is flat rather than rippled, so the motion band passes without ripple [55], [56]. A fourth-order filter with a 3 hertz cutoff runs twice over each run, forward and backward [57], and the second pass cancels the phase shift of the first. Cancelling phase delay does not undo attenuation: the two passes multiply their amplitude gains, so a movement with components near 3 hertz comes out smaller. A peak whose components are attenuated unevenly can shift by a sample. Running a filter backward needs the samples that follow the current one, so this candidate exists only offline."),

(P, "The Savitzky-Golay smoother fits a second-order polynomial by least squares to a sliding window of 9 samples and takes its value at the centre [46]. Fitting a curve rather than averaging preserves a peak that a moving average would flatten, which Figure F.1(a) shows, but the fit gives every sample in the window a weight, so a single far-off sample still pulls the result."),

(P, "The rolling median replaces each sample by the median of a 9-sample window. Because the median takes the middle value of the window, an isolated extreme sample is discarded rather than averaged in, although it can still move the median by one rank. A steady ramp passes through a centred window unchanged, since the middle value of a monotone run is the sample at its centre. The cost falls on a peak, where the median takes a sample below the top and flattens it."),

(P, 'The One Euro filter is the causal candidate, the one the real-time pipeline of Chapter 8 can run, because it needs only the samples already received [40]. It is a first-order low-pass whose cutoff is not fixed. The filter keeps a smoothed estimate of the rate of change of the signal and raises its cutoff in proportion to it, through a minimum cutoff of 0.05 hertz and a coefficient of 1.0. Jitter is therefore suppressed while the signal holds still, and the filter follows while it changes. Being causal, it cannot avoid lag, which Figure F.1(a) shows as the late arrival of its output at the end of the move.'),


(IMG, FIG / "appF_fig_filters.png", 6.3),
(CAP, "Figure F.1. (a) The four candidates on one constructed test signal, not recorded data: a profile that holds still, moves and holds again, with broadband noise and one single-frame spike added. (b) The amplitude gain of the two linear candidates against frequency, including both passes of the Butterworth filter, with the 3 hertz cutoff marked and the band of deliberate upper-body motion shaded."),

(P, "Table F.1 compares phase and response to spikes. Centred offline filtering avoids a systematic causal delay, but smoothing can change peaks and event boundaries. A spike missed by the Hampel stage is rejected by the median or spread across neighbouring samples by the other candidates, so despiking runs first. Offline availability and lag determine which smoother the pipeline can use."),

(TBL, [["Candidate", "What it computes", "Phase", "Single-frame spike", "Parameters used"],
       ["Butterworth, zero phase", "low-pass, applied forward and backward", "none", "spread over the filter length", "order 4, cutoff 3 Hz"],
       ["Savitzky-Golay", "least-squares polynomial over a sliding window", "none", "spread over the window", "window 9, order 2"],
       ["rolling median", "median of a sliding window", "none", "removed", "window 9"],
       ["One Euro", "first-order low-pass with a rate-dependent cutoff", "lag", "partly passed through", "minimum cutoff 0.05 Hz, coefficient 1.0"]]),
(CAP, "Table F.1. The four candidate smoothers, with the parameters used in Figure F.1(a). The Butterworth, Savitzky-Golay and rolling median smoothers need the whole recording or a window around each sample; only the One Euro filter runs on the samples already received."),




# ==========================================================================
# D2/M06: describes the frozen implementation without changing its policy.
# Sources: v1/kinematics/occlusion_ext.py:219-299,397-674,735-750;
# eval/failure/recovery_core.py:181-234; v2/person/v2_person.py:201-202.
(H1, "Appendix G: Hip Depth Preparation"),

(P, "Section 2.6 introduces the preparation used before the body solve. This appendix gives its rules for the hips and shoulders, whose depth can come from an object in front of the body even when the landmark pixel remains usable. The rules operate in the y-up camera frame {Camera'}; flipping the camera y axis leaves the camera-depth coordinate z unchanged."),

(H2, "G.1 Inputs and Initial State"),

(P, "The state contains calibrated hip and shoulder widths, distances from each torso landmark to the opposite pair's midpoint, and remembered pair directions and depths. Offline evaluation supplies the arm segment lengths and the shoulder width. Every remaining length is fixed at the median of its first 60 available observations, offline and live alike. These lengths include the hip width and the four distances to the opposite pair midpoint. Direction and depth memories start empty. Accepted pair directions are normalized after blending 70 percent of the previous direction with 30 percent of the current one."),

(P, "For a point p with camera-depth coordinate p_z greater than 5 centimetres, its retained ray is p divided by p_z. If z_star is the chosen replacement depth, the repaired point p_star is"),
(EQ, sub(r("p"), nor("star")) + r(" = ") + sub(r("z"), nor("star")) + r(" ") + frac(r("p"), sub(r("p"), r("z"))), "G.1"),
(P, "This scaling preserves the direction from the camera. A missing point or unusable ray cannot receive this repair."),

(H2, "G.2 Rejection Order"),

(P, "The preparation first checks each pair's width. A deviation greater than 30 percent from its calibrated width triggers an endpoint check. Each endpoint's distance to the opposite pair midpoint is compared with its own calibrated distance, using the same tolerance. If exactly one endpoint fails, that endpoint is rejected and its ray retained. An ambiguous width violation alone leaves both endpoints in place."),

(P, "The next check compares the directed hip and shoulder lines when both remain available. With both direction memories seeded, disagreement above 20 degrees activates the check. The pair farther from its remembered direction is selected, with a tie assigned to the hips. That pair stays selected until disagreement falls to 10 degrees or below. The preparation applies the endpoint-distance check described for width deviations and rejects a single point when it can. Otherwise, a known pair width permits rejection of both endpoints, and the remembered pair direction is moved toward the surviving line. The solver's 30 percent width tolerance and 20 degree line trigger are not the thresholds of the offline detectors of Section 5.2, whose torso width ratio detector fires at 20 percent and whose torso line angle detector fires above 22 degrees."),

(P, "The depth check next rejects each remaining torso point whose depth differs from its memory by more than 10 centimetres. If both endpoints fail, the width and direction are known, and no whole-pair repair has already been scheduled, the pair is rebuilt together. Otherwise, each rejected endpoint is considered separately."),

(P, "The final check rejects a remaining hip more than 15 centimetres nearer the camera than the input shoulder midpoint. That midpoint is saved before any checks and is not recomputed after shoulder rejection. This check can operate before depth memory exists, but assumes the input shoulders have usable depths and the trunk does not lean substantially. The earlier depth checks take priority because this check only sees hips they left available."),

(H2, "G.3 Memory Updates and Replacement"),

(P, "After rejection and before replacement, each surviving torso sample updates its depth memory. Its first accepted depth seeds the memory; later updates blend 70 percent previous depth with 30 percent new depth. Repaired points never update this memory. It has no expiry, so a persistent rejected interval can continue to use an old accepted depth."),

(P, "A single point rejected by the width, line or depth checks uses its remembered depth in equation (G.1). A hip rejected by the last check instead uses the saved shoulder-midpoint depth. Neither individual placement enforces pair width. The latter placement assumes an approximately upright trunk in camera depth."),

(P, "For a scheduled whole-pair repair, both rays are first placed at their remembered depths. Their average gives the repaired midpoint. Their line direction is blended into the remembered pair direction with the same 70/30 weighting and normalized; following a line-angle rejection, this is an additional update after steering toward the other pair. The endpoints are then placed half the calibrated width on either side of that midpoint. The final endpoints therefore need not lie on their original rays. If either ray or depth memory is unavailable, or the rebuilt pair has zero width, the fallback keeps the input midpoint and rebuilds around it using the remembered direction and calibrated width. That fallback cannot remove a common depth error in the midpoint."),

(P, "If one endpoint is still missing, a separate fallback places it one calibrated width from its available partner along the remembered pair direction. This fallback is limited to 45 consecutive unmeasured frames for that endpoint. The limit does not expire depth memory or the whole-pair repair. When the required geometry remains unavailable, the dependent joint group keeps its last solved value."),

(H2, "G.4 Output States and Scope"),

# Method names follow the rebuilt Chapter 7 (D-029) and match the wording
# of Chapter 2 Section 2.6 and Chapter 5 Section 5.1 word for word:
# Section 7.4.1 compares hold-last, direction memory and object-assisted
# recovery, and Section 7.4.2 compares the plain solve, hold-last and
# object-assisted recovery. Chapter 7 has no "masked solve" name any more.
# The claim that hold-last does not rest on this preparation was
# re-verified against the variant definitions: the solving variants run
# RobustChainSolver, which carries the preparation, while hold-last runs
# ChainFallbackSolver, which has none (eval/failure/run_recovery.py variant
# list; eval/failure/recovery_core.py lines 193 and 227-230;
# eval/failure/compare_hip_hold.py, "the hold-last solve, no preparation").
(P, "A successfully solved joint group that consumes a repaired point is tagged constrained and loses its live bit (Section 6.1). This includes the root when a hip, or the shoulder used as its tilt reference, was repaired. A group that cannot be solved is held. The plain, direction-memory and object-assisted solves of Chapter 7 all rest on this preparation. The hold-last solve keeps the last accepted joint angles and does not use it. Offline pelvis translation uses the repaired hip midpoint. The live path runs the same preparation before its solve but still publishes the original hip midpoint as the pelvis translation."),

(P, "The preparation depends on correct landmark rays and accepted calibration samples. A desk-depth sample can seed memory if it passes the checks, and genuine lean or body translation can trigger a depth rejection. Starting with hidden hips is not evaluated: the shoulder-depth check is available, but success still depends on usable hip pixels and shoulder depths. Section 9.4 discusses the observed effect without treating it as independent anatomical validation."),

# ==========================================================================
# Moved here from Chapter 7, Section 7.5 at commit 347f873, after D-029
# removed solver verification from the evaluation chapter. Prose, table
# and source comments are the committed ones; only the framing sentences
# were changed and the chapter roadmap sentence dropped. Old Table 7.5 is
# Table H.1 here. The reference to Section 7.4 still resolves: the
# rebuilt Chapter 7 grades recovery in Section 7.4.
(H1, "Appendix H: Kinematic Solver Verification"),

(P, "Chapter 3 derives the solve that turns metric landmarks into a root pose and four angles for each arm. This appendix is the record of the checks made on that mathematics, first against injected angle trajectories and then on measured input. It verifies the equations of Chapter 3 and nothing else, so it sits here rather than among the experimental results of Chapter 7."),

(H2, "H.1 Synthetic Datasets"),

(P, "Before any recording ran through the kinematic solver of Chapter 3, the six synthetic datasets of Table H.1 checked it. Each dataset injects a known angle trajectory, synthesizes landmark positions in the camera frame {Camera}, pushes them through the complete pipeline, and compares the decoded angles with the injected ones. The injected trajectory is a synthetic exact reference, known before the landmarks exist."),

# per-dataset worst decode errors: v1/KINEMATIC_MODEL.md and
# v1/kinematics/dataset/README.md
(TBL, [["Synthetic dataset", "Motion"],
       ["Root sweeps", "torso yaw, pitch, roll, then combined"],
       ["Shoulder only", "shoulder swing and twist, torso still"],
       ["Shoulder with torso", "the same shoulder motion, torso moving"],
       ["Elbow only", "flexion, then a rotating flexion plane, then both"],
       ["Full arm", "root, shoulder, and elbow together"],
       ["Both arms", "torso and both arms, eleven angles nonzero"]]),
(CAP, "Table H.1. The six synthetic datasets. On each, the decoded angles agree with the injected ones over every frame and every angle, to the precision of the arithmetic."),

(P, "The residual is floating-point rounding, so the solver inverts the geometric construction exactly. The gimbal-lock case behaves as derived: a pose at the singularity decodes without error, as a different set of numbers for the same physical pose."),

(H2, "H.2 Measured Input"),

# rail sweep: writing/v7/scripts/ch3_numbers.py on recording_20260831_065553
# (77 of 900 frames with all eight landmarks; elbow ranges right -30.83 to
# -14.75, left -40.40 to -15.54)
(P, "On measured input, where no injected trajectory exists, the solver ran on the 77 frames of the rail recording, out of its 900, in which all eight landmarks are measured, all of them inside the slide. On the worst of those frames the angles it returns reconstruct the measured upper-arm direction on either arm to the precision of the arithmetic. Elbow flexion stays inside its anatomical range throughout, between -30.8 and -14.8 degrees on the right arm and between -40.4 and -15.5 on the left."),

(H2, "H.3 Scope of the Checks"),

(P, "These checks verify the kinematic mathematics that was tested, on the trajectories that were tested. They do not validate the landmark detector, the depth the sensor returns, the marker pose, or the anatomical accuracy of the reconstructed pose. They also say nothing about the recovery of a missing observation, which Section 7.4 grades separately."),

]

# --------------------------------------------------------------------------
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

# D-106: unnumbered display ordinals continuing into the following clause.
DISPLAY_COMMAS = {1, 2, 4, 6}
math_display_index = 0

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == STEP:
        p = doc.add_paragraph(item[1], style="List Number")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    elif kind == EQ:
        add_display_eq(doc, item[1], item[2])
    elif kind == MATH:
        math_display_index += 1
        add_display_math(doc, item[1],
                         punctuation="," if math_display_index in DISPLAY_COMMAS else ".")
    elif kind == CODE:
        for line in item[1].split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line if line else " ")
            run.font.name = "Courier New"
            run.font.size = Pt(9)
    elif kind == IMG:
        path = Path(item[1])
        assert path.exists(), f"missing figure: {path}"
        doc.add_picture(str(path), width=Inches(item[2]))
        if path.name == "appD_fig_alignment.png":
            # D-085: the actual PDF review found this caption detached.
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row_ in enumerate(rows):
            for j, c in enumerate(row_):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True
        if rows[0][:2] == ["Marker", "Role"]:
            # D-076: Table E.1's long header must remain complete and attached
            # to its first data row when the table moves across a page break.
            header = t.rows[0]
            props = header._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            props.append(OxmlElement("w:tblHeader"))
            for cell in header.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True

out = REPO / "writing" / "v8" / "condensed" / "Appendices.docx"
doc.save(out)
print("saved", out, "| items:", len(content))
