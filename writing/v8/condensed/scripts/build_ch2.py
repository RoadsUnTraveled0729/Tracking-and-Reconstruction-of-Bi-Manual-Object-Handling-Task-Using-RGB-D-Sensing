#!/usr/bin/env python3
"""Build the condensed Chapter 2: Experimental Setup and Measurements.

2026-09-07: approved D2/M06 preparation in Section 2.6 and Appendix G;
minimal marker/filter corrections and C38 rounding recorded in
notes_followup_ch2.md. Sections 2.1 and 2.2 condensed further.

2026-09-06 (MATH_LOGIC_REVIEW.md M10): Sections 2.2.1 and 2.2.2 state
both mounting assumptions of the gravity reference, the plumb wall and
the upright marker; Sections 2.3 to 2.5 untouched (user verbatim).

Chapter 2 "Experimental Setup and Measurements" (title per user decision
2026-09-05, supervisor round 4): 2.1 The Recording Setup, 2.2 Coordinate
Frames and Scene Calibration (new; C45/C46: the frame list, Figure 2.4
camera frame moved from Chapter 3, Figure 2.5 scene frames with the T
arrows, Table 2.2 transformations, and the static scene calibration
moved from Chapter 4 Section 4.1 with equation (2.1), Table 2.3 and
Figure 2.6), 2.3 Pose Landmarks from MediaPipe, 2.4 Object Pose from
ArUco Markers, 2.5 Measurement Filtering (C43 rename). Sections 2.3-2.5
are the user's own text (commit 669cff0); every edit to them is logged in
CH2_CONCISION_LOG.md. Environment only (user
direction 2026-09-01): the route the author measured lives in Chapter 7,
and the rail dimension figure (ch2_fig_rail_dims.png) moved there with it.
After the Chapter 7 restructure of D-029 that figure and the drawn waypoint
figures no longer exist; Section 7.2 names the four physical waypoint roles
W1 to W4 (D-036) and defines the endpoint-length and fitted-line
comparisons, and Chapter 7 defines no reference path. The rail lies
FLAT (a 2-by-4: 3.8 cm tall, top face 8.8 cm wide; E-024). Marker size:
the desk and object markers are 45 mm and the pipeline uses 45 mm
(E-009a; the provenance of that number is the comment above Table 2.1).
Rail-only round (2026-09-01, user direction): the thesis is
based on the rail recording recording_20260831_065553 (900 frames,
30 s) alone; the R5 loop recording appears only in Section 7.4.2, the
natural-occlusion evaluation, and is not mentioned here. All example figures
(setup, MediaPipe, ArUco, filters) come from rail frame 533 and the
rail landmark CSVs. Filter parameters per the run metadata: hampel
window 7, k 3, floor 2 cm, gaps to 5 frames, Butterworth 3 Hz order 4
zero-phase; edge backfill disabled.
Supervisor comments addressed: C4 (setup photographs + user study
objective), C6 (filters named, justified, referenced, results shown),
C8 (filter panel split into separate larger figures), C43, C45, C46
(round 4, 2026-09-05; C44 keeps the user's sentences as they are).
Natural-voice pass 2026-09-05 (user: no 'participant', no 'X, the Y of
Z' appositives, active voice, natural joins): 2.1 and 2.2 redrafted,
2.3-2.5 touched only where the rule required; log entries 131-169.
Humanizer pass 2026-09-05 (writing/WRITING_SKILL.md = blader/humanizer):
2.1 and 2.2 redrafted again, 2.3-2.5 touched in four places; log
entries 170-173 and the 2.1/2.2 redraft block.
Figures: make_ch2_frames_fig.py (Figure 2.5), make_ch2_fig1.py (2.3).
Figures: writing/v8/figures/. Citations: writing/v7/references.md
numbering ([1]-[52] frozen, [53]-[57] appended).

Round 1 fix pass of 2026-09-11 (writing/reviews/final_2026-09-11/round1/
Chapter_2.txt and Thesis_wide.txt, under FIX_DECISIONS.md G1 and G6; log in
FIXES_ch1-3.md). Section 2.6 no longer says "Every arm result of Chapter 7
rests on it", which Appendix G.4 denies for the hold-last baseline; it now
carries the G6 sentence, worded identically in Chapter 5 Section 5.1. The
same paragraph names the causal path of Chapter 8 instead of "the live
angle solve" and "the live pelvis translation" (G1). The Figure 2.5 caption
calls Chapter 4's product the cleaned object track, not a path. Section
2.2.2 reserves "physical reference" for a tape or ruler measurement and
names the antecedent of the gravity step. Sections 2.3 to 2.5 are untouched
in this pass; the two source comments added for the 1.3 metre working
distance and the 2.9 centimetre wrist spike are builder comments and change
no printed word.

Follow-up pass of 2026-09-12, after the Chapter 7 restructure completed
Section 7.2 (D-044, D-045):
 - Section 2.1 gains the task phase definition of D-044, renamed under D-051:
   object acquisition, carry, lift,
   slide and release for the rail task, and the same physical phases for the
   handover recording, whose slide additionally contains a right-hand
   segment, a handover interval and a left-hand segment. It is prose, it is a
   task definition only, and it carries no accuracy, measurement or outcome.
   The deleted Chapter 7 protocol figure is not reinstated.
 - Section 2.1 no longer promises that "Chapter 7 reports the physically
   measured route". Section 7.2 compares reconstructed endpoint separations
   with tape measurements of the physical setup, and it registers no physical
   route into the world frame, so no physical path or waypoint coordinate is
   drawn or reported. The sentence now says that, and names no quantity.
 - D-045: the 3.8 centimetre rail height survives only as the apparatus
   dimension of Table 2.1 (387 by 88 by 38 mm). No accuracy claim of this
   chapter rests on the dropped rail-height cross-check.
No number, figure, table or equation of this chapter changed in this pass.

Notation pass of 2026-09-12 (D-054, D-056, D-058, D-061, D-063; the apply
rows are FRAME_INVENTORY.md section 8.9.1, sixteen rows for this file and
one for make_ch2_frames_fig.py):
 - Section 2.2.1 and Table 2.2 now carry the whole frame inventory of the
   thesis under the semantic names, with the origin, the axes, the
   handedness and the parent of each frame, and the transformations
   between them in Craig form, reference frame in the left superscript and
   described frame in the left subscript. The five gaps the notation audit
   found are closed: the wall marker frame has a name, one marker rule
   generates the three marker frames instead of a generic marker frame,
   the root, shoulder and elbow frames are named after their landmarks
   with their parents and their sides given, the depth imager is recorded
   as not a frame of this system, and the two levelled systems are placed
   in Section 6.4 under D-061 rather than here.
 - The flip F and the swap S are declared once, after Table 2.2, by source
   basis, destination basis, matrix form, determinant and self-inverseness
   (D-054 class 4). Neither carries frame names and neither composes.
 - Figure 2.4(b)'s caption describes the coordinate system of the finished
   scene instead of the Unity frame at the root of the avatar rig (D-063),
   and the body sentence that repeated it now separates the two systems.
   The figure file is unchanged and was not re-rendered.
 - Figure 2.5's four arrow labels were re-rendered from
   writing/v8/scripts/make_ch2_frames_fig.py in the same Craig form.
No number, section, figure, table or equation number changed in this pass,
and no non-frame quantity was renamed: the singular-vector matrix keeps U,
the segment lengths and the landmark names keep L.
"""
# D-073: current frame labels are Camera and Camera'; historical notes
# above retain the terminology of their original decisions. No numeric map changes.
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches

from eqn import r, nor, sub, sup, d, hat, mat, add_display_eq, add_inline_math

H1, H2, H3, P = "h1", "h2", "h3", "p"
CAP, TBL, LIN, IMG, EQ, PM = "cap", "tbl", "lin", "img", "eq", "pm"


def T(t):
    return ("t", t)


def X(f):
    return ("m", f)


def M(f):
    """A table cell holding inline math instead of text."""
    return ("m", f)


# --- Craig frame notation (D-054, D-058, FRAME_INVENTORY.md sections 2, 3
#     and 8.9.1) ---
# The reference frame is the left superscript and the described frame the
# left subscript, so ^A_B T carries a point of {B} into {A}. Frame names are
# set upright. An absent script is a blank run: an empty <m:e> draws a
# placeholder box in some renderers, a single space draws nothing.
BLANK = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'


def pre(base, lsub, lsup):
    """Craig pre-sub-superscript: left superscript over left subscript."""
    # An empty OMML script slot renders as a placeholder box in the PDF
    # render path, so an absent frame carries a blank run instead.
    _BLANK = '<m:r><m:t xml:space="preserve"> </m:t></m:r>'
    return ('<m:sPre><m:sub>' + (lsub or _BLANK) + '</m:sub><m:sup>'
            + (lsup or _BLANK) + '</m:sup><m:e>' + base + '</m:e></m:sPre>')


def cT(ref, desc):
    """Class-2 homogeneous frame transform, both frames named."""
    return pre(r("T"), nor(desc), nor(ref))


def cR(ref, desc):
    """Class-1 frame-relative proper rotation, both frames named."""
    return pre(r("R"), nor(desc), nor(ref))


def inf(ref, sym):
    """A quantity expressed in one frame: left superscript only."""
    return pre(sym, BLANK, nor(ref))


# --- equation-fragment shorthands (Table 2.2, Table 2.3 and Section 2.2) ---
TSW = cT("Camera", "World")
TSwall = cT("Camera", "Wall")
TSO = cT("Camera", "Object")
TWO = cT("World", "Object")
TWS = cT("World", "Camera")
TSceneUnity = cT("Scene", "Unity")
TWwall = cT("World", "Wall")
TAB = cT("A", "B")
TPL24 = cT("Camera'", "L24")
TL24L12 = cT("L24", "L12")
TL24L11 = cT("L24", "L11")
TL12L14 = cT("L12", "L14")
TL11L13 = cT("L11", "L13")
TSWinv = sup(d(TSW), r("−1"))
PA = inf("A", r("P"))
PB = inf("B", r("P"))
gW = inf("World", hat(r("g")))
RAB = cR("A", "B")
RSW = cR("Camera", "World")
RSwall = cR("Camera", "Wall")
RWO = cR("World", "Object")
# The swap S exchanges the second and third coordinates (FRAME_INVENTORY.md
# section 6; v1/aruco/frames.py lines 38 to 40).
Smat = mat([[r("1"), r("0"), r("0")],
            [r("0"), r("0"), r("1")],
            [r("0"), r("1"), r("0")]])


def bar(e):
    """Overbar accent (mean matrix); the shared equation helpers have none."""
    return ('<m:acc><m:accPr><m:chr m:val="\u0304"/></m:accPr>'
            f'<m:e>{e}</m:e></m:acc>')


Mbar = bar(r("M"))
MbarSW = pre(Mbar, nor("World"), nor("Camera"))
MbarSwall = pre(Mbar, nor("Wall"), nor("Camera"))

REPO = Path(__file__).resolve().parents[4]
FIG = REPO / "writing" / "v8" / "figures"

# Rail-only round: the rail recording is the only recording of the
# chapter. The loop recording belongs to
# Chapter 7's occlusion evaluation and is introduced there.
F_RAILS, F_SETUP, F_FLOW, F_CAM, F_FRAMES = 1, 2, 3, 4, 5
F_MP, F_DESP, F_GAP, F_SMOOTH = 6, 7, 8, 9

content = [
(H1, "Chapter 2: Experimental Setup and Measurements"),

(P, f"Chapter 1 asked what can be reconstructed of a person and a manipulated object from one consumer-grade RGB-D viewpoint. Answering it starts from what the sensor delivers, and this chapter builds that measurement layer. The camera supplies aligned colour and depth, and the intrinsics turn pixels into metric points. A static calibration puts every observation in one calibrated metric scene. Its world frame is on the desk marker, and its vertical comes from the wall marker. Section 2.1 describes the recording setup, Section 2.2 the coordinate frames and the scene calibration, Sections 2.3 and 2.4 the landmark and the marker observations, Section 2.5 the filtering of the landmark signals, and Section 2.6 the preparation of the hip depth. A landmark is a measured point with a depth and a visibility score and not a joint angle; Chapter 3 builds the kinematic state from those points."),

(H2, "2.1 The Recording Setup"),

# 2.1 rewritten on the user's outline (2026-09-02): controlled-lab
# opening, one table of the scene objects, the 2 m range argument from
# the Intel D400 tuning guide [58], no field-of-view sentence, and no
# Appendix D pointer here (appendix order = first appearance in body;
# Appendix D was first cited in the MediaPipe section, user decision
# 2026-09-02; since round 4 the moved calibration text of Section 2.2.2
# cites Appendix D and Appendix E before Section 2.3 cites Appendix C,
# which the appendix order rule must absorb - flagged to the main agent).
(P, f"The recording took place in a controlled lab to limit lighting and background interference. The subject moves a cube along the desk and a wooden rail, with one hand in the rail recording of this chapter and with a handover between the hands in the handover recording of Chapter 7. Chapter 7 compares the reconstructed separations between the endpoints of that route with physical references measured on the setup with a tape, and it draws the reconstructed trajectories. The physical route is not registered into the calibrated world frame, so no physical path is drawn. Figure 2.{F_RAILS} shows the setup from three viewpoints, and Table 2.1 lists its objects."),

# Task phase definition (D-044). D-029 removed the four-step protocol
# drawing and the colour frame strip from Chapter 7, so the phases the
# later chapters name were defined nowhere. The five rail phases and the
# handover segmentation of the slide are the author's wording of
# 2026-09-12, recorded in D-044. This is a task definition only: it
# reports no accuracy, no measurement and no outcome, and the deleted
# protocol figure is not reinstated.
(P, "The task runs in five phases, and the later chapters refer to them by these names. The hand acquires the cube at its initial position in the object acquisition, and the carry translates the cube from that position toward the rail. The lift raises the cube from the desk level to the rail level, the slide translates it along the rail, and in the release the hand leaves the cube at the far end. The handover recording has the same physical phases, and its slide also contains a right-hand segment, a handover interval and a left-hand segment."),

(IMG, FIG / "ch2_fig_rail_setup.png", 6.3),
(CAP, f"Figure 2.{F_RAILS}. The recording setup. (a) The working area drawn on grid paper, with the rail, the cube at its start position, and the desk marker. (b) The scene from behind the RGB-D camera. (c) A colour frame of the recording."),

# src: marker sizes -> eval/common/marker_size.py, E-009a. The desk and
#      object markers are 45 mm and the pipeline uses 45 mm (user
#      measurement 2026-09-01; locked fact 1 of the revision brief of
#      2026-09-11). Provenance of the pinned chain: it implements the
#      45 mm size as the detector's 50 mm default multiplied by 0.9
#      (eval/reports/r6b_marker_scale.json and r7_marker_scale.json record
#      assumed_black_square_mm 45, scale_applied 0.9), and for a planar
#      pose solver the translation scales linearly with the size, so that
#      is exactly the 45 mm result. D-026 (author, 2026-09-11): the earlier
#      declared sizes are legacy values, not current thesis dimensions;
#      the 50 mm default is retained only as historical calculation provenance.
#      Rail 38.7 x 8.8 x 3.8 cm lying flat (E-024).
(TBL, [["Object", "Role", 'ArUco identifier', "Dictionary", "Printed size (mm)", "Physical dimensions"],
       ["Wall marker", "Static; gravity reference", "0", "5 by 5", "150", "-"],
       ["Desk marker", "Static; world origin", "2", "5 by 5", "45", "-"],
       ["Object marker", "Dynamic; tracked cube", "1", "5 by 5", "45", "Cube: 70 mm per side"],
       ["Wooden rail", "Reference object", "-", "-", "-", "387 by 88 by 38 mm, lying flat"]]),
(CAP, f"Table 2.1. The objects in the scene. All markers are from the OpenCV 5 by 5 dictionary, and the size is the side of the printed black square. The translation a planar marker solver returns scales with the size it is given, so every object-derived distance in this thesis rests on it."),

# src: working distance, author setup measurement; CH2_CONCISION_LOG.md
#      entry 7, kept at one decimal as a setup measurement and not a
#      computed result (entry 100). The 2 metre range argument follows the
#      Intel D400 tuning guide [58].
(P, f"Figure 2.{F_SETUP} labels the parts of the scene in one colour frame from the sensor. The camera is an Intel RealSense D435 on a fixed support at the edge of the desk, about 1.3 metres from the subject. The subject stays within 2 metres of it for the whole experiment, because the depth error of the sensor grows with the square of the distance [37], [38], [58]."),

(IMG, FIG / "ch2_fig_setup.png", 6.3),
(CAP, f"Figure 2.{F_SETUP}. The experimental setup seen from the RGB-D camera, with the subject, the rail, the carried cube, and the three ArUco markers labelled."),

(P, f"The recording is saved as a bag file, the format of the Intel RealSense software. Each frame holds a colour image and a depth image aligned to it, and a shared frame index says which person and object measurements belong to the same moment. Appendix A describes the recording and the file layout."),

(P, f"Figure 2.{F_FLOW} follows the recording through parallel landmark and marker pipelines. Pipeline A combines landmark detection, depth, filtering and the body solve of Chapters 3 and 5. Pipeline B measures the scene and object poses of Sections 2.2 and 2.4 and Chapter 4; these also support recovery in Chapter 5. Chapter 6 pairs the two outputs at a common render time for Unity. Appendix B lists the workstation."),

(IMG, FIG / "ch2_fig_flow.png", 6.3),
(CAP, f"Figure 2.{F_FLOW}. Overall data flow. Pipeline A (top) tracks the person, Pipeline B (bottom) the scene and object, and each block names the chapter that develops it."),

(H2, "2.2 Coordinate Frames and Scene Calibration"),

(P, f"Each position and orientation is expressed in a coordinate frame, an origin with three perpendicular axes. The transformations between them follow standard robotics notation [42]. Every frame of this thesis carries a name, and a name in braces is the label that frame takes in the superscript and subscript slots of a transformation. Table 2.2 gives the origin, the axes, the handedness and the parent of each frame, and then the transformations between them."),

(H3, "2.2.1 The Frames"),

# Semantic frame names (D-056, D-058) and the Craig two-frame rule (D-054).
# The frames, their origins, axes, handedness and parents are the verified
# inventory of FRAME_INVENTORY.md, refined by D-073 and D-082.
# Scene is the final display frame; Section 6.4 defines its full mapping.
# Gravity alignment and floor placement are operations, not extra frames.
(P, f"The raw camera frame {{Camera}} belongs to the RGB-D camera, and it is the frame of the colour camera to which each depth image is aligned. Its origin is at the colour camera lens, the x axis points to the right of the image, the y axis points down, and the z axis points forward into the scene [10], which makes it right-handed. The depth of a point is its z coordinate in that frame. Every raw measurement of both pipelines starts out there, and Figure 2.{F_CAM}(a) draws it on the sensor, where it carries the letter C."),

# D-065: the previous asset carried a baked-in panel title naming the frame
# as U and placing its origin at the avatar rig root, both false. Regenerated
# from the same two source images by make_ch2_sensor_scene_fig.py with that
# title reduced to the panel letter; the caption carries the description.
(IMG, REPO / "writing" / "v8" / "condensed" / "figures" / "ch2_fig_sensor_scene.png", 6.3),
# Figure 2.4(b) caption under D-063: the figure is accurate and is not
# re-rendered, and the caption describes the coordinate system of the
# finished scene without naming the formal frame symbol, which Section 6.4
# introduces beside the levelling rotation and the floor drop that define
# it. The four caption defects it repairs are recorded in FRAME_INVENTORY.md
# section 8.1.1: the old wording named the Unity frame U, which is 32.04
# degrees and 71.6 cm away from what the screenshot shows; it placed the
# origin at the root of the avatar rig, a node that is re-translated on
# every frame; and it named a levelling parent node the screenshot does not
# contain.
(CAP, f"Figure 2.{F_CAM}. (a) The sensor frame, drawn as C, on a front photograph of the D435 sensor (product photograph: Intel [59]; axes added by the author). (b) The coordinate system of the finished Unity scene, left-handed with y up and its origin on the drawn floor. The axes are drawn at that origin, where the avatar rig spawns in the T-pose, the reference configuration of the model. Section 6.4 builds this system from the display axes of Section 6.2."),

# One rule generates the three marker frames rather than a generic marker
# frame symbol (FRAME_INVENTORY.md section 2.4; the same rule is stated in
# Appendix E.2). The marker centre as the origin is D-053, traced to the
# symmetric corner layout of v1/aruco/extract_aruco_poses.py.
(P, f"Every printed marker carries a frame of its own. Its origin is at the centre of the printed square, its x and y axes lie in the printed plane and its z axis points out of the printed face, so each marker of Table 2.1 defines one right-handed frame. The world frame is the desk marker's, and it is fixed to the desk. Every object pose ends up in it, and Chapter 6 brings the reconstructed body into it as well. The wall marker's frame is fixed on a plumb wall behind the working area and is never used as an origin. The up axis in the plane of that marker gives the direction of gravity, on the assumptions Section 2.2.2 states, and the calibration uses it for that purpose only. The object frame rides the marker on the cube, and its pose is measured on every frame where the marker is detected."),

# The chain frames take the names of the landmarks at their origins
# (D-058), which names the parents the earlier draft left as "their parent
# frame" and distinguishes the two sides, which it never did.
(P, f"The y-up camera frame {{Camera'}} has the same optical centre as the raw camera frame {{Camera}} and differs only by flipping the y axis upward, which makes it left-handed. It is not person-centred; it is the coordinate convention used for body kinematics. Chapter 3 builds the body model in it as a chain of frames, each named after the landmark at its origin, so that the name also gives the side. The root frame sits at the right hip landmark L24 and carries the orientation of the torso. A shoulder frame sits at each shoulder landmark, L12 on the right and L11 on the left, and inherits the orientation of the root. An elbow frame sits at each elbow landmark, L14 on the right and L13 on the left, and carries the basis of the root turned by the solved shoulder rotation."),

# The body sentence that used to read "Figure 2.4(b) shows U at the root of
# the avatar rig" conflated two systems 32.04 degrees and 71.6 cm apart
# (FRAME_INVENTORY.md section 8.1.1). It now separates them and names no
# final Scene frame explicitly under D-082, derived in Section 6.4.
(P, f"Chapter 6 first converts the calibrated world coordinates to the left-handed display convention {{Unity}} by exchanging the second and third coordinates. This convention retains the desk-marker origin; its second axis follows the desk-marker normal and is not generally vertical. Gravity alignment then rotates these coordinates, and a recording-specific translation places the drawn floor at zero height. The result is the final {{Scene}} frame shown in Figure 2.{F_CAM}(b). Section 6.4 combines those two operations into one transformation from {{Unity}} to {{Scene}}, without defining an intermediate frame."),

(PM, [T("The pose of a frame B seen in a frame A is written "), X(TAB),
      T(", a 4 by 4 matrix that holds a rotation and a translation, and Chapter 3 gives its form in equation (3.2). The same matrix carries a point of B into A, so that "),
      X(mat([[PA], [r("1")]]) + r(" = ") + TAB + r(" ") + mat([[PB], [r("1")]])),
      T(f". A pair of them composes when the subscript of the first matches the superscript of the second, and the shared name drops out of the product. Figure 2.{F_FRAMES} draws the frames of the scene on a colour frame and on a top view, with an arrow for each transformation. Table 2.2 lists them and says which are fixed by the calibration and which are measured on every frame.")]),

# D-066: the condensed thesis uses a Craig-labelled copy of Figure 2.5,
# generated by make_ch2_frames_craig_fig.py. The shared generator and asset
# under writing/v8 are untouched, so the non-condensed v8 thesis keeps arrow
# labels matching its own prose.
(IMG, REPO / "writing" / "v8" / "condensed" / "figures" / "ch2_fig_frames_craig.png", 6.3),
(CAP, f"Figure 2.{F_FRAMES}. The scene and its coordinate frames. (a) The colour frame used in the worked examples, with {{Wall}}, {{World}} on the desk marker, and {{Object}} on the cube projected through the colour intrinsics. The raw camera frame {{Camera}} is the viewpoint of the image itself, and each arrow is one transformation of Table 2.2. Arrows point from the reference frame to the described frame; the labelled transform maps coordinates in the opposite direction. (b) The same frames in a top view of {{Scene}}, with the cleaned object track of Chapter 4 in grey. The common floor translation affects only height, so it does not change this x-z projection."),

# Table 2.2 is the frame inventory of the thesis and the transformations
# between the frames, in one table. The earlier version listed the
# transformations only, and the frames it drew on were incomplete in five
# ways (CRAIG_NOTATION_AUDIT_ch2_ch4.md section 2): the wall marker frame
# had no label, the generic marker frame had none, the root, shoulder and
# elbow frames were named in words with their parents unnamed and their
# sides never distinguished, and two further systems were missing. The
# marker family rule of Section 2.2.1 removes the need for a generic marker
# frame, Appendix D.2's depth imager is not a frame of this system because
# only the aligned colour frame is ever used downstream. D-082 includes
# the final Scene frame and its full transform, derived in Section 6.4.
(TBL, [["Frame", "Origin", "Axes and handedness", "Parent frame", "Also called"],
       ["{Camera}", "the colour camera lens", "x to the image right, y down, z forward into the scene; right-handed", "none; every measurement starts here", "the raw camera frame; drawn as C in Figure 2.4(a)"],
       ["{Camera'}", "the colour camera lens, shared with {Camera}", "x right, y up, z forward; left-handed", "{Camera}, through the flip F", "y-up camera frame for body kinematics"],
       ["{World}", "the desk marker centre", "x and y in the printed plane, z out of the printed face; right-handed", "{Camera}, fixed by the calibration", "the world frame"],
       ["{Wall}", "the wall marker centre", "the same marker convention; right-handed", "{Camera}, fixed by the calibration", "the wall marker frame"],
       ["{Object}", "the object marker centre", "the same marker convention; right-handed", "{Camera}, measured on every frame", "the object frame"],
       ["{Unity}", "the desk marker centre", "the axes of {World} with the second and third coordinates exchanged; left-handed", "{World}, through the swap S", "the display axes of Chapter 6; U"],
       ["{Scene}", "the drawn floor origin", "gravity-aligned display axes, y up; left-handed", "{Unity}, through gravity alignment and floor translation", "the final Unity scene of Section 6.4"],
       ["{L24}", "landmark L24, the right hip", "x to the right along the hip line, y up the trunk, z out of the chest; left-handed", "{Camera'}", "the root frame, the torso frame"],
       ["{L12}, {L11}", "landmark L12, the right shoulder, and landmark L11, the left", "the orientation of {L24}; left-handed", "{L24}", "the shoulder frames"],
       ["{L14}, {L13}", "landmark L14, the right elbow, and landmark L13, the left", "the basis of {L24} turned by the solved shoulder rotation; left-handed", "{L12}, {L11}", "the elbow frames"],
       ["Transformation", "Carries a point from", "Into", "Fixed or measured", "Derived in"],
       [M(TSW), "{World}", "{Camera}", "fixed by the calibration", "Section 2.2.2, Chapter 4"],
       [M(TSwall), "{Wall}", "{Camera}", "fixed by the calibration", "Section 2.2.2"],
       [M(TSO), "{Object}", "{Camera}", "measured on every frame", "Section 2.4"],
       # Split into two inline expressions so the cell can wrap between the
       # two factors; a single expression overflows the column.
       [[X(TWO + r(" = ") + TWS), T(" "), X(TSO)], "{Object}", "{World}", "computed on every frame", "Chapter 4"],
       [[X(TWS), T(" = "), X(TSWinv)], "{Camera}", "{World}", "fixed by the calibration", "Chapter 4"],
       [M(r("F")), "{Camera}", "{Camera'}", "fixed, the y flip; reverses handedness", "Chapter 3, equation (3.1)"],
       [M(TPL24), "{L24}", "{Camera'}", "measured on every frame", "Chapter 3"],
       [M(TL24L12 + r(", ") + TL24L11), "{L12}, {L11}", "{L24}", "measured on every frame", "Chapter 3"],
       [M(TL12L14 + r(", ") + TL11L13), "{L14}, {L13}", "{L12}, {L11}", "measured on every frame", "Chapter 3"],
       [M(r("S")), "{World}", "{Unity}", "fixed, the second and third coordinates exchanged; reverses handedness", "Chapter 6, equation (6.1)"],
       [M(TSceneUnity), "{Unity}", "{Scene}", "fixed per recording; gravity rotation and floor translation", "Chapter 6, equation (6.6)"]],
      (0, 11)),
(CAP, f"Table 2.2. The coordinate frames of the thesis and the transformations between them. A transformation written with two frame names carries a point from the frame in the subscript into the frame in the superscript, and two of them compose when the inner names match. The flip F and the swap S are not transformations of that kind, because each reverses handedness, so neither carries frame names and neither composes in that way."),

# Class-4 declarations (D-054): source basis, destination basis, matrix
# form, determinant and self-inverseness, for both maps and in one place.
# Content verified in FRAME_INVENTORY.md section 6 against v1/kinematics/
# root_frame.py lines 9 and 12 to 14 for F and v1/aruco/frames.py lines 38
# to 40 for S, and numerically on eval/output/scene_calibration_r6bc.json,
# which gives det S = −1 and det F = −1. The thesis stated the
# self-inverseness of the swap and never of the flip, which is what makes
# the bare F of equations (6.2) and (6.3) harmless; it is now stated for
# both. Code-to-thesis naming, recorded here and not printed because no
# prose of this thesis names a source module: the implementation calls the
# swap S by the name P and the flip F by the name D (eval/offset/carry.py
# lines 24 and 25; eval/unity_check/check_unity_log.py lines 88 and 89).
(PM, [T("Table 2.2 has two entries that reverse handedness, and they are declared here once. The flip "),
      X(r("F")),
      T(" takes the right-handed sensor frame to the left-handed y-up camera frame, and as a matrix it is "),
      X(nor("diag") + d(r("1, −1, 1"))),
      T(". The swap "), X(r("S")),
      T(" takes the right-handed world frame to the left-handed Unity frame by exchanging the second and third coordinates, so it is "),
      X(Smat),
      T(". Each has determinant −1, so neither is a rotation. Each is its own inverse, so applying either one twice returns the original coordinates. Neither carries frame names in the superscript and subscript slots, and the composition rule above never applies through one of them. Where an orientation crosses one of them, Chapter 6 prints the change of basis as a conjugation, with the map on both sides of the rotation, rather than as a multiplication on one side.")]),

(H3, "2.2.2 Scene Calibration"),

# Moved from Chapter 4 Section 4.1 (round 4, C46; build_ch4.py at commit
# d3561d2 lines 94-131). Long sentences re-patterned to the chapter's
# voice; each change is logged in CH2_CONCISION_LOG.md.
(P, f"The desk and wall markers are assumed fixed. Their poses are calibrated at the start and held constant throughout the recording."),

# C2-11 and C2-12: the mean, the decomposition and equation (2.1) stay
# generic, and the instantiations carry the frame pair. U keeps its letter:
# it is the left singular-vector matrix and not a frame, and the semantic
# frame names dissolve the collision it used to have with the Unity frame
# (D-056).
(PM, [T("Calibration averages the first ten detections of each static marker. Translation uses the arithmetic mean. The element-wise rotation mean generally loses orthonormality, so it is projected onto the nearest proper rotation [45], giving the chordal mean. Let "),  # src: eval/output/scene_calibration_r6bc.json (calib_frames = 10); default in v1/aruco/calibrate_scene.py
      X(Mbar), T(" be the element-wise mean of the rotation blocks of "), X(TAB),
      T(" across the window. Its singular value decomposition is "),
      X(Mbar + r(" = ") + r("U") + r(" ") + r("Σ") + r(" ") + sup(r("V"), r("T"))),
      T(", with the singular values in the middle factor. The projection is")]),

(EQ, RAB + r(" = ") + r("U") + r(" ") + nor("diag") +
     d(r("1, 1, ") + nor("det") + d(r("U") + sup(r("V"), r("T")))) + r(" ") +
     sup(r("V"), r("T")), "2.1"),

(PM, [T("where the determinant factor excludes reflections and the result is orthonormal. The calibration runs it on "),
      X(MbarSW), T(" for the desk marker and on "), X(MbarSwall),
      T(" for the wall marker, which gives "), X(RSW), T(" and "), X(RSwall),
      T(". Chapter 4 uses the same projection to smooth the object rotations "), X(RWO), T(".")]),

(P, f"Freezing the anchor prevents later desk-marker noise from rotating the reconstructed scene. An anchor orientation error displaces each point in proportion to its distance from the anchor. The frozen calibration can still be biased, but it adds no per-frame anchor wobble."),

(P, f"Calibration also measures the geometry of the scene around the markers, on two physical assumptions. The wall carrying the wall marker is plumb, and the marker is mounted with its printed up axis vertical. Under these assumptions, the up axis in the plane of the calibrated wall marker gives the direction of gravity. A plumb wall alone fixes only the normal of the marker plane, and the marker could still be turned within that plane. The upright mounting was not measured, so an error in how the marker hangs enters every gravity-referenced height and tilt of this thesis."),

(P, f"The axes of the world frame are not lined up with gravity because the desk marker card sits on a small stand that tilts it toward the camera. Gravity is therefore taken from the wall marker."),

(P, f"A coarser estimate of the same direction comes from a plane fitted to the depth pixels in a ring around that card, after the deprojection of Section 2.3 has turned those pixels into metric points. Its normal fixes the sign of the gravity direction and settles the two-solution ambiguity of the planar pose estimation that Section 2.4 describes. The depth fit also supplies a point on the tabletop. For height measurements, the implementation uses a horizontal tabletop model through that point, with its normal set by the calibrated gravity direction."),

(P, f"The remaining calibrated quantities follow from those pieces, and Table 2.3 lists the whole set with the measurement each one comes from."),

(TBL, [["Quantity", "How it is obtained"],
       ["World anchor pose", "mean translation and chordal-mean rotation over the calibration window"],
       ["Wall marker pose in the world", [T("the calibrated wall pose carried through the inverted anchoring transformation, "), X(TWwall + r(" = ") + TWS + r(" ") + TSwall), T(", where the shared name drops out")]],
       ["Camera pose in the world", [T("the inverse of the anchor, "), X(TWS), T(" of Table 2.2")]],
       [[T("Gravity direction "), X(gW)], "up axis in the plane of the calibrated wall marker, signed by the depth plane"],
       ["Tabletop model", "horizontal plane through the depth-derived tabletop point, with calibrated gravity as its normal"],
       ["World origin height above the tabletop", "the origin's signed offset along gravity from the horizontal tabletop model"],
       ["Object cube edge", "twice the marker centre height above the tabletop while the cube rests, or the measured edge"],
       ["Sensor field of view", "the colour intrinsics of the recording"]]),
(CAP, f"Table 2.3. The frozen scene description produced by the static calibration."),

(P, f"The pipeline hands this description to the reconstruction once, before any per-frame streaming begins, and Chapter 6 builds the Unity scene from it."),

(H2, "2.3 Pose Landmarks From MediaPipe"),

(P, f"Pipeline A starts with MediaPipe Pose. It is an open-source detector from the MediaPipe framework [28], and it runs in real time on ordinary hardware. For each colour frame, it first locates the person and then regresses 33 body landmarks from the face to the feet. Landmark indices are fixed. The pipeline can therefore select the eight required upper-body points directly: shoulders L11 and L12, elbows L13 and L14, wrists L15 and L16, and hips L23 and L24. Every result includes image coordinates, an appearance-derived depth estimate, and a visibility score between 0 and 1. Appendix C gives the detector architecture, the steps that extract the landmarks, and a numerical example."),

(P, f"Because MediaPipe estimates depth from two-dimensional (2D) appearance rather than physical sensor measurements, the pipeline discards those values and uses the depth readings from the aligned RealSense stream instead. It then uses the camera intrinsics to back-project each landmark pixel into metric 3D coordinates (Appendix D). The pipeline treats landmarks with visibility scores below 0.50 as missing data, so the system does not report coordinates for occluded points. Figure 2.{F_MP} illustrates the full detector output on a sample recording frame."),

(IMG, FIG / "ch2_fig_mediapipe.png", 6.3),
(CAP, f"Figure 2.{F_MP}. MediaPipe Pose output on a mid-task frame of the recording. The eight numbered landmarks (shoulders, elbows, wrists, hips) are the ones this system uses. The leg landmarks below the desk edge are guesses for body parts the desk hides, and the visibility gate removes them."),

(H2, "2.4 Object Pose From ArUco Markers"),

# C2-13: the last two sentences carry the frame names of Table 2.2. The
# rest of this section is the author's own text and is otherwise untouched.
(PM, [T("Pipeline B gets its measurements from ArUco fiducial markers [29], which have a black square border around a unique binary identifier (ID) grid. Once supplied with the marker's physical side length, it solves the planar Perspective-n-Point problem [52] to resolve both position and orientation in a single pass. The result is the pose of the marker in the sensor frame. For the object marker that pose is "),
      X(TSO), T(" of Table 2.2, and it is measured on every frame.")]),

# The marker table lives in Section 2.1 (Table 2.1) since 2026-09-02;
# the old duplicate table here is removed.
(P, f"The wall marker (ID 0, 150 mm) is mounted on a plumb wall. It is printed larger only because it sits farthest from the sensor. The wall marker's printed up axis is assumed vertical and defines the gravity reference. Appendix E covers corner detection, planar ambiguity, and rotation matrix conversion."),

(H2, "2.5 Measurement Filtering"),

(P, f"The raw 3D landmark trajectories carry three kinds of error:"),

(LIN, f"Case A, spikes: an isolated sample jumps far off the local trajectory for one frame, from a depth speckle or a momentary false detection."),
(LIN, f"Case B, gaps: a few consecutive samples are missing, because the visibility gate of Section 2.3 blocked them or the depth sample was invalid."),
(LIN, f"Case C, jitter: every sample carries small broadband noise from detection and depth measurement."),

(P, f"A single smoothing filter cannot solve all three kinds of errors. If a filter is strong enough to remove a sharp spike, it also distorts fast motion. No smoothing filter can fill in missing data gaps either. The pipeline therefore handles each case in a separate stage."),

# src: 2.94 cm off the local trend, right-wrist filtered track of the rail
#      recording (the series Figure 2.7 plots); rounded to 2.9 under C38,
#      notes_followup_ch2.md. Entry 46 of CH2_CONCISION_LOG.md is where
#      2.94 entered.
(P, f"A rolling Hampel test handles case A. The Hampel identifier [53] spots outliers by checking how far a point drifts from the local median. It works well on time-series data because a sudden spike will not pull the median off track the way it pulls a mean [54]. The test uses a 7-frame window and a threshold of three times a robust scale. On each coordinate, the pipeline subtracts each sample's own centred rolling median and takes the absolute residual. A second centred 7-frame rolling median over those residuals, multiplied by about 1.48, supplies the scale (Appendix F). A baseline floor of 2 centimetres prevents small sensor noise from triggering the test when the subject is standing still. The pipeline discards any spike the test flags and marks it as missing data, which then passes into case B. Figure 2.{F_DESP} shows the test flagging a wrist spike that jumped 2.9 centimetres off track in a single frame."),

(IMG, FIG / "ch2_fig_filter_despike.png", 6.0),
(CAP, f"Figure 2.{F_DESP}. Despiking on the right wrist of the recording: per-axis raw and filtered trajectories, with a mark on the sample flagged by the Hampel test."),

(P, f"Linear interpolation handles case B. The interpolation fills interior gaps no longer than 5 frames, a sixth of a second, along a straight line between the two valid endpoints. Longer gaps remain missing, and the pose recovery of Chapter 5 handles them. Figure 2.{F_GAP} shows bridged wrist dropouts from the recording."),

(IMG, FIG / "ch2_fig_filter_gap.png", 6.0),
(CAP, f"Figure 2.{F_GAP}. Gap bridging on the right wrist of the recording: the short gaps filled by interpolation, with the replaced samples marked."),

(P, f"A zero-phase Butterworth low-pass filter then handles case C. The filter follows the design procedure of [55]. The Butterworth family has a flat passband, so it passes the motion without ripple [56]. A fourth-order filter with a 3 Hz cutoff runs forward and backward over the recording [57]. The two passes cancel the phase delay. The forward and backward low-pass filter reduces jitter but also attenuates motion near its cutoff. The backward pass needs the whole recording, so the filter can only run offline. The real-time alternative is the One Euro filter [40]. It is a causal low-pass filter that adapts its cutoff to how fast the signal is changing. Chapter 8 compares the causal filter with the offline one on the same recording (Section 8.3). Appendix F compares the candidate smoothers side by side. Figure 2.{F_SMOOTH} shows the raw jitter against the filtered track."),

(IMG, FIG / "ch2_fig_filter_smooth.png", 6.0),
(CAP, f"Figure 2.{F_SMOOTH}. Smoothing on the right-wrist depth of the recording: the raw and filtered tracks over an eight-second span, with the shaded window enlarged below."),

(P, f"Every repaired sample keeps a flag, which separates measured data from filled-in data for later stages. Chapter 3 solves the joint angles from these filtered landmarks."),


(H2, "2.6 Hip Depth Preparation"),

# D2/M06: the existing solver's torso preprocessing, not a new method.
# Source: v1/kinematics/occlusion_ext.py:397-674; exact branches Appendix G.
(P, "A visible landmark can take the depth of the desk or rail in front of the body. Before solving the body angles, the pipeline checks hip and shoulder widths against calibrated values, then compares the directions of the hip line and shoulder line. It next rejects a remaining torso depth more than 10 centimetres from its remembered value. A final check rejects a hip more than 15 centimetres nearer the camera than the input shoulder midpoint. These checks run in that order, so the remembered-depth repair takes priority over the last check."),

(P, "Depth memory starts empty. The first sample that survives the checks seeds it; later accepted samples update it with 70 percent previous depth and 30 percent new depth. A repaired sample does not update it, and the memory does not expire. A single rejected point is placed on its retained camera ray at the remembered depth, or at the shoulder-midpoint depth for the last check. Rebuilding a whole pair also restores its calibrated width, so its final endpoints can leave their original rays."),

# Method names follow the rebuilt Chapter 7 (D-029): Section 7.4.1 compares
# hold-last, direction memory and object-assisted recovery, and Section 7.4.2
# compares the plain solve, hold-last and object-assisted recovery. The claim
# that hold-last does not rest on this preparation was re-verified against the
# variant definitions: the solving variants run RobustChainSolver, which
# carries the preparation, while hold-last runs ChainFallbackSolver, which has
# none (eval/failure/run_recovery.py variant list; eval/failure/recovery_core.py
# lines 193 and 227-230; eval/failure/compare_hip_hold.py, "the hold-last
# solve, no preparation"). Appendix G.4 states the same.
(P, "Appendix G gives the branching, fallbacks and output states. The preparation assumes usable landmark pixels and, for the shoulder-depth check, limited trunk lean. It does not establish anatomical hip depth. Chapter 5 uses the prepared torso for arm recovery. The plain, direction-memory and object-assisted solves of Chapter 7 all rest on it, and the hold-last solve, which keeps the last accepted joint angles, does not. The causal path's angle solve includes this preparation, but its pelvis translation still uses the original hip midpoint."),

(P, "The chapter has established the measurement layer: aligned colour and depth from one camera, the intrinsics and the frames, and the static calibration that fixes the world frame and the direction of gravity. In that scene, the pipeline observes the body as landmarks and the object as marker poses. It filters the landmark signals and prepares the hip depth. It leaves three things unresolved. The landmarks are metric points rather than a kinematic state. The object pose is a per-frame observation rather than a track. Nothing so far links the body to the object. Chapter 3 takes the first of them, Chapter 4 the second and Chapter 5 the third."),

]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == H3:
        doc.add_heading(item[1], level=3)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == PM:
        p = doc.add_paragraph()
        for typ, val in item[1]:
            if typ == "t":
                p.add_run(val)
            else:
                add_inline_math(p, val)
    elif kind == EQ:
        add_display_eq(doc, item[1], item[2])
    elif kind == LIN:
        p = doc.add_paragraph(item[1])
        p.paragraph_format.left_indent = Inches(0.3)
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == IMG:
        path = Path(item[1])
        assert path.exists(), f"missing figure: {path}"
        doc.add_picture(str(path), width=Inches(item[2]))
    elif kind == TBL:
        rows = item[1]
        # Optional third element: the row indices to set in bold. Table 2.2
        # carries two header rows, one for the frames and one for the
        # transformations between them.
        hdr = item[2] if len(item) > 2 else (0,)
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        if rows[0][0] == "Frame":
            # D-082 visual QA: keep each inventory entry on one page.
            for row in t.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j)
                if isinstance(c, tuple):
                    add_inline_math(cell.paragraphs[0], c[1])
                    continue
                if isinstance(c, list):
                    # A cell that mixes text and inline mathematics.
                    par = cell.paragraphs[0]
                    for typ, val in c:
                        if typ == "t":
                            run = par.add_run(val)
                            run.font.size = Pt(10)
                        else:
                            add_inline_math(par, val)
                    continue
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i in hdr:
                            run.font.bold = True

out = REPO / "writing" / "v8" / "condensed" / "Chapter_2_Experimental_Setup.docx"
doc.save(out)
print("saved", out, "| items:", len(content))
