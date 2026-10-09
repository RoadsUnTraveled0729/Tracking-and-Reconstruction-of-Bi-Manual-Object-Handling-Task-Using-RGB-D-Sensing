#!/usr/bin/env python3
"""Build writing/v8/condensed/Chapter_8_Real_Time_Feasibility.docx.

Revision of 2026-09-11 (Phase 1B of REVISION_2026-09-11_BRIEF.md). The
chapter stays "Chapter 8: Real-Time Feasibility" and is not turned into a
discussion. Its question: can the offline reconstruction architecture be
transferred to causal processing, where no future sample is available.
Sections: 8.1 From Offline to Causal Processing (the question and the six
offline operations that read the future), 8.2 Causal and Real-Time
Architecture (the live structure, the causal replacements, Figure 8.1 and
Table 8.1), 8.3 Offline-Causal Comparison (Figures 8.2 and 8.3 and the
corrected pairing statistics of the rerun), 8.4 Runtime, Buffering, and
Delay (the replay rate and the three separate timing quantities),
8.5 Feasibility and Remaining Validation Gap.

Previous structure (build of 2026-09-07, D-008): 8.1 The Live Structure,
8.2 Changes in the Transfer, 8.3 Losses of the Live Structure, 8.4 Cost.
The 8.1 and 8.2 material is now 8.2, the 8.3 comparison is now 8.3, the
8.3 open items are now 8.5, and 8.4 is now 8.4 with the delay quantities
separated as MATH_LOGIC_REVIEW.md M06 requires.

Numbers printed here and their pinned sources:
 - replay: 900 frames consumed, none skipped, 29.0 fps sustained, steady
   merger ticks -> v2/dataset/r6b_probe_baseline.txt (E-026); written
   "about 29 frames per second" in prose.
 - branch budgets: landmark p99 13.0 ms, marker p99 16.7 ms against the
   33.3 ms budget at 30 frames per second -> v2/dataset/m03_rail_corrected.txt
   (also parked Table 8.3 in writing/v8/CH8_PARKED.md).
 - configured delay of two frames -> v2/integration/v2_integrate.py.
 - offline/causal angle differences in the three windows (42.3 degrees
   twist, written "about 40"; 24.1 degrees elbow) ->
   writing/v8/condensed/figures/ch8_compare.json.
 - corrected pairing of the rerun: lag 5 frames, 174 paired frames, per
   angle Rel_y 0.41/1.18, Rsh_z 1.45/2.30, Rsh_tau 7.04/15.09, Rsh_y
   16.42/26.94, worst group 3.90/23.25 degrees -> v2/dataset/
   m03_rail_corrected.json (rerun recorded in MATH_LOGIC_REVIEW.md M03,
   inputs v2/dataset/m03_rail_inputs.zip, reference rebuilt by
   writing/v8/condensed/audit_evidence/followup/m03_reconstruct.py);
   printed to 0.1 degree.
 - 45 mm desk and object markers -> REVISION_2026-09-11_BRIEF.md section 3
   fact 1 (provenance: eval/common/marker_size.py, E-009a).

Sources for the structure and the causal replacements: v2/V2.md,
v2/broker/capture_broker.py, v2/common/shm_ring.py, v2/person/v2_person.py,
v2/object/v2_object.py, v2/common/object_link.py (SKEW_MAX 3),
v2/integration/v2_integrate.py (delay 2 frames, INTERP/HELD/BLEND),
v2/common/display_lpf.py, eval/failure/grip_state.py (EXIT_FRAMES 5),
v2/reports/r5_realtime_probe.md (feature preservation).

Chapter 7 restructure of 2026-09-12 (D-029, D-037; CH7_INTEGRATION_FIXLIST.md
items 8-1 to 8-3). Section 8.3 no longer says that Chapter 7 grades the
offline pass against "physical, manual and synthetic references": it names
the measured landmarks of Section 7.3 and the manual wrist labels of
Section 7.4.2, and the synthetic exact references now point at Appendix H,
where D-037 put the solver verification. Section 8.5 takes the rebuilt
chapter's vocabulary, synthetic landmark removal and manual wrist labels.
The Table 8.1 offline-only row drops the coarse rail-height cross-check,
which the restructure removed from Chapter 7 and which is printed nowhere
in the thesis. No number, figure, table or equation of this chapter changed,
and no quantity of Section 7.2 is named while D-036 is being computed.

Follow-up pass of 2026-09-12, after Section 7.2 was completed (D-044, D-045):
 - Section 8.1 names Section 2.1 as the source of the task phase vocabulary,
   which D-044 defines there. That is the single pointer of the chapter,
   placed before the first use, the slide of Table 8.1.
 - Section 8.3 carries the physical reference again. Sections 7.2.1 and 7.2.2
   now compare reconstructed marker-centre segment lengths with the author's
   tape measurements of the physical route, so the clause dropped under D-036
   is restored and the sentence is split in two for the 45-word limit. Still
   no quantity of Section 7.2 is named here.
 - D-045: the rail-height cross-check stays out. The Table 8.1 offline-only
   row still does not name it, and no claim of this chapter rests on it.
No number, figure, table or equation of this chapter changed in this pass.

Parked verbatim in writing/v8/CH8_PARKED.md: the previous Section 8.2 with
Tables 8.2 and 8.3 and the latency figure, and the check-suite paragraphs
of the previous Section 8.3. Restored from it in this revision: only the
two branch p99 values and the frame budget, in Section 8.4.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

H1, H2, P, CAP, TBL, IMG = "h1", "h2", "p", "cap", "tbl", "img"

REPO = Path(__file__).resolve().parents[4]
FIGC = REPO / "writing" / "v8" / "condensed" / "figures"

F_SYS, F_TRACE, F_OVER = 1, 2, 3
T_FEAT = 1

content = [
(H1, "Chapter 8: Real-Time Feasibility"),

(P, "Every result of Chapter 7 came from a finished recording, and a clinical or robotic use needs the reconstruction while the motion happens. This chapter asks whether the reconstruction architecture of Chapters 2 to 6 can be transferred to causal processing, in which no sample from the future is available. It names the offline operations that read the future, gives the causal architecture that replaces them, compares the two processing paths on one recording, reports the runtime, and states what the experiment leaves open."),

(H2, "8.1 From Offline to Causal Processing"),

(P, 'A stage is causal when its output for a frame depends only on that frame and the frames before it [55]. The offline pass breaks that condition in six places, each a point where the reconstruction reads a sample a live system would not yet hold. The landmark smoother of Section 2.5 is a fourth-order Butterworth low-pass run forward and backward over the whole signal, and the two passes cancel the phase delay only because the backward pass starts at the end of the recording [55], [57]. The despiking test of the same section compares a sample with the median of a window centred on it, so half of that window lies ahead of the sample. Gap bridging needs the sample on the far side of a gap.'),

(P, 'The hand-object offset of Chapter 5 warms up on the mean over a whole grip episode. The same tracker confirms an episode exit over a run of five frames and then clears those frames, so the state of a frame changes after the frame has passed. The cleaning rule of Chapter 4 reads the object track as a complete time series, after the track has been smoothed.'),

# The task phase names (D-044) are defined in Chapter 2 Section 2.1. This
# is the single pointer of the chapter, placed before the first use of the
# vocabulary, which is "the line fitted to the slide" in Table 8.1; later
# uses, among them the lift of Figure 8.2 and the slide of Section 8.3, are
# not tagged again.
(P, "The causal structure must replace each of those operations with one that reads only the past, or do without it. This chapter tests the exchange by running the causal structure on the rail recording of Chapter 2 at its recorded pace and comparing its output with the offline pass frame by frame. The task phase names used in this chapter, among them the lift and the slide, are those of Section 2.1. Replaying one recording keeps the scene fixed, so a difference between the two outputs follows from the processing path and not from the motion, the lighting or the camera placement. Exchanging the recording for the sensor changes only the capture process."),

# the five-frame exit run and the clearing of those frames: Chapter 5
# Section 5.3 (exit on five consecutive frames of rest or of a cleanly
# measured wrist beyond the release radius, then the five frames of the
# exit run are cleared), v2/common/object_link.py EXIT_FRAMES; the whole
# episode still enters the sixth operation through the warm-up mean above
# (Chapter_8.txt SHOULD FIX 4).
(H2, "8.2 Causal and Real-Time Architecture"),

(P, f"Figure 8.{F_SYS} sets the offline pass and the causal structure side by side, each causal stage under the offline stage it replaces. Offline, every script opens the recording itself and writes a table that the analysis tools of Chapter 7 read afterwards. Live, one capture process owns the sensor or the recording, aligns depth to colour once, stamps each frame with the session clock, and publishes the pair into a shared frame buffer (Section 6.1). The landmark branch and the marker branch read that buffer as separate processes and write their records. Every record carries the frame index and the session time of the frame it came from, and that pair is the only synchronization between the branches. The recovery of Chapter 5 inside the landmark branch also reads the object record the marker branch publishes, so it can use the object pose."),

(P, f"The merger pairs the two records at a render time that trails capture by a configured delay of two frames. It interpolates within that delay and holds across a longer gap as Section 6.1.2 describes, and writes the combined record that the Unity receiver of Chapter 6 draws. The per-joint output states of Chapter 5 travel in the streamed packet, so a constrained joint group is distinguishable on screen from a held one."),

(IMG, FIGC / "ch8_fig_system.png", 6.4),
(CAP, f"Figure 8.{F_SYS}. The offline pass (above) and the causal structure (below), each causal stage under the offline stage it replaces. Green boxes in the landmark lane and peach boxes in the marker lane keep the same construction as offline, amber boxes are causal replacements, and violet boxes are parts that exist only in the causal structure. Grey boxes with a dashed edge are analysis tools with no causal counterpart, yellow boxes are records in shared memory, and the blue-grey box at the left of each band is the frame source."),

(P, f"The structure rests on three choices. A single process owns the device, because two processes cannot open the same camera. The merger renders behind capture by a fixed, configured delay rather than at the newest sample, so it can interpolate between two samples it already holds instead of guessing forward. That delay is a buffering choice, distinct from the lag of the causal filter that Section 8.3 shows and Section 8.4 separates from it. The display smoother of Section 6.3 acts on the output record only, so what the branches compute is what their records hold."),

(P, f"Table 8.{T_FEAT} sorts the stages into three groups, unchanged, adapted, and analysis tools with no live form. The largest group needs no change. The kinematic solve of Chapter 3 is closed form and holds no state across frames. The recovery layer of Chapter 5 does carry state, the direction memories, the offset estimate, the holding state, the twist hold and the rate limit, but every update reads the current frame and what earlier frames left behind, at a bounded cost per frame. The marker detector and the planar pose solver read a single image."),

# three rejections in a row: MAX_REJECTS = 3 in
# v1/realtime/person/realtime_person.py (the causal despike class used by
# v2/person/v2_person.py)
(P, "The causal structure replaces each of the three landmark operations of Section 8.1. The despiking window trails the sample instead of surrounding it, and a rejected sample becomes a missing sample. After the filter rejects three samples in a row, it accepts the next one and rebuilds its window around it, so that a trailing window does not go stale when the landmark moves. Gap bridging moves into the merger, which bridges only between two samples it already holds within its delay and holds the last pose when the gap runs past that delay. The zero-phase Butterworth smoother gives way to the One Euro filter on each landmark coordinate [40]. Its cutoff frequency increases as the coordinate changes faster, so it smooths a still pose and follows a moving one."),

(P, 'The marker branch keeps its mechanism and changes its order. The cleaning rule of Chapter 4 discards a sample farther from the last kept sample than the slow task can produce. Live it runs on each detected sample as it arrives, ahead of the causal filter, where offline it runs on the smoothed track.'),

(P, 'The grip tracker of Chapter 5 changes in two places. The hand-object offset keeps its slow update, and its warm-up becomes a running mean over the first clean frames of the episode, fitted in the frame the solver works in. The causal tracker decides episode entry and exit after a limited number of frames. The test for a carried object takes the first object samples as its rest reference and gates entry only, so a pause with the object held does not split an episode. Exit is confirmed only after several consecutive frames of a cleanly measured wrist beyond the release radius.'),

# Table 8.1 spans Chapters 2 to 5: three rows are the filtering of
# Section 2.5, row 1 carries the hip depth preparation of Section 2.6 and
# one row is the cleaning rule of Chapter 4 (Chapter_8.txt MUST FIX 2).
# The offline-only row claims no comparison against the physical route
# (Chapter_8.txt MUST FIX 3, locked fact 3). Refreshed for the Chapter 7
# restructure (D-029) and D-035/D-036: the coarse rail-height cross-check
# was dropped with the old Section 7.2 and is printed nowhere in the
# thesis, so this row no longer names it. Under D-036 the ArUco marker
# centre is the physical and reconstructed endpoint, so Section 7.2
# compares endpoint segment lengths with the tape lengths without a
# registered route; those results are still being computed, so no
# quantity from Section 7.2 is named here.
(TBL, [["capability", "offline form", "live form", "relation"],
       ["Failure detectors, hip depth preparation, twist hold and angle rate limit", "per frame", "same construction", "unchanged"],
       ["Wrist from the object pose and two-link elbow solution", "per frame", "same construction", "unchanged"],
       ["Object pose read by the landmark branch", "the cleaned track as a table", "newest published marker record, rebuilt in the camera frame {Camera}", "new data path"],
       ["Per-joint output states carried to the reconstruction", "written into the record", "packed into the streamed packet and carried through the merger's buffer", "new data path"],
       ["Landmark despiking", "window centred on the sample", "trailing window; a rejection becomes a missing sample", "adapted"],
       ["Interpolation across a gap", "over the whole signal", "only between samples the merger holds, then the last pose is held", "adapted"],
       ["Landmark smoothing", "forward and backward Butterworth", "One Euro filter on each coordinate", "adapted"],
       ["Cleaning rule on the object track", "applied to the smoothed track", "applied to the detected sample ahead of the causal filter", "adapted"],
       ["Hand-object offset and grip episodes", "warm-up over the whole episode, then the slow update", "warm-up over the first clean frames, then the same update, with entry and exit decided after a limited number of frames", "adapted"],
       ["Comparison against the line fitted to the slide, the failure mask as a reference, the within-recording comparison under synthetic landmark removal", "over a finished recording", "no live counterpart, beyond the solver's own gates and the grip-plausibility detector", "offline only"]]),
(CAP, f"Table 8.{T_FEAT}. The parts of the offline pass of Chapters 2 to 5 that the causal structure keeps; the replacements of this section are the adapted rows."),

# the object record read by the landmark branch: Chapter 6 Section 6.1
# describes the one-way link, and the transformation it inverts is the
# axis swap of Section 6.2 with the calibrated anchor of Section 2.2.2
# (Chapter 6 Section 6.1.2). The skew bound is three frames, SKEW_MAX = 3
# in v2/common/object_link.py, printed as Chapter 6 prints it
# (Chapter_8.txt SHOULD FIX 3b and SHOULD FIX 6).
(P, "The two capabilities marked as a new data path needed no new mathematics. Offline the recovery reads the cleaned object track as a table. Live it reads the newest record the marker branch published and rebuilds the object pose in the camera frame {Camera}. It undoes the axis swap of Section 6.2, then composes the calibrated pose of the world frame in the camera frame {Camera}, established in Section 2.2.2, with the object pose in the world frame. It rejects a record whose frame index sits more than three frames from the frame being solved. The offline-only group has no live form beyond the solver's own gates, which do the work of the failure mask, and the grip-plausibility detector of Section 5.2 inside the live grip tracker."),

(H2, "8.3 Offline-Causal Comparison"),

# The physical reference was dropped from this list while Section 7.2 was
# Data Required under D-036. Sections 7.2.1 and 7.2.2 now compare
# reconstructed marker-centre segment lengths with the author's tape
# measurements of the physical route, so the clause is restored. The
# sentence is split in two to stay inside the 45-word limit of
# check_style.py. No quantity of Section 7.2 is named here.
(P, f"The transfer keeps the mathematics of Chapters 3 to 5 and changes what each stage may look at. Setting the two outputs against each other is a within-recording comparison, one output of the pipeline beside another output of the same recording. Neither path is the reference for the other, and a difference between them is a processing-path difference rather than an error against physical truth. Chapter 7 evaluates the offline pass against physical references, measured landmarks and manual wrist labels, and against its own outputs where no independent reference exists. Section 7.2 sets the reconstructed object segment lengths against tape measurements of the physical route, and Appendix H verifies the kinematic solve against synthetic exact references."),

(P, f"This section measures the distance between the two processing paths. Figure 8.{F_TRACE} draws the right elbow flexion and the right shoulder twist of both paths on the same frame axis in three windows of the recording. Figure 8.{F_OVER} draws both chains on the colour frame at one moment of each window. The angles of the recovery solve of Section 7.4 place the offline chain and those of the causal path place the causal chain, both built in the y-up camera frame {{Camera'}} on the measured shoulder of the frame."),

(IMG, FIGC / "ch8_fig_traces.png", 6.4),
(CAP, f"Figure 8.{F_TRACE}. The right elbow flexion (upper row) and the right shoulder twist (lower row) of the offline pass and the causal path on the same frames. The windows are the start of the recording, the lift onto the rail just before the wrist gap opens, and the wrist gap. The grey band shows the frames on which the offline pass rebuilt the elbow, the striped band those on which the causal path did."),

(P, f"The smoother loses its view of the future. A zero-phase filter sees both sides of every sample; the causal filter follows a moving joint late. In the lift onto the rail of Figure 8.{F_TRACE} the causal curve trails the offline one by a few frames and rounds its turns. The lag is small on this slow task and grows with the speed of the motion."),

# twist difference: figures/ch8_compare.json, window "start of the take",
# Rsh_tau max_abs_diff_deg 42.3 at frame 32 (written "about 40 degrees").
# ch8_compare.json stores only the counts for that window
# (offline_constrained_frames 18, live_constrained_frames 0), so the range
# "frames 7 to 24" is recomputed from the per-frame group states:
# eval/output/recovery_r6b/angles_recovery.csv has tag_1, tag_2 and tag_3
# (the right swing, twist and elbow groups) all constrained on frames 7 to
# 24, eighteen frames, and v2/output/v2_person_dump_r6b_full.csv has no
# such frame in the window, which is the live count of 0.
# Wrist separation at frame 18: figures/ch8_compare.json overlay_frames,
# wrist_offline_vs_live_cm 12.0.
# the twist hold releases above 25 degrees of elbow flexion and holds
# below 15 (Chapter 5 Section 5.5); the release radius is the wrist to
# object distance of Section 5.3 and plays no part here, which is what
# Chapter_8.txt MUST FIX 1 corrects. Figure 8.2 upper right shows the
# flexion passing 30 degrees at frame 550 and both twists leaving their
# frozen values on the same frame.
(P, f"The states carried by the solver depend on the starting frames. The twist hold of Section 5.5, which acts while the elbow is nearly straight, the direction memories and the hand-object offset keep whatever the earliest frames gave them. On this recording the offline failure mask rebuilt the right arm on frames 7 to 24 while the live gates held it. The two paths therefore froze shoulder twists about 40 degrees apart and kept them until the elbow bent far enough inside the gap to release the twist hold (Figure 8.{F_TRACE}, lower row). At frame 18 the offline pass rebuilt the arm and the causal path held it, and the rendered wrists sit 12.0 centimetres apart (first panel of Figure 8.{F_OVER}). The frozen twists keep the two wrists apart on the later frames where every landmark is measured. Runs of the same structure agree on the same frames only when they start from the same first frames."),

# elbow difference inside the gap: figures/ch8_compare.json, window
# "natural wrist gap", Rel_y max_abs_diff_deg 24.1 at frame 594; wrist
# separation 1.9 cm at frame 595 in the same file
(P, f"The hand-object offset differs as well. Offline it is the mean over the whole episode, and live it is a running mean over the first clean frames of the episode. The two paths therefore rebuild the wrist onto different points of the object. Just before the gap opens, the offline pass still follows the tracker's wrist, which has already drifted onto the object, while the causal path has begun to rebuild from the object (second panel of Figure 8.{F_OVER}). Inside the gap the rebuilt elbow-flexion angles differ by as much as 24.1 degrees (Figure 8.{F_TRACE}, right), although the wrists nearly coincide (third panel of Figure 8.{F_OVER}). The recovery also starts and ends on different frames along the two paths, because the live gates and the offline mask do not fire on the same frames."),

(IMG, FIGC / "ch8_fig_overlays.png", 6.4),
(CAP, f"Figure 8.{F_OVER}. The offline chain (blue) and the causal chain (orange) drawn on the colour image. Each chain is built in the y-up camera frame {{Camera'}} and carried back into the camera frame {{Camera}} by the fixed axis flip of Section 3.1 before the colour intrinsics place it on the image. At frame 18, the start of the recording, the offline pass rebuilt the arm and the causal path held it. At frame 548, just before the wrist gap opens, the offline pass still follows the drifting tracker wrist and the causal path has begun to rebuild from the object. At frame 595, inside the gap, both paths rebuild and the elbows differ more than the wrists."),

# gated hips of the causal run: v2/dataset/r6b_probe_baseline.txt, the
# [robust] line of the landmark branch, gated {'right_hip': 858,
# 'left_hip': 847, 'right_wrist': 61} of the 900 frames consumed; the
# larger of the two hip counts is printed. Round-1 review, Chapter_8.txt
# SHOULD FIX 5: 858 is the count of frames on which the checks REJECTED a
# measured right hip; the recovered count of the same line is 867, since a
# whole-pair repair also replaces a hip the gates did not reject. The verb
# now names the gated quantity that the printed number measures.
# corrected pairing of the rerun: v2/dataset/m03_rail_corrected.json
# (lag_frames 5, paired_frames 174, per_angle Rel_y 0.41/1.18, Rsh_z
# 1.45/2.30, Rsh_tau 7.04/15.09, Rsh_y 16.42/26.94, worst group R_shoulder
# 3.90/23.25); rerun and its inputs recorded in MATH_LOGIC_REVIEW.md M03.
# Printed to 0.1 degree (measurement-precision-reporting.md).
(P, "A rerun of the live validator on the archived records of the same replay puts numbers on the comparison. Each causal frame is paired with the offline frame that the alignment places beside it, and the arm must be cleanly measured on both sides of the pair. The alignment that brings the two angle series closest together is five frames, and the pairing keeps 174 frames. At that alignment the elbow flexion of the two paths differs by a median of 0.4 degrees, the shoulder elevation by 1.4, and the shoulder twist by 7.0."),

(P, "The shoulder azimuth differs by a median of 16.4 degrees and carries the shoulder group to a 95th percentile of 23.3 degrees. Section 3.4.2 gives the reason: the arm hangs close to straight down through the slide, near the configuration in which the azimuth is undefined. The comparison covers the arm alone, because the hip depth preparation of Section 2.6 rejects a measured hip on 858 of the 900 frames of this recording, so the root is rarely a measured group on both sides at once."),

# the merger's blend is two ticks, the count Section 6.1.2 defines
# (v2/integration/v2_integrate.py BLEND); Chapter 6 Section 6.1.2 stays
# primary for the gap mechanism and this chapter points at it
# (Chapter_8.txt SHOULD FIX 1).
(P, "The edges of an episode and of a gap differ too. The causal exit rule keeps the wrist anchored to the object during the frames that confirm a release, while the offline episode logic clears those frames once it sees the release. The merger blends a return from a long gap over two ticks instead of snapping to it (Section 6.1.2). On this recording the merger never had to hold a person record mid-run."),

(H2, "8.4 Runtime, Buffering, and Delay"),

# replay: v2/dataset/r6b_probe_baseline.txt (900 frames consumed, 0
# skipped by latest-wins, 31.0 s, 29.0 fps sustained, 929 merger ticks);
# written "about 29 frames per second". Branch budgets restored from the
# parked Table 8.3: v2/dataset/m03_rail_corrected.txt (person compute p99
# 13.0 ms vs 33.3 ms, object compute p99 16.7 ms). Detector share and the
# per-stage costs: parked Table 8.2 (writing/v8/CH8_PARKED.md), whose
# profile is the same replay in v2/dataset/r6b_probe_baseline.txt. On the
# run with the recovery layer on, the stages this thesis adds have 99th
# percentiles of 0.63 ms (landmark deprojection and causal filtering),
# 0.58 (kinematic solve and recovery), 0.01 (record write), 0.22 (planar
# pose estimation) and 0.24 (marker causal filtering and world mapping),
# against detection at 12.15 and 16.33 ms; 0.63 ms is the largest, and it
# is printed as the bound.
(P, "On the workstation of Appendix B the causal structure consumed all 900 frames of the recording at its recorded pace, sustaining about 29 frames per second, with none dropped, and the merger emitted its output at a steady interval. Detection takes most of each frame's budget on both branches, and no stage this thesis adds exceeds 0.63 milliseconds at the 99th percentile. Each branch also finishes inside the frame budget of 33.3 milliseconds at the 99th percentile, the landmark branch within 13.0 milliseconds and the marker branch within 16.7. A lighter detector variant is the documented way to reduce the cost [43]."),

# configured delay 2 frames: v2/integration/v2_integrate.py. The three
# quantities are kept separate on MATH_LOGIC_REVIEW.md M06: the two-frame
# output delay, the causal filter's lag, the estimated offline/live frame
# alignment, and no verified physical latency measurement exists.
(P, f"This chapter reports three timing quantities, and they are separate. The configured two-frame delay is a buffering and render delay: the merger renders the world two frame intervals behind capture so that it can interpolate between two samples it already holds. It is not a measured end-to-end latency from the sensor to the screen, which would also carry the capture, the transport and the display's own refresh, and which this thesis does not measure. The second quantity is the lag of the causal filter, which depends on the One Euro filter and the speed of the motion. In the lift, Figure 8.{F_TRACE} shows the causal output following the offline output a few frames later. The five-frame alignment of Section 8.3 is the third, an estimate of the shift that brings the two angle series into closest agreement on this recording. None of the three bounds the others."),

(H2, "8.5 Feasibility and Remaining Validation Gap"),

(P, "The causal structure ran a whole recording at its recorded pace on a fixed scene, so every difference from the offline pass follows from the structure rather than from the conditions. The stages of Chapters 3 to 5 needed no new mathematics, and the operations that read the future were each replaced by one that reads only the past. The reconstruction still followed the task: the elbow flexion of the two paths stays within a degree at the median, and the causal path rebuilt the wrist from the object through the natural gap as the offline pass did."),

# marker size: REVISION_2026-09-11_BRIEF.md section 3, fact 1 (45 mm desk
# and object markers; provenance eval/common/marker_size.py, E-009a).
(P, "The experiment does not demonstrate a working live system. No complete live-camera session was validated, and the step from the recording to the device leaves several items untested. A scene calibration built from a live stream needs the printed size of each marker before the first detection, including the 45 millimetre desk and object markers of Table 2.1. The calibration of this thesis was computed from a recording. The rest reference of the grip tracker assumes that the object begins at rest and in view, which a recording guarantees and a live session does not. The capture process has been exercised on a file and not on the sensor, since no session with the camera in place exists. The evaluation itself has no live form. The synthetic landmark removal of Section 7.4.1, which is evaluated against the unmasked reconstruction, and the manual wrist labels of Section 7.4.2 both read a finished recording, so a live run can report only what its own gates decided."),

(P, "The result is therefore feasibility evidence for the architecture under causal processing, and not a validated real-time system. Section 9.6 discusses the differences in Section 8.3 and the remaining validation gap in relation to the rest of the thesis."),
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == IMG:
        path = Path(item[1])
        assert path.exists(), f"missing figure: {path}"
        doc.add_picture(str(path), width=Inches(item[2]))
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

out = REPO / "writing" / "v8" / "condensed" / "Chapter_8_Real_Time_Feasibility.docx"
doc.save(out)
print("saved", out, "| items:", len(content))
