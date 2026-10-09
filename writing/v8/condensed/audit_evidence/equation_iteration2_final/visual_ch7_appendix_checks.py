"""Pin the final Ch7/Appendix assembly and the manually reviewed PDF pages.

Run with --render to reproduce the rasters before reviewing them. This script
checks artifact identity; its PASS does not replace the manual visual review.
It writes only the adjacent evidence directory.
"""
import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

from docx import Document
from lxml import etree

OUT = Path(__file__).resolve().parent
REPO = OUT.parents[4]
V8 = REPO / 'writing/v8'
PDF = V8 / 'Thesis_V8_Condensed.pdf'
ASSEMBLED = V8 / 'Thesis_V8_Condensed.docx'
RASTERS = OUT / 'visual_ch7_appendix_pages'
PAGES = {
    114: 'Equations 7.1-7.4',
    115: 'Table 7.1',
    116: 'Figure 7.1, Scene axes and floor elevation',
    117: 'Table 7.2',
    118: 'Figure 7.2, Scene axes and floor elevation',
    120: 'Equation 7.5 and its frame/provenance labels',
    121: 'Table 7.3 and Figure 7.3',
    122: 'Table 7.4',
    123: 'Figure 7.4',
    124: 'Bare-model scope and Section 7.3.3 spatial medians',
    125: 'Section 7.3.4 qualitative-example scope',
    126: 'Figure 7.5, all three paired rail frames and caption',
    127: 'Figure 7.6, all three paired handover frames and caption',
    128: 'Table 7.5 and Equation 7.6',
    129: 'Tables 7.6 and 7.7',
    130: 'Table 7.8',
    131: 'Table 7.9',
    167: 'Appendix C landmark indices and Table C.1',
    170: 'Appendix C wrist equation array and Table C.2',
    171: 'Equation D.1 and Table D.1',
    172: 'Nominal FOV qualification and depth-alignment prose',
    173: 'Figure D.1 image with its caption and Equation D.2',
    174: 'Appendix D object-depth equation array',
    179: 'Equation E.1 and rotation-operator scope',
    180: 'Table E.1, rotation matrix, trace-angle, axis and skew matrix',
    181: 'Squared skew matrix and recomposition matrix',
    182: 'Appendix F two rolling medians and two-pass amplitude prose',
    183: 'Figure F.1 image and amplitude-gain caption',
    185: 'Equation G.1 and positive-depth domain',
    188: 'Table H.1 eleven-nonzero-angle row',
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def math(doc):
    return [etree.tostring(node, method='c14n') for node in doc.element.xpath('.//m:oMath')]


def tables(doc):
    return [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]


def unique_start(whole, part):
    matches = [i for i in range(len(whole) - len(part) + 1)
               if whole[i:i + len(part)] == part]
    assert len(matches) == 1, matches
    return matches[0]


parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--render', action='store_true')
args = parser.parse_args()
pdf_hash_before = sha(PDF)
pdf_info = subprocess.run(['pdfinfo', str(PDF)], check=True, capture_output=True,
                          text=True).stdout
page_count = int(re.search(r'^Pages:\s+(\d+)$', pdf_info, re.M).group(1))
assembled = Document(ASSEMBLED)
assembly_math, assembly_tables = math(assembled), tables(assembled)
assembly_paragraphs = [p.text for p in assembled.paragraphs]
report = {
    'status': 'PASS',
    'scope': 'Artifact identity; page content was assessed manually in visual_ch7_appendix.md.',
    'pdf': {'path': PDF.relative_to(REPO).as_posix(), 'sha256': pdf_hash_before,
            'pages': page_count},
    'assembled_docx': {'path': ASSEMBLED.relative_to(REPO).as_posix(),
                       'sha256': sha(ASSEMBLED), 'total_math_objects': len(assembly_math)},
    'parts': {},
}
for name in ['Chapter_7_Evaluation', 'Appendices']:
    path = V8 / 'condensed' / (name + '.docx')
    part = Document(path)
    part_math, part_tables = math(part), tables(part)
    mi, ti = unique_start(assembly_math, part_math), unique_start(assembly_tables, part_tables)
    first_heading = part.paragraphs[0].text
    body_starts = [i for i, p in enumerate(assembled.paragraphs)
                   if p.style.name == 'Heading 1' and p.text == first_heading]
    assert len(body_starts) == 1
    body_start = body_starts[0]
    if name == 'Chapter_7_Evaluation':
        body_end = next(i for i in range(body_start + 1, len(assembled.paragraphs))
                        if assembled.paragraphs[i].style.name == 'Heading 1')
    else:
        body_end = len(assembled.paragraphs)
    captions = [p.text for p in part.paragraphs
                if re.match(r'^(?:Figure|Table) (?:7|[A-H])\.\d+\. ', p.text)]
    caption_positions = []
    for caption in captions:
        hits = [i for i in range(body_start, body_end) if assembly_paragraphs[i] == caption]
        assert len(hits) == 1, caption
        caption_positions.append(hits[0])
    assert caption_positions == sorted(caption_positions)
    report['parts'][name] = {
        'sha256': sha(path), 'canonical_math_objects': len(part_math),
        'math_block_start_one_based': mi + 1, 'math_block_end_one_based': mi + len(part_math),
        'tables': len(part_tables), 'table_block_start_one_based': ti + 1,
        'identical_ordered_captions': len(captions),
        'canonical_math_sha256': [hashlib.sha256(value).hexdigest() for value in part_math],
        'status': 'PASS',
    }

RASTERS.mkdir(exist_ok=True)
if args.render:
    for page in PAGES:
        subprocess.run(['pdftoppm', '-f', str(page), '-l', str(page), '-scale-to', '1600',
                        '-png', '-singlefile', str(PDF), str(RASTERS / f'p{page}')], check=True)
    (RASTERS / 'source_pdf.json').write_text(json.dumps({'sha256': pdf_hash_before}, indent=2) + '\n')
assert json.loads((RASTERS / 'source_pdf.json').read_text())['sha256'] == pdf_hash_before, \
    'Page rasters belong to a different PDF; rerender and review them.'
report['reviewed_pages'] = []
for page, content in PAGES.items():
    raster = RASTERS / f'p{page}.png'
    assert raster.is_file(), f'Missing {raster}; run with --render and review it.'
    report['reviewed_pages'].append({
        'physical_pdf_page': page, 'printed_page': page - 20, 'content': content,
        'raster': raster.relative_to(OUT).as_posix(), 'sha256': sha(raster),
    })
assert sha(PDF) == pdf_hash_before, 'PDF changed during evidence collection.'
report['render_command'] = ('pdftoppm -f PAGE -l PAGE -scale-to 1600 -png -singlefile '
                            'writing/v8/Thesis_V8_Condensed.pdf OUTPUT')
(OUT / 'visual_ch7_appendix_checks.json').write_text(json.dumps(report, indent=2) + '\n')
print('PASS: ordered math, tables and captions match the reviewed parts.')
print('PDF_SHA256', pdf_hash_before)
print('PAGES', page_count, 'REVIEW_RASTERS', len(PAGES))
