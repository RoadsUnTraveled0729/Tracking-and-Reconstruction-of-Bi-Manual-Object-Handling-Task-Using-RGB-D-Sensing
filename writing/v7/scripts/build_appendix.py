#!/usr/bin/env python3
"""Build writing/v7/Appendices.docx (V7 rewrite round).

Six appendices against the V7 table of contents, lettered in the order
they are first referenced in the body (Chapter 1, Section 1.4 lists them
in this order, and Chapters 2 and 3 reference them in place):

  A  RGB-D Data Acquisition        A.1 procedure, A.2 session file
                                   structure, A.3 replay and inspection
  B  Hardware and Software Platform
  C  Pose Landmark Detection       C.1 BlazePose, C.2 extraction,
                                   C.3 worked example
  D  Depth Sensing and Deprojection D.1 calibration, D.2 alignment,
                                   D.3 pixel-to-world, D.4 worked example
  E  Fiducial Marker Detection and Pose Recovery
                                   E.1 detection, E.2 planar pose and the
                                   two-solution ambiguity, E.3 Rodrigues,
                                   E.4 worked example
  F  Characteristics of the Candidate Filters

Appendices are method and reference material only: procedures,
definitions, derivations and worked examples. Validation, coverage and
accuracy figures belong to the evaluation chapter and are not repeated
here (skill_set/thesis-structure-rules.md, user rule 2026-08-28).

Every number printed here is produced by
writing/v7/scripts/appendix_numbers.py from the frozen artifacts of the
rail recording (recording_20260831_065553: 900 frames, 29.99 s,
640 by 480 at 30 frames per second; rail-only round 2026-09-01, which
replaced the R5 loop recording and its frame 270). The worked-example
frame is 533, the Chapter 3 worked-example frame, which carries all
three marker detections and all eight landmarks. Figures come from
make_appA_figs.py, make_appC_figs.py, make_appD_fig.py,
make_appE_figs.py and make_appF_fig.py.
Citations follow writing/v7/references.md ([1]-[57]).
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

from eqn import r, nor, sub, sup, frac, mat, d, eqArr, \
    add_display_eq, add_display_math

REPO = Path(__file__).resolve().parents[3]
FIG = REPO / "writing" / "v7" / "figures"

H1, H2, P, IMG, CAP, TBL, EQ = "h1", "h2", "p", "img", "cap", "tbl", "eq"
CODE = "code"   # monospace listing
MATH = "mth"    # displayed, unnumbered worked-example mathematics

# --- equation-fragment shorthands ---
fx, fy = sub(r("f"), r("x")), sub(r("f"), r("y"))
cx, cy = sub(r("c"), r("x")), sub(r("c"), r("y"))
wx = sub(d(r("w"), "[", "]"), r("×"))
wx2 = sup(sub(d(r("w"), "[", "]"), r("×")), r("2"))

content = [

# ==========================================================================
(H1, "Appendix A: RGB-D Data Acquisition"),

(H2, "A.1 Acquisition Procedure"),

(P, "The recording of this thesis begins as one continuous take captured with the RealSense D435 sensor [10] on the workstation of Appendix B. The physical arrangement is prepared first, in a fixed order. The working area is drawn on grid paper on the desk, the rail is placed inside it and set level, the wall marker is fixed at the back of the scene, the desk marker is placed at the world origin on the desk surface, and the object marker is mounted on one face of the cube. The camera then goes on its support at the desk edge, at about chest height. It is aimed so that the participant, the desk surface, the rail and all three markers stay in view for the whole take."),

(P, "The capture program requests two streams from the sensor: the colour image in the 8-bit-per-channel bgr8 format and the depth image in the 16-bit z16 format, both at 640 by 480 pixels and 30 frames per second. It directs the pipeline to write everything the device delivers into a single session file, the recording container described in Section A.2. Listing A.1 reproduces the essential lines. One check runs before any frame is kept: the profile the device negotiates is compared against the profile that was requested, and the capture stops if they differ. A camera attached to a slower bus quietly falls back to a lower resolution or frame rate, and a recording made under a fallback profile would not match the profile the rest of this thesis assumes."),

(CODE, """import sys
import pyrealsense2 as rs

pipeline = rs.pipeline()
config = rs.config()

config.enable_stream(rs.stream.color, 640, 480, rs.format.bgr8, 30)
config.enable_stream(rs.stream.depth, 640, 480, rs.format.z16, 30)
config.enable_record_to_file(str(bag_path))

profile = pipeline.start(config)

# the negotiated profile must equal the requested one, per stream
ok = True
for stream in (rs.stream.color, rs.stream.depth):
    sp = profile.get_stream(stream).as_video_stream_profile()
    got = (sp.width(), sp.height(), sp.fps())
    print(f"[info] negotiated {stream}: {got[0]}x{got[1]} @ {got[2]} fps")
    if got != (640, 480, 30):
        print(f"[ERROR] wanted 640x480 @ 30 fps for {stream} -- "
              f"check the camera is on a USB 3.x port")
        ok = False
if not ok:
    sys.exit(1)

while recording:
    frames = pipeline.wait_for_frames()   # written to the session file

pipeline.stop()                           # closes and finalizes the file"""),
(CAP, "Listing A.1. The core of the capture path. The pipeline writes every frame it delivers into the session file. The profile check runs right after the streams start, and it names the stream that fell back before stopping, so the operator can see which one to fix."),

(P, "With the streams running, the participant takes position at the desk, the recording starts, and the task is performed once: the participant lifts the cube from the desk onto the rail and slides it to the far end. Recording stops after the participant has stepped back from the desk. The recording of Chapter 2 holds 900 frames, the last of them timestamped 29.99 seconds after the first. Nothing is trimmed afterwards. The approach and retreat spans stay in the file, so the later stages segment the take from the recorded data."),

(H2, "A.2 Structure of the Recorded Bag File"),

(P, "The sensor stores each session as a single bag file, named for its container format: a rosbag container, the logging format of the Robot Operating System, in which a set of named topics carries messages held in compressed chunks. The layout separates description from data, and Figure A.1 draws it. Written once at the start are the file version, a description of the device, one description per sensor together with the value of every sensor option at record time, such as exposure, gain, laser power and the depth unit, and the calibration of each stream: the camera intrinsics of Appendix D and the pose of the stream relative to the device. After that the file is a sequence of per-frame messages. For every captured frame of each stream there is an image message carrying the pixel data and a metadata message carrying that frame's timestamps and counters."),

(IMG, FIG / "appA_fig_bag.png", 6.3),
(CAP, "Figure A.1. Structure of the recorded session file, read from the recording: the one-time description topics on the left, the per-frame image and metadata messages on the right."),

(P, "Replay reads the same structure back. The intrinsics used throughout this thesis come from these calibration topics, so a recording carries its own calibration with it and no separate calibration file has to be kept alongside it. Table A.1 lists the contents of the recording's file. The two streams together hold about 1.38 GB of raw pixel data across the 900 frames, and the compressed chunks bring the file on disk to 0.82 GB."),

(TBL, [["Property", "Value, read from the recording"],
       ["container", "rosbag 2.0, compressed chunks"],
       ["streams", "depth (z16, 16-bit) and colour (bgr8, 24-bit)"],
       ["resolution and rate", "640 by 480 at 30 frames per second, both streams"],
       ["duration", "29.99 seconds, first to last frame"],
       ["frames delivered on replay", "900 depth and 900 colour"],
       ["raw pixel payload", "1.38 GB"],
       ["file size on disk", "0.82 GB"],
       ["one-time topics", "file version, device and sensor descriptions, sensor options, per-stream intrinsics and pose"],
       ["per-frame topics", "image data and frame metadata, per stream"]]),
(CAP, "Table A.1. Contents of the recording's session file."),

(H2, "A.3 Replay and Visual Inspection"),

(P, "The vendor's viewer application opens the same session file and replays it as though a live sensor were attached, which makes it the quickest way to audit a recording before any processing runs. Its two-dimensional view shows the recorded colour and depth streams side by side. Its three-dimensional view deprojects every depth pixel through the recorded intrinsics into a metric point cloud, the mapping of equation (D.2) applied to the whole image rather than to eight landmark pixels, and lets the viewpoint orbit freely. Figure A.2 places the camera's own view of one moment of the recording beside the point cloud of the same moment in two colouring modes. In panel (b) the colour stream is draped over the measured geometry, so the desk, the participant and the markers appear as coloured surfaces standing in space. In panel (c) the same cloud is coloured by distance, which exposes the metric structure directly: the desk plane nearest the camera, the participant behind it, and the room surfaces furthest away. The flat image carries no explicit geometry, while the point cloud shows these surfaces from a viewpoint the camera never occupied, and that is the information the depth stream adds."),

(IMG, FIG / "appA_fig_cloud.png", 6.3),
(CAP, "Figure A.2. (a) The recorded colour frame of the worked-example moment. (b) The same moment deprojected into a metric point cloud with the colour stream draped over it. (c) The same cloud coloured by distance instead, near surfaces dark and far surfaces bright."),

(P, "A second, automated pass reads the raw outputs of the two extractors and reports what the recording contains before any modelling starts: the span of frames in which the participant is present, the segmentation of the take into approach, manipulation and retreat, the frames on which each marker and each landmark is available, and an inventory of the occlusion episodes with their start and end frames. One condition in that pass is a hard gate rather than a report. The pass runs on the extracted landmark and marker tables rather than on the images, and the desk marker defines the world origin, so if that marker is missing on any frame the pass stamps the recording as rejected and every stage after it is stopped."),

# ==========================================================================
(H1, "Appendix B: Hardware and Software Platform"),

(P, "Every processing result and every timing in this thesis was produced on one workstation, recorded here once; the cost measurements of Chapter 8 in particular should be read against this configuration. The division of labour is fixed throughout. The pose detector's neural network inference runs on the graphics card through the framework's GPU delegate, and everything else, including the depth sampling, the marker pose recovery, all of the kinematic mathematics and the filtering, runs on the processor in a single thread."),

(TBL, [["Component", "Value"],
       ["processor", "AMD Ryzen 9 7950X, 16 cores, 32 threads"],
       ["memory", "64 GB"],
       ["graphics card", "NVIDIA GeForce RTX 3080, 10 GB"],
       ["sensor", "Intel RealSense D435 RGB-D camera [10]"],
       ["operating system", "Ubuntu 24.04 LTS"],
       ["Python", "3.11.9"],
       ["RealSense library", "2.55.1"],
       ["MediaPipe", "0.10.15"],
       ["OpenCV", "4.10.0"],
       ["NumPy", "1.26.4"],
       ["SciPy", "1.13.1"],
       ["Unity editor", "6000.3.19f1"]]),
(CAP, "Table B.1. The workstation and the software versions behind every result in this thesis."),

# ==========================================================================
(H1, "Appendix C: Pose Landmark Detection"),

(H2, "C.1 MediaPipe BlazePose Overview"),

(P, "Learning-based pose detectors share a common structure. A detector network first locates the person, or a region of interest containing the person, in the image, and a regression network then predicts a fixed set of landmark coordinates inside that region. Systems differ mainly in the backbone networks, in the training data, and in whether they return 2D points, 3D points or a full body mesh. The requirements here are specific: per-joint image coordinates that can be paired with a depth buffer, a per-joint confidence signal that the occlusion handling can use, and real-time operation on ordinary hardware. The BlazePose model [43], deployed through the MediaPipe framework [28], meets all three. It runs on the device, it returns individual landmarks rather than an opaque mesh, and it attaches a visibility score to every landmark."),

(P, "BlazePose follows the two-stage pattern: a lightweight person detector localizes a region of interest, and a landmark network regresses 33 anatomical keypoints inside it, covering the face, the torso, the arms and the legs [43]. For each landmark the model returns normalized image coordinates in the range 0 to 1, a depth estimate inferred from appearance, and a visibility score between 0 and 1 expressing the model's confidence that the landmark is present and unoccluded. This work uses the heavy variant of the model in video mode, which tracks the person across frames rather than detecting from scratch on each one, with a single person expected in the scene and the detection, presence and tracking confidence thresholds of the tracker left at 0.30. The appearance-derived depth is not used. It carries substantial uncertainty during fast motion and object interaction [3], [13], which are the conditions of the recorded task. Metric depth comes instead from the aligned depth buffer of Appendix D, which is a measurement rather than an inference."),

(H2, "C.2 Landmark Coordinate Extraction"),

(P, "Of the 33 landmarks the model returns, this thesis uses eight, covering the torso and the two arms. Table C.1 lists them. Anatomical left and right always refer to the participant; because the participant faces the sensor, the participant's right side appears on the left of the image."),

(TBL, [["ID", "Landmark", "ID", "Landmark"],
       ["L11", "left shoulder", "L12", "right shoulder"],
       ["L13", "left elbow", "L14", "right elbow"],
       ["L15", "left wrist", "L16", "right wrist"],
       ["L23", "left hip", "L24", "right hip"]]),
(CAP, "Table C.1. The eight MediaPipe landmarks this work consumes, with the indices the model assigns them."),

(P, "Extraction proceeds in three steps per landmark and per frame. First the normalized coordinates are scaled to the colour image and truncated to whole pixels, so a normalized pair becomes the pixel u equal to the integer part of the horizontal coordinate times 640 and v the integer part of the vertical coordinate times 480. Second, the aligned depth image is sampled at that pixel by the 5 by 5 median of Section D.3. Third, the pixel and its depth are back-projected to a metric position in the camera frame by equation (D.2). Figure C.1 shows the eight landmarks drawn on the worked-example frame at the pixels the extractor sampled, and Figure C.2 shows the same eight as position vectors from the sensor origin, which is the form the back-projection produces."),

(IMG, FIG / "appC_fig_landmarks.png", 5.2),
(CAP, "Figure C.1. The eight upper-body landmarks on the worked-example frame of the recording, drawn at the pixels the extractor sampled, with the segments between them. The participant stands at the desk with the right hand on the cube on the rail."),

(IMG, FIG / "appC_fig_vectors.png", 6.0),
(CAP, "Figure C.2. The same eight landmarks as metric position vectors from the sensor origin, with the camera axes drawn at the origin: x to image right, y down, z forward into the scene."),

(P, "One property of the detector shapes the whole design. The model returns coordinates for all 33 landmarks whether it can see them or not, so an occluded wrist still receives a position, and that position is a guess. The visibility score is the only signal that separates a measurement from a guess. The extractor therefore gates on it: a landmark whose visibility falls below 0.50 is written as an empty record rather than as a coordinate. The same applies when the 5 by 5 depth window holds no valid return. Every exported landmark carries a code recording which of three cases produced it, a measured sample, a sample blocked by low visibility, or a sample blocked by missing depth, so that later stages can tell a gap from a value. Converting wrong data into declared missing data, and then handling missing data explicitly, is a rule the whole system follows: the recovery of Chapter 5 and the marker gap handling of Chapter 4 both depend on gaps being declared rather than filled at the source."),

(H2, "C.3 Worked Example"),

(P, "Take the right wrist, landmark L16, on frame 533 of the recording, at 17.780 seconds into the take. The frame is the one drawn in Figure C.1, part way along the slide, with the right hand on the cube. The detector reports the landmark with a visibility of 0.5778, above the 0.50 gate. The extractor stores the sampled pixel and the metric position rather than the normalized pair, so the normalized coordinates are given here as the interval that truncates to the stored pixel:"),
(MATH, eqArr(
    sub(r("x"), r("n")) + r(" ∈ [0.3266, 0.3281)   →   u = 209"),
    sub(r("y"), r("n")) + r(" ∈ [0.6417, 0.6438)   →   v = 308"))),
(P, "The 5 by 5 median depth at pixel (209, 308) is 1.060 metres. Back-projection through equation (D.2), with the intrinsics of Table D.1, gives"),
(MATH, eqArr(
    r("x") + r(" = 1.060 · ") + frac(r("209 − 323.93756"), r("607.56128")) + r(" = −0.2005"),
    r("y") + r(" = 1.060 · ") + frac(r("308 − 248.01741"), r("607.01508")) + r(" = +0.1047"))),
(P, "so the right wrist sits at (-0.201, 0.105, 1.060) metres in the camera frame: about 20 centimetres to the left of the optical axis in the image, 10 centimetres below it, at a depth of 1.06 metres. Table C.2 carries the same three steps through the other two landmarks of the right arm on the same frame. The upper arm and the forearm read as plausible segments from these three positions, and Section 3.5 continues from the filtered form of these numbers into the kinematic solve."),

(TBL, [["Landmark", "Visibility", "Pixel", "Depth (m)", "Camera-frame position (m)"],
       ["L12 right shoulder", "0.9987", "(248, 93)", "1.311", "(-0.164, -0.335, 1.311)"],
       ["L14 right elbow", "0.7031", "(222, 225)", "1.196", "(-0.201, -0.045, 1.196)"],
       ["L16 right wrist", "0.5778", "(209, 308)", "1.060", "(-0.201, 0.105, 1.060)"]]),
(CAP, "Table C.2. The right-arm landmarks of the worked-example frame, from the visibility score through the sampled pixel and depth to the back-projected position. These are the raw extractor positions, before the filtering of Section 2.4; the filtered values the kinematic solver consumes are listed in Table 3.1."),

# ==========================================================================
(H1, "Appendix D: Depth Sensing and Deprojection"),

(P, "The Intel RealSense D435 is a stereo-depth camera that produces a per-pixel depth measurement alongside a colour image [10]. Depth is computed by triangulating between two infrared imagers separated by a known baseline, assisted by an infrared projector that adds texture to surfaces too uniform for stereo matching to work on. Both streams are recorded at 640 by 480 pixels and 30 frames per second. The participant stands at a working distance of about 1.3 metres, and the scene spans depths from about 0.53 metres at the desk marker to about 3.2 metres at the wall marker, inside the sensor's usable range. Two properties of stereo depth carry through everything the main body does with it: the depth noise grows roughly with the square of the distance [37], [38], and one pixel of image displacement corresponds to a lateral distance on the object that also grows with distance."),

(H2, "D.1 RGB Camera Calibration"),

(P, "Calibration establishes the relationship between a 3D point in the scene and the pixel it lands on. The pinhole model describes that relationship through four numbers: the focal lengths in pixels along the horizontal and vertical image axes, and the principal point, the pixel where the optical axis meets the image [51]. Figure D.1(a) draws the geometry. A point p with coordinates x, y and z in the camera frame projects to the pixel"),
(EQ, r("u") + r(" = ") + fx + r(" ") + frac(r("x"), r("z")) + r(" + ") + cx + r(",        ") +
     r("v") + r(" = ") + fy + r(" ") + frac(r("y"), r("z")) + r(" + ") + cy, "D.1"),
(P, "Reading the first of the two in words: the point's horizontal offset is divided by its depth, which captures the fact that the same offset looks smaller when it is farther away; the ratio is scaled by the focal length to turn a geometric angle into a pixel count; and the principal point shifts the result from optics-centred coordinates onto the pixel grid, whose origin sits at the top-left corner of the image. The vertical equation works the same way."),

(P, "Calibration normally estimates these numbers by imaging a planar checkerboard from several viewpoints and fitting the model to the observed corners, following Zhang's method [51]. No manual calibration was performed in this work. The vendor's library exposes the factory calibration held in each individual device, and the recording carries those values in its own calibration topic, so the intrinsics used are the ones the sensor itself reports for the colour stream at 640 by 480. Table D.1 lists them."),

(TBL, [["Parameter", "Value", "Meaning"],
       ["fx", "607.56128 px", "focal length along the horizontal image axis"],
       ["fy", "607.01508 px", "focal length along the vertical image axis"],
       ["cx", "323.93756 px", "principal point, horizontal (image centre 320)"],
       ["cy", "248.01741 px", "principal point, vertical (image centre 240)"],
       ["distortion model", "inverse Brown-Conrady", "the model the stream declares"],
       ["distortion coefficients", "all zero", "the recorded values for this stream"]]),
(CAP, "Table D.1. Factory colour-camera intrinsics of the sensor, read from the calibration topic of the recording."),

(P, "Three observations connect the table to the model. The horizontal focal length is 607.56 pixels and the vertical 607.02, so the two are nearly equal and the pixels are very nearly square. The principal point sits a few pixels away from the geometric image centre, and the model carries it as a parameter of its own, so the deprojection of Section D.3 measures every pixel offset from that point rather than from the centre of the image. And the distortion coefficients the stream reports are all zero, so the general deprojection routine of the vendor's library reduces to the closed form of equation (D.2) for this data. Combining the focal lengths with the image size gives the field of view of the colour camera, 55.55 degrees across and 43.15 degrees vertically, the number the camera placement of Chapter 2 was chosen against."),

(H2, "D.2 Depth-to-RGB Alignment"),

(P, "The depth imager and the colour camera sit at a small physical offset inside the sensor housing, and each has its own intrinsics [10]. A raw depth pixel at a given location therefore does not observe the same scene point as the colour pixel at the same location. The vendor's library resolves this with a depth-to-colour alignment step: each depth pixel is lifted to a 3D point through the depth camera's intrinsics, moved into the colour camera's frame through the known transformation between the two imagers, and projected back onto the colour image grid through the colour intrinsics of Table D.1. Figure D.1(b) draws the arrangement."),

(IMG, FIG / "appD_fig_alignment.png", 6.3),
(CAP, "Figure D.1. (a) The pinhole geometry the intrinsics describe: a scene point, the ray from it to the camera centre, and the pixel where that ray crosses the image plane. (b) Depth-to-colour alignment: the depth buffer is reprojected into the colour camera's geometry so that both buffers share one pixel grid."),

(P, "Both buffers are 640 by 480 here, so alignment produces a depth image on exactly the same pixel grid as the colour image. After it, the colour value and the depth value at one pixel describe the same physical point. That aligned pair is the only input the rest of the pipeline sees. Landmark detection and marker detection both happen in colour pixel coordinates, and the same coordinates index the aligned depth buffer directly, with no further correspondence step anywhere downstream."),

(H2, "D.3 Pixel-to-World Mapping"),

(P, "Given the aligned depth image and the colour intrinsics, a pixel with a valid depth is lifted to a 3D point in the camera frame by running the projection backwards:"),
(EQ, r("x") + r(" = ") + r("z") + r(" ") + frac(r("u") + r(" − ") + cx, fx) + r(",      ") +
     r("y") + r(" = ") + r("z") + r(" ") + frac(r("v") + r(" − ") + cy, fy) + r(",      ") +
     r("p") + r(" = ") + d(r("x, y, z")), "D.2"),
(P, "The mapping is exact because the depth supplies the one piece of information the forward projection throws away. The resulting camera frame is right-handed, with x to image right, y downward and z forward into the scene along the optical axis, and the depth measurement is the z coordinate itself. The depth image stores integer counts rather than metres, and the recording carries the conversion in its sensor options: one count is one millimetre for this device. Converting the camera-frame point into the left-handed convention Unity uses is a separate fixed axis flip, developed in Section 3.1."),

(P, "One practical detail matters as much as the equation. Reading depth at a single pixel fails in two ways: the pixel can be a hole, a zero return, which is common at object edges and on dark or reflective surfaces, and it can carry single-pixel speckle. The pipeline therefore never samples one pixel. It takes the median of the nonzero depths in a 5 by 5 window centred on the target pixel, so up to half the window may be outliers before the median moves, and a zero return counts as an absent sample rather than as a depth of zero metres. When the whole window holds no valid return the sample is reported as missing, never guessed."),

(H2, "D.4 Worked Example"),

(P, "Take the object marker on frame 533 of the recording, the frame of Appendix C. The marker detector of Appendix E places the centre of the marker at pixel (200, 341), and the 5 by 5 median depth at that pixel is 0.986 metres. Substituting into equation (D.2) with the intrinsics of Table D.1:"),
(MATH, eqArr(
    r("x") + r(" = 0.986 · ") + frac(r("200 − 323.93756"), r("607.56128")) + r(" = 0.986 · (−0.20399) = −0.2011"),
    r("y") + r(" = 0.986 · ") + frac(r("341 − 248.01741"), r("607.01508")) + r(" = 0.986 · (+0.15318) = +0.1510"))),
(P, "so the marker centre sits at (-0.201, 0.151, 0.986) metres in the camera frame: about 20 centimetres to the left of the optical axis in the image, 15 centimetres below it, at just under a metre of depth. Running the result forward through equation (D.1) returns the pixel (200.0, 341.0), the pixel the detection started from: the two equations are inverses. The same two lines run for every landmark of Appendix C and for every marker centre of Appendix E, on every frame of the recording."),

# ==========================================================================
(H1, "Appendix E: Fiducial Marker Detection and Pose Recovery"),

(P, "The landmark detections of Appendix C measure the participant. The objectives of Chapter 1 also require a measurement of the manipulated object, expressed in a frame that does not move with the camera, and a scene geometry against which the reconstruction can be read. Fiducial markers supply both. They are printed patterns whose geometry is known exactly, so their pose follows from a single image with no training data and no appearance model. The markers here play two roles: two static markers define the world frame and the gravity direction, and one marker rides on the carried cube and supplies the object trajectory. Among the available fiducial systems, ArUco markers combine detection robustness, low computational cost and built-in identity encoding [29], and the relationship between marker size, viewing distance and recovered pose is characterized in [30], which informed the marker sizes of Chapter 2. The markers used here come from the 5 by 5 OpenCV dictionary, which encodes each identity in a grid of 25 black and white cells inside a black border."),

(H2, "E.1 Detection and Sub-Pixel Corners"),

(P, "Detection runs in four stages. Adaptive thresholding first turns the image into binary regions by comparing each pixel with the mean of its own neighbourhood, which keeps the detector working under uneven lighting. Contour extraction then finds candidate quadrilaterals. The interior of each candidate is perspective-corrected and its cell pattern matched against the dictionary, which either rejects the candidate or returns a marker identity. Finally the four corner positions are refined below the pixel grid. Figure E.1 shows the result on the worked-example frame: all three markers of the scene are detected, each outlined with its four corners in detection order and its recovered axes drawn at its centre."),

(IMG, FIG / "appE_fig_detect.png", 5.6),
(CAP, "Figure E.1. Marker detection on the worked-example frame of the recording. (a) The full frame with the three detected markers outlined and the recovered axes drawn at each centre. (b) to (d) Each marker enlarged, with its four refined corners numbered in detection order."),

(P, "Detection itself is standard and is not a contribution of this work [29]. The corner refinement step is a design decision. The pose mathematics of the next section consumes exactly four corner positions per marker, so what the corners can resolve bounds what the pose can resolve. With the detector's default settings the corner coordinates land on whole pixels, and a marker that sits still in the image is then reported at identical corners frame after frame, which looks like perfect stability and is in fact the detector re-snapping to the same grid positions. The pipeline therefore enables the corner refinement based on the AprilTag method [48], which fits straight lines to the marker's four edges in the local image and takes each corner as the intersection of the two lines that meet there, so the corner lands between the grid points. Every frame of the recording uses the refined corners."),

(P, "One further quantity is recorded at detection time: the median of a 5 by 5 patch of aligned depth values at each marker's centre pixel. The pose recovery of the next section never reads this number. It is kept precisely because it comes from the other measurement principle, the stereo depth stream rather than the marker geometry, and Chapter 9 uses the pair. Frames on which a marker is not detected are written as empty records and never interpolated at this stage, the same convention the landmark branch follows."),

(H2, "E.2 Planar Pose Recovery and the Two-Solution Ambiguity"),

(P, "A detected marker gives four corner points in the image, and the printed marker gives the same four corners in its own frame. For a printed side length L the corners sit at (-L/2, L/2, 0), (L/2, L/2, 0), (L/2, -L/2, 0) and (-L/2, -L/2, 0), listed in the order the detector returns them, all in the marker plane where the third coordinate is zero. Recovering the pose means finding the rotation R and the translation t that carry these four known points onto the four detected pixels through the projection of equation (D.1). The detected corners carry noise, so no pose lands on them exactly, and the problem is posed as a minimization: choose R and t so that the sum over the four corners of the squared image distance between the detected corner and the projected corner is as small as possible. Counting unknowns shows why four corners are enough. The pose has six unknowns, three for the rotation and three for the translation, and each corner contributes two equations, one per image axis, which gives eight equations for six unknowns."),

(P, "A planar target has a structural weakness that a target with depth does not have, and that weakness shaped the pipeline, so it is developed here in full. Because the four corners lie in one plane, the image observation constrains only the mapping between that plane and the image. A plane seen nearly face-on can be tilted slightly toward the camera or slightly away from it by the same amount and produce very nearly the same four pixels. The two readings are mirror images of the marker normal about the viewing ray. The solver used here, the infinitesimal plane-based method [52], makes the situation explicit: it returns both locally optimal poses, called lobes below, rather than silently choosing one. Figure E.2 draws the geometry in panel (a) and, in panel (b), the pair the solver returns for one recorded detection, the wall marker on the worked-example frame. The marker is small in the image and seen nearly face-on, so the branch records its near-degenerate flag there, as on almost every frame of the recording. The two normals of the pair differ by 36.37 degrees, and both readings reproject onto the four detected corners with a largest corner gap of 0.22 pixels on sides averaging 30.03 pixels."),

(IMG, FIG / "appE_fig_lobes.png", 6.3),
(CAP, "Figure E.2. (a) A nearly face-on plane admits two pose readings, tilted toward the camera or away from it, that are mirror images about the viewing ray, drawn here at the tilt of the recorded pair of panel (b). (b) Both readings the solver returns for the wall marker on the worked-example frame, reprojected into image pixels over the detection they come from, with the corner where they separate most magnified in the inset. The outlines nearly coincide, so the four detected corners alone barely separate the readings."),

(P, "How far apart the two lobes sit is itself a conditioning signal, and the pipeline switches its policy on that separation rather than on how well either lobe reprojects. A marker close to the camera shows visible perspective across its own width, the two lobes then sit far apart, and the wrong one lands its corners visibly away from where they were detected, so the closer-fitting lobe is taken. A marker whose projection is nearly an affine copy of the printed square gives almost no perspective cue, the two lobes collapse toward each other, and the four corners no longer separate them. The threshold between the two regimes is a lobe separation of 40 degrees. The pair drawn in Figure E.2 sits just below it, so its sub-pixel outline gap leaves the corner fit without a usable preference and the discriminators below decide instead."),

(P, "Below that separation the discriminator becomes physical and temporal instead. On the first such frame the lobe is chosen by plumbness: the wall is vertical, so the lobe whose in-plane up axis better agrees with an independent estimate of the gravity direction is taken. That estimate is seeded from the plane fitted to the depth returns around the desk marker, and it only has to be right to within about 10 degrees, because the two lobes sit far enough apart for a coarse direction to separate them. On every later frame the lobe closest to a reference pose is kept. The reference updates only when the chosen pose lies within 30 degrees of it, so a single corrupted detection, an arm clipping the marker's edge for one frame, cannot poison the frames that follow, and after five consecutive frames far from the reference the tracker re-seeds on the current detection. A per-frame flag records how each pose was resolved, so later stages can tell a lobe decided by fit from one decided by continuity."),

(H2, "E.3 From the Rodrigues Form to a Rotation Matrix"),

(P, "The pose solver returns its rotation in the compact Rodrigues form: a single 3-vector whose direction is the rotation axis and whose length is the rotation angle in radians. The rest of the system works in rotation matrices, so the vector is converted on receipt, and the conversion is written out here because the same formula reappears in the object-track cleaning of Chapter 4."),

(P, "Write w for the unit axis and theta for the angle, so the Rodrigues vector is theta times w. The first ingredient is the skew-symmetric matrix of w, written with the cross subscript, built from the components of w as the matrix with rows (0, -w3, w2), (w3, 0, -w1) and (-w2, w1, 0). That matrix implements the cross product: multiplying it by any vector v gives w cross v. Applying it twice removes from v the part that lies along the axis and reverses the rest, which is the second ingredient the formula needs. The conversion is Rodrigues' rotation formula [42]:"),
(EQ, r("R") + r(" = ") + r("I") + r(" + ") + nor("sin") + d(r("θ")) + r(" ") + wx +
     r(" + ") + d(r("1 − ") + nor("cos") + d(r("θ"))) + r(" ") + wx2, "E.1"),
(P, "Reading the three terms in words: the identity term starts from the vector unchanged; the sine term adds the part perpendicular to the axis turned a quarter turn, scaled by the sine of the angle, which produces the circular sweep; and the third term pulls that perpendicular part toward its rotated position while leaving the part along the axis untouched, so points on the axis do not move at all. At an angle of zero both factors vanish and R is the identity, as it must be. The formula holds for every angle, uses no small-angle approximation, and produces a proper rotation matrix by construction. Its inverse, recovering the axis and the angle from a matrix, is the matrix logarithm the gap bridging of Chapter 4 uses, and the two together let the pipeline move between the compact form and the matrix form without introducing quaternions anywhere in between."),

(H2, "E.4 Worked Example: Frame 533 Detections"),

(P, "Frame 533 of the recording carries all three markers, the frame drawn in Figure E.1. Table E.1 collects what the marker branch records for each of them: the printed side length, the centre pixel, the recovered translation in the camera frame, the range that translation implies, and the independent 5 by 5 median depth at the same centre pixel. The wall marker is printed at 150 millimetres because it sits farthest from the camera. The black square of the desk and object markers is printed at 45 millimetres, while the size declared to the detector is 50 millimetres. Declaring a marker larger than it is printed scales every recovered translation by the same ratio, so every recovered translation is corrected by a factor of 0.90 before any pose is used. Table E.1 lists the scaled values. The object marker occupies one face of the 70 millimetre cube."),

# src: frame 533 of eval/output/recording_20260831_065553_aruco_raw_scaled.csv
#      (the scaled track every later stage reads); ranges are the norms of
#      the tabulated translations
(TBL, [["Marker", "Role", "Black square (mm)", "Centre pixel", "Translation t (m)", "Range (m)", "Depth sample (m)"],
       ["id 0", "wall", "150", "(119, 155)", "(-1.031, -0.470, 3.055)", "3.258", "3.185"],
       ["id 2", "desk", "45", "(350, 455)", "(0.023, 0.187, 0.552)", "0.583", "0.532"],
       ["id 1", "object", "45", "(200, 341)", "(-0.209, 0.156, 1.023)", "1.056", "0.986"]]),
(CAP, "Table E.1. The three marker detections of the worked-example frame, in the order of the panels of Figure E.1. The translation and the range come from the corner geometry alone and carry the scaling of the printed black-square size, so they are the values the rest of the pipeline reads; the depth sample in the last column comes from the stereo depth stream and is recorded alongside without entering the pose recovery."),

(P, "Take the object marker through the rotation conversion. The pose recovery returns the rotation"),
(MATH, r("R") + r(" = ") + mat([
    [r("0.998259"), r("−0.008185"), r("0.058415")],
    [r("−0.005833"), r("−0.999170"), r("−0.040308")],
    [r("0.058697"), r("0.039897"), r("−0.997478")]])),
(P, "whose determinant is one, so it is a proper rotation and not a reflection. The angle follows from the trace, which is the sum of the three diagonal entries:"),
(MATH, eqArr(
    nor("tr") + r(" R = 0.998259 − 0.999170 − 0.997478 = −0.998389"),
    r("θ") + r(" = ") + nor("acos") + d(frac(r("−0.998389 − 1"), r("2"))) + r(" = 177.7007°"))),
(P, "and the axis follows from the skew-symmetric part of the matrix, each component being one off-diagonal difference divided by twice the sine of the angle:"),
(MATH, r("w") + r(" = (+0.9996, −0.0035, +0.0293)")),
(P, "The marked face of the cube is turned almost 180 degrees about an axis close to the camera's x axis. The cube sits on the rail with its marked face square to the camera. Building the skew-symmetric matrix from those three components gives"),
(MATH, wx + r(" = ") + mat([
    [r("0.000000"), r("−0.029302"), r("−0.003507")],
    [r("0.029302"), r("0.000000"), r("−0.999564")],
    [r("0.003507"), r("0.999564"), r("0.000000")]])),
(P, "and multiplying it by itself gives the second ingredient,"),
(MATH, wx2 + r(" = ") + mat([
    [r("−0.000871"), r("−0.003506"), r("0.029290")],
    [r("−0.003506"), r("−0.999988"), r("−0.000103")],
    [r("0.029290"), r("−0.000103"), r("−0.999141")]])),
(P, "The two scalar factors of equation (E.1) at this angle are the sine, 0.0401, and one minus the cosine, 1.9992. Assembling the three terms:"),
(MATH, r("R") + r(" = ") + r("I") + r(" + 0.0401 ") + wx + r(" + 1.9992 ") + wx2 + r(" = ") + mat([
    [r("0.998259"), r("−0.008185"), r("0.058415")],
    [r("−0.005833"), r("−0.999171"), r("−0.040308")],
    [r("0.058697"), r("0.039897"), r("−0.997479")]])),
(P, "which reproduces the matrix the conversion started from to within 0.00001 in every entry (two entries differ in the sixth decimal). The axis is hardest to read back out of a matrix at an angle this close to 180 degrees, because the sine in the denominator, 0.0401 here, becomes small. The round trip still returns the input with no special handling. The third column of the recovered rotation is the direction of the marker's own outward normal expressed in the camera frame, (0.058, -0.040, -0.997), a vector pointing almost straight back at the camera, which matches the marked face of the cube facing the camera in Figure E.1. That pose is the input the world anchoring of Chapter 4 carries forward."),

# ==========================================================================
(H1, "Appendix F: Characteristics of the Candidate Filters"),

(P, "The filtering stage of Section 2.4 can select any of four smoothers, and it runs the same two stages ahead of all four. This appendix sets out what each candidate computes and how each behaves, so that the choice made in Chapter 2 and the real-time comparison of Chapter 8 can be read against one description. Every parameter quoted here is the one the filter run of the recording used, and the sampling rate throughout is the recorded frame rate of 29.98 samples per second."),

(P, "The two stages ahead of the smoothers are fixed. A rolling Hampel test flags a sample that deviates from the median of a 7-frame window centred on it by more than three robust standard deviations, the robust standard deviation being the median absolute deviation of the same window scaled by the usual factor of 1.4826, with a floor of 2 centimetres below which nothing is flagged [53], [54]. A flagged sample becomes a gap. Gaps of at most 5 frames are then filled by linear interpolation between the valid samples that bound them, and longer gaps stay empty and are passed on as declared gaps. Leading and trailing gaps are not backfilled. Each smoother then runs on each unbroken run of valid samples separately, so no filter ever smooths across a gap it cannot see into."),

(P, "The zero-phase Butterworth low-pass is the candidate the results use. The Butterworth family is the standard low-pass design of the signal processing literature, chosen for a passband that is flat rather than rippled, so the motion band passes with its shape intact [55], [56]. A fourth-order filter with a 3 hertz cutoff is applied twice over each run, once forward and once backward [57]. The second pass cancels the phase shift the first pass introduces, so no event moves in time, and the squared gain of the pair is the curve drawn in Figure F.1(b). Running a filter backward requires the samples that follow the current one, so this candidate exists only offline."),

(P, "The Savitzky-Golay smoother fits a second-order polynomial by least squares to a sliding window of 9 samples and takes the value of that polynomial at the centre of the window [46]. Fitting a curve rather than averaging preserves a peak that a moving average would flatten, which the shape following in Figure F.1(a) shows, and the same fit gives the derivative at no extra cost. The trade is that a polynomial fit gives every sample in the window a weight, so a single far-off sample still pulls the result."),

(P, "The rolling median replaces each sample by the median of a 9-sample window. Because the median ignores the ordering of the values and follows the middle one, a single-frame spike inside the window has no effect at all, which is the behaviour Figure F.1(a) shows at the spike. The same property costs it on smooth motion: the output holds a value while the window's middle sample holds, then steps, so a steady move comes out as a staircase rather than a ramp."),

(P, "The One Euro filter is the causal candidate, and it is the one the real-time pipeline of Chapter 8 can run, because it needs only the samples already received [40]. It is a first-order low-pass whose cutoff is not fixed: the filter keeps a smoothed estimate of the rate of change of the signal and raises its cutoff in proportion to that estimate, through a minimum cutoff of 0.05 hertz and a coefficient of 1.0. Holding still, the cutoff stays low and the jitter is heavily suppressed; while the signal is changing, the cutoff rises and the filter follows. Being causal, it cannot avoid lag, which Figure F.1(a) shows as the late arrival of its output at the end of the move."),

(P, "Figure F.1(a) puts the four side by side on a constructed test signal rather than on recorded data: a position profile that holds still, moves, and holds again, with broadband noise and one single-frame spike added. Constructing the signal makes the comparison readable, because the motion underneath is then known exactly and can be drawn beside the four outputs, which no recording allows. The filter parameters in the panel are listed in Table F.1."),

(IMG, FIG / "appF_fig_filters.png", 6.3),
(CAP, "Figure F.1. (a) The four candidates on one constructed test signal, not recorded data: a position profile that holds still, moves, and holds again, with broadband noise and one single-frame spike added. Shape following, spike handling and lag are all visible in the one panel. (b) The gain against frequency of the two linear candidates at the recorded frame rate, with the 3 hertz cutoff marked and the band that carries deliberate upper-body motion shaded."),

(TBL, [["Candidate", "What it computes", "Phase", "Single-frame spike", "Parameters used"],
       ["Butterworth, zero phase", "low-pass, applied forward and backward", "none", "spread over the filter length", "order 4, cutoff 3 Hz"],
       ["Savitzky-Golay", "least-squares polynomial over a sliding window", "none", "spread over the window", "window 9, order 2"],
       ["rolling median", "median of a sliding window", "none", "removed", "window 9"],
       ["One Euro", "first-order low-pass with a rate-dependent cutoff", "lag", "partly passed through", "minimum cutoff 0.05 Hz, coefficient 1.0"]]),
(CAP, "Table F.1. The four candidate smoothers, with their parameters. The first three need the whole recording or a window around each sample; only the last runs on the samples already received."),

(P, "Two properties separate the four for this application. The first is phase: the three offline candidates leave event timing where it is, either because the forward and backward passes cancel or because the window is centred, while the causal candidate cannot. The second is what happens to a spike that the Hampel stage did not catch, which the median removes outright and the three others spread across their window. The despiking stage runs ahead of all four for exactly that reason, so the smoother is asked only to handle broadband noise, and the choice among the four then rests on phase and on whether the pipeline runs offline or live."),

]

# --------------------------------------------------------------------------
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == EQ:
        add_display_eq(doc, item[1], item[2])
    elif kind == MATH:
        add_display_math(doc, item[1])
    elif kind == CODE:
        for line in item[1].split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line if line else " ")
            run.font.name = "Courier New"
            run.font.size = Pt(9)
    elif kind == IMG:
        path = Path(item[1])
        assert path.exists(), f"missing figure: {path}"
        doc.add_picture(str(path), width=Inches(item[2]))
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row_ in enumerate(rows):
            for j, c in enumerate(row_):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

out = REPO / "writing" / "v7" / "Appendices.docx"
doc.save(out)
print("saved", out, "| items:", len(content))
