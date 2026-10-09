#!/usr/bin/env python3
"""Build writing/v7/Thesis_V7_TOC.docx: the V7 rewrite table of contents.

Rewrite round of 2026-08-26 (after the R5 experiment freeze). Structure:
the supervisor-shaped Chapters 1-3 and the appendix set are kept; the
middle and back are restructured around the completed R5 work:
  - New Chapter 5 "Pose Recovery During Tracking Failure" (the
    contribution; was only future-work item 8.3.5 in the old TOC).
  - Evaluation rebuilt around the designed 8-waypoint path (replaces
    the Lifting/Rotating/Passing scenario sections).
  - Old Ch 8 + 9 merged into one closing chapter; "Data Fusion" and all
    honesty-framed sections dropped; old 2.2 comms preliminaries moved
    into the integration chapter.
  - Appendices reordered by first appearance in the body (provisional
    lettering, verified against the built text before shipping - see
    skill_set/thesis-structure-rules.md).
  - No question-style titles; no speed information anywhere (the task
    is defined as slow movement).
"""
from docx import Document
from docx.shared import Pt, Inches

TITLE = "TRACKING AND RECONSTRUCTION OF BI-MANUAL OBJECT HANDLING TASK USING RGB-D SENSING"

# (level, text): 0 = top-level line (bold), 1 = section, 2 = subsection
entries = [
(0, "Front Matter"),
(1, "Title Page"),
(1, "Approval"),
(1, "Abstract"),
(1, "Acknowledgments"),
(1, "Table of Contents"),
(1, "List of Figures"),
(1, "List of Tables"),
(1, "List of Equations"),

(0, "Chapter 1: Introduction"),
(1, "1.1 Motivation"),
(1, "1.2 Problem Statement and Literature Review"),
(2, "1.2.1 Hand and Object Tracking Approaches"),
(2, "1.2.2 Bimanual Manipulation in Motion Analysis"),
(2, "1.2.3 Graphical Reconstruction in Robotics and VR"),
(2, "1.2.4 Gap Analysis and Motivation for This Work"),
(1, "1.3 Research Objectives"),
(1, "1.4 Organization of the Thesis"),

(0, "Chapter 2: Experimental Setup and Data Acquisition"),
(1, "2.1 The Recording Setup"),
(1, "2.2 Pose Landmarks from MediaPipe"),
(1, "2.3 Object Pose from ArUco Markers"),
(1, "2.4 Signal Filtering"),

(0, "Chapter 3: Kinematic Modeling"),
(1, "3.1 Coordinate Systems and Transformations"),
(1, "3.2 Torso Reference Frame"),
(2, "3.2.1 Landmark Selection"),
(2, "3.2.2 Derivation of the Torso Frame"),
(1, "3.3 The Arm Model"),
(2, "3.3.1 Kinematic Schematic and Solved Angles"),
(2, "3.3.2 Definition of Shoulder, Elbow, and Wrist"),
(2, "3.3.3 Establishing Local Arm Frames"),
(1, "3.4 Joint Angle Computation"),
(2, "3.4.1 Torso Pose"),
(2, "3.4.2 Shoulder Orientation"),
(2, "3.4.3 Elbow Flexion and Extension"),
(2, "3.4.4 The Wrist as a Tracked Point"),
(1, "3.5 Worked Example"),

(0, "Chapter 4: Object Tracking and Scene Reconstruction"),
(1, "4.1 Static Scene Calibration"),
(1, "4.2 World Anchoring and Camera Placement Invariance"),
(1, "4.3 Filtering and Cleaning the Object Track"),
(1, "4.4 Worked Example"),

(0, "Chapter 5: Pose Recovery During Tracking Failure"),
(1, "5.1 Landmark Tracking Failures"),
(1, "5.2 Grip Episodes and the Hand-Object Offset"),
(1, "5.3 Wrist Recovery from the Object Pose"),
(1, "5.4 Elbow Recovery by Two-Link Inverse Kinematics"),
(1, "5.5 Torso Repair on Camera Rays"),
(2, "5.5.1 The Failure Model"),
(2, "5.5.2 The Ray Repair and the Depth Memory"),
(2, "5.5.3 Gate One: Rigidity of the Torso"),
(2, "5.5.4 Gate Two: Depth Jump"),
(2, "5.5.5 Gate Three: Shoulder Depth"),
(2, "5.5.6 Rebuilding the Pair"),
(2, "5.5.7 What the Repair Cannot Do"),
(1, "5.6 Output States of the Solver"),

(0, "Chapter 6: System Integration and Graphical Reconstruction"),
(1, "6.1 Communication Architecture"),
(1, "6.2 Coordinate System Conversion"),
(1, "6.3 Avatar Reconstruction in Unity"),
(1, "6.4 Scene Reconstruction in Unity"),

(0, "Chapter 7: Evaluation"),
(1, "7.1 Evaluation Protocol and the Recordings"),
(2, "7.1.1 The Recordings and Their Reference Paths"),
(2, "7.1.2 Solver Validation Against Synthetic Ground Truth"),
(2, "7.1.3 Tracking Failures in the Recordings"),
(1, "7.2 Object Trajectory Against the Reference Paths"),
(1, "7.3 Pose Recovery Accuracy"),
(2, "7.3.1 Synthetic Masking on the Rail Recording"),
(2, "7.3.2 Synthetic Masking on the Loop Recording"),
(2, "7.3.3 Recovery Accuracy on Labelled Natural Failure Windows"),
(1, "7.4 Torso Stability"),
(1, "7.5 Qualitative Results"),

(0, "Chapter 8: Real-Time Feasibility"),
(1, "8.1 Causality of the Pipeline"),
(1, "8.2 Computational Cost and Latency"),
(1, "8.3 Offline and Real-Time Comparison"),

(0, "Chapter 9: Discussion, Future Work, and Conclusion"),
(1, "9.1 Limitations"),
(1, "9.2 Future Directions"),
(1, "9.3 Conclusion"),

(0, "References"),

(0, "Appendix A: RGB-D Data Acquisition"),
(1, "A.1 Acquisition Procedure"),
(1, "A.2 Structure of the Recorded Bag File"),
(1, "A.3 Replay and Visual Inspection"),

(0, "Appendix B: Hardware and Software Platform"),

(0, "Appendix C: Pose Landmark Detection"),
(1, "C.1 MediaPipe BlazePose Overview"),
(1, "C.2 Landmark Coordinate Extraction"),
(1, "C.3 Worked Example"),

(0, "Appendix D: Depth Sensing and Deprojection"),
(1, "D.1 RGB Camera Calibration"),
(1, "D.2 Depth-to-RGB Alignment"),
(1, "D.3 Pixel-to-World Mapping"),
(1, "D.4 Worked Example"),

(0, "Appendix E: Fiducial Marker Detection and Pose Recovery"),
(1, "E.1 Detection and Sub-Pixel Corners"),
(1, "E.2 Planar Pose Recovery and the Two-Solution Ambiguity"),
(1, "E.3 From the Rodrigues Form to a Rotation Matrix"),
(1, "E.4 Worked Example: Frame 533 Detections"),

(0, "Appendix F: Characteristics of the Candidate Filters"),
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

t = doc.add_paragraph()
r = t.add_run(TITLE)
r.font.bold = True
t.alignment = 1

s = doc.add_paragraph()
r = s.add_run("Table of Contents (V7 rewrite)")
r.font.italic = True
s.alignment = 1

doc.add_paragraph()

for level, text in entries:
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 0:
        run.font.bold = True
        p.paragraph_format.space_before = Pt(10)
    else:
        p.paragraph_format.left_indent = Inches(0.3 * level)
    p.paragraph_format.space_after = Pt(2)

out = "/home/luo/Desktop/New_SandBox/writing/v7/Thesis_V7_TOC.docx"
doc.save(out)
print("saved", out, "| entries:", len(entries))
