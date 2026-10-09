#!/usr/bin/env python3
"""Build Chapter_2_Change_Log.docx: a standalone change log for Chapter 2
only (user request 2026-09-02), for the supervisor.

Two parts: what changed in the chapter between Version 6 (the version his
2026-08-11 comments were written on) and the current Version 7 build, and
what changed in answer to each of his comments that targets Chapter 2
(C4, C6, C7, C8 of writing/v7/PROF_COMMENTS_ROUND2.md).

Style rules as for make_change_summary.py: plain statements, student
voice, no emojis, no decorative Unicode, no dashes as separators, no
internal file names, no decision identifiers.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[3]
OUT = REPO / "writing" / "v7" / "Chapter_2_Change_Log.docx"

INTRO = (
    "This document lists the changes made to Chapter 2 between Version 6, "
    "the version your comments of August 11 were written on, and the "
    "current Version 7 build. The first part gives the new structure and "
    "the changes section by section. The second part goes through each of "
    "your comments that concerns Chapter 2 and states what was done in "
    "answer to it and where it can be found."
)

TOC_ROWS = [
    ("Version 6", "Version 7"),
    ("Chapter 2: Background", "Chapter 2: Experimental Setup and Data Acquisition"),
    ("2.1 Data Acquisition and Preprocessing", "2.1 The Recording Setup"),
    ("2.1.1 Experimental Setup and Recordings", "2.2 Pose Landmarks from MediaPipe"),
    ("2.1.2 Pose Landmarks from MediaPipe", "2.3 Object Pose from ArUco Markers"),
    ("2.1.3 Fiducial Markers (ArUco)", "2.4 Signal Filtering"),
    ("2.1.4 Recorded Session Replay and Frame Extraction", ""),
    ("2.1.5 Signal Smoothing", ""),
    ("2.2 Unity-Python Communication Preliminaries", ""),
    ("2.2.1 UDP Protocol Basics and Packet Structure", ""),
    ("2.2.2 Frame Synchronization Overview", ""),
    ("2.2.3 Unity Humanoid Rig Fundamentals", ""),
]

STRUCTURE_BULLETS = [
    "The chapter is retitled from Background to Experimental Setup and "
    "Data Acquisition and now holds only the physical experiment and the "
    "data it produces. The three-level numbering of Version 6 is replaced "
    "by four flat sections.",
    "The Unity and transport preliminaries of the old Section 2.2 leave "
    "the chapter. The UDP link they described is gone: the streaming "
    "stage now passes the data to Unity through shared memory, and that "
    "is described in Chapter 6 with the system integration.",
    "The old Section 2.1.4 on session replay and frame extraction is "
    "reduced to one paragraph in Section 2.1, and the acquisition "
    "procedure and the file structure move to Appendix A.",
    "The reference path of the task, the waypoints and the rail "
    "measurements move to the start of Chapter 7, where the evaluation "
    "uses them. Chapter 2 keeps only the environment.",
]

SECTION_BULLETS = [
    "Section 2.1 opens with the reason for the controlled laboratory "
    "setting, the task in one sentence, and the whole-setup photograph as "
    "Figure 2.1, before every other image: the working area from above the "
    "desk with the rail, the cube, the desk marker and the spirit level, "
    "then the scene from behind the sensor, then a colour frame of the "
    "recording.",
    "A new Table 2.1 lists every object in the scene in one place: the "
    "three ArUco markers with their identifiers, dictionary, measured "
    "printed size and the size declared to the detector, the cube, and the "
    "wooden rail with its dimensions. The caption states that the desk and "
    "object markers are printed at 45 millimetres, that the detector was "
    "given 50, and that every recovered object translation is corrected by "
    "the factor 0.9 before use.",
    "Section 2.1 states the sensor placement and the reason the "
    "participant is kept within 2 metres of the camera: the depth error of "
    "the RealSense grows with the square of the distance, with the Intel "
    "tuning guide added as a reference. Figure 2.2 labels the parts of the "
    "scene in the camera's own view, and Figure 2.3 is a new data-flow "
    "figure of the two pipelines, the person branch and the marker branch, "
    "with the chapter that develops each block written on it.",
    "Section 2.2 keeps the MediaPipe description, states which part of "
    "its output the pipeline uses and why the appearance-based depth is "
    "replaced by the aligned RealSense depth, and adds Figure 2.4, the "
    "detector output on a frame of the recording with the eight consumed "
    "landmarks numbered. The detector architecture and a worked example "
    "move to Appendix C, the deprojection to Appendix D.",
    "Section 2.3 keeps the ArUco description and now describes the three "
    "markers by role, with Figure 2.5 showing the detections on the same "
    "frame as Figure 2.4 and each marker enlarged with its corners. The "
    "corner detection, the planar ambiguity and the rotation conversion "
    "move to Appendix E.",
    "Section 2.4 replaces the old Signal Smoothing section. It names the "
    "three kinds of error in the landmark signals, then the filter used "
    "for each, the reason it is chosen, and its reference, and shows a "
    "before and after figure for each stage (Figures 2.6 to 2.8). The "
    "candidate filters are compared side by side in Appendix F.",
    "All example figures of the chapter come from the same recording, "
    "which is the recording evaluated in Chapter 7.",
]

COMMENT_ROWS = [
    ("Your comment", "What changed", "Where"),
    ("In chapter 2, first show the image of the whole setup using your "
     "phone. This should include the ground truth wire that the person "
     "need to follow, the Aruco markers, the RGB-D sensor and you. This "
     "should come first before other images and explain the overall "
     "objective of the user study.",
     "The whole-setup photograph now opens the chapter as Figure 2.1, "
     "before every other image, with the objective of the study stated "
     "beside it. The ground-truth path is now the wooden rail on the desk, "
     "and the photograph shows the rail, the cube, the markers and the "
     "spirit level. The wire recording is used only in the occlusion part "
     "of Chapter 7. I still owe you a retake that also has the sensor and "
     "me in the same frame; for now panel (b) shows the scene from behind "
     "the sensor.",
     "Section 2.1, Figure 2.1"),
    ("When you talk about data filtering, you need to mention which one "
     "you have used and why. Then, refer the reader to a reference "
     "textbook (or wherever you have borrowed them from). You should also "
     "show the results after filtering.",
     "Section 2.4 names each filter actually used and why: the Hampel "
     "test for single-frame spikes, linear interpolation for short gaps, "
     "and the zero-phase Butterworth low-pass for jitter, each with its "
     "reference (Hampel 1974 and Pearson et al. 2016; Oppenheim and "
     "Schafer, Butterworth 1930, Gustafsson 1996). Each stage has its own "
     "before and after figure on the right wrist of the recording. The "
     "object track filter is described in Chapter 4.",
     "Section 2.4, Figures 2.6 to 2.8"),
    ("Not sure if we should add an appendix highlighting the features of "
     "each filter.",
     "I drafted Appendix F with the candidate filters side by side "
     "(Butterworth, Savitzky-Golay, rolling median, One Euro), and Section "
     "2.4 points to it once. It is a single deletion if you would rather "
     "leave it out.",
     "Section 2.4, Appendix F"),
    ("Figure 2.5 is not clear. You may want to break it down to make a "
     "bigger. Also label the Figures and add captions for each set.",
     "The old combined filtering panel is gone. It is now three separate "
     "larger figures, one per stage, each labelled and captioned: "
     "despiking, gap bridging, and smoothing.",
     "Figures 2.6, 2.7 and 2.8"),
]

OPEN_ITEMS = [
    "The retake of the setup photograph with the sensor and me in the "
    "same frame.",
    "Your decision on keeping or removing Appendix F.",
]


def style_table(table):
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(widths))
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            table.rows[i].cells[j].text = text
    for row in table.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Inches(w)
    style_table(table)
    return table


def add_heading(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.add_run(text).font.bold = True
    return h


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(4)


doc = Document()
for sname in ("Normal", "List Bullet"):
    st = doc.styles[sname]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Chapter 2 Change Log: Version 6 to Version 7")
r.font.size = Pt(14)
r.font.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Lanqing Luo, September 2026").font.size = Pt(11)

intro = doc.add_paragraph(INTRO)
intro.paragraph_format.space_after = Pt(10)

add_heading(doc, "1. Structure")
doc.add_paragraph(
    "The structure of Chapter 2 in Version 6 and in Version 7, side by "
    "side.").paragraph_format.space_after = Pt(6)
add_table(doc, TOC_ROWS, (3.0, 3.0))
doc.add_paragraph().paragraph_format.space_after = Pt(2)
doc.add_paragraph("The structural changes and the reason for each:")
add_bullets(doc, STRUCTURE_BULLETS)

add_heading(doc, "2. Changes by Section")
add_bullets(doc, SECTION_BULLETS)

add_heading(doc, "3. Your Comments on Chapter 2")
doc.add_paragraph(
    "Each comment of August 11 that concerns Chapter 2, what was done in "
    "answer to it, and where it is in the chapter.").paragraph_format.space_after = Pt(6)
add_table(doc, COMMENT_ROWS, (2.2, 2.9, 1.1))

add_heading(doc, "4. Open Items")
add_bullets(doc, OPEN_ITEMS)

doc.save(OUT)
print("saved", OUT)
