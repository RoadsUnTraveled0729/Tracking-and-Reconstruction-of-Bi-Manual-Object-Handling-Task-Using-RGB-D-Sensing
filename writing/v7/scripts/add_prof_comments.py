"""Restore the supervisor's V7-round comments into the assembled thesis,
each with a short reply stating what changed.

Comment sources:
  C1-C13  : the supervisor's message of 2026-08-11 on V6
            (writing/v7/PROF_COMMENTS_ROUND2.md; C14 was the user's own
            language directive and is not a supervisor comment)
  C15-C21 : the supervisor's message of 2026-08-31 on the assembled V7
            Chapter 3 (writing/v7/PROF_COMMENTS_ROUND3.md)

Each site carries:
  anchor : a substring of the CURRENT Thesis_V7.docx text at the spot the
           fix landed (asserted unique across the document)
  prof   : the supervisor's comment, verbatim or lightly trimmed
  reply  : what changed, one or two short sentences, student voice

Mechanism identical to writing/v6/scripts/add_prof_comments.py: comments
authored "Shahram Payandeh", replies "Lanqing Luo", threaded through a
commentsExtended part (w15:paraIdParent), the mechanism Word itself uses.

Run from writing/v7/scripts after build_thesis.py. Produces
Thesis_V7_commented.docx (the clean build is not modified) plus
PROF_REPLIES_V7.md.
"""
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn

BASE = str(Path(__file__).resolve().parents[1]) + "/"

PROF = "Shahram Payandeh"
ME = "Lanqing Luo"

COMMENTS = [
# ---- round of 2026-08-11 (on V6): C1-C13 ------------------------------------
dict(tag="C1",
     anchor="Examining Committee:",
     prof="The title page should follow the format of the thesis template. "
          "There should be other members listed. I do not see the name of "
          "other committee members.",
     reply="Rebuilt the title and approval pages in the template order and "
           "added the committee block. The other members' names are marked "
           "placeholders until I can confirm them with you."),
dict(tag="C2",
     anchor="This thesis develops a system that tracks a person's upper body",
     prof="The abstract needs to be rewritten more aligned with your "
          "language. Also, make sure you mention that the tracking results "
          "are compared with the ground truth which is the actual path you "
          "define in the lab.",
     reply="Rewrote the abstract in my own words. It now says the path "
           "defined in the laboratory is the ground truth and the tracked "
           "trajectory is compared against it."),
dict(tag="C3",
     anchor="A thesis submitted in partial fulfillment",
     prof="Make each chapter starts on a new page. Make sure you follow "
          "the eng. Sci thesis format template.",
     reply="Every chapter, the references and each appendix now start on "
           "a new page, and the whole file runs on one font set per the "
           "template."),
dict(tag="C4",
     anchor="shows the arrangement in one photograph taken from above the desk",
     prof="In chapter 2, first show the image of the whole setup using "
          "your phone. This should include the ground truth wire that the "
          "person need to follow, the Aruco markers, the RGB-D sensor and "
          "you. This should come first before other images and explain the "
          "overall objective of the user study.",
     reply="The whole-setup photo now opens Chapter 2 as Figure 2.1, with "
           "the study objective stated beside it. I still owe you a retake "
           "that also has the sensor and me in the frame."),
dict(tag="C5",
     anchor="Averaging the ten translations is an ordinary arithmetic mean",
     prof="Somehow the equations in the Word file is messed up. Equations "
          "in section 4.2.1 are not showing correctly.",
     reply="The equations are rebuilt as native Word math throughout. "
           "Please tell me if any still look broken in your Word and I "
           "will switch that section to rendered images."),
dict(tag="C6",
     anchor="The pipeline therefore handles each case in a separate stage",
     prof="When you talk about data filtering, you need to mention which "
          "one you have used and why. Then, refer the reader to a "
          "reference textbook (or wherever you have borrowed them from). "
          "You should also show the results after filtering.",
     reply="Section 2.5 now names each filter used, why it is chosen, and "
           "its textbook reference, and shows a before and after figure "
           "for every stage."),
dict(tag="C7",
     anchor="Appendix F compares the candidate smoothers side by side",
     prof="Not sure if we should add an appendix highlighting the features "
          "of each filter.",
     reply="I drafted Appendix F with the candidate filters side by side. "
           "It is a single deletion if you would rather leave it out."),
dict(tag="C8",
     anchor="shows the test in action, catching a wrist spike",
     prof="Figure 2.5 is not clear. You may want to break it down to make "
          "a bigger. Also label the Figures and add captions for each set.",
     reply="The combined panel is now three separate larger figures, one "
           "per stage, each labelled and captioned."),
dict(tag="C9",
     anchor="The arm model borrows its shape from robotics",
     prof="In the joint angle calculation section, show a kinematic model "
          "of an arm (similar to what we draw in robotics course) and the "
          "label which joint angle is which. This clear figure should come "
          "before section 3.4.1. For example, draw some schematic similar "
          "to PUMA robot for the shoulder joint and connected elbow joint "
          "and the wrist point. In there, label the joints which you are "
          "solving for.",
     reply="Added Section 3.3.1 before the angle computation: Figure 3.7 "
           "gives the standard seven-degree-of-freedom description with "
           "its PUMA-style chain, and Figure 3.8 draws the four solved "
           "joints on the body with every angle labelled."),
dict(tag="C10",
     anchor="The wrist is not a solved joint of this model",
     prof="Section 3.4.4, we do not solve for wrist joints. We are solving "
          "for the shoulder and elbow joints. The wrist would be just a 3D "
          "point which we can then track.",
     reply="Section 3.4.4 is retitled The Wrist as a Tracked Point and now "
           "presents the wrist exactly that way."),
dict(tag="C11",
     anchor="A single measured frame now carries the whole chain",
     prof="In the example you are solving, you need to directly refer to "
          "the kinematic diagram of PUMA type model which I mention above "
          "and show the corresponding angles that you are solving for.",
     reply="The worked example now points every solved value at its label "
           "in the schematic at the step where it is solved."),
dict(tag="C12",
     anchor="The rotation obtained that way is called the chordal mean",
     prof="What is chordal-mean anchor?",
     reply="Defined it here at first use: the element-wise mean of the "
           "calibration rotations projected back onto the nearest proper "
           "rotation, then frozen as the anchor."),
dict(tag="C13",
     anchor="[UPDATE FIELDS IN WORD: F9]",
     prof="Prepare the table of content of the thesis as we have agreed "
          "and based on the thesis format. Sent it to me so we can call it "
          "finalized.",
     reply="This is the finalized structure we agreed on. The contents and "
           "list pages are self-updating Word fields; pressing F9 fills in "
           "the page numbers."),
# ---- round of 2026-08-31 (on V7 Chapter 3): C15-C21 -------------------------
dict(tag="C15",
     anchor="The two combine into one representation, the 4 by 4 homogeneous transformation matrix T",
     prof="You should also show the T matrix which combines both R and t "
          "in one representation. For example, equation 3.2 is valid when "
          "the origin of both child and parent are the same location. You "
          "should show the corresponding T matrix here as well.",
     reply="Equation (3.2) now shows T with R and t as blocks and its "
           "action on a point. The rotation-only mapping is scoped to "
           "directions and shared origins, with the inverse T form beside "
           "it in equation (3.4)."),
dict(tag="C16",
     anchor="Each neighbouring pair of frames",
     prof="Referring to Figure 3.3, you should also define the T matrices "
          "between the various frames. For example, T of the shoulder "
          "about the root (torso) frame, elbow with respect to the "
          "shoulder and .. This makes it clear that each frame is "
          "different from the other in terms of relative rotation and "
          "position of their origin.",
     reply="Equation (3.5) now defines the three T matrices of the chain "
           "right after the chain figure, each with its own rotation and "
           "origin, and equation (3.6) composes them."),
dict(tag="C17",
     anchor="These relative transformations are the plan of the whole chapter",
     prof="Then it is clear when you mention that when you measure the "
          "landmark position of say the shoulder with respect to the "
          "sensor frame, you can use these relative T matrices to obtain "
          "for example, the position of the shoulder landmark with respect "
          "to the torso (root frame). As you have it, you do not make this "
          "clear in explaining the logic before you go over the details.",
     reply="Added this logic paragraph before any construction: each "
           "landmark is measured with respect to the sensor and the "
           "inverse T matrices carry it into the frame where its joint is "
           "solved."),
dict(tag="C18",
     anchor="the first transformation announced in equation",
     prof="Same for equation 3.7. You should define the T matrix here as "
          "well which defines the relative position and orientation of the "
          "root frame with respect to the sensor frame.",
     reply="Equation (3.12) now shows the root T matrix right where the "
           "rotation and origin are assembled."),
dict(tag="C19",
     anchor="this fills in the second transformation",
     prof="Same goes in the section about the shoulder T matrix with "
          "respect to the root.",
     reply="Equations (3.13) and (3.14) define the shoulder and elbow T "
           "matrices in the section that places those frames."),
dict(tag="C20",
     anchor="Appending the origin",
     prof="Then in the numerical section, you should show how each of "
          "these T matrices look like.",
     reply="The worked example now prints all three T matrices as numbers "
           "and closes with the chain check: their product lands exactly "
           "on the measured elbow."),
dict(tag="C21",
     anchor="Figure 3.1 draws the information flow",
     prof="I also think that at the beginning of this chapter, you should "
          "draw a visual information flow diagram. This way, the reader "
          "can clearly see how each sections come together. For example, "
          "you have raw/calibrated information from the landmarks of torso "
          "and arms coming in. The torso landmark are used to establish "
          "the T matrix of the torso frame with respect to the sensor "
          "frame. ... So use AI tools to create this visual data flow "
          "where the reader can then get a better idea of the process.",
     reply="Figure 3.1 is the new information-flow diagram at the start of "
           "the chapter. Each box names the section that builds it and the "
           "arrows carry the T matrices between the stages."),
]

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
CT_COMMENTS_EX = ("application/vnd.openxmlformats-officedocument"
                  ".wordprocessingml.commentsExtended+xml")
RT_COMMENTS_EX = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"


def find_anchor_runs(doc, anchor):
    """The paragraph runs covering `anchor`; asserts exactly one match."""
    hits = []
    for p in doc.paragraphs:
        joined = "".join(r.text for r in p.runs)
        if anchor in joined:
            hits.append(p)
    assert len(hits) == 1, f"anchor matched {len(hits)} paragraphs: {anchor!r}"
    p = hits[0]
    runs, texts = p.runs, [r.text for r in p.runs]
    joined = "".join(texts)
    i = joined.find(anchor)
    j = i + len(anchor)
    span, pos = [], 0
    for r, t in zip(runs, texts):
        if pos < j and pos + len(t) > i:
            span.append(r)
        pos += len(t)
    assert span, f"no runs cover anchor: {anchor!r}"
    return span


def thread_replies(doc, pairs):
    """Give every comment paragraph a w14:paraId and write a commentsExtended
    part that nests each reply under its professor comment."""
    comments_part = doc.part.part_related_by(RT.COMMENTS)
    comments_el = comments_part._element
    para_id = {}
    counter = 0x2B3C0001
    for c in comments_el.findall(qn("w:comment")):
        cid = c.get(qn("w:id"))
        last_p = None
        for p in c.findall(qn("w:p")):
            p.set(f"{{{W14}}}paraId", f"{counter:08X}")
            last_p = p
            counter += 1
        para_id[cid] = last_p.get(f"{{{W14}}}paraId")

    lines = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
             f'<w15:commentsEx xmlns:w14="{W14}" xmlns:w15="{W15}" '
             'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
             'mc:Ignorable="w14">']
    for prof_id, reply_id in pairs:
        lines.append(f'<w15:commentEx w15:paraId="{para_id[prof_id]}" w15:done="0"/>')
        lines.append(f'<w15:commentEx w15:paraId="{para_id[reply_id]}" '
                     f'w15:paraIdParent="{para_id[prof_id]}" w15:done="0"/>')
    lines.append("</w15:commentsEx>")
    blob = "\n".join(lines).encode("utf-8")
    part = Part(PackURI("/word/commentsExtended.xml"), CT_COMMENTS_EX,
                blob, doc.part.package)
    doc.part.relate_to(part, RT_COMMENTS_EX)


def add_comments(src, dst):
    doc = Document(src)
    pairs = []
    for entry in COMMENTS:
        span = find_anchor_runs(doc, entry["anchor"])
        c_prof = doc.add_comment(span, text=entry["prof"], author=PROF, initials="SP")
        c_reply = doc.add_comment(span, text=entry["reply"], author=ME, initials="LL")
        pairs.append((str(c_prof.comment_id), str(c_reply.comment_id)))
    thread_replies(doc, pairs)
    doc.save(dst)
    print(f"{dst}: {len(COMMENTS)} comment sites, {2 * len(COMMENTS)} comments")


def write_md():
    out = [
        "# Supervisor comments restored with replies (V7)",
        "",
        "The comments of the 2026-08-11 round (C1-C13) and the 2026-08-31",
        "Chapter 3 round (C15-C21), restored as threaded Word comments in",
        "Thesis_V7_commented.docx, each with a reply stating what changed.",
        "The anchor column is where the comment sits in the built text, for",
        "adding the same threads to an online copy.",
        "",
    ]
    for n, e in enumerate(COMMENTS, 1):
        out.append(f"## {n}. ({e['tag']}) anchored at: \"{e['anchor']}\"")
        out.append(f"- **Comment (S.P.):** {e['prof']}")
        out.append(f"- **Reply (L.L.):** {e['reply']}")
        out.append("")
    with open(BASE + "PROF_REPLIES_V7.md", "w") as f:
        f.write("\n".join(out))
    print(f"PROF_REPLIES_V7.md: {len(COMMENTS)} sites")


if __name__ == "__main__":
    add_comments(BASE + "Thesis_V7.docx", BASE + "Thesis_V7_commented.docx")
    write_md()
