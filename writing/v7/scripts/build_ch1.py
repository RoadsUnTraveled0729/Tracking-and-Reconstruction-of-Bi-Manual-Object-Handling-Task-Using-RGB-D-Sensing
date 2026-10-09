#!/usr/bin/env python3
"""Build writing/v7/Chapter_1_Introduction.docx (V7 rewrite round, English).

Source text: the consolidated post-Turnitin V6 Chapter 1 (writing/v6/thesis.md
lines 75-164) with the consolidated citation numbers [1]-[39] preserved.
Changes in this round (see writing/v7/DECISIONS.md D9, D11):
  - 1.2.1 closing sentence: the recovery strategy is object-conditioned,
    matching the new Chapter 5, not interpolation-based.
  - Objective 4: recovery from the tracked object and two-link inverse
    kinematics; the Bayesian extension is future work in Chapter 9.
  - Objective 5: evaluation against the designed object path (the ground
    truth defined in the laboratory), replacing the dropped
    lifting/rotating/passing scenario framing.
  - 1.4 rewritten against the V7 table of contents.
Everything else is carried over verbatim.
"""
from docx import Document
from docx.shared import Pt, Inches

H1, H2, H3, P = "h1", "h2", "h3", "p"
CAP, TBL, LIN = "cap", "tbl", "lin"

content = [
(H1, "Chapter 1: Introduction"),

(P, "Human motion is useful to robots and clinicians only when it can be measured reliably in three dimensions. Robots learn manipulation from recorded demonstrations, and clinicians need objective evidence of how a patient moves. Both applications often rely on specially equipped laboratories, however. This thesis examines a harder practical case: tracking and graphically reconstructing bimanual object manipulation with a single consumer-grade RGB-D camera."),

(H2, "1.1 Motivation"),

(P, "Human hands are the primary means through which people interact with the physical world. Coordinated two-handed tasks such as lifting, rotating, transferring, and assembling objects appear in nearly every part of daily life. Capturing these interactions in three dimensions has direct uses in several fields. Robotics uses recorded human demonstrations as training data for imitation learning [1], [2]. In virtual and augmented reality, skeletal and object pose data animate avatars and simulate physical tasks [3], [4]."),

(P, "A vision-based system that records and reconstructs a patient's upper-body motion in three dimensions would let a physician review, at any convenient time, how the patient reaches for an object, transfers it between hands, or completes a routine daily task. An objective and repeatable record of this kind could assist in assessing progressive conditions such as Parkinson's disease, stroke-related loss of upper-limb function, or age-related decline in motor coordination. Brief clinical observations may miss occasional abnormalities in these conditions [5], [6]."),

(P, "Reconstruction also protects patient privacy. The recorded motion is played back on a graphical avatar, so the patient's face, body, and home environment never leave the capture device."),

(P, "These requirements narrow the choice of sensor. A standard RGB camera provides only two-dimensional information, so metric 3D measurement requires either a depth camera or multiple synchronized viewpoints, such as the two-camera arrangement of FreeMoCap [7]. Multi-camera motion capture systems are accurate, but they demand calibrated camera arrays, physical infrastructure, and expert operation [8]. Wearable sensors avoid the cameras at a different cost: they must be attached and calibrated for each participant, they restrict natural movement, and they raise hygiene concerns in patient-facing settings [9]. A single consumer-grade RGB-D camera avoids both drawbacks. It is inexpensive, places no hardware on the participant, and captures registered colour and metric depth from one fixed viewpoint [10]."),

(P, "Reconstructing bimanual motion from that single sensor is difficult. When the two hands occlude each other or the object, landmark detectors fail at the moments that matter clinically [11], [12]. Direct depth readings also become ambiguous or invalid when landmarks overlap in the image or approach the edge of the reliable depth range [13]. Keeping all tracked segments and objects in one consistent coordinate frame requires calibration, and prior single-camera systems address this only partially [14]. The torso also moves slowly while the hands move fast, so a single low-pass filter with one cutoff frequency applied to every joint signal over-smooths some segments and under-smooths others [15]. These problems are not solved by better sensor resolution alone. They call for a structured kinematic framework that encodes the relationships between body segments and stays physically consistent when measurements are missing or unreliable. A wider trend in computing supports this route: as transistor scaling reaches its physical limits, future capability is expected to come from algorithms and system organization rather than from faster parts [16]. This thesis, written within a systems engineering program, takes that as a working assumption and returns to it in Chapter 9."),

(P, "The Networked Robotics and Sensing Laboratory at Simon Fraser University has produced earlier work that this thesis builds on. A hierarchical kinematic model of the human hand, constructed from spatially tracked RGB-D landmarks, was developed in [17]. A subsequent study [18] formulated hand joint estimation as a Gaussian Bayesian Network problem, recovering anatomically consistent joint configurations when landmark measurements are occluded or corrupted by noise. The present thesis builds on both contributions, extending the scope from a single hand to the full bimanual upper body with concurrent object tracking and Unity-based visualization."),

(H2, "1.2 Problem Statement and Literature Review"),

(P, "Each challenge identified in Section 1.1 reflects a broader research problem. This section reviews the literature in three areas: hand and object tracking approaches, bimanual manipulation in motion analysis, and graphical reconstruction in robotics and virtual reality. Table 1.1 at the end of the section summarizes the reviewed work and the limitation of each approach family for the setting of this thesis."),

(H3, "1.2.1 Hand and Object Tracking Approaches"),

(P, "Hand tracking research has produced two broad families of methods: appearance-based approaches, which learn a direct mapping from image features to pose, and model-based approaches, which fit a parameterized hand model to sensor observations [19]. Early wearable systems encoded finger configuration directly in the sensor signal, through instrumented colour gloves [20] and, in earlier work in our laboratory, marker-based visual tracking combined with a vibrotactile glove [21]. The hardware these systems place on the user makes them impractical for clinical use. Consumer depth cameras then shifted the field toward markerless tracking: full-hand articulation was demonstrated with Particle Swarm Optimization over a geometric hand model [22], with an articulated Iterative Closest Point solver [23], and with a sphere-based model fitted to depth point clouds [24]. These methods established depth-based hand tracking, but they target a single hand, need explicit geometric models, and must be initialized for each user."),

(P, "Deep learning replaced hand-crafted models with learned representations, from multi-view convolutional networks over depth images [25] to transformer-based mesh recovery from monocular RGB [26] and graph-guided state space models [27]. These methods reach high benchmark accuracy, but they need large annotated training sets, they do not encode biomechanical constraints on joint motion, and they can behave unpredictably on data unlike their training distribution. MediaPipe [28] represents the current mainstream for lightweight landmark detection without specialized hardware: a two-stage pipeline that first detects a body region and then regresses keypoint positions within it. Its per-landmark depth is inferred from image appearance and carries uncertainty during fast motion, contact, and occlusion [3]. Back-projecting the 2D detections through a calibrated depth camera resolves this ambiguity, an approach set out in the hierarchical hand framework of [17] and independently validated in [13]. This thesis adopts the same fusion strategy."),

(P, "For object tracking, fiducial markers give reliable pose estimation without appearance models or training data. A comparative evaluation of marker systems concluded that ArUco markers balance robustness and efficiency well for single-camera setups [29], and the relationship between marker size, viewing distance, and pose error is documented in [30]. Each marker also encodes an integer identity, so a single detection gives both the object identity and its pose. Occlusion between the hand and the object remains the most persistent failure mode for single-camera systems. One line of work restores each occluded hand with an inpainting framework before estimating its pose [11]; another models hand and object jointly through mutual attention [12]. Both show that occlusion needs either explicit recovery or contextual constraints, and this observation motivates the object-conditioned recovery strategy of this thesis."),

(H3, "1.2.2 Bimanual Manipulation in Motion Analysis"),

(P, "Bimanual manipulation goes beyond running two hand trackers in parallel. The hands overlap in the image, they interact with the same object, and they are kinematically coupled through whatever they hold. An error in one hand therefore spreads into the reconstructed scene [19]. One study defines bimanual task categories, distinguishing symmetric from asymmetric coordination, and the two produce different patterns of landmark dropout [31]. The infrastructure behind large bimanual datasets shows how demanding the capture problem is: InterHand2.6M was collected with between 80 and 140 synchronized cameras [32], and OakInk2 combined multi-camera rigs with wearable sensors [8]. On the reconstruction side, RGB2Hands recovers two strongly interacting hands from a single RGB video [33], but it does not track objects. It relies on a learned model that is hard to interpret in terms of physical joint constraints."),

(P, "In clinical practice, objective kinematic measures of hand motion asymmetry capture deficits that observation scores may miss [5], and a structured review of upper-limb assessment in amyotrophic lateral sclerosis found existing clinical tools too insensitive to track gradual functional change [6]. Both findings support a low-overhead, marker-free reconstruction system. The output must also stay physically possible, since clinical assessment depends on anatomically consistent motion even when sensor data is incomplete."),

(H3, "1.2.3 Graphical Reconstruction in Robotics and VR"),

(P, "Unity has become the dominant platform for avatar-driven motion replay. Human motion has been reconstructed in Unity from a monocular camera using kinematic retargeting [34], a Kinect-based rehabilitation system showed that avatar-driven feedback supports clinical assessment of exercise quality [35]. The RePose system streams MediaPipe-derived joint angles to a Unity patient avatar for rehabilitation monitoring [36], the same end-to-end architecture adopted here. A persistent engineering issue in such systems is the coordinate mismatch between sensor and engine. The RealSense frame is right-handed with the y axis pointing down, and Unity's frame is left-handed with the y axis pointing up. A point expressed in the first frame therefore cannot be placed directly in the second, and without an explicit conversion the streamed poses accumulate systematic spatial error [14]."),

(P, "A single fixed change of coordinates moves landmark positions between the two environments. The avatar, however, is not driven by points in space. It is driven kinematically, joint by joint, through rotations, so the relative rotation angles of the body segments must be resolved from the measured landmark positions. That problem forms the modeling core of this thesis, and its foundations come from the laboratory's hand kinematics work: the hierarchical RGB-D hand model of [17] and its Bayesian estimation extension [18], with the RealSense depth noise characterized as Gaussian consistent with [37], [38]."),

(TBL, [["Approach family", "Representative work", "Sensing", "Limitation for this setting"],
       ["Instrumented gloves and wearables", "[20], [21]", "worn sensors", "hardware on the user; per-user calibration; hygiene"],
       ["Model-based depth tracking", "[22], [23], [24]", "one depth camera", "single hand; explicit geometric models; per-user initialization"],
       ["Learning-based pose estimation", "[25], [26], [27]", "depth or RGB", "large training sets; no biomechanical constraints"],
       ["Landmark detection with depth fusion", "[28], [3], [13], [17]", "RGB-D", "prior work covers a single hand, not the bimanual upper body"],
       ["Fiducial object tracking", "[29], [30]", "RGB", "tracks marked objects only; no body tracking"],
       ["Hand-object occlusion handling", "[11], [12]", "RGB, depth", "single hand and object; learned models"],
       ["Bimanual capture datasets", "[32], [8]", "rigs of 80 or more cameras; wearables", "laboratory infrastructure far beyond clinical settings"],
       ["Two-hand reconstruction", "[33]", "monocular RGB", "no object tracking; no metric depth"],
       ["Avatar replay in Unity", "[34], [35], [36]", "monocular or Kinect", "single tracked stream; no concurrent object in a shared frame"]]),
(CAP, "Table 1.1. Summary of the reviewed literature: each approach family, its representative work, its sensing arrangement, and its limitation for single-sensor bimanual capture."),

(H3, "1.2.4 Gap Analysis and Motivation for This Work"),

(P, "The reviewed work shows progress in hand tracking, object pose estimation, bimanual motion analysis, and graphical reconstruction taken individually. No single-sensor system combines these capabilities while remaining physically consistent under real bimanual manipulation. This thesis addresses five gaps."),

(LIN, "1. Existing single-camera systems either fail silently when landmark confidence drops or smooth all joints uniformly, without a recovery strategy that respects the different kinematics of the torso, upper arm, forearm, and hand [19]."),
(LIN, "2. MediaPipe's monocular depth estimates are unreliable during manipulation [3]. Depth fusion with a calibrated sensor resolves this, but no reported work integrates it with a full upper-body kinematic chain in a bimanual context."),
(LIN, "3. No published single-camera system reconstructs two hands and a manipulated object simultaneously in one calibrated world frame. Bimanual methods omit the object [33], and hand-object methods cover only one hand [39]."),
(LIN, "4. The sensor-to-Unity coordinate transformation is documented for the single-hand case [14], but a calibration that keeps a full upper-body system and a concurrently tracked object in one world frame has not been addressed."),
(LIN, "5. A kinematic framework valued for staying physically consistent and recovering under occlusion, dropout, and noisy depth, rather than for raw accuracy under ideal unoccluded conditions, has not been established for this setting."),

(P, "The kinematic and Bayesian estimation foundations developed in [17], [18] provide the mathematical basis for such a framework at the level of a single hand. The present thesis extends and integrates these foundations into a system that addresses all five gaps together. The same structure lays the groundwork for future systems in which better sensors, denser landmark sets, and stronger estimation methods can be incorporated without redesigning the architecture."),

(H2, "1.3 Research Objectives"),

(P, "The main goal of this thesis is to develop and evaluate a structured framework for tracking and reconstructing bimanual object manipulation from a single RGB-D sensor, with attention to physical consistency under the occlusion and measurement noise conditions that arise in practice. The thesis pursues this goal through five objectives."),

(LIN, "1. Design and implement a calibrated RGB-D sensing pipeline on the Intel RealSense D435 that aligns the depth and colour streams, establishes the camera intrinsics, and expresses all observations in a consistent world coordinate frame."),
(LIN, "2. Develop an upper-body pose estimation module that fuses MediaPipe Pose landmark detections with registered depth data, producing metric 3D joint positions for the torso, arms, and hands from a single fixed viewpoint."),
(LIN, "3. Implement concurrent object tracking from ArUco fiducial markers, synchronized frame by frame with the skeletal stream in the same world frame. The object's orientation is recovered and carried through filtering, streaming, and evaluation, but no second measurement chain validates it. Robust tracking through all orientations of the object remains an extension of this framework."),
(LIN, "4. Recover the pose of a holding hand during landmark tracking failures by conditioning on the concurrently tracked object, restoring the wrist from the object pose and the elbow by two-link inverse kinematics. Validate the recovered trajectories against reference data. A chain-level Bayesian recovery scheme following [18], in which a Gaussian network corrects each joint estimate from the estimates of its neighbours, is discussed as future work in Chapter 9."),
(LIN, "5. Integrate the complete pipeline with a Unity-based visualization environment. Evaluate the reconstruction against the original recordings of a designed bimanual object-handling task, comparing the tracked object trajectory with the ground-truth path defined in the laboratory."),

(H2, "1.4 Organization of the Thesis"),

(P, "Chapter 2 describes the experimental setup and the data it produces: the recording arrangement, the designed object path that serves as the ground-truth reference, the pose landmarks obtained from MediaPipe, the object poses obtained from the ArUco markers, and the filtering applied to the recorded landmark signals. The kinematic model of the person is developed in Chapter 3, which constructs the torso and arm coordinate frames from measured landmarks, computes the torso pose and the arm joint angles, and closes with a worked numerical example. Chapter 4 develops the object branch: static scene calibration, the world anchoring that makes the measurements independent of where the camera is placed, and how the object track is cleaned. The pose recovery strategy of Chapter 5 carries the reconstruction through landmark tracking failures, restoring the wrist from the object pose, the elbow by two-link inverse kinematics, and the torso on camera rays."),

(P, "Chapter 6 integrates the two branches and streams the combined person and object record into the Unity reconstruction. The evaluation follows in Chapter 7, comparing the tracked object trajectory with the designed path and grading the recovery under synthetic masking. Chapter 8 examines real-time feasibility and measures the computational cost of the causal pipeline, and Chapter 9 states the limitations, outlines future work, and concludes. The appendices collect the supporting internals. Appendix A describes how the RGB-D recordings are acquired, Appendix B the hardware and software platform, and Appendix C pose landmark detection. Appendix D covers depth sensing and deprojection, Appendix E fiducial marker detection and pose recovery, and Appendix F the characteristics of the candidate filters."),

]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

for item in content:
    kind = item[0]
    if kind == H1:
        doc.add_heading(item[1], level=1)
    elif kind == H2:
        doc.add_heading(item[1], level=2)
    elif kind == H3:
        doc.add_heading(item[1], level=3)
    elif kind == P:
        doc.add_paragraph(item[1])
    elif kind == LIN:
        p = doc.add_paragraph(item[1])
        p.paragraph_format.left_indent = Inches(0.3)
    elif kind == CAP:
        p = doc.add_paragraph(item[1])
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
    elif kind == TBL:
        rows = item[1]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j)
                cell.text = c
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

out = "/home/luo/Desktop/New_SandBox/writing/v7/Chapter_1_Introduction.docx"
doc.save(out)
print("saved", out, "| items:", len(content))
