"""Build Thesis_V7_Changes.docx: the change summary that accompanies the
full V7 delivery.

Delivery protocol (user decision 2026-08-05): every full round ships
Thesis_Vx.docx, Thesis_Vx_commented.docx, and Thesis_Vx_Changes.docx. Scope of
a change summary is the changes since the last version SENT to the supervisor.

Baseline: V6, the version the supervisor's 2026-08-11 comments were written on
(his comments name Section 3.4.4 "Wrist Rotation", Section 4.2.1 and Figure
2.5, which are V6 numbers). Target: the complete assembled V7. The Chapter 3
section also covers his 2026-08-31 comments on the assembled V7 (the T
matrices and the information-flow diagram), which were applied to V7 in place.

Style rules: plain statements, student voice, no emojis, no decorative
Unicode, no dashes as separators, no internal file names, no decision
identifiers.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/home/luo/Desktop/New_SandBox/writing/v7/Thesis_V7_Changes.docx"

INTRO = (
    "This document lists the changes of Version 7 against Version 6, the "
    "version your comments of August 11 were written on. The first section "
    "gives the new structure side by side with the old one and the reason for "
    "every structural change. The sections after it go through the front "
    "matter and each chapter. The Chapter 3 section also covers your comments "
    "of August 31 on the assembled Version 7, which are applied in this same "
    "file. The last section lists every comment with its status.")

TOC_ROWS = [
    ("Version 6", "Version 7"),
    ("Chapter 1: Introduction", "Chapter 1: Introduction"),
    ("Chapter 2: Background",
     "Chapter 2: Experimental Setup and Data Acquisition"),
    ("Chapter 3: Kinematic Modeling", "Chapter 3: Kinematic Modeling"),
    ("Chapter 4: Offline Processing Pipeline System",
     "Chapter 4: Object Tracking and Scene Reconstruction"),
    ("No equivalent chapter",
     "Chapter 5: Pose Recovery During Tracking Failure"),
    ("Chapter 5: System Integration and Graphical Reconstruction",
     "Chapter 6: System Integration and Graphical Reconstruction"),
    ("Chapter 6: Experimental Design and Evaluation", "Chapter 7: Evaluation"),
    ("Chapter 7: Real-Time Feasibility and Trade-off Analysis",
     "Chapter 8: Real-Time Feasibility"),
    ("Chapter 8: Discussion and Future Work; Chapter 9: Conclusion",
     "Chapter 9: Discussion, Future Work, and Conclusion"),
    ("Appendix A: RGB-D Data Acquisition", "Appendix A: RGB-D Data Acquisition"),
    ("Appendix E: Hardware and Software Platform",
     "Appendix B: Hardware and Software Platform"),
    ("Appendix C: Pose Landmark Detection", "Appendix C: Pose Landmark Detection"),
    ("Appendix B: Depth Sensing and Deprojection",
     "Appendix D: Depth Sensing and Deprojection"),
    ("Appendix D: Fiducial Marker Detection and Pose Recovery",
     "Appendix E: Fiducial Marker Detection and Pose Recovery"),
    ("No equivalent appendix",
     "Appendix F: Characteristics of the Candidate Filters"),
]

TOC_BULLETS = [
    "A new Chapter 5 collects the pose recovery work. In Version 6 the "
    "handling of missing landmarks was one interpolation step inside the "
    "pipeline chapter. It has since become a method with its own parts: "
    "detecting the tracking failures, finding the grip episodes and the fixed "
    "offset between hand and object, restoring the wrist from the object pose, "
    "solving the elbow by two-link inverse kinematics, repairing the torso on "
    "the camera rays, and reporting which of these produced each output frame. "
    "That does not fit as a subsection.",
    "The evaluation is now built around the object paths defined in the "
    "laboratory. Section 7.2 compares the tracked object trajectory against "
    "the designed path, which is the ground truth you asked to be named in "
    "the abstract. The Version 6 scenario sections on lifting, rotating, and "
    "passing are dropped, because the recorded task is one continuous path "
    "through eight waypoints rather than three separate scenarios.",
    "The evaluation chapter is Chapter 7, and Chapters 1 to 5 now present the "
    "method only. Validation results, error figures, and coverage counts are "
    "collected in the evaluation chapter instead of appearing next to the "
    "derivations that produce them.",
    "The appendices are lettered in the order they are first referenced in the "
    "body. The hardware and software platform moves to Appendix B, depth "
    "sensing becomes Appendix D, and marker detection becomes Appendix E. The "
    "content of each is unchanged by the relettering.",
    "Appendix F on the characteristics of the candidate filters is added, "
    "following your suggestion. It can be removed as one block if you decide "
    "against it.",
    "Chapter 2 is retitled from Background to Experimental Setup and Data "
    "Acquisition, and the communication preliminaries move to the integration "
    "chapter, so the setup chapter holds the setup and the data it produces.",
    "The former discussion chapter and the former conclusion chapter are "
    "merged into Chapter 9.",
]

FRONT_BULLETS = [
    "The title page and the Approval page follow the template order. The "
    "Approval page carries the committee block; the names of the chair and "
    "the other members are marked placeholders until I can confirm them.",
    "The abstract is rewritten in my own words. It states that the drawn path "
    "in the laboratory is the ground truth and that the tracked object "
    "trajectory is compared with it, and it quotes only numbers printed in "
    "the evaluation and feasibility chapters.",
    "The table of contents and the lists of figures and tables are "
    "self-updating fields; pressing F9 in Word fills in the page numbers.",
    "Every chapter, the references, and each appendix start on a new page, "
    "and the whole file runs on a single font set.",
]

CH1_BULLETS = [
    "The chapter is in English, as is the rest of the thesis; the Chinese "
    "version sent in the last round is set aside.",
    "The opening paragraph is rewritten as a shorter statement of what the "
    "thesis does, and the whole chapter is reworded in plainer language.",
    "Objective 4 describes the recovery actually implemented: the pose of a "
    "holding hand is recovered by conditioning on the tracked object, the "
    "wrist from the object pose and the elbow by two-link inverse "
    "kinematics. The Bayesian scheme stays as future work in Chapter 9.",
    "Objective 5 states that the reconstruction is evaluated by comparing "
    "the tracked object trajectory with the ground-truth path defined in "
    "the laboratory.",
    "Section 1.4 walks through the new chapter structure and introduces the "
    "six appendices in their new lettering.",
]

CH2_BULLETS = [
    "The chapter opens with the study objective and the whole-setup "
    "photograph as Figure 2.1, before every other image. The study now uses "
    "two recordings, and the chapter presents the rail recording first: a "
    "straight wooden rail on the desk, measured and levelled, along which "
    "the cube is slid with one hand, so the ground-truth path is built into "
    "the setup. Figure 2.1 shows that arrangement from above the desk, from "
    "behind the sensor, and from the sensor's own viewpoint, followed by the "
    "annotated camera view and the data-flow figure of the two pipelines. "
    "The example figures of the chapter (the MediaPipe output, the marker "
    "detections, and the three filter stages) are drawn from the rail "
    "recording.",
    "Section 2.2 gives the rail's measured dimensions and then presents the "
    "loop recording: the eight stations, the laboratory photograph of the "
    "wire and the drawn path, and the designed loop in the levelled world "
    "frame. The evaluation separates the two by occlusion: the rail "
    "recording is the scenario without it, the loop recording with its two "
    "handovers is the scenario with it, and both are evaluated in Section "
    "7.2.",
    "Section 2.5 names each filter actually used with the reason it is "
    "chosen and its textbook reference, and shows a before and after figure "
    "for each stage: despiking, gap bridging, and smoothing. The old "
    "combined panel is gone; the three figures are larger, labelled, and "
    "captioned. Section 2.5 points to Appendix F for the candidate filters "
    "side by side.",
    "Section 2.1 now carries one table of the scene objects: the three "
    "markers with their measured printed size of 45 millimetres (150 for "
    "the wall marker) beside the 50 millimetres declared to the detector, "
    "the cube, and the rail. The caption states that a factor of 0.9 "
    "corrects the declared size on every object pose.",
]

CH3_BULLETS = [
    "A new Section 3.3.1 comes before all the joint angle sections and holds "
    "two figures: the seven rotations of the human arm with their "
    "serial-chain equivalent in the PUMA convention, drawn after the "
    "reference figures you gave, and the schematic of the modelled arm as a "
    "chain of revolute-joint cylinders with every solved angle labelled. "
    "This answers your comment asking for a kinematic model of the arm "
    "before the joint angle sections.",
    "Section 3.4.4 is retitled The Wrist as a Tracked Point and rewritten: "
    "the chain solves the shoulder and the elbow, and the wrist is a "
    "tracked 3D point. This answers your comment that the wrist joints are "
    "not solved.",
    "The worked example is consolidated into one Section 3.5 on a single "
    "frame of the recording, and every solved value is read against its "
    "label in the schematic. This answers your comment that the example "
    "must refer to the kinematic diagram.",
    "The chain figure is now a photograph of the recording with the torso "
    "frame and the arm chains drawn on the person; the Version 6 mannequin "
    "drawing is removed.",
    "The Version 6 sections on the kinematic foundations are merged into "
    "one Section 3.1, and the Unity mapping section moves to Chapter 6, so "
    "the chapter holds the model of the person and nothing else.",
]

CH3_ROUND3_BULLETS = [
    "The homogeneous transformation matrix T is now shown wherever a "
    "rotation and a translation appear together. Equation (3.2) gives the "
    "block form and its action on a point; equation (3.4) gives the inverse "
    "that carries a point from the parent frame into the child frame, and "
    "the rotation-only mapping is scoped to directions and to frames that "
    "share an origin.",
    "After the chain figure, equation (3.5) defines the three relative "
    "transformations of the chain, the root with respect to person space, "
    "the shoulder with respect to the root, and the elbow with respect to "
    "the shoulder, each with its own rotation and origin, and equation "
    "(3.6) composes them. A paragraph before any construction states the "
    "logic: every landmark is measured with respect to the sensor, and the "
    "inverse transformations carry it into the frame where its joint is "
    "solved.",
    "The root, shoulder, and elbow transformations are each defined in the "
    "section that builds them, as equations (3.12), (3.13), and (3.14).",
    "The worked example prints all three transformations as numbers on the "
    "measured frame and closes with a chain check: their product lands on "
    "the measured elbow position exactly, and the wrist carried through the "
    "inverse chain reproduces the forearm length and the flexion input.",
    "A visual information-flow diagram opens the chapter as Figure 3.1: the "
    "landmarks enter from the sensing stage, the torso landmarks establish "
    "the root transformation, each arm landmark is carried through the "
    "relative transformations into the frame where its joint is solved, and "
    "the thirteen angles per frame leave for the avatar. Every box names "
    "the section that builds it. The other figures and the equations "
    "renumber accordingly.",
    "Section 3.1 also states why the kinematics are computed in the "
    "left-handed frame: the angles drive the Unity animation system, points "
    "cross a handedness change with one sign flip while rotations need a "
    "conversion of every sign convention, so the flip is paid once per "
    "point at the entry.",
]

CH4_BULLETS = [
    "The chapter is rewritten from the Version 6 pipeline chapter into the "
    "object branch proper: static scene calibration, world anchoring, "
    "filtering and cleaning of the object track, and a worked example on "
    "the same frame as the Chapter 3 example.",
    "The chordal mean is defined in plain words at its first use: the "
    "element-wise mean of the calibration rotations projected back onto the "
    "nearest proper rotation, then frozen as the anchor. This answers your "
    "question about the chordal-mean anchor.",
    "The world frame is anchored at the desk marker, so the reconstruction "
    "does not depend on where the camera stands; the chapter shows the "
    "invariance directly.",
    "The object-track filter and the cleaning stage are described with "
    "their actual behaviour on the recording, including how an interior "
    "gap in the detections is held at the previous valid pose instead of "
    "being smoothed across.",
]

CH5_BULLETS = [
    "This chapter is new. It collects the pose recovery that Version 6 "
    "handled as interpolation into a method of its own: how tracking "
    "failures are detected, how grip episodes and the hand-to-object offset "
    "are found, how the wrist is restored from the object pose during a "
    "grip, how the elbow is solved by two-link inverse kinematics, how the "
    "torso is repaired on the camera rays when the hip depth corrupts, and "
    "which source produced each output frame.",
]

CH6_BULLETS = [
    "The integration chapter now holds the communication architecture, the "
    "coordinate conversion, and the avatar and scene reconstruction in "
    "Unity. The mapping from the solved angles onto the avatar, which sat "
    "in Chapter 3 in Version 6, lives here.",
]

CH7_BULLETS = [
    "The evaluation opens with its protocol: accuracy is quoted only where "
    "the correct answer is known independently, either injected in a "
    "synthetic dataset or measured on frames the tracker handled and then "
    "masked. The solver is validated against six synthetic datasets, where "
    "its worst decoding error is at floating-point rounding size.",
    "Section 7.2 compares the tracked object trajectory with the designed "
    "path over the loop traversal, with a per-segment table, and explains "
    "the two outlier segments structurally. The same section closes with "
    "the no-occlusion rail scenario: the rail line is fitted to the slide "
    "and the deviations are reported, and the two scenarios are read "
    "against each other.",
    "Section 7.3 grades the recovery under synthetic masking, on still and "
    "moving windows, against holding the last value; a subsection defines "
    "the protocol for manual wrist labels on the natural failure windows, "
    "whose measurements are pending until I place the labels.",
    "Section 7.4 examines the torso repair over the corrupt-depth windows, "
    "and Section 7.5 shows the reconstruction on the frames that carry no "
    "independent measurement, which are shown rather than scored.",
]

CH8_BULLETS = [
    "The feasibility chapter reports which parts of the offline chain are "
    "causal, the measured computational cost and latency of each stage, and "
    "a live run of the same feature set at 29.6 frames per second, compared "
    "with the offline configuration.",
]

CH9_BULLETS = [
    "The discussion and the conclusion are merged. The limitations name "
    "what the recordings actually show, including the landmark detector "
    "losing low-visibility wrists and the marker being hidden in the "
    "handover. The future directions keep the Bayesian recovery scheme and "
    "the sensor extensions.",
]

APP_BULLETS = [
    "Appendices A to E keep their content under the new lettering, and "
    "their worked examples run on the same frame of the recording as the "
    "Chapter 3 and Chapter 4 examples, so one moment of the recording can "
    "be followed through the landmark, depth, and marker chains.",
    "Appendix F is new: the candidate smoothing filters side by side, with "
    "a comparison figure and table.",
]

STATUS_AUG11 = [
    ("Comment", "Status"),
    ("Title page format and committee members",
     "Done with placeholders. The pages follow the template; the committee "
     "names are marked in bold brackets and I need the names to fill them."),
    ("Abstract rewritten, naming the ground-truth path", "Done."),
    ("Each chapter starts on a new page, format template", "Done."),
    ("Whole-setup photo first in Chapter 2",
     "Done. Figure 2.1 opens the chapter with the rail setup from above the "
     "desk, from behind the sensor, and from the sensor's own view with me "
     "in the frame."),
    ("Equations broken in the Word file",
     "Open on your side. The equations are native Word math throughout; "
     "please tell me if any still look broken in your copy and I will "
     "switch that section to images."),
    ("Data filtering: which filter, why, reference, results", "Done."),
    ("Optional appendix on filter features",
     "Done as Appendix F; one deletion if you decide against it."),
    ("Figure 2.5 unclear, break it up and caption it", "Done."),
    ("Kinematic model of the arm before the joint angle sections", "Done."),
    ("The wrist is not a solved joint", "Done."),
    ("Worked example must refer to the kinematic diagram", "Done."),
    ("What is a chordal-mean anchor", "Done, defined at first use in 4.1."),
    ("Finalized table of contents",
     "Done; the changes since your sign-off are in the first section."),
]

STATUS_AUG31 = [
    ("Comment", "Status"),
    ("Show the T matrix combining R and t; T at equation 3.2", "Done."),
    ("Define the T matrices between the frames of the chain figure", "Done."),
    ("State the mapping logic before the details", "Done."),
    ("T of the root with respect to the sensor at equation 3.7", "Done."),
    ("The shoulder T matrix with respect to the root", "Done."),
    ("Show each T matrix numerically", "Done."),
    ("Information-flow diagram at the start of Chapter 3", "Done."),
]


def style_table(table):
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
    for row in table.rows:
        row.cells[0].width = Inches(widths[0])
        row.cells[1].width = Inches(widths[1])
    style_table(table)
    return table


def add_heading(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.add_run(text).font.bold = True
    return h


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(4)


doc = Document()
for sname in ("Normal", "List Bullet"):
    st = doc.styles[sname]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Summary of Changes: Version 7")
r.font.size = Pt(14)
r.font.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Lanqing Luo, August 2026").font.size = Pt(11)

intro = doc.add_paragraph(INTRO)
intro.paragraph_format.space_after = Pt(10)

add_heading(doc, "1. Structure")
doc.add_paragraph(
    "The structure of Version 6 and the structure of Version 7, side by "
    "side.").paragraph_format.space_after = Pt(6)
add_table(doc, TOC_ROWS, (3.0, 3.0))
doc.add_paragraph().paragraph_format.space_after = Pt(2)
doc.add_paragraph("The structural changes and the reason for each:")
add_bullets(doc, TOC_BULLETS)

add_heading(doc, "2. Front Matter")
add_bullets(doc, FRONT_BULLETS)

add_heading(doc, "3. Chapter 1: Introduction")
add_bullets(doc, CH1_BULLETS)

add_heading(doc, "4. Chapter 2: Experimental Setup and Data Acquisition")
add_bullets(doc, CH2_BULLETS)

add_heading(doc, "5. Chapter 3: Kinematic Modeling")
add_bullets(doc, CH3_BULLETS)
doc.add_paragraph(
    "Your comments of August 31 on the assembled Version 7 are applied in "
    "this same file:").paragraph_format.space_before = Pt(4)
add_bullets(doc, CH3_ROUND3_BULLETS)

add_heading(doc, "6. Chapter 4: Object Tracking and Scene Reconstruction")
add_bullets(doc, CH4_BULLETS)

add_heading(doc, "7. Chapter 5: Pose Recovery During Tracking Failure")
add_bullets(doc, CH5_BULLETS)

add_heading(doc, "8. Chapter 6: System Integration and Graphical Reconstruction")
add_bullets(doc, CH6_BULLETS)

add_heading(doc, "9. Chapter 7: Evaluation")
add_bullets(doc, CH7_BULLETS)

add_heading(doc, "10. Chapter 8: Real-Time Feasibility")
add_bullets(doc, CH8_BULLETS)

add_heading(doc, "11. Chapter 9: Discussion, Future Work, and Conclusion")
add_bullets(doc, CH9_BULLETS)

add_heading(doc, "12. Appendices")
add_bullets(doc, APP_BULLETS)

add_heading(doc, "13. Status of the Comments")
doc.add_paragraph(
    "Your comments of August 11, in the order they were "
    "given.").paragraph_format.space_after = Pt(6)
add_table(doc, STATUS_AUG11, (2.2, 3.8))
doc.add_paragraph(
    "Your comments of August 31 on Chapter 3 of the assembled "
    "thesis.").paragraph_format.space_before = Pt(8)
add_table(doc, STATUS_AUG31, (2.2, 3.8))

doc.save(OUT)
print("saved", OUT)
