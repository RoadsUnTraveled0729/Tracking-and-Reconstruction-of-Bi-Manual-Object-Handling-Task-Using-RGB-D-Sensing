"""Build journal/KINEMATIC_ANGLES_pandoc.docx from journal/KINEMATIC_ANGLES.md with pandoc.

The pandoc route, an alternative to the python-docx route of
journal/scripts/build_docx.py (journal/DECISIONS.md J-017).

Pandoc install (author permission 2026-10-07): pandoc 3.12 from conda-forge
in its own conda env "pandoc", binary /home/luo/anaconda3/envs/pandoc/bin/pandoc;
the base environment is untouched (journal/DECISIONS.md J-017).

Steps:
  1. Copy the Markdown and strip every \\tag{n} from the $$ displays; the
     numbers are kept in display order (pandoc's texmath drops \\tag).
  2. Reference docx: pandoc's default reference.docx restyled with the
     thesis house style (writing/v9/scripts/build_frontmatter.py
     apply_house_style), plus Word's Normal Table and Table Grid styles.
  3. pandoc -f markdown+tex_math_dollars-smart-implicit_figures -t docx.
     smart is off so Camera' keeps its ASCII prime; implicit_figures is off
     because every figure already has its own "Figure N." paragraph.
  4. Post-process with python-docx: each numbered display goes into the
     thesis's borderless three-column equation table
     (writing/v9/scripts/eqn.py add_display_eq), each unnumbered display is
     centred with eqn.EQ_SPACE_PT around it, data tables take Table Grid
     with a bold header, "Figure N." and "Table N." paragraphs take the
     caption styles, and writing/v9/scripts/submission_layout.apply runs
     last, as in build_thesis.py.

Run: /home/luo/anaconda3/bin/python journal/scripts/build_docx_pandoc.py
"""
import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree

REPO = Path(__file__).resolve().parents[2]
JOURNAL = REPO / "journal"
SRC = JOURNAL / "KINEMATIC_ANGLES.md"
OUT = JOURNAL / "KINEMATIC_ANGLES_pandoc.docx"
PANDOC = "/home/luo/anaconda3/envs/pandoc/bin/pandoc"
# Code blocks only: the longest code line in the source is 73 characters; a
# monospace advance of 0.6 em at 9 pt gives 73 * 5.4 = 394 pt, inside the
# 432 pt (6 in) text width, so no line wraps. Prose and inline code keep the
# thesis minimum of 12 pt (submission_layout.apply).
CODE_PT = 9
READER = "markdown+tex_math_dollars-smart-implicit_figures"

sys.path.insert(0, str(REPO / "writing" / "v9" / "scripts"))
import eqn  # noqa: E402
import submission_layout  # noqa: E402
from build_frontmatter import apply_house_style, define_style  # noqa: E402

DISPLAY_RE = re.compile(r"^\$\$\n(.*?)\n\$\$$", re.S | re.M)
TAG_RE = re.compile(r"\s*\\tag\{([^}]*)\}")
IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
CAPTION_RE = re.compile(r"^(Figure|Table) \d+\. ")
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
ZWSP = "\u200b"


def preprocess(text):
    """Strip \\tag{n} from every display; return the text and the numbers."""
    numbers = []

    def strip(match):
        body = match.group(1)
        tags = TAG_RE.findall(body)
        assert len(tags) <= 1, f"display with {len(tags)} tags: {body[:60]}"
        numbers.append(tags[0] if tags else None)
        return "$$\n" + TAG_RE.sub("", body) + "\n$$"

    return DISPLAY_RE.sub(strip, text), numbers


def tidy_inline(text):
    """Trim whitespace inside each inline $...$ span.

    Pandoc rejects a closing $ that follows a space (source line
    "$180^\\circ - $ Method A"); the dollars of a paragraph pair in order.
    Displays and fenced code are left alone. Returns the text and the
    number of spans changed."""
    parts = re.split(r"(^\$\$\n.*?\n\$\$$|^```.*?^```$)", text, flags=re.S | re.M)
    changed = 0
    for i in range(0, len(parts), 2):
        paras = re.split(r"(\n[ \t]*\n)", parts[i])
        for j in range(0, len(paras), 2):
            pieces = paras[j].split("$")
            assert len(pieces) % 2 == 1, f"odd $ count: {paras[j][:80]}"
            for k in range(1, len(pieces), 2):
                if pieces[k] != pieces[k].strip():
                    pieces[k] = pieces[k].strip()
                    changed += 1
            paras[j] = "$".join(pieces)
        parts[i] = "".join(paras)
    return "".join(parts), changed


def reference_doc(path):
    """Pandoc's reference.docx in the thesis house style."""
    data = subprocess.run([PANDOC, "--print-default-data-file", "reference.docx"],
                          check=True, capture_output=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    apply_house_style(doc)
    # One Heading 1 (the title) in this document: no page break before it.
    for el in doc.styles["Heading 1"].element.iter(qn("w:pageBreakBefore")):
        el.getparent().remove(el)
    for name in ("Body Text", "First Paragraph", "Block Text"):
        define_style(doc, name, size=12, line=1.5, before=0, after=6)
    define_style(doc, "Compact", size=12, line=1.0, before=0, after=0)
    define_style(doc, "Title", size=16, bold=True, line=1.0, before=0, after=14)
    for level in (4, 5, 6):
        define_style(doc, f"Heading {level}", size=12, bold=True, line=1.0,
                     before=12, after=6, keep_next=True)
    for name in ("Caption", "Image Caption", "Table Caption", "Figure Caption"):
        define_style(doc, name, size=12, line=1.0, before=4, after=14)
    # Word's Normal Table (borderless, the default) and Table Grid, taken
    # from python-docx's default template; pandoc's "Table" stays for nothing.
    styles = doc.styles.element
    for st in styles.findall(qn("w:style")):
        if st.get(qn("w:type")) == "table":
            st.attrib.pop(qn("w:default"), None)
    word = Document().styles
    for name in ("Normal Table", "Table Grid"):
        styles.append(copy.deepcopy(word[name].element))
    doc.save(str(path))


def run_pandoc(md, ref, raw):
    subprocess.run([PANDOC, str(md), "-f", READER, "-t", "docx",
                    "--resource-path", str(JOURNAL), "--reference-doc", str(ref),
                    "-o", str(raw)], check=True)


def no_borders(table_el):
    """Borderless cells, also against a table style's first-row rule."""
    for tc in table_el.iter(qn("w:tc")):
        tcPr = tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement("w:" + side)
            b.set(qn("w:val"), "nil")
            borders.append(b)
        tcPr.append(borders)


def _mq(tag):
    return qn("m:" + tag)


def _mtext(el):
    return "".join(t.text or "" for t in el.iter(_mq("t")))


def _slot():
    # build_ch3.py pre(): a one-space run fills an absent left index.
    return etree.fromstring(f'<m:r xmlns:m="{M_NS}"><m:t xml:space="preserve"> </m:t></m:r>')


def fix_math(root):
    """Rewrite pandoc's OMML into the thesis's constructs.

    texmath writes Craig's {}^{A}_{B}R as a script on an empty base (a
    zero-width space) followed by R; the thesis writes m:sPre with R as its
    base (build_ch3.py pre). texmath writes Camera' as Camera with a
    superscript prime; the thesis writes the run Camera' (build_ch3.py FPER).
    texmath splits \\mathrm words into one run per letter; adjacent runs with
    the same properties are merged. Returns (prescripts, primes, merges,
    trailing scripts attached to the element before them)."""
    counts = [0, 0, 0, 0]
    for el in list(root.iter(_mq("sSup"), _mq("sSub"), _mq("sSubSup"))):
        e = el.find(_mq("e"))
        if not (len(e) == 1 and e[0].tag == _mq("r") and _mtext(e) in ("", ZWSP)):
            continue
        base = el.getnext()
        if base is None or not _mtext(base).strip():
            # A trailing script texmath could not attach (f'_y^2, or a
            # degree sign after "\ "): its base is the element before it.
            prev = el.getprevious()
            assert prev is not None, "script with no base on either side"
            if not _mtext(prev).strip() and _mtext(el) == ZWSP + "\u2218":
                # "\ {}^\circ": a degree sign with nothing to sit on.
                deg = _slot()
                deg.find(_mq("t")).text = "\u00b0"
                el.addprevious(deg)
                el.getparent().remove(el)
                counts[3] += 1
                continue
            e.remove(e[0])
            e.append(prev)
            counts[3] += 1
            continue
        assert base.tag.startswith("{%s}" % M_NS)
        pre = etree.SubElement(el.getparent(), _mq("sPre"))
        el.addprevious(pre)
        for slot in ("sub", "sup"):
            part = el.find(_mq(slot))
            new = etree.SubElement(pre, _mq(slot))
            new.extend(list(part) if part is not None else [_slot()])
        etree.SubElement(pre, _mq("e")).append(base)
        el.getparent().remove(el)
        counts[0] += 1
    for el in list(root.iter(_mq("sSup"))):
        sup, e = el.find(_mq("sup")), el.find(_mq("e"))
        if _mtext(sup) != "\u2032" or any(c.tag != _mq("r") for c in e):
            continue
        prime = copy.deepcopy(e[-1])
        prime.find(_mq("t")).text = "'"
        for c in list(e) + [prime]:
            el.addprevious(c)
        el.getparent().remove(el)
        counts[1] += 1
    for parent in list(root.iter()):
        prev = None
        for child in list(parent):
            if (prev is not None and child.tag == _mq("r") and prev.tag == _mq("r")
                    and len(child.findall(_mq("t"))) == 1
                    and len(prev.findall(_mq("t"))) == 1
                    and [etree.tostring(x) for x in child if x.tag != _mq("t")]
                    == [etree.tostring(x) for x in prev if x.tag != _mq("t")]):
                t = prev.find(_mq("t"))
                t.text = (t.text or "") + (child.find(_mq("t")).text or "")
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                parent.remove(child)
                counts[2] += 1
                continue
            prev = child
    return tuple(counts)


def place_equations(doc, numbers):
    body = doc.element.body
    paras = list(body.iter(qn("m:oMathPara")))
    assert len(paras) == len(numbers), (len(paras), len(numbers))
    eq_tables = []
    for omp, number in zip(paras, numbers):
        p = omp.getparent()
        assert p.tag == qn("w:p") and p.getparent() is body
        assert not "".join(t.text or "" for t in p.iter(qn("w:t"))
                           if t.getparent().getparent() is p).strip()
        maths = omp.findall(qn("m:oMath"))
        assert len(maths) == 1
        if number is None:
            omp.addprevious(maths[0])
            omp.getparent().remove(omp)
            par = Paragraph(p, doc._body)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            eqn._space(par)
            continue
        fragments = "".join(etree.tostring(c, encoding="unicode") for c in maths[0])
        table = eqn.add_display_eq(doc, fragments, number, punctuation="")
        no_borders(table._tbl)
        p.addprevious(table._tbl)
        body.remove(p)
        eq_tables.append(table._tbl)
    return eq_tables


def style_tables(doc, eq_tables):
    eq_ids = {id(t) for t in eq_tables}
    n = 0
    for tbl in doc.element.body.iter(qn("w:tbl")):
        if id(tbl) in eq_ids:
            continue
        n += 1
        tblPr = tbl.find(qn("w:tblPr"))
        tblPr.find(qn("w:tblStyle")).set(qn("w:val"), "TableGrid")
        first = tbl.find(qn("w:tr"))
        for r in first.iter(qn("w:r")):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                r.insert(0, rPr)
            if rPr.find(qn("w:b")) is None:
                rPr.insert(0, OxmlElement("w:b"))
    return n


def style_captions(doc):
    n = 0
    for par in doc.paragraphs:
        m = CAPTION_RE.match(par.text)
        if m:
            par.style = doc.styles["Figure Caption" if m.group(1) == "Figure"
                                   else "Table Caption"]
            n += 1
    return n


def code_size(doc):
    """Code-block runs at CODE_PT, after submission_layout's 12 pt floor."""
    for par in doc.paragraphs:
        if par.style is not None and par.style.style_id == "SourceCode":
            for run in par.runs:
                run.font.size = Pt(CODE_PT)


def check(doc, md_text):
    body = doc.element.body
    text = "".join(t.text or "" for t in body.iter(qn("w:t")))
    for bad in ("$", "\\tag", "pmatrix"):
        assert bad not in text, f"left in run text: {bad!r}"
    omath = len(list(body.iter(qn("m:oMath"))))
    media = [p for p in doc.part.package.iter_parts()
             if str(p.partname).startswith("/word/media/")]
    images = IMAGE_RE.findall(md_text)
    assert len(media) == len(images), (len(media), len(images))
    return omath, len(media)


def main():
    md_text = SRC.read_text(encoding="utf-8")
    pre, numbers = preprocess(md_text)
    pre, trimmed = tidy_inline(pre)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "in.md").write_text(pre, encoding="utf-8")
        reference_doc(tmp / "ref.docx")
        run_pandoc(tmp / "in.md", tmp / "ref.docx", tmp / "raw.docx")
        doc = Document(str(tmp / "raw.docx"))
    fixes = fix_math(doc.element.body)
    eq_tables = place_equations(doc, numbers)
    n_tables = style_tables(doc, eq_tables)
    n_captions = style_captions(doc)
    submission_layout.apply(doc)
    code_size(doc)
    omath, media = check(doc, md_text)
    body = doc.element.body
    n_pre = len(list(body.iter(_mq("sPre"))))
    empty = sum(1 for e in body.iter(_mq("e")) if _mtext(e) in ("", ZWSP))
    assert empty == 0, f"{empty} empty math bases remain"
    doc.save(str(OUT))
    numbered = sum(1 for n in numbers if n is not None)
    print(f"saved {OUT}")
    print(f"inline spans trimmed of inner whitespace: {trimmed}")
    print(f"m:sPre {n_pre} (converted {fixes[0]}); primes {fixes[1]}; "
          f"runs merged {fixes[2]}; trailing scripts attached {fixes[3]}; "
          f"empty bases left {empty}")
    print(f"displays {len(numbers)} (numbered {numbered}, unnumbered "
          f"{len(numbers) - numbered}); m:oMath total {omath}, inline "
          f"{omath - len(numbers)}; data tables {n_tables}; captions "
          f"{n_captions}; media {media}")


if __name__ == "__main__":
    main()
