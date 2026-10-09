from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import json,re,math,posixpath

B=Path(__file__).resolve().parent;W=B.parents[1]
N={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
speech=json.loads((B/'main_speech.json').read_text());entries={}
with ZipFile(B/'Thesis_Defence_personal_review_candidate.pptx') as z:
 relns='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
 pres=E.fromstring(z.read('ppt/presentation.xml'))
 rels={r.get('Id'):r.get('Target') for r in E.fromstring(z.read('ppt/_rels/presentation.xml.rels'))}
 slideparts=['ppt/'+rels[r.get('{'+relns+'}id')] for r in pres.find('p:sldIdLst',N)]
 for n,part in enumerate(slideparts,1):
  slide=E.fromstring(z.read(part))
  title=''
  for s in slide.findall('p:cSld/p:spTree/p:sp',N):
   q=s.find('.//p:cNvPr',N)
   if q is not None and (q.get('id')=='2' or (not title and int(q.get('id'))>=1000)):
    title=''.join(s.xpath('.//a:t/text()',namespaces=N))
  if n==1:title='Thesis defence'
  sr=E.fromstring(z.read(posixpath.dirname(part)+'/_rels/'+posixpath.basename(part)+'.rels'))
  nt=next(r.get('Target') for r in sr if r.get('Type','').endswith('/notesSlide'))
  notes=E.fromstring(z.read(posixpath.normpath(posixpath.join(posixpath.dirname(part),nt))))
  body=next(s for s in notes.findall('.//p:sp',N) if s.find('.//p:ph',N) is not None and s.find('.//p:ph',N).get('type')=='body')
  txt=[''.join(p.xpath('.//a:t/text()',namespaces=N)) for p in body.findall('p:txBody/a:p',N)]
  cue=' '.join(t[1:-1] for t in txt if t.startswith('['))
  spoken=' '.join(t for t in txt if t and not t.startswith(('Sources:','[')))
  entries[n]={'title':title,'cue':cue,'spoken':spoken}
  if n<=46:assert spoken==speech[str(n)],(n,spoken,speech[str(n)])

wordcount=lambda t:len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b",t))
words=sum(wordcount(s) for s in speech.values())
extra={4:4,13:2,14:2,15:2,20:3,23:6,25:7,26:7,30:10,34:3,35:3,38:10}
budgets={n:math.ceil(wordcount(speech[str(n)])/110*60)+2+extra.get(n,0) for n in range(1,47)}
assert sum(budgets.values())<=1560,sum(budgets.values())
budgets[38]+=1560-sum(budgets.values())
elapsed=0;timing={}
for n,seconds in budgets.items():
 timing[n]={'start_s':elapsed,'end_s':elapsed+seconds,'seconds':seconds,'words':wordcount(speech[str(n)])};elapsed+=seconds
assert elapsed==1560
fmt=lambda v:f'{v//60:02d}:{v%60:02d}'
doc=Document(W/'outputs/Thesis_Defence_2026_Speaker_Script.docx')
for e in list(doc.element.body):
 if e.tag!=qn('w:sectPr'):doc.element.body.remove(e)
for name in ['Normal','Title','Heading 1','Heading 2','Spoken script','Stage cue','Rehearsal timing']:
 st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'DejaVu Sans');st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'DejaVu Sans')
for name in ['Heading 2','Stage cue','Rehearsal timing']:
 doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Spoken script'].paragraph_format.keep_together=True
doc.add_paragraph('Thesis Defence Speaker Script','Title')
doc.add_paragraph('Lanqing Luo')
doc.add_paragraph('Tracking and Reconstruction of Bi Manual Object Handling Task Using RGB D Sensing')
doc.add_paragraph(f'The main presentation contains {words:,} spoken words and follows slides 1–46. Its planned duration is 26 minutes, leaving one minute before the 27 minute limit. The timing allows about 110 words per minute, slide changes and the viewing pauses marked below. Speak while the demonstrations play; do not add their full duration again. The schedule is a rehearsal guide, not a measured delivery time.')
doc.add_paragraph(f'Checkpoints: finish motivation and goal at {fmt(timing[4]["end_s"])}; finish the offline system at {fmt(timing[30]["end_s"])}; finish results and next steps at 26:00. References are advanced without narration. Discussion pages are outside the main presentation budget.')
blocks={1:'Block 1 Motivation and goal',5:'Block 2 Offline reconstruction system',31:'Block 3 Results and next steps'}
for n in range(1,47):
 if n in blocks:doc.add_paragraph(blocks[n],'Heading 1')
 e=entries[n];title=re.sub(r'[^\w\s]',' ',e['title']);title=re.sub(r'\s+',' ',title).strip()
 doc.add_paragraph(f'Slide {n} {title}','Heading 2')
 t=timing[n];doc.add_paragraph(f'{fmt(t["start_s"])} to {fmt(t["end_s"])}  |  {t["seconds"]} seconds','Rehearsal timing')
 cue=e['cue']
 if extra.get(n):cue+=f' Allow {extra[n]} seconds for viewing or pointing within this slot.'
 doc.add_paragraph(cue,'Stage cue')
 if e['spoken']:doc.add_paragraph(e['spoken'],'Spoken script')
 else:doc.add_paragraph('No narration.','Stage cue')
p=doc.add_paragraph('Supporting discussion pages','Heading 1');p.paragraph_format.page_break_before=True
doc.add_paragraph('Use these pages only when a related question arises. Slides 47–77 contain the QA material. Slide 78 contains the remaining IEEE references. None is included in the 26 minute main schedule.')
for n in range(47,78):
 e=entries[n];title=re.sub(r'[^\w\s]',' ',e['title']);title=re.sub(r'\s+',' ',title).strip()
 doc.add_paragraph(f'Slide {n} {title}','Heading 2');doc.add_paragraph('On demand','Rehearsal timing')
 doc.add_paragraph(e['cue'],'Stage cue');doc.add_paragraph(e['spoken'],'Spoken script')
doc.add_paragraph('Slide 78 Additional references','Heading 2')
doc.add_paragraph('On demand. Reference list only; no prepared narration.','Stage cue')
for tree in [doc.element,doc.styles.element]:
 for node in list(tree.xpath('.//w:pBdr')):node.getparent().remove(node)
doc.core_properties.title='Thesis Defence Speaker Script';doc.core_properties.author='Lanqing Luo'
doc.save(B/'Thesis_Defence_2026_Speaker_Script.docx')
(B/'timing_plan.json').write_text(json.dumps({'words':words,'assumed_wpm':110,'planned_seconds':1560,'reserved_seconds_to_27min':60,'actual_timed_rehearsal':False,'slides':timing},indent=2))

# Retain the requested two-page Chinese guide and update its three-block tree.
FONT='/System/Library/Fonts/STHeiti Medium.ttc'
im=Image.new('RGB',(2800,1620),'white');d=ImageDraw.Draw(im)
rootfont=ImageFont.truetype(FONT,49);branchfont=ImageFont.truetype(FONT,49);leaf=ImageFont.truetype(FONT,43)
branches=[
 ('1 引言与目标',180,[(100,'现有技术：可穿戴、实验室光学动捕、视觉方法'),(180,'重建目标：单 RGB-D、运动学模型、低成本'),(260,'输出：双手操作中的人体上肢与物体，共同场景')]),
 ('2 离线重建系统',715,[(400,'总架构：人体与物体两条分支，完整数据流'),(490,'预处理：跳变／Hampel；短缺／插值；抖动／Butterworth'),(580,'运动学：链模型、T-pose、L24 root、肩与肘坐标系'),(670,'ArUco：静态前 10 次有效检测平均；动态物体追踪'),(760,'局部恢复：A 方向记忆；B 物体约束；C 双连杆几何'),(850,'无可靠支持：保持上一有效角度，保留状态标签'),(940,'系统整合：shared memory、跳帧与覆盖、时间对齐'),(1030,'Unity 映射：人体角度；物体、桌面和墙面')]),
 ('3 结果与后续方向',1360,[(1170,'结果：物体；人体无遮挡；人体遮挡'),(1260,'实时转换：未来信息的因果替换；准确度与时延权衡'),(1350,'演示：实际 RGB 在左，Unity 重建在右'),(1440,'局限：评估范围、物体路径参考、录制重放'),(1530,'未来：保留可信像素；运动学约束结合贝叶斯估计')])]
spine=310;branchx=365;join=970;leafx=1030;line='#888888'
d.line((spine,180,spine,1360),fill=line,width=4)
d.multiline_text((5,665),'双手操作任务\n检测与重建',font=rootfont,fill='black',spacing=18)
d.line((290,735,spine,735),fill=line,width=4)
for label,y,children in branches:
 d.line((spine,y,branchx-14,y),fill=line,width=4);d.text((branchx,y-30),label,font=branchfont,fill='black')
 tw=d.textlength(label,font=branchfont);d.line((branchx+tw+20,y,join,y),fill=line,width=4)
 d.line((join,children[0][0],join,children[-1][0]),fill=line,width=4)
 for cy,t in children:
  assert d.textlength(t,font=leaf)<2800-leafx-10,(t,d.textlength(t,font=leaf))
  d.line((join,cy,leafx-15,cy),fill=line,width=4);d.text((leafx,cy-25),t,font=leaf,fill='black')
fig=B/'outline_tree.png';im.save(fig)
groups=[
 ('坐标系间转换','47','旋转与原点；逆变换；链式变换'),
 ('左手与右手人体流程','48','8 个输入点；5 个输出旋转；两条流程对比'),
 ('ArUco 与 AprilTag 比较','49–50','文献比较；角点误差；检测时间'),
 ('传输方案与帧包','51–52','shared memory、TCP、UDP；传输测量范围'),
 ('滤波方案','53–57','Savitzky–Golay、median、Butterworth、One Euro'),
 ('人体运动学推导','58–65','Camera 与链；root；肩摆动与扭转；肘坐标'),
 ('物体位姿推导','66','Camera 到 World；位姿组合；样本接受'),
 ('遮挡恢复推导','67–70','抓握偏移；方向记忆；手腕目标；两球交圆'),
 ('Unity 映射与场景标定','71–74','轴转换；rig 与 rest axes；重力和地面'),
 ('评估指标推导','75–76','长度误差；直线散布；关节位置与参照'),
 ('实时管线日志','77','帧号；处理时间；处理速率；丢帧与 merger ticks')]
guide=Document();sec=guide.sections[0]
sec.page_width=Inches(11);sec.page_height=Inches(8.5)
sec.left_margin=sec.right_margin=Inches(.45);sec.top_margin=sec.bottom_margin=Inches(.4);sec.footer_distance=Inches(.2)
for name,size in [('Normal',11.5),('Title',20),('Heading 1',18)]:
 st=guide.styles[name];st.font.name='Arial Unicode MS';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=name!='Normal'
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.1
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
guide.add_paragraph('答辩结构树','Title')
guide.add_paragraph('正文 1–46 页  ·  三个大块，先总览再展开  ·  计划 26 分钟，预留 1 分钟余量')
p=guide.add_paragraph();p.paragraph_format.space_after=Pt(0);p.add_run().add_picture(str(fig),width=Inches(10.1))
for q in guide.element.xpath('.//wp:docPr'):q.set('descr','引言与目标、离线重建系统、结果与后续方向组成的三块式答辩结构树')
p=guide.add_paragraph('隐藏问答索引','Heading 1');p.paragraph_format.page_break_before=True
guide.add_paragraph('按最终 PPT 页码查找  ·  47–77 页，共 31 页 QA  ·  坐标转换与左右手流程优先')
t=guide.add_table(rows=1,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;widths=[2.65,.8,6.65]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for c,txt in zip(t.rows[0].cells,['主题','页码','关键词']):c.text=txt
for vals in groups:
 for c,txt in zip(t.add_row().cells,vals):c.text=txt
for ri,row in enumerate(t.rows):
 for ci,cell in enumerate(row.cells):
  cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr()
  borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   q=OxmlElement('w:'+edge)
   for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:q.set(qn('w:'+k),v)
   borders.append(q)
  pr.append(borders);margins=OxmlElement('w:tcMar')
  for edge in ['top','bottom','left','right']:
   q=OxmlElement('w:'+edge);q.set(qn('w:w'),'100');q.set(qn('w:type'),'dxa');margins.append(q)
  pr.append(margins);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if ri==0 else ('F4F4F4' if ri%2==0 else 'FFFFFF'));pr.append(sh)
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15;p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==1 else WD_ALIGN_PARAGRAPH.LEFT
   for r in p.runs:r.font.name='Arial Unicode MS';r.font.size=Pt(11.5);r.font.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
repeated=OxmlElement('w:tblHeader');repeated.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeated)
guide.add_paragraph('参考文献：42–44 页；其余条目位于隐藏的第 78 页。QA 不计入正文演讲时间。')
for tree in [guide.element,guide.styles.element]:
 for node in list(tree.xpath('.//w:pBdr')):node.getparent().remove(node)
guide.core_properties.title='答辩结构与隐藏问答索引';guide.core_properties.author='Lanqing Luo'
guide.save(B/'Thesis_Defence_2026_Outline_and_QA_Index.docx')
print(json.dumps({'spoken_words':words,'planned_seconds':elapsed,'checkpoint_after_slide4':fmt(timing[4]['end_s']),'checkpoint_after_slide30':fmt(timing[30]['end_s']),'docx_count':2},indent=2))
