from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from collections import Counter
from docx import Document
from pypdf import PdfReader
import json,hashlib,posixpath,re
B=Path(__file__).resolve().parent;W=B.parents[1]
SOURCE=W/'outputs/Thesis_Defence_2026_RESTRUCTURED_FIXED.pptx'
CAND=B/'Thesis_Defence_personal_review_candidate.pptx'
N={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
sha=lambda b:hashlib.sha256(b).hexdigest()
def readzip(p):
 with ZipFile(p)as z:
  assert z.testzip() is None
  assert len(z.namelist())==len(set(z.namelist()))
  return {n:z.read(n) for n in z.namelist()}
s=readzip(SOURCE);c=readzip(CAND);checks=[]
def check(name,ok,detail=None):
 checks.append({'name':name,'passed':bool(ok),'detail':detail})
 assert ok,(name,detail)
def tree(d,p):return E.fromstring(d[p])
def rels(d,p):
 rp=posixpath.dirname(p)+'/_rels/'+posixpath.basename(p)+'.rels'
 return {r.get('Id'):r for r in tree(d,rp)} if rp in d else {}
def ordered(d):
 r=rels(d,'ppt/presentation.xml')
 return [posixpath.normpath('ppt/'+r[q.get('{'+N['r']+'}id')].get('Target')) for q in tree(d,'ppt/presentation.xml').find('p:sldIdLst',N)]
sp=ordered(s);cp=ordered(c)
check('Existing slide order retained; one hidden bibliography page appended',cp[:77]==sp and len(cp)==78)
check('46 shown pages and 32 hidden pages',sum(tree(c,p).get('show')!='0' for p in cp)==46 and all(tree(c,p).get('show')=='0' for p in cp[46:]))
for p in c:
 if p.endswith(('.xml','.rels')):tree(c,p)
check('All XML parts parse',True)
for part in cp:
 r=tree(c,part);shape_ids={q.get('id') for q in r.findall('.//p:cNvPr',N)}
 refs=[q.get('spid') for q in r.iter() if q.get('spid') is not None]
 check('Animation and build targets resolve to existing shapes',all(x in shape_ids for x in refs),{'part':part,'missing':sorted(set(refs)-shape_ids)})

links=0
for p in c:
 if not p.endswith('.rels'):continue
 base='' if p=='_rels/.rels' else posixpath.dirname(posixpath.dirname(p))
 seen=[]
 for r in tree(c,p):
  seen.append(r.get('Id'))
  if r.get('TargetMode')=='External':continue
  target=r.get('Target');resolved=posixpath.normpath(posixpath.join(base,target)).lstrip('/')
  check('Relationship target exists',resolved in c,{'part':p,'target':target});links+=1
 check('Relationship IDs unique',len(set(seen))==len(seen),p)
for p in cp:
 root=tree(c,p);ids=root.xpath('.//p:cNvPr/@id',namespaces=N)
 check('Shape IDs unique',len(ids)==len(set(ids)),p)
 targets=root.xpath('.//p:spTgt/@spid',namespaces=N)
 check('Animation shape targets resolve',all(i in ids for i in targets),p)
 for el in root.iter():
  for k,v in el.attrib.items():
   if k.startswith('{'+N['r']+'}') and v:check('Slide relationship reference resolves',v in rels(c,p),{'part':p,'id':v})
media=[p for p in s if p.startswith('ppt/media/')]
changed=[p for p in media if c.get(p)!=s[p]]
check('Only rejected p38 movie and poster replaced',set(changed)=={'ppt/media/fable_repaired_causal38.mp4','ppt/media/fable_repaired_causal38.png'},changed)
check('No media removed',all(p in c for p in media))
check('p23 original recording remains embedded',sha(c['ppt/media/media11.mp4'])=='74231478b2b7a0df6f56ca89a7c14d216e0de6bb943d275ffb0c4a567a05f66d')
check('p23 synchronized movie and poster embedded',all(c['ppt/media/'+n]==(B/n).read_bytes() for n in ['p23_synchronized_geometry.mp4','p23_synchronized_geometry.png']))
check('p23 media links use synchronized animation',all(rels(c,'ppt/slides/slide23.xml')[i].get('Target')=='../media/p23_synchronized_geometry.mp4' for i in ['rId5','rId6']))
for p in sp:
 before=tree(s,p);after=tree(c,p)
 def pics(root):
  result={}
  for q in root.findall('.//p:pic',N):
   id=q.find('.//p:cNvPr',N).get('id')
   q=E.fromstring(E.tostring(q))
   result[id]=E.tostring(q,method='c14n')
  return result
 check('All original picture and formula shapes retained unchanged',pics(before)==pics(after),p)
 for tag in ['timing','transition']:
  a=before.find('p:'+tag,N);b=after.find('p:'+tag,N)
  if p=='ppt/slides/slide23.xml' and tag=='timing':
   check('p23 original video autoplay retained',E.tostring(a.find('.//p:video',N))==E.tostring(b.find('.//p:video',N)))
   check('p23 geometry and video visible without clicks',not b.findall('.//p:seq',N) and not b.findall('.//p:set',N))
   continue
  check('Playback timing retained',E.tostring(a) if a is not None else None == (E.tostring(b) if b is not None else None),p) if a is None else check('Playback timing retained',E.tostring(a)==E.tostring(b),p)
origfonts=set();newfonts=set()
for d,parts,fonts in [(s,sp,origfonts),(c,cp,newfonts)]:
 for p in parts:
  for q in tree(d,p).iter():
   if q.get('typeface'):fonts.add(q.get('typeface'))
check('No new slide typefaces',newfonts<=origfonts,sorted(newfonts-origfonts))
slidesizes=lambda d:dict(tree(d,'ppt/presentation.xml').find('p:sldSz',N).attrib)
check('Slide dimensions unchanged',slidesizes(s)==slidesizes(c))
citations=[]
for i,p in enumerate(cp,1):
 text=' '.join(tree(c,p).xpath('.//a:t/text()',namespaces=N))
 check('No audience question marks',not ('?'in text or '？'in text),i)
 if i not in [42,43,44,78]:citations += [int(x) for x in re.findall(r'\[(\d+)\]',text)]
check('Citations resolve to deck reference list',set(citations)<=set(range(1,24)),sorted(set(citations)))
check('23 numbered IEEE entries',sum(len(re.findall(r'^\[\d+\]$',t)) for i in [42,43,44,78] for t in tree(c,cp[i-1]).xpath('.//a:t/text()',namespaces=N))==23)
doc=Document(B/'Thesis_Defence_2026_Speaker_Script.docx');speech={};n=None
for p in doc.paragraphs:
 if p.style.name=='Heading 2':
  m=re.match(r'Slide (\d+) ',p.text);n=int(m[1]) if m else None
 elif p.style.name=='Spoken script' and n:speech[n]=p.text
for i,p in enumerate(cp,1):
 nr=next(q for q in rels(c,p).values() if q.get('Type','').endswith('/notesSlide'))
 np=posixpath.normpath(posixpath.join(posixpath.dirname(p),nr.get('Target')))
 note=tree(c,np);body=next(q for q in note.findall('.//p:sp',N) if q.find('.//p:ph',N) is not None and q.find('.//p:ph',N).get('type')=='body')
 paragraphs=[''.join(q.xpath('.//a:t/text()',namespaces=N)) for q in body.findall('p:txBody/a:p',N)]
 spoken=' '.join(t for t in paragraphs if t and not t.startswith(('Sources:','[')))
 check('Word narration matches logical PPT page',speech.get(i,'')==spoken,i)
plan=json.loads((B/'timing_plan.json').read_text())
check('Main schedule is 26 minutes with 60-second reserve',plan['planned_seconds']==1560 and plan['reserved_seconds_to_27min']==60)
check('Spoken-word count matches main script',sum(len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b",v)) for n,v in speech.items() if n<=46)==plan['words'])
g=Document(B/'Thesis_Defence_2026_Outline_and_QA_Index.docx')
check('Outline is Letter landscape',all(x.page_width==10058400 and x.page_height==7772400 for x in g.sections))
covered=[]
for row in g.tables[0].rows[1:]:
 value=row.cells[1].text;lo,*hi=value.split('–');covered.extend(range(int(lo),int(hi[0] if hi else lo)+1))
check('QA index covers 47–77 exactly once',covered==list(range(47,78)))
check('Guide renders to exactly two pages',len(PdfReader(next((B/'outline_render').glob('*.pdf'))).pages)==2)
report={'passed':True,'checks_passed':len(checks),'checks_failed':0,'source_sha256':sha(SOURCE.read_bytes()),'candidate_sha256':sha(CAND.read_bytes()),'relationship_targets':links,'slide_count':78,'shown':46,'hidden':32,'media_count':len(media),'main_spoken_words':plan['words'],'planned_seconds':1560,'actual_timed_rehearsal':False,'native_powerpoint_current_file_verified':False,'formulas_and_other_media_preserved':True,'checks':checks}
(B/'delivery_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='checks'},indent=2))
