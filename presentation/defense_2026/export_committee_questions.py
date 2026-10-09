#!/usr/bin/env python3
"""Export the reviewed committee bank verbatim, with whole-document checks."""

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "COMMITTEE_QUESTIONS.md"
INDEX = HERE / "review/question_index.json"
DOCX = HERE / "COMMITTEE_QUESTIONS.docx"
PDF = HERE / "COMMITTEE_QUESTIONS.pdf"
REPORT = HERE / "review/COMMITTEE_EXPORT_CHECK.json"
HEADER = "Thesis defence | Committee questions"
FIELDS = {
    "Question": "question",
    "Short answer": "short_answer",
    "Supporting detail": "supporting_detail",
    "Sources": "sources",
    "Limits and evidence": "limits",
    "Supporting slide": "supporting_slide",
}
# D-230: counts are the accepted M1 bank, not inferred from an export.
EXPECTED_QUESTIONS = 134
EXPECTED_BASIC = 16


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


# Copied verbatim from export_defense_script.py at 87e87ae (deleted, D-385).
def plain_inline(value):
    """Strip presentation markup while retaining every authored word."""
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", value)
    return value.replace("**", "").replace("`", "")


# Copied verbatim from export_defense_script.py at 87e87ae (deleted, D-385).
def add_page_field(paragraph):
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def blocks(text):
    """Join Markdown paragraph wraps without changing authored words."""
    return [" ".join(part.splitlines()).strip()
            for part in re.split(r"\n\s*\n", text.strip()) if part.strip()]


def validate_source():
    index = json.loads(INDEX.read_text())
    if index["bank_sha256"] != digest(SOURCE):
        raise AssertionError("The reviewed index does not describe this bank.")
    questions = []
    group = None
    groups = []
    for block in blocks(SOURCE.read_text()):
        if block.startswith("## "):
            group = block[3:]
            groups.append(group)
        elif block.startswith("### "):
            questions.append({"id": block[4:], "group": group, "fields": {}})
        elif questions:
            label, separator, value = block.partition(": ")
            if not separator or label not in FIELDS:
                raise AssertionError("Unexpected question field: " + block[:100])
            if label in questions[-1]["fields"]:
                raise AssertionError("Repeated question field: " + label)
            questions[-1]["fields"][label] = value
    if len(questions) != EXPECTED_QUESTIONS:
        raise AssertionError("Question count differs from the accepted bank.")
    if len({q["id"] for q in questions}) != len(questions):
        raise AssertionError("Repeated question ID.")
    if groups[0] != "Basic questions":
        raise AssertionError("Basic questions must be the first separate section.")
    basic = [q for q in questions if q["group"] == "Basic questions"]
    if len(basic) != EXPECTED_BASIC or questions[:EXPECTED_BASIC] != basic:
        raise AssertionError("The separate first Basic section must contain 16 questions.")
    if len(index["questions"]) != len(questions):
        raise AssertionError("Index count differs from bank.")
    for observed, accepted in zip(questions, index["questions"]):
        if observed["id"] != accepted["id"] or observed["group"] != accepted["group"]:
            raise AssertionError("Question order/group differs from the reviewed index.")
        if accepted["independent_review_verdict"] != "ACCEPT":
            raise AssertionError("Question lacks independent source acceptance.")
        if list(observed["fields"]) != list(FIELDS):
            raise AssertionError("Missing, extra or reordered answer field: " + observed["id"])
        for label, key in FIELDS.items():
            if observed["fields"][label] != accepted[key] or not accepted[key].strip():
                raise AssertionError("Bank/index field mismatch: " + observed["id"] + "/" + label)
    return questions, groups


def make_document():
    doc = Document()
    section = doc.sections[0]
    # D-230: the accepted script's Letter page and margin geometry is reused.
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.widow_control = True
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    header = section.header.paragraphs[0]
    header.text = HEADER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)
    for paragraph in (header, footer):
        for run in paragraph.runs:
            run.font.name = "DejaVu Sans"
            run.font.size = Pt(9)
    doc.core_properties.title = "Committee questions for rehearsal"
    doc.core_properties.author = "Lanqing Luo"
    doc.core_properties.created = datetime(2026, 9, 29)
    doc.core_properties.modified = datetime(2026, 9, 29)
    doc.core_properties.comments = "Verbatim reviewed COMMITTEE_QUESTIONS.md; D-230."
    expected = []
    for block in blocks(SOURCE.read_text()):
        heading = re.fullmatch(r"(#{1,3}) (.+)", block)
        if heading:
            value = plain_inline(heading.group(2))
            style = {1: "Title", 2: "Heading 1", 3: "Heading 2"}[len(heading.group(1))]
            paragraph = doc.add_paragraph(value, style)
        else:
            if block.startswith(("|", "```", "- ", "* ")):
                raise ValueError("Unsupported bank markup: " + block[:100])
            value = plain_inline(block)
            paragraph = doc.add_paragraph(style="Normal")
            label, separator, rest = value.partition(": ")
            if separator and label in FIELDS:
                paragraph.add_run(label + ": ").bold = True
                paragraph.add_run(rest)
                if label == "Question":
                    paragraph.paragraph_format.keep_with_next = True
            else:
                paragraph.add_run(value)
        expected.append(value)
    doc.save(DOCX)
    observed = [p.text for p in Document(DOCX).paragraphs if p.text]
    if observed != expected:
        raise AssertionError("DOCX changed bank wording or order.")
    return expected


def normalized(value):
    """Character sequence; compare_body_text separately checks word boundaries."""
    return "".join(value.split())


def whitespace_boundaries(value):
    """Map an internal whitespace gap to its non-whitespace character offset."""
    boundaries = {}
    offset = 0
    for part in re.findall(r"\s+|\S+", value.strip()):
        if part.isspace():
            boundaries[offset] = part
        else:
            offset += len(part)
    return boundaries


def compare_body_text(expected, observed):
    expected_text = normalized(expected)
    observed_text = normalized(observed)
    if observed_text != expected_text:
        position = next((i for i, (a, b) in enumerate(zip(expected_text, observed_text)) if a != b),
                        min(len(expected_text), len(observed_text)))
        raise AssertionError("Whole PDF text differs at normalized character %d: expected %r, got %r" %
                             (position, expected_text[position:position+100], observed_text[position:position+100]))
    authored = whitespace_boundaries(expected)
    extracted = whitespace_boundaries(observed)
    missing = sorted(set(authored) - set(extracted))
    additions = sorted(set(extracted) - set(authored))
    unjustified = [offset for offset in additions
                   if not any(character in extracted[offset] for character in "\n\r\f")]
    if missing or unjustified:
        raise AssertionError("PDF word boundaries differ: missing authored gaps %r; added same-line gaps %r" %
                             (missing, unjustified))
    return {"authored_whitespace_boundaries": len(authored),
            "missing_authored_boundaries": len(missing),
            "additional_extraction_linewrap_boundaries": len(additions),
            "unjustified_same_line_boundaries": len(unjustified)}


def verify_pdf(expected, questions, groups):
    extracted = subprocess.run(["pdftotext", "-layout", str(PDF), "-"],
                               check=True, capture_output=True, text=True).stdout
    cleaned = "\n".join(line for line in extracted.splitlines()
                        if line.strip() != HEADER and not re.fullmatch(r"Page\s+\d+", line.strip()))
    boundary_check = compare_body_text("\n".join(expected), cleaned)
    expected_text = normalized("\n".join(expected))
    ids = re.findall(r"^\s*(Q(?:[AB]|\d+)\.(?:N)?\d+)\s*$", cleaned, re.MULTILINE)
    if ids != [q["id"] for q in questions]:
        raise AssertionError("PDF question headings differ from the accepted order.")
    found_groups = [line.strip() for line in cleaned.splitlines() if line.strip() in groups]
    if found_groups != groups:
        raise AssertionError("PDF group headings differ from bank.")
    return {"pdf": PDF.name, "pdf_sha256": digest(PDF),
            "pdf_pages": len([page for page in extracted.split("\f") if page.strip()]),
            "pdf_whole_text_preserved": True, "pdf_question_ids": ids,
            "pdf_text_normalization": "Markdown emphasis/code delimiters omitted and links expanded to label (target); repeated page header and Page N footer removed. All authored whitespace boundaries must remain. Added gaps are allowed only when their extracted whitespace contains a line break. Same-line word splitting and missing word gaps are rejected. All non-whitespace characters and their order must match the complete DOCX body.",
            "pdf_word_boundary_check": boundary_check,
            "normalized_body_sha256": hashlib.sha256(expected_text.encode()).hexdigest()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", action="store_true", help="Also export and verify the PDF with LibreOffice.")
    args = parser.parse_args()
    questions, groups = validate_source()
    expected = make_document()
    report = {"status": "PASS", "decision": "D-230", "source": SOURCE.name,
              "source_sha256": digest(SOURCE), "index": str(INDEX.relative_to(HERE)),
              "index_sha256": digest(INDEX), "exporter_sha256": digest(Path(__file__)),
              "docx": DOCX.name, "docx_sha256": digest(DOCX), "paragraphs": len(expected),
              "docx_wording_and_order_match": True, "questions": len(questions),
              "basic_questions": EXPECTED_BASIC, "groups": groups,
              "bank_index_field_comparisons": len(questions) * len(FIELDS),
              "independent_source_verdicts": {"ACCEPT": len(questions)}}
    if args.pdf:
        with tempfile.TemporaryDirectory(prefix="qa-") as folder:
            scratch = Path(folder)
            snapshot = scratch / DOCX.name
            shutil.copyfile(DOCX, snapshot)
            environment = dict(os.environ, GSETTINGS_BACKEND="memory")
            subprocess.run(["soffice", "-env:UserInstallation=" + (scratch / "lo").as_uri(),
                            "--headless", "--convert-to", "pdf", "--outdir", str(scratch), str(snapshot)],
                           check=True, env=environment)
            exported = scratch / PDF.name
            if not exported.exists():
                raise RuntimeError("LibreOffice did not produce the committee PDF.")
            shutil.copyfile(exported, PDF)
        report.update(verify_pdf(expected, questions, groups))
    REPORT.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps({key: value for key, value in report.items() if key != "pdf_question_ids"}, indent=2))


if __name__ == "__main__":
    main()
