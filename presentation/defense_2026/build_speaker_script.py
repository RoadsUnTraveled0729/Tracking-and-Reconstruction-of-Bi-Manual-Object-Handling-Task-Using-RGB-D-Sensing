#!/usr/bin/env python3
"""Build and check the defence speaker script from a tracked text source.

Source of truth: presentation/defense_2026/speaker_script.json.
Style template: the current Thesis_Defence_2026_Speaker_Script.docx (body
cleared, styles kept), the technique of the author's builder
review/personal_review_20261005/build_documents.py.

Modes
  --extract   read the current docx, write speaker_script.json
  (default)   build the docx from the JSON
  --check     lint the JSON (slides 1-46 word counts, timing cap, style hits)

Timing constants (source: review/personal_review_20261005/build_documents.py
lines 41-44): 110 words per minute, 2 seconds per slide overhead, per-slide
extra pause, remainder to the cap assigned to slide 38, cap read from timing.cap_s of the JSON (1620 seconds since
D-388) for slides 1-46.
"""
import argparse
import json
import math
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DEFAULT_DOCX = HERE / "Thesis_Defence_2026_Speaker_Script.docx"
DEFAULT_JSON = HERE / "speaker_script.json"

WPM = 110
OVERHEAD_S = 2
CAP_S = 1620  # default only; the JSON timing.cap_s is authoritative (D-388)
REMAINDER_SLIDE = 38
TIMED_LAST = 46
EXTRA_SUFFIX = " Allow {n} seconds for viewing or pointing within this slot."
EXTRA_RE = re.compile(r" Allow (\d+) seconds for viewing or pointing within this slot\.$")
WORD_RE = re.compile(r"\b[\w]+(?:[-'][\w]+)*\b")

SECTION_MAP = {
    "Block 1": [("Motivation and goal", 1, 4)],
    "Block 2": [("Data preprocessing", 6, 10), ("Kinematic model", 11, 15),
                ("Object tracking", 16, 18), ("Partial occlusion", 19, 24),
                ("From modules to one system", 25, 30)],
    "Block 3": [("Results", 31, 38),
                ("Limitations, future directions and summary", 39, 41),
                ("References and close", 42, 46)],
}


def wordcount(text):
    return len(WORD_RE.findall(text or ""))


def fmt(v):
    return "%02d:%02d" % (v // 60, v % 60)


def base_seconds(spoken):
    return math.ceil(wordcount(spoken) / WPM * 60) + OVERHEAD_S


def compute_timing(data):
    """Return {n: (start_s, end_s, seconds)} for timed slides."""
    slides = data["slides"]
    secs = {}
    for n in range(1, TIMED_LAST + 1):
        s = slides[str(n)]
        secs[n] = base_seconds(s["spoken"]) + int(s.get("timing_extra_s", 0))
    total = sum(secs.values())
    cap = int(data.get("timing", {}).get("cap_s", CAP_S))
    if total > cap:
        raise SystemExit("ERROR: slides 1-%d need %d s, cap is %d s" % (TIMED_LAST, total, cap))
    secs[REMAINDER_SLIDE] += cap - total
    out, t = {}, 0
    for n in range(1, TIMED_LAST + 1):
        out[n] = (t, t + secs[n], secs[n])
        t += secs[n]
    return out


# ---------------------------------------------------------------- extract
def extract(docx_path, json_path):
    import docx
    d = docx.Document(str(docx_path))
    P = d.paragraphs
    title = P[0].text
    assert P[0].style.name == "Title"
    i = 1
    header = []
    while P[i].style.name == "Normal":
        header.append(P[i].text)
        i += 1
    blocks, slides = [], {}
    cur = None
    pending = {}
    while i < len(P):
        p = P[i]
        st = p.style.name
        if st == "Heading 1":
            pb = bool(p.paragraph_format.page_break_before)
            cur = {"title": p.text, "slides": [None, None], "intro": None,
                   "page_break_before": pb, "sections": []}
            blocks.append(cur)
            i += 1
            if P[i].style.name == "Normal":
                cur["intro"] = P[i].text
                i += 1
            continue
        m = re.match(r"Slide (\d+) (.*)$", p.text) if st == "Heading 2" else None
        if not m:
            raise SystemExit("unexpected paragraph %d (%s): %r" % (i, st, p.text[:60]))
        n = int(m.group(1))
        entry = {"title": m.group(2), "cue": None, "spoken": None, "timing_extra_s": 0}
        timing_line = P[i + 1]
        assert timing_line.style.name == "Rehearsal timing", (n, timing_line.style.name)
        pending[n] = timing_line.text
        i += 2
        cues = []
        while i < len(P) and P[i].style.name == "Stage cue":
            cues.append(P[i].text)
            i += 1
        assert 1 <= len(cues) <= 2, (n, cues)
        cue = cues[0]
        if len(cues) == 2:
            assert cues[1] == "No narration.", (n, cues)
        elif cues == ["No narration."]:
            pass
        mm = EXTRA_RE.search(cue)
        if mm:
            entry["timing_extra_s"] = int(mm.group(1))
            cue = cue[:mm.start()]
        entry["cue"] = cue
        if i < len(P) and P[i].style.name == "Spoken script":
            entry["spoken"] = P[i].text
            i += 1
        slides[str(n)] = entry
        if cur["slides"][0] is None:
            cur["slides"][0] = n
        cur["slides"][1] = n

    # Recover timing_extra_s from the timing lines versus the formula.
    unresolved = []
    shown = {}
    for n in range(1, TIMED_LAST + 1):
        m = re.match(r"(\d\d):(\d\d) to (\d\d):(\d\d)  \|  (\d+) seconds$", pending[n])
        assert m, (n, pending[n])
        a = int(m.group(1)) * 60 + int(m.group(2))
        b = int(m.group(3)) * 60 + int(m.group(4))
        sec = int(m.group(5))
        assert b - a == sec, (n, pending[n])
        shown[n] = (a, b, sec)
    for n in range(1, TIMED_LAST + 1):
        e = slides[str(n)]
        diff = shown[n][2] - base_seconds(e["spoken"])
        if n == REMAINDER_SLIDE:
            continue
        if diff < 0:
            unresolved.append((n, diff))
        if diff != e["timing_extra_s"]:
            # the Allow-suffix in the cue disagrees with the timing line
            unresolved.append((n, "cue suffix %d vs timing diff %d" % (e["timing_extra_s"], diff)))
        e["timing_extra_s"] = max(diff, 0)
    # slide 38: extra is the cue suffix; the remainder is derived by the build
    for n in range(TIMED_LAST + 1, 78):
        if n in slides:
            assert pending[n] == "On demand", (n, pending[n])

    sections_map = {}
    for b in blocks:
        for key, val in SECTION_MAP.items():
            if b["title"].startswith(key + " "):
                b["sections"] = [{"title": t, "slides": [a, z]} for t, a, z in val]
    data = {
        "title": title,
        "header": header,
        "emit_section_headings": False,
        "timing": {"wpm": WPM, "overhead_s": OVERHEAD_S, "cap_s": CAP_S,
                   "remainder_slide": REMAINDER_SLIDE, "timed_slides": [1, TIMED_LAST]},
        "blocks": blocks,
        "slides": slides,
    }
    # verify reproduction of every timing line
    rebuilt = compute_timing(data)
    bad = [(n, shown[n], rebuilt[n]) for n in rebuilt if shown[n] != rebuilt[n]]
    json_path.write_text(json.dumps(data, indent=1, ensure_ascii=True) + "\n", encoding="ascii")
    print("extracted %d slides, %d blocks -> %s" % (len(slides), len(blocks), json_path))
    print("timing lines reproduced exactly: %d of %d" % (len(rebuilt) - len(bad), len(rebuilt)))
    for n, a, b in bad:
        print("  TIMING MISMATCH slide %d: docx %s rebuilt %s" % (n, a, b))
    for u in unresolved:
        print("  UNRESOLVED extra: slide %s: %s" % u)
    return 0 if not bad and not unresolved else 1


# ------------------------------------------------------------------ build
def ensure_heading3(doc):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    st = doc.styles["Heading 3"]
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.font.size = Pt(12)  # one step below Heading 2 (13 pt)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "DejaVu Sans")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    st.paragraph_format.keep_with_next = True


def build(json_path, out_path, template_path):
    import docx
    from docx.oxml.ns import qn
    data = json.loads(json_path.read_text(encoding="ascii"))
    timing = compute_timing(data)
    doc = docx.Document(str(template_path))
    for e in list(doc.element.body):
        if e.tag != qn("w:sectPr"):
            doc.element.body.remove(e)
    emit_sections = bool(data.get("emit_section_headings"))
    if emit_sections:
        ensure_heading3(doc)
    doc.add_paragraph(data["title"], "Title")
    for line in data["header"]:
        doc.add_paragraph(line)
    slides = data["slides"]
    for b in data["blocks"]:
        p = doc.add_paragraph(b["title"], "Heading 1")
        if b.get("page_break_before"):
            p.paragraph_format.page_break_before = True
        if b.get("intro"):
            doc.add_paragraph(b["intro"])
        first, last = b["slides"]
        starts = {}
        if emit_sections:
            for s in b.get("sections", []):
                if s.get("title"):
                    starts[s["slides"][0]] = s["title"]
        for n in range(first, last + 1):
            if n in starts:
                doc.add_paragraph(starts[n], "Heading 3")
            s = slides[str(n)]
            doc.add_paragraph("Slide %d %s" % (n, s["title"]), "Heading 2")
            if n in timing:
                a, z, sec = timing[n]
                doc.add_paragraph("%s to %s  |  %d seconds" % (fmt(a), fmt(z), sec), "Rehearsal timing")
            else:
                doc.add_paragraph("On demand", "Rehearsal timing")
            cue = s["cue"]
            if s.get("timing_extra_s"):
                cue += EXTRA_SUFFIX.format(n=s["timing_extra_s"])
            doc.add_paragraph(cue, "Stage cue")
            if s["spoken"]:
                doc.add_paragraph(s["spoken"], "Spoken script")
            else:
                doc.add_paragraph("No narration.", "Stage cue")
    doc.core_properties.title = data["title"]
    doc.core_properties.author = "Lanqing Luo"
    doc.save(str(out_path))
    print("built %s" % out_path)
    return 0


# ------------------------------------------------------------------ check
NUMWORDS = ("one two three four five six seven eight nine ten eleven twelve thirteen "
            "fourteen fifteen sixteen seventeen eighteen nineteen twenty hundred thousand").split()
FILLERS = ["as we can see", "basically", "in other words", "very ", "quite ", "really ",
           "Thesis:", "/home/", "New_SandBox"]
R_RE = re.compile(r"\bR[1-9]\b")


def sentences(text):
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]


def norm(s):
    return " ".join(re.findall(r"[a-z0-9]+", s.lower()))


def check(json_path):
    data = json.loads(json_path.read_text(encoding="ascii"))
    slides = data["slides"]
    timing = compute_timing(data)
    hits = []
    seen = {}
    total_words = 0
    print("slide  words  seconds")
    for n in range(1, TIMED_LAST + 1):
        s = slides[str(n)]
        txt = s["spoken"] or ""
        w = wordcount(txt)
        total_words += w
        print("%5d  %5d  %7d" % (n, w, timing[n][2]))
        if ";" in txt:
            hits.append((n, "semicolon"))
        low = txt.lower()
        for f in FILLERS:
            hay = txt if f == "Thesis:" or f.startswith("/") or f == "New_SandBox" else low
            if f.lower() == "very " or f.lower() == "quite " or f.lower() == "really ":
                if re.search(r"\b" + f.strip() + r"\b", low):
                    hits.append((n, "filler %r" % f.strip()))
            elif f in hay:
                hits.append((n, "filler %r" % f))
        if R_RE.search(txt):
            hits.append((n, "internal name %s" % R_RE.search(txt).group(0)))
        for sent in sentences(txt):
            first = re.match(r"[\w']+", sent)
            if first:
                fw = first.group(0).lower()
                if fw[0].isdigit() or fw in NUMWORDS:
                    hits.append((n, "sentence starts with number: %r" % sent[:50]))
            ns = norm(sent)
            if len(ns.split()) >= 8:
                seen.setdefault(ns, set()).add(n)
    for ns, ns_slides in seen.items():
        if len(ns_slides) > 1:
            hits.append((sorted(ns_slides)[0], "sentence repeated in slides %s: %r"
                         % (sorted(ns_slides), ns[:60])))
    total_s = timing[TIMED_LAST][1]
    cap_s = int(data.get("timing", {}).get("cap_s", CAP_S))
    ok_cap = total_s <= cap_s
    print("total words (slides 1-%d): %d" % (TIMED_LAST, total_words))
    print("total seconds: %d, cap %d: %s" % (total_s, cap_s, "PASS" if ok_cap else "FAIL"))
    # header consistency (stated word count and checkpoints)
    hdr = " ".join(data["header"])
    m = re.search(r"contains ([\d,]+) spoken words", hdr)
    if m and int(m.group(1).replace(",", "")) != total_words:
        hits.append((0, "header states %s words, computed %d" % (m.group(1), total_words)))
    for label, n in (("motivation and goal", 4), ("the offline system", 30)):
        m = re.search(label + r" at (\d\d:\d\d)", hdr)
        if m and m.group(1) != fmt(timing[n][1]):
            hits.append((0, "header checkpoint %r %s, computed %s" % (label, m.group(1), fmt(timing[n][1]))))
    print("hits: %d" % len(hits))
    for n, what in sorted(hits, key=lambda h: (h[0], h[1])):
        print("  slide %d: %s" % (n, what))
    return 0 if ok_cap and not hits else 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--extract", action="store_true", help="docx -> JSON")
    g.add_argument("--check", action="store_true", help="lint the JSON")
    ap.add_argument("--json", type=Path, default=DEFAULT_JSON)
    ap.add_argument("--template", type=Path, default=DEFAULT_DOCX,
                    help="docx used as extraction source and as style template")
    ap.add_argument("--out", type=Path, default=DEFAULT_DOCX, help="output docx for the build")
    a = ap.parse_args()
    if a.extract:
        return extract(a.template, a.json)
    if a.check:
        return check(a.json)
    return build(a.json, a.out, a.template)


if __name__ == "__main__":
    sys.exit(main())
